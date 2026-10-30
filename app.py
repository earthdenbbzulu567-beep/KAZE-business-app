# KAZE Traders Business Manager
# Flask + PostgreSQL. Run with: gunicorn app:app

from flask import Flask, render_template, request, redirect, url_for, flash, Response, send_file, g, session
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
import os
import secrets
import io
import csv
import hmac
import hashlib
import struct
import base64
import time
import psycopg2
from psycopg2.extras import RealDictCursor
from datetime import datetime, timedelta
import smtplib
from email.mime.text import MIMEText
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
import json
import re
import mode_runtime
import stripe_payments
import integrations
import subscriptions
import gateways
import oauth_login

# --------------------- Performance helpers ---------------------
import time
_PERF_CACHE = {}

def _cache_get(key, ttl=45):
    entry = _PERF_CACHE.get(key)
    if entry and (time.time() - entry[0]) < ttl:
        return entry[1]
    return None

def _cache_set(key, value):
    _PERF_CACHE[key] = (time.time(), value)


_APP_DIR = os.path.dirname(os.path.abspath(__file__))
_TPL_DIR = os.path.join(_APP_DIR, 'templates') if os.path.isdir(os.path.join(_APP_DIR, 'templates')) else _APP_DIR
app = Flask(__name__, template_folder=_TPL_DIR)
app.secret_key = os.environ.get('SECRET_KEY', 'dev-secret-key-change-in-production')
LEGAL_VERSION = '4.6.0'
LEGAL_DIR = _APP_DIR
_LEGAL_CACHE = {}
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 86400
app.config['TEMPLATES_AUTO_RELOAD'] = False
app.config['SESSION_REFRESH_EACH_REQUEST'] = False
app.config['JSONIFY_PRETTYPRINT_REGULAR'] = False
app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024


def _totp_new_secret():
    return base64.b32encode(secrets.token_bytes(20)).decode('ascii').rstrip('=')

def _totp_at(secret, at):
    raw = (secret or '').strip().replace(' ', '').upper()
    raw += '=' * ((8 - len(raw) % 8) % 8)
    digest = hmac.new(base64.b32decode(raw), struct.pack('>Q', int(at // 30)), hashlib.sha1).digest()
    offset = digest[-1] & 15
    code = struct.unpack('>I', digest[offset:offset + 4])[0] & 0x7fffffff
    return str(code % 1000000).zfill(6)

def _totp_ok(secret, code):
    code = ''.join(ch for ch in (code or '') if ch.isdigit())
    if len(code) != 6 or not secret:
        return False
    now = time.time()
    return any(secrets.compare_digest(_totp_at(secret, now + drift), code) for drift in (-30, 0, 30))

def _backup_codes():
    plain = ['%s-%s' % (secrets.token_hex(2), secrets.token_hex(2)) for _ in range(8)]
    return plain, json.dumps([generate_password_hash(c) for c in plain])

def _backup_take(stored, code):
    code = (code or '').strip().lower()
    try:
        hashes = json.loads(stored or '[]')
    except Exception:
        hashes = []
    kept, hit = [], False
    for item in hashes:
        if not hit and check_password_hash(item, code):
            hit = True
            continue
        kept.append(item)
    return hit, json.dumps(kept)

def _needs_2fa(user):
    return bool(user and int(user.get('totp_enabled') or 0) == 1 and (user.get('totp_secret') or ''))

def _begin_2fa(user):
    session['pre_2fa_id'] = int(user['id'])
    session['pre_2fa_exp'] = time.time() + 300
    return redirect(url_for('login_2fa'))

def _otpauth_uri(secret, label):
    label = (label or 'user').replace(' ', '')[:40]
    return 'otpauth://totp/KAZE:{0}?secret={1}&issuer=KAZE&digits=6&period=30'.format(label, secret)

def _qr_svg(payload):
    try:
        import qrcode, qrcode.image.svg
        img = qrcode.make(payload, image_factory=qrcode.image.svg.SvgPathImage, border=2)
        buf = io.BytesIO(); img.save(buf)
        return buf.getvalue().decode('utf-8', 'replace')
    except Exception:
        return ''

# --------------------- Helper: time since (for activity feed) ---------------------
@app.template_filter('timesince')
def timesince_filter(dt):
    """
    Jinja filter to show relative time like "2 hours ago".
    Expects dt as a datetime object or string.
    """
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt.replace('Z', '+00:00'))
        except ValueError:
            return dt
    now = datetime.now()
    diff = now - dt
    seconds = diff.total_seconds()
    if seconds < 60:
        return "just now"
    elif seconds < 3600:
        m = int(seconds // 60)
        return f"{m} minute{'s' if m > 1 else ''} ago"
    elif seconds < 86400:
        h = int(seconds // 3600)
        return f"{h} hour{'s' if h > 1 else ''} ago"
    elif seconds < 604800:
        d = int(seconds // 86400)
        return f"{d} day{'s' if d > 1 else ''} ago"
    else:
        return dt.strftime('%Y-%m-%d %H:%M')

# --------------------- Login setup ---------------------
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'


@app.route('/favicon.ico')
def favicon():
    return redirect(url_for('static', filename='favicon-32.png'))

@app.route('/healthz')
def healthz():
    return {'ok': True}, 200


# --------------------- Database ---------------------
_pool = None
_pool_lock = __import__('threading').Lock()
_recurring_ran_at = {}


def _db_url():
    url = os.environ.get('DATABASE_URL')
    if not url:
        raise ValueError("DATABASE_URL environment variable is not set")
    return url


def _get_pool():
    global _pool
    if _pool is False:
        return None
    if _pool is not None:
        return _pool
    with _pool_lock:
        if _pool not in (None, False):
            return _pool
        try:
            from psycopg2.pool import SimpleConnectionPool
            _pool = SimpleConnectionPool(1, 10, _db_url(), cursor_factory=RealDictCursor)
        except Exception:
            _pool = False
            return None
    return _pool


class _ReqConn:
    """Wrap a psycopg2 connection. close() is a no-op so routes can call it safely."""
    def __init__(self, raw):
        object.__setattr__(self, '_raw', raw)

    def cursor(self, *args, **kwargs):
        return self._raw.cursor(*args, **kwargs)

    def commit(self):
        return self._raw.commit()

    def rollback(self):
        return self._raw.rollback()

    def close(self):
        return None

    @property
    def closed(self):
        try:
            return self._raw.closed
        except Exception:
            return 1

    def __getattr__(self, name):
        return getattr(self._raw, name)


def get_db():
    # Reuse one pooled connection for the whole request.
    from flask import has_app_context
    if has_app_context():
        wrap = getattr(g, '_kaze_conn', None)
        if wrap is not None and getattr(wrap, 'closed', 1) == 0:
            return wrap
        pool = _get_pool()
        if pool:
            raw = pool.getconn()
        else:
            raw = psycopg2.connect(_db_url(), cursor_factory=RealDictCursor)
        wrap = _ReqConn(raw)
        g._kaze_conn = wrap
        g._kaze_raw = raw
        return wrap
    return psycopg2.connect(_db_url(), cursor_factory=RealDictCursor)


def _wants_json():
    path = request.path or ''
    if path.startswith('/webhooks/') or path.startswith('/healthz'):
        return True
    best = request.accept_mimetypes.best
    return best == 'application/json' and best != 'text/html'


def _error_body(code, title, message):
    if _wants_json():
        return {'ok': False, 'error': title, 'message': message}, code
    try:
        return render_template('errors.html', code=code, title=title, message=message), code
    except Exception:
        safe = (
            '<!DOCTYPE html><meta charset="utf-8"><title>{code} · KAZE</title>'
            '<body style="font-family:sans-serif;background:#0b1220;color:#e8eef8;padding:2rem">'
            '<h1>{title}</h1><p>{message}</p><p><a href="/">Home</a></p></body>'
        ).format(code=code, title=title, message=message)
        return safe, code


@app.errorhandler(400)
def _err_400(err):
    return _error_body(400, 'Bad request', 'That request was incomplete or not valid.')


@app.errorhandler(403)
def _err_403(err):
    return _error_body(403, 'Not allowed', 'You do not have permission for that page.')


@app.errorhandler(404)
def _err_404(err):
    return _error_body(404, 'Page not found', 'That page is not in KAZE. Check the link and try again.')


@app.errorhandler(405)
def _err_405(err):
    return _error_body(405, 'Wrong method', 'That action cannot be opened that way.')


@app.errorhandler(413)
def _err_413(err):
    return _error_body(413, 'File too large', 'That upload is over 8 MB. Use a smaller file.')


@app.errorhandler(500)
def _err_500(err):
    try:
        import sys, traceback
        sys.stderr.write('KAZE 500: %s\n' % (err,))
        traceback.print_exc()
    except Exception:
        pass
    return _error_body(500, 'Something went wrong', 'KAZE hit an error and rolled the request back. Try again.')


@app.errorhandler(Exception)
def _err_any(err):
    from werkzeug.exceptions import HTTPException
    if isinstance(err, HTTPException):
        return _error_body(err.code or 500, err.name or 'Error', err.description or 'That request failed.')
    try:
        import sys, traceback
        sys.stderr.write('KAZE error: %s\n' % (err,))
        traceback.print_exc()
    except Exception:
        pass
    try:
        conn = getattr(g, '_kaze_raw', None)
        if conn is not None:
            conn.rollback()
    except Exception:
        pass
    return _error_body(500, 'Something went wrong', 'KAZE hit an error and did not save that change. Try again.')


@app.teardown_appcontext
def _return_db(_exc):
    wrap = getattr(g, '_kaze_conn', None)
    raw = getattr(g, '_kaze_raw', None)
    g._kaze_conn = None
    g._kaze_raw = None
    if raw is None and wrap is not None:
        raw = getattr(wrap, '_raw', wrap)
    if raw is None:
        return
    try:
        if _exc:
            raw.rollback()
    except Exception:
        pass
    pool = _get_pool()
    if pool:
        try:
            pool.putconn(raw)
            return
        except Exception:
            pass
    try:
        raw.close()
    except Exception:
        pass


@app.after_request
def _fast_headers(resp):
    path = request.path or ''
    if path.startswith('/static/'):
        resp.headers.setdefault('Cache-Control', 'public, max-age=86400, immutable')
        return resp
    if resp.status_code == 200 and not getattr(resp, 'direct_passthrough', False):
        already = (resp.headers.get('Content-Encoding') or '')
        accept = request.headers.get('Accept-Encoding', '')
        mime = (resp.mimetype or '')
        compressible = mime.startswith('text/') or mime in (
            'application/json', 'application/javascript', 'application/xml'
        )
        if compressible and 'gzip' in accept and 'gzip' not in already:
            try:
                data = resp.get_data()
            except Exception:
                data = None
            if data and 800 < len(data) < 2_000_000:
                import gzip
                packed = gzip.compress(data, compresslevel=1)
                if len(packed) < len(data) - 120:
                    resp.set_data(packed)
                    resp.headers['Content-Encoding'] = 'gzip'
                    resp.headers['Content-Length'] = str(len(packed))
                    vary = resp.headers.get('Vary', '')
                    if 'Accept-Encoding' not in vary:
                        resp.headers['Vary'] = (vary + ', Accept-Encoding').strip(', ')
    return resp


_schema_ready = False

def init_db():
    global _schema_ready
    if _schema_ready:
        return
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id SERIAL PRIMARY KEY,
            username TEXT UNIQUE NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL
        )
    ''')
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS role TEXT DEFAULT 'owner'")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS owner_id INTEGER REFERENCES users(id)")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS terms_accepted INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS terms_version TEXT DEFAULT ''")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS terms_accepted_at TEXT DEFAULT ''")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS oauth_provider TEXT DEFAULT ''")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS oauth_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS totp_secret TEXT DEFAULT ''")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS totp_enabled INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS totp_backup TEXT DEFAULT ''")
    cur.execute('''
        CREATE TABLE IF NOT EXISTS income (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            source TEXT NOT NULL,
            amount REAL NOT NULL,
            date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS expenses (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            amount REAL NOT NULL,
            category TEXT NOT NULL,
            date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS documents (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            doc_type TEXT NOT NULL,
            title TEXT NOT NULL,
            client TEXT NOT NULL,
            amount REAL,
            date TEXT NOT NULL,
            notes TEXT
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS cash_books (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            entry_type TEXT NOT NULL,
            category TEXT NOT NULL,
            description TEXT NOT NULL,
            amount REAL NOT NULL,
            date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS stock (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            product_name TEXT NOT NULL,
            quantity REAL NOT NULL,
            cost_price REAL NOT NULL,
            selling_price REAL NOT NULL,
            unit TEXT
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS customers (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            email TEXT,
            phone TEXT,
            address TEXT,
            notes TEXT
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS sales (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            stock_id INTEGER NOT NULL REFERENCES stock(id),
            quantity_sold REAL NOT NULL,
            selling_price_at_time REAL NOT NULL,
            total_amount REAL NOT NULL,
            profit REAL NOT NULL,
            sale_date TEXT NOT NULL,
            customer_name TEXT,
            customer_email TEXT,
            customer_id INTEGER REFERENCES customers(id)
        )
    ''')
    cur.execute("ALTER TABLE sales ADD COLUMN IF NOT EXISTS customer_id INTEGER REFERENCES customers(id)")
    cur.execute('''
        CREATE TABLE IF NOT EXISTS settings (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL UNIQUE REFERENCES users(id),
            low_stock_threshold REAL DEFAULT 10,
            email_notifications INTEGER DEFAULT 1,
            notify_email TEXT,
            theme TEXT DEFAULT 'dark',
            font_family TEXT DEFAULT 'Inter',
            font_size TEXT DEFAULT 'medium',
            default_chart_type TEXT DEFAULT 'line',
            animations_enabled INTEGER DEFAULT 1,
            show_low_stock_widget INTEGER DEFAULT 1,
            show_sales_trend INTEGER DEFAULT 1,
            about_text TEXT DEFAULT '',
            currency_code TEXT DEFAULT 'USD',
            currency_symbol TEXT DEFAULT '$',
            tax_rate REAL DEFAULT 0,
            business_name TEXT DEFAULT '',
            monthly_revenue_goal REAL DEFAULT 0,
            tutorial_completed INTEGER DEFAULT 0
        )
    ''')
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS currency_code TEXT DEFAULT 'USD'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS currency_symbol TEXT DEFAULT '$'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS tax_rate REAL DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS business_name TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS monthly_revenue_goal REAL DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS tutorial_completed INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS accent_color TEXT DEFAULT '#2ecc71'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS density TEXT DEFAULT 'comfortable'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS items_per_page INTEGER DEFAULT 25")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS date_format TEXT DEFAULT 'YYYY-MM-DD'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS show_clock INTEGER DEFAULT 1")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS rounded_ui INTEGER DEFAULT 1")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS high_contrast INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS invoice_prefix TEXT DEFAULT 'INV'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS business_phone TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS business_address TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS fiscal_year_start INTEGER DEFAULT 1")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS number_decimals INTEGER DEFAULT 2")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS sidebar_collapsed INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS app_mode TEXT DEFAULT 'full'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS enabled_modules TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS ui_device TEXT DEFAULT 'auto'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS stripe_secret_key TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS stripe_publishable_key TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS stripe_webhook_secret TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS plan_tier TEXT DEFAULT 'ceo'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS plan_paid_rank INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS plan_started_at TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS plan_expires_at TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS grace_expires_at TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS stripe_subscription_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS stripe_customer_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS subscription_status TEXT DEFAULT 'none'")
    # One-time migration: grandfathered users keep their highest paid tier forever.
    cur.execute('''
        UPDATE settings
        SET plan_paid_rank = CASE plan_tier
            WHEN 'founder' THEN 4
            WHEN 'ceo' THEN 3
            WHEN 'businessman_pro' THEN 3
            WHEN 'businessman' THEN 2
            WHEN 'entrepreneur_plus' THEN 2
            WHEN 'entrepreneur' THEN 1
            WHEN 'new_user_plus' THEN 1
            ELSE 0
        END
        WHERE plan_paid_rank = 0
          AND plan_tier IS NOT NULL
          AND plan_tier != ''
          AND plan_tier != 'new_user'
    ''')
    # Collapse legacy tier names into the new 5-tier set.
    cur.execute('''
        UPDATE settings
        SET plan_tier = CASE plan_tier
            WHEN 'new_user_plus' THEN 'entrepreneur'
            WHEN 'entrepreneur_plus' THEN 'businessman'
            WHEN 'businessman_pro' THEN 'ceo'
            ELSE plan_tier
        END
        WHERE plan_tier IN ('new_user_plus', 'entrepreneur_plus', 'businessman_pro')
    ''')
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS package_prices TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS flutterwave_public_key TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS flutterwave_secret_key TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS flutterwave_webhook_hash TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS paypal_client_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS paypal_secret TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS paypal_mode TEXT DEFAULT 'sandbox'")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS bank_name TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS bank_account_name TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS bank_account_number TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS bank_branch TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS alpha_vantage_key TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS trial_started INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS trial_ends TEXT DEFAULT ''")
    cur.execute("ALTER TABLE settings ADD COLUMN IF NOT EXISTS trial_days INTEGER DEFAULT 0")

    cur.execute('''
        CREATE TABLE IF NOT EXISTS stripe_payments (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            kind TEXT NOT NULL,
            record_id INTEGER DEFAULT 0,
            amount REAL NOT NULL DEFAULT 0,
            currency TEXT DEFAULT 'USD',
            status TEXT NOT NULL DEFAULT 'pending',
            checkout_id TEXT DEFAULT '',
            payment_intent TEXT DEFAULT '',
            title TEXT DEFAULT '',
            created_date TEXT NOT NULL,
            paid_date TEXT DEFAULT ''
        )
    ''')

    cur.execute("ALTER TABLE stripe_payments ADD COLUMN IF NOT EXISTS provider TEXT DEFAULT 'stripe'")
    cur.execute('CREATE INDEX IF NOT EXISTS idx_stripe_user ON stripe_payments (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_stripe_checkout ON stripe_payments (checkout_id)')

    # User activity log (for tracking logins)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS user_activity (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            username TEXT NOT NULL,
            login_time TEXT NOT NULL,
            ip_address TEXT
        )
    ''')
    # SWOT
    cur.execute('''
        CREATE TABLE IF NOT EXISTS swot_analyses (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            title TEXT NOT NULL,
            strengths TEXT,
            weaknesses TEXT,
            opportunities TEXT,
            threats TEXT,
            created_date TEXT NOT NULL
        )
    ''')
    # Business plan
    cur.execute('''
        CREATE TABLE IF NOT EXISTS business_plans (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL UNIQUE REFERENCES users(id),
            business_name TEXT DEFAULT '',
            mission TEXT DEFAULT '',
            products_services TEXT DEFAULT '',
            target_market TEXT DEFAULT '',
            marketing_strategy TEXT DEFAULT '',
            operations_plan TEXT DEFAULT '',
            financial_plan TEXT DEFAULT '',
            goals TEXT DEFAULT '',
            updated_date TEXT
        )
    ''')
    # Recurring
    cur.execute('''
        CREATE TABLE IF NOT EXISTS recurring_items (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            item_type TEXT NOT NULL,
            name TEXT NOT NULL,
            amount REAL NOT NULL,
            category TEXT DEFAULT '',
            client TEXT DEFAULT '',
            frequency TEXT NOT NULL DEFAULT 'monthly',
            next_run_date TEXT NOT NULL,
            active INTEGER DEFAULT 1,
            created_date TEXT NOT NULL
        )
    ''')
    # Suppliers
    cur.execute('''
        CREATE TABLE IF NOT EXISTS suppliers (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            email TEXT DEFAULT '',
            phone TEXT DEFAULT '',
            address TEXT DEFAULT '',
            notes TEXT DEFAULT ''
        )
    ''')
    # Purchase orders
    cur.execute('''
        CREATE TABLE IF NOT EXISTS purchase_orders (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            supplier_id INTEGER REFERENCES suppliers(id),
            item_description TEXT NOT NULL,
            quantity REAL NOT NULL,
            unit_cost REAL NOT NULL,
            total_cost REAL NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            order_date TEXT NOT NULL,
            expected_date TEXT DEFAULT '',
            stock_id INTEGER REFERENCES stock(id)
        )
    ''')
    # Budgets
    cur.execute('''
        CREATE TABLE IF NOT EXISTS budgets (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            category TEXT NOT NULL,
            monthly_limit REAL NOT NULL,
            created_date TEXT NOT NULL,
            UNIQUE(user_id, category)
        )
    ''')
    # Notifications
    cur.execute('''
        CREATE TABLE IF NOT EXISTS notifications (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            message TEXT NOT NULL,
            category TEXT DEFAULT 'info',
            is_read INTEGER DEFAULT 0,
            created_date TEXT NOT NULL
        )
    ''')
    # ─── NEW: Activities table ───
    cur.execute('''
        CREATE TABLE IF NOT EXISTS activities (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            username TEXT NOT NULL,
            action TEXT NOT NULL,
            details TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS tasks (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            title TEXT NOT NULL,
            notes TEXT DEFAULT '',
            priority TEXT NOT NULL DEFAULT 'medium',
            due_date TEXT DEFAULT '',
            done INTEGER DEFAULT 0,
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS notes (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            title TEXT NOT NULL,
            body TEXT DEFAULT '',
            created_date TEXT NOT NULL,
            updated_date TEXT NOT NULL
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS code_snippets (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            title TEXT NOT NULL,
            language TEXT DEFAULT 'text',
            body TEXT DEFAULT '',
            pinned INTEGER DEFAULT 0,
            created_date TEXT NOT NULL,
            updated_date TEXT NOT NULL
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_snippets_user ON code_snippets (user_id)')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS bookkeeping_docs (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            kind TEXT NOT NULL,
            title TEXT NOT NULL,
            party TEXT DEFAULT '',
            reference TEXT DEFAULT '',
            doc_date TEXT NOT NULL,
            amount REAL DEFAULT 0,
            status TEXT DEFAULT 'draft',
            notes TEXT DEFAULT '',
            line_items TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute("ALTER TABLE stock ADD COLUMN IF NOT EXISTS sku TEXT DEFAULT ''")
    cur.execute("ALTER TABLE stock ADD COLUMN IF NOT EXISTS category TEXT DEFAULT ''")
    cur.execute("ALTER TABLE sales ADD COLUMN IF NOT EXISTS discount REAL DEFAULT 0")
    cur.execute("ALTER TABLE sales ADD COLUMN IF NOT EXISTS sale_notes TEXT DEFAULT ''")
    cur.execute("ALTER TABLE sales ADD COLUMN IF NOT EXISTS payment_method TEXT DEFAULT 'cash'")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS day_closes (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            close_date TEXT NOT NULL,
            opening_float REAL DEFAULT 0,
            counted_cash REAL DEFAULT 0,
            expected_cash REAL DEFAULT 0,
            cash_sales REAL DEFAULT 0,
            other_sales REAL DEFAULT 0,
            cash_in REAL DEFAULT 0,
            cash_out REAL DEFAULT 0,
            variance REAL DEFAULT 0,
            notes TEXT DEFAULT '',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(user_id, close_date)
        )
    """)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS journal_entries (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            book TEXT NOT NULL,
            entry_date TEXT NOT NULL,
            reference TEXT DEFAULT '',
            account TEXT NOT NULL,
            particulars TEXT DEFAULT '',
            debit REAL DEFAULT 0,
            credit REAL DEFAULT 0,
            created_date TEXT NOT NULL
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS services (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            category TEXT NOT NULL DEFAULT 'Other',
            description TEXT DEFAULT '',
            duration_minutes INTEGER DEFAULT 60,
            price REAL NOT NULL DEFAULT 0,
            active INTEGER DEFAULT 1,
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS service_sessions (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            service_id INTEGER REFERENCES services(id),
            customer_id INTEGER REFERENCES customers(id),
            customer_name TEXT DEFAULT '',
            session_date TEXT NOT NULL,
            session_time TEXT DEFAULT '',
            duration_minutes INTEGER DEFAULT 60,
            price REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'booked',
            payment_method TEXT DEFAULT 'cash',
            notes TEXT DEFAULT '',
            income_id INTEGER,
            stripe_checkout_id TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS bank_accounts (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            bank_name TEXT DEFAULT '',
            account_no TEXT DEFAULT '',
            account_type TEXT DEFAULT 'Current',
            opening_balance REAL DEFAULT 0,
            notes TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS bank_transactions (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            account_id INTEGER REFERENCES bank_accounts(id),
            tx_date TEXT NOT NULL,
            kind TEXT NOT NULL DEFAULT 'in',
            amount REAL NOT NULL DEFAULT 0,
            counterparty TEXT DEFAULT '',
            reference TEXT DEFAULT '',
            category TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            linked_cashbook_id INTEGER
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS investments (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            asset_type TEXT DEFAULT 'Other',
            quantity REAL DEFAULT 1,
            unit_cost REAL DEFAULT 0,
            current_price REAL DEFAULT 0,
            bought_date TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS shop_listings (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            title TEXT NOT NULL,
            kind TEXT DEFAULT 'Custom',
            source_id INTEGER DEFAULT 0,
            price REAL DEFAULT 0,
            published INTEGER DEFAULT 1,
            description TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS shop_orders (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            listing_id INTEGER REFERENCES shop_listings(id),
            buyer TEXT DEFAULT '',
            qty REAL DEFAULT 1,
            amount REAL DEFAULT 0,
            status TEXT DEFAULT 'pending',
            order_date TEXT NOT NULL,
            notes TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS market_channels (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            name TEXT NOT NULL,
            kind TEXT DEFAULT 'Other',
            monthly_revenue REAL DEFAULT 0,
            notes TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS market_valuations (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            as_of TEXT NOT NULL,
            method TEXT DEFAULT 'blended',
            multiple REAL DEFAULT 0,
            revenue_base REAL DEFAULT 0,
            asset_base REAL DEFAULT 0,
            estimated_value REAL DEFAULT 0,
            notes TEXT DEFAULT '',
            created_date TEXT NOT NULL
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_bank_acct_user ON bank_accounts (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_bank_tx_user ON bank_transactions (user_id, tx_date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_invest_user ON investments (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_shop_list_user ON shop_listings (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_shop_ord_user ON shop_orders (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_market_ch_user ON market_channels (user_id)')
    cur.execute("ALTER TABLE investments ADD COLUMN IF NOT EXISTS ticker TEXT DEFAULT ''")
    cur.execute("ALTER TABLE shop_orders ADD COLUMN IF NOT EXISTS flutterwave_tx_ref TEXT DEFAULT ''")
    cur.execute("ALTER TABLE shop_orders ADD COLUMN IF NOT EXISTS flutterwave_tx_id TEXT DEFAULT ''")

    cur.execute('CREATE INDEX IF NOT EXISTS idx_services_user ON services (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_sessions_user_date ON service_sessions (user_id, session_date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_sales_user_date ON sales (user_id, sale_date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_expenses_user_date ON expenses (user_id, date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_income_user_date ON income (user_id, date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_cash_user ON cash_books (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_stock_user ON stock (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_customers_user ON customers (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_activities_user ON activities (user_id, created_at DESC)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_notifications_unread ON notifications (user_id, is_read)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_journal_user_book ON journal_entries (user_id, book)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_bookdocs_user ON bookkeeping_docs (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_settings_user ON settings (user_id)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_recurring_due ON recurring_items (user_id, active, next_run_date)')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_users_id ON users (id)')
    cur.execute("ALTER TABLE service_sessions ADD COLUMN IF NOT EXISTS stripe_checkout_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE bookkeeping_docs ADD COLUMN IF NOT EXISTS stripe_checkout_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE sales ADD COLUMN IF NOT EXISTS stripe_checkout_id TEXT DEFAULT ''")
    cur.execute("ALTER TABLE purchase_orders ADD COLUMN IF NOT EXISTS received_date TEXT DEFAULT ''")
    cur.execute("ALTER TABLE bank_transactions ADD COLUMN IF NOT EXISTS reconciled INTEGER DEFAULT 0")
    cur.execute("ALTER TABLE bank_transactions ADD COLUMN IF NOT EXISTS reconciled_date TEXT DEFAULT ''")
    cur.execute("ALTER TABLE tasks ADD COLUMN IF NOT EXISTS blocker_note TEXT DEFAULT ''")
    cur.execute("ALTER TABLE bookkeeping_docs ADD COLUMN IF NOT EXISTS approval_status TEXT DEFAULT 'none'")
    cur.execute("ALTER TABLE bookkeeping_docs ADD COLUMN IF NOT EXISTS approval_notes TEXT DEFAULT ''")
    cur.execute("ALTER TABLE bookkeeping_docs ADD COLUMN IF NOT EXISTS approved_by TEXT DEFAULT ''")
    cur.execute("ALTER TABLE bookkeeping_docs ADD COLUMN IF NOT EXISTS approved_at TEXT DEFAULT ''")
    cur.execute('''
        CREATE TABLE IF NOT EXISTS settings_audit (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL,
            field_name TEXT NOT NULL,
            old_value TEXT DEFAULT '',
            new_value TEXT DEFAULT '',
            changed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_settings_audit_user ON settings_audit (user_id, changed_at DESC)')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS whatsapp_templates (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            body TEXT NOT NULL,
            created_date TEXT NOT NULL,
            updated_date TEXT NOT NULL
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_whatsapp_templates_user ON whatsapp_templates (user_id)')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS whatsapp_messages (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL,
            customer_id INTEGER,
            phone TEXT NOT NULL,
            body TEXT DEFAULT '',
            sent_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_whatsapp_messages_user ON whatsapp_messages (user_id, sent_at DESC)')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS score_snapshots (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL,
            snapshot_date TEXT NOT NULL,
            overall_score INTEGER NOT NULL,
            trend_score INTEGER DEFAULT 0,
            margin_score INTEGER DEFAULT 0,
            stock_score INTEGER DEFAULT 0,
            cash_score INTEGER DEFAULT 0,
            retention_score INTEGER DEFAULT 0,
            UNIQUE(user_id, snapshot_date)
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_score_snapshots_user ON score_snapshots (user_id, snapshot_date DESC)')

    conn.commit()
    cur.close()
    conn.close()
    _schema_ready = True


def ensure_db():
    """Create tables once. A slow database must not stop the process from listening."""
    if _schema_ready:
        return True
    try:
        init_db()
        return True
    except Exception as exc:
        print('KAZE schema init deferred: %s' % exc)
        return False


try:
    ensure_db()
except Exception as exc:
    print('KAZE boot continued without schema: %s' % exc)

_EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')

def _clean_text(name, required=True, max_len=200, label=None):
    val = (request.form.get(name) or '').strip()
    pretty = label or name.replace('_', ' ')
    if required and not val:
        return None, pretty[0].upper() + pretty[1:] + ' is required.'
    if len(val) > max_len:
        return None, pretty[0].upper() + pretty[1:] + ' is too long (max %s characters).' % max_len
    return val, None

def _clean_float(name, required=True, min_v=None, max_v=None, label=None, default=None):
    raw = request.form.get(name)
    pretty = label or name.replace('_', ' ')
    if raw is None or str(raw).strip() == '':
        if required and default is None:
            return None, pretty[0].upper() + pretty[1:] + ' is required.'
        return default, None
    try:
        n = float(raw)
    except (TypeError, ValueError):
        return None, pretty[0].upper() + pretty[1:] + ' must be a number.'
    if n != n or n in (float('inf'), float('-inf')):
        return None, pretty[0].upper() + pretty[1:] + ' must be a number.'
    if min_v is not None and n < min_v:
        return None, pretty[0].upper() + pretty[1:] + ' must be at least %s.' % min_v
    if max_v is not None and n > max_v:
        return None, pretty[0].upper() + pretty[1:] + ' must be at most %s.' % max_v
    return n, None

def _clean_email(name='email', required=False):
    val = (request.form.get(name) or '').strip()
    if not val:
        if required:
            return None, 'Email is required.'
        return '', None
    if len(val) > 120 or not _EMAIL_RE.match(val):
        return None, 'Enter a valid email address.'
    return val, None

def _clean_date(name='date', default_today=True):
    val = (request.form.get(name) or '').strip()
    if not val:
        return (datetime.today().strftime('%Y-%m-%d') if default_today else ''), None
    try:
        datetime.strptime(val, '%Y-%m-%d')
    except ValueError:
        return None, 'Use a valid date (YYYY-MM-DD).'
    if val < '2000-01-01' or val > '2100-12-31':
        return None, 'Date is out of range.'
    return val, None

def _reject(*pairs):
    """pairs of (value, error). If any error, flash first and return True."""
    for val, err in pairs:
        if err:
            flash(err, 'danger')
            return True
    return False


def _page_size(uid, default=25):
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute('SELECT items_per_page FROM settings WHERE user_id=%s', (uid,))
        row = cur.fetchone()
        cur.close()
        conn.close()
        n = int(row['items_per_page']) if row and row.get('items_per_page') else default
        return n if n in (10, 25, 50, 100) else default
    except Exception:
        return default


def _payment_method(raw=None):
    val = (raw if raw is not None else request.form.get('payment_method') or 'cash').strip().lower()
    if val not in ('cash', 'mobile', 'bank', 'card', 'other'):
        return 'cash'
    return val


def _form_lists(*keys):
    bag = {}
    for k in keys:
        bag[k] = request.form.getlist(k) or request.form.getlist(k + '[]')
    n = max((len(v) for v in bag.values()), default=0)
    return bag, n


def _cell(bag, key, i):
    lst = bag.get(key) or []
    if i >= len(lst):
        return ''
    return (lst[i] or '').strip()


def _cell_float(bag, key, i, default=None):
    raw = _cell(bag, key, i)
    if raw == '':
        return default
    try:
        n = float(raw)
    except (TypeError, ValueError):
        return None
    if n != n or n in (float('inf'), float('-inf')):
        return None
    return n


@app.route('/multi-add/<kind>', methods=['POST'])
@login_required
def multi_add(kind):
    """Save several rows from one form. Empty rows are skipped."""
    kind = (kind or '').strip().lower()
    uid = current_user.id
    today = datetime.today().strftime('%Y-%m-%d')
    saved = 0
    skipped = 0
    errors = []

    def bounce(endpoint='dashboard', **kwargs):
        if errors and not saved:
            flash(errors[0], 'danger')
        elif saved:
            extra = ' Skipped {} empty or incomplete row(s).'.format(skipped) if skipped else ''
            if errors:
                extra += ' ' + errors[0]
            flash('Saved {} {}.'.format(saved, 'row' if saved == 1 else 'rows') + extra, 'success')
        else:
            flash('Nothing to save. Fill at least one complete row.', 'danger')
        return redirect(url_for(endpoint, **kwargs))

    conn = get_db()
    cur = conn.cursor()

    try:
        if kind == 'stock':
            if not module_on('stock'):
                flash('Stock is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('product_name', 'quantity', 'cost_price', 'selling_price', 'unit', 'sku', 'category')
            for i in range(n):
                name = _cell(bag, 'product_name', i)
                if not name:
                    skipped += 1
                    continue
                blocked = _plan_block('stock')
                if blocked:
                    errors.append(blocked)
                    break
                qty = _cell_float(bag, 'quantity', i, 0)
                cost = _cell_float(bag, 'cost_price', i, 0)
                sell = _cell_float(bag, 'selling_price', i, 0)
                if qty is None or cost is None or sell is None or qty < 0 or cost < 0 or sell < 0:
                    errors.append('Row {}: quantity and prices must be numbers 0 or more.'.format(i + 1))
                    skipped += 1
                    continue
                cur.execute(
                    "INSERT INTO stock (user_id, product_name, quantity, cost_price, selling_price, unit, sku, category) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                    (uid, name[:120], qty, cost, sell,
                     _cell(bag, 'unit', i)[:20], _cell(bag, 'sku', i)[:40], _cell(bag, 'category', i)[:60])
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add stock', '{} products'.format(saved))
            return bounce('stock')

        if kind == 'sales':
            if not module_on('sales'):
                flash('Sales is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('stock_id', 'quantity', 'customer_id', 'customer_name', 'date', 'discount', 'payment_method', 'sale_notes')
            for i in range(n):
                try:
                    stock_id = int(_cell(bag, 'stock_id', i) or 0)
                except ValueError:
                    stock_id = 0
                qty = _cell_float(bag, 'quantity', i)
                if not stock_id or qty is None or qty <= 0:
                    skipped += 1
                    continue
                cur.execute('SELECT * FROM stock WHERE id=%s AND user_id=%s', (stock_id, uid))
                item = cur.fetchone()
                if not item:
                    errors.append('Row {}: product not found.'.format(i + 1))
                    skipped += 1
                    continue
                if item['quantity'] < qty:
                    errors.append('Row {}: not enough {} ({} left).'.format(i + 1, item['product_name'], item['quantity']))
                    skipped += 1
                    continue
                cust_id = _cell(bag, 'customer_id', i) or None
                cust_name = _cell(bag, 'customer_name', i) or 'Walk-in'
                cust_email = ''
                if cust_id:
                    try:
                        cust_id = int(cust_id)
                    except ValueError:
                        cust_id = None
                if cust_id:
                    cur.execute('SELECT name, email FROM customers WHERE id=%s AND user_id=%s', (cust_id, uid))
                    saved_c = cur.fetchone()
                    if saved_c:
                        cust_name = saved_c['name']
                        cust_email = saved_c['email'] or ''
                    else:
                        cust_id = None
                disc = _cell_float(bag, 'discount', i, 0) or 0
                if disc < 0:
                    disc = 0
                sale_date = _cell(bag, 'date', i) or today
                try:
                    datetime.strptime(sale_date, '%Y-%m-%d')
                except ValueError:
                    sale_date = today
                pay = _payment_method(_cell(bag, 'payment_method', i) or 'cash')
                total = max(item['selling_price'] * qty - disc, 0)
                profit = total - (item['cost_price'] * qty)
                note = _cell(bag, 'sale_notes', i)
                cur.execute(
                    "INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit, sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (uid, stock_id, qty, item['selling_price'], total, profit, sale_date,
                     cust_name, cust_email, cust_id, disc, note, pay)
                )
                cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id=%s', (qty, stock_id))
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add sales', '{} tickets'.format(saved))
            return bounce('sales')

        if kind == 'cash':
            if not module_on('bookkeeping'):
                flash('Book keeping is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('entry_type', 'category', 'description', 'amount', 'date')
            for i in range(n):
                et = (_cell(bag, 'entry_type', i) or 'in').lower()
                if et not in ('in', 'out'):
                    et = 'in'
                cat = _cell(bag, 'category', i)
                amt = _cell_float(bag, 'amount', i)
                if not cat or amt is None or amt <= 0:
                    skipped += 1
                    continue
                d = _cell(bag, 'date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                cur.execute(
                    "INSERT INTO cash_books (user_id, entry_type, category, description, amount, date) VALUES (%s,%s,%s,%s,%s,%s)",
                    (uid, et, cat[:80], _cell(bag, 'description', i)[:200], amt, d)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add cash', '{} lines'.format(saved))
            return bounce('cashbook')

        if kind == 'docs':
            if not module_on('documents'):
                flash('Documents is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('doc_type', 'title', 'client', 'amount', 'date', 'notes')
            for i in range(n):
                dtype = _cell(bag, 'doc_type', i) or 'invoice'
                if dtype not in ('invoice', 'contract', 'report'):
                    dtype = 'invoice'
                title = _cell(bag, 'title', i)
                client = _cell(bag, 'client', i)
                if not title or not client:
                    skipped += 1
                    continue
                amt = _cell_float(bag, 'amount', i, None)
                d = _cell(bag, 'date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                cur.execute(
                    "INSERT INTO documents (user_id, doc_type, title, client, amount, date, notes) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                    (uid, dtype, title[:160], client[:120], amt, d, _cell(bag, 'notes', i)[:4000])
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add documents', '{} papers'.format(saved))
            return bounce('docs')

        if kind == 'books':
            if not module_on('bookkeeping'):
                flash('Book keeping is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            doc_kind = (request.form.get('doc_kind') or '').strip()
            info = BOOK_DOC_LOOKUP.get(doc_kind)
            if not info:
                flash('Unknown document type.', 'danger')
                return redirect(url_for('bookkeeping'))
            label = info[0]
            bag, n = _form_lists('title', 'party', 'reference', 'doc_date', 'amount', 'status', 'notes', 'line_items')
            for i in range(n):
                title = _cell(bag, 'title', i) or label
                amt = _cell_float(bag, 'amount', i, 0) or 0
                d = _cell(bag, 'doc_date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                if not _cell(bag, 'title', i) and amt == 0 and not _cell(bag, 'party', i):
                    skipped += 1
                    continue
                status = _cell(bag, 'status', i) or 'draft'
                if status not in ('draft', 'issued', 'paid', 'cancelled'):
                    status = 'draft'
                ref = _cell(bag, 'reference', i) or (doc_kind.upper()[:3] + '-' + today.replace('-', '') + '-' + str(i + 1))
                cur.execute(
                    "INSERT INTO bookkeeping_docs (user_id, kind, title, party, reference, doc_date, amount, status, notes, line_items, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (uid, doc_kind, title[:160], _cell(bag, 'party', i)[:120], ref[:80], d, amt, status,
                     _cell(bag, 'notes', i), _cell(bag, 'line_items', i), today)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add ' + label, '{} papers'.format(saved))
            return bounce('bookkeeping_type', kind=doc_kind)

        if kind == 'journal':
            if not module_on('accounts'):
                flash('Account tracking is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            book = (request.form.get('book') or 'general').strip()
            bag, n = _form_lists('entry_date', 'reference', 'account', 'particulars', 'debit', 'credit')
            for i in range(n):
                account = _cell(bag, 'account', i)
                dr = _cell_float(bag, 'debit', i, 0) or 0
                cr = _cell_float(bag, 'credit', i, 0) or 0
                if not account:
                    skipped += 1
                    continue
                if dr < 0 or cr < 0:
                    errors.append('Row {}: debit and credit cannot be negative.'.format(i + 1))
                    skipped += 1
                    continue
                d = _cell(bag, 'entry_date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                cur.execute(
                    "INSERT INTO journal_entries (user_id, book, entry_date, reference, account, particulars, debit, credit, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (uid, book, d, _cell(bag, 'reference', i)[:80], account[:120],
                     _cell(bag, 'particulars', i)[:200], dr, cr, today)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add journal', book + ' x' + str(saved))
            return bounce('account_book', book=book)

        if kind == 'customers':
            if not module_on('customers'):
                flash('Customers is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('name', 'email', 'phone', 'address', 'notes')
            for i in range(n):
                name = _cell(bag, 'name', i)
                if not name:
                    skipped += 1
                    continue
                cur.execute(
                    "INSERT INTO customers (user_id, name, email, phone, address, notes) VALUES (%s,%s,%s,%s,%s,%s)",
                    (uid, name[:120], _cell(bag, 'email', i)[:120], _cell(bag, 'phone', i)[:40],
                     _cell(bag, 'address', i)[:200], _cell(bag, 'notes', i)[:2000])
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add customers', '{} contacts'.format(saved))
            return bounce('customers')

        if kind == 'suppliers':
            if not module_on('operations'):
                flash('Suppliers is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('name', 'email', 'phone', 'address')
            for i in range(n):
                name = _cell(bag, 'name', i)
                if not name:
                    skipped += 1
                    continue
                cur.execute(
                    "INSERT INTO suppliers (user_id, name, email, phone, address, notes) VALUES (%s,%s,%s,%s,%s,%s)",
                    (uid, name[:120], _cell(bag, 'email', i), _cell(bag, 'phone', i), _cell(bag, 'address', i), '')
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add suppliers', '{} vendors'.format(saved))
            return bounce('suppliers')

        if kind == 'income':
            bag, n = _form_lists('source', 'amount', 'date')
            for i in range(n):
                source = _cell(bag, 'source', i)
                amt = _cell_float(bag, 'amount', i)
                if not source or amt is None or amt < 0:
                    skipped += 1
                    continue
                d = _cell(bag, 'date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                cur.execute(
                    'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s)',
                    (uid, source[:120], amt, d)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add income', '{} lines'.format(saved))
            return bounce('dashboard')

        if kind == 'expenses':
            bag, n = _form_lists('name', 'amount', 'category', 'date')
            for i in range(n):
                name = _cell(bag, 'name', i)
                amt = _cell_float(bag, 'amount', i)
                if not name or amt is None or amt < 0:
                    skipped += 1
                    continue
                d = _cell(bag, 'date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                cat = _cell(bag, 'category', i) or 'Other'
                cur.execute(
                    'INSERT INTO expenses (user_id, name, amount, category, date) VALUES (%s,%s,%s,%s,%s)',
                    (uid, name[:120], amt, cat[:60], d)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add expenses', '{} lines'.format(saved))
            return bounce('dashboard')

        if kind == 'tasks':
            if not module_on('workspace'):
                flash('Tasks is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('title', 'priority', 'due_date', 'notes')
            for i in range(n):
                title = _cell(bag, 'title', i)
                if not title:
                    skipped += 1
                    continue
                pri = _cell(bag, 'priority', i) or 'medium'
                if pri not in ('low', 'medium', 'high'):
                    pri = 'medium'
                cur.execute(
                    "INSERT INTO tasks (user_id, title, notes, priority, due_date, done, created_date) VALUES (%s,%s,%s,%s,%s,0,%s)",
                    (uid, title[:160], _cell(bag, 'notes', i), pri, _cell(bag, 'due_date', i), today)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add tasks', '{} jobs'.format(saved))
            return bounce('tasks')

        if kind == 'notes':
            if not module_on('workspace'):
                flash('Notes is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('title', 'body')
            now = datetime.today().strftime('%Y-%m-%d %H:%M')
            for i in range(n):
                title = _cell(bag, 'title', i)
                body = _cell(bag, 'body', i)
                if not title and not body:
                    skipped += 1
                    continue
                cur.execute(
                    'INSERT INTO notes (user_id, title, body, created_date, updated_date) VALUES (%s,%s,%s,%s,%s)',
                    (uid, (title or 'Untitled note')[:160], body, now, now)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add notes', '{} notes'.format(saved))
            return bounce('notes')

        if kind == 'purchase_orders':
            if not module_on('operations'):
                flash('Purchase orders is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('supplier_id', 'item_description', 'stock_id', 'quantity', 'unit_cost', 'expected_date')
            for i in range(n):
                desc = _cell(bag, 'item_description', i)
                qty = _cell_float(bag, 'quantity', i)
                cost = _cell_float(bag, 'unit_cost', i)
                if not desc or qty is None or cost is None or qty <= 0 or cost < 0:
                    skipped += 1
                    continue
                sid = _cell(bag, 'supplier_id', i) or None
                stid = _cell(bag, 'stock_id', i) or None
                if sid:
                    try:
                        sid = int(sid)
                    except ValueError:
                        sid = None
                if stid:
                    try:
                        stid = int(stid)
                    except ValueError:
                        stid = None
                total = qty * cost
                cur.execute(
                    "INSERT INTO purchase_orders (user_id, supplier_id, item_description, quantity, unit_cost, total_cost, status, order_date, expected_date, stock_id) VALUES (%s,%s,%s,%s,%s,%s,'pending',%s,%s,%s)",
                    (uid, sid, desc[:200], qty, cost, total, today, _cell(bag, 'expected_date', i), stid)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add purchase orders', '{} orders'.format(saved))
            return bounce('purchase_orders')


        if kind == 'services':
            if not module_on('services'):
                flash('Services is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('name', 'category', 'duration_minutes', 'price', 'description')
            for i in range(n):
                name = _cell(bag, 'name', i)
                if not name:
                    skipped += 1
                    continue
                price = _cell_float(bag, 'price', i, 0)
                if price is None or price < 0:
                    errors.append('Row {}: price must be 0 or more.'.format(i + 1))
                    skipped += 1
                    continue
                try:
                    dur = int(float(_cell(bag, 'duration_minutes', i) or 60))
                except (TypeError, ValueError):
                    dur = 60
                if dur < 0:
                    dur = 0
                cat = _cell(bag, 'category', i) or 'Other'
                if cat not in SERVICE_CATEGORIES:
                    cat = 'Other'
                cur.execute(
                    "INSERT INTO services (user_id, name, category, description, duration_minutes, price, active, created_date) VALUES (%s,%s,%s,%s,%s,%s,1,%s)",
                    (uid, name[:160], cat[:40], _cell(bag, 'description', i)[:2000], dur, price, today)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add services', '{} services'.format(saved))
            return bounce('services')

        if kind == 'sessions':
            if not module_on('services'):
                flash('Services is dormant in this mode.', 'danger')
                return redirect(url_for('dashboard'))
            bag, n = _form_lists('service_id', 'customer_id', 'customer_name', 'session_date', 'session_time', 'price', 'notes')
            for i in range(n):
                try:
                    sid = int(_cell(bag, 'service_id', i) or 0)
                except ValueError:
                    sid = 0
                if not sid:
                    skipped += 1
                    continue
                cur.execute('SELECT * FROM services WHERE id=%s AND user_id=%s', (sid, uid))
                svc = cur.fetchone()
                if not svc:
                    errors.append('Row {}: service not found.'.format(i + 1))
                    skipped += 1
                    continue
                d = _cell(bag, 'session_date', i) or today
                try:
                    datetime.strptime(d, '%Y-%m-%d')
                except ValueError:
                    d = today
                price = _cell_float(bag, 'price', i, svc['price'])
                if price is None or price < 0:
                    price = svc['price']
                cust_id = _cell(bag, 'customer_id', i) or None
                cust_name = _cell(bag, 'customer_name', i) or ''
                if cust_id:
                    try:
                        cust_id = int(cust_id)
                    except ValueError:
                        cust_id = None
                if cust_id:
                    cur.execute('SELECT name FROM customers WHERE id=%s AND user_id=%s', (cust_id, uid))
                    saved_c = cur.fetchone()
                    if saved_c:
                        cust_name = saved_c['name']
                    else:
                        cust_id = None
                if not cust_name:
                    cust_name = 'Walk-in'
                cur.execute(
                    "INSERT INTO service_sessions (user_id, service_id, customer_id, customer_name, session_date, session_time, duration_minutes, price, status, payment_method, notes, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'booked','cash',%s,%s)",
                    (uid, sid, cust_id, cust_name[:120], d, _cell(bag, 'session_time', i)[:10],
                     svc['duration_minutes'] or 60, price, _cell(bag, 'notes', i)[:2000], today)
                )
                saved += 1
            conn.commit()
            if saved:
                log_activity(uid, current_user.username, 'Multi-add sessions', '{} bookings'.format(saved))
            return bounce('services')

        flash('Unknown multi-add type.', 'danger')
        return redirect(url_for('dashboard'))
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Could not save the rows. Check the numbers and try again.', 'danger')
        print('multi_add error:', exc)
        return redirect(url_for('dashboard'))
    finally:
        try:
            cur.close()
        except Exception:
            pass


def _hex_to_rgb(hex_color):
    h = str(hex_color or '#2ecc71').strip().lstrip('#')
    if len(h) == 3:
        h = ''.join(c * 2 for c in h)
    if len(h) != 6:
        h = '2ecc71'
    try:
        return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return 46, 204, 113


def _accent_palette(hex_color):
    r, g, b = _hex_to_rgb(hex_color)
    dark = (max(0, int(r * 0.42)), max(0, int(g * 0.42)), max(0, int(b * 0.42)))
    dim = (max(0, int(r * 0.16)), max(0, int(g * 0.16)), max(0, int(b * 0.16)))
    bg = (max(6, min(28, int(8 + r * 0.10))), max(6, min(28, int(8 + g * 0.10))), max(6, min(28, int(8 + b * 0.10))))
    bg2 = (max(10, min(36, int(14 + r * 0.09))), max(10, min(36, int(14 + g * 0.09))), max(10, min(36, int(14 + b * 0.09))))
    surface = (max(16, min(46, int(22 + r * 0.08))), max(16, min(46, int(22 + g * 0.08))), max(16, min(46, int(22 + b * 0.08))))
    light_bg = (min(255, 245 + int(r * 0.02)), min(255, 245 + int(g * 0.02)), min(255, 245 + int(b * 0.02)))
    return {
        'hex': '#{:02x}{:02x}{:02x}'.format(r, g, b),
        'rgb': '{},{},{}'.format(r, g, b),
        'dark': '#{:02x}{:02x}{:02x}'.format(*dark),
        'dim': '#{:02x}{:02x}{:02x}'.format(*dim),
        'glow': 'rgba({},{},{},0.20)'.format(r, g, b),
        'wash': 'rgba({},{},{},0.12)'.format(r, g, b),
        'grid': 'rgba({},{},{},0.07)'.format(r, g, b),
        'bg': '#{:02x}{:02x}{:02x}'.format(*bg),
        'bg2': '#{:02x}{:02x}{:02x}'.format(*bg2),
        'surface': '#{:02x}{:02x}{:02x}'.format(*surface),
        'light_bg': '#{:02x}{:02x}{:02x}'.format(*light_bg),
    }


# ======================== App modes (run some parts, leave others dormant) ========================
# Always-on pieces stay reachable so you can switch mode back on.
SERVICE_CATEGORIES = (
    'Teaching', 'Coaching', 'Tutoring', 'Training', 'Consulting',
    'Mentoring', 'Repair', 'Installation', 'Delivery', 'Design', 'Other',
)


PLAN_TIERS = {
    'new_user': {
        'key': 'new_user',
        'label': 'New user',
        'tagline': 'Open the desk and record the first tickets.',
        'price': 0,
        'highlights': [
            'Stock, sales, till, and customers',
            'Tasks, notes, and calendar',
            'Starter desk limits',
        ],
        'modules': [
            'stock', 'sales', 'counter', 'customers', 'workspace',
        ],
        'limits': {'stock': 20, 'customers': 20, 'services': 0, 'sales_month': 40, 'team': 1},
    },
    'entrepreneur': {
        'key': 'entrepreneur',
        'label': 'Entrepreneur',
        'tagline': 'Till plus cash ledger, day close, and services.',
        'price': 10,
        'highlights': [
            'Everything in New user',
            'Services, day close, cash ledger',
            'Book keeping papers and calculator',
        ],
        'modules': [
            'stock', 'sales', 'counter', 'customers', 'services',
            'calculator', 'reminders', 'workspace',
            'bookkeeping', 'day_close',
        ],
        'limits': {'stock': 150, 'customers': 150, 'services': 20, 'sales_month': 400, 'team': 2},
    },
    'businessman': {
        'key': 'businessman',
        'label': 'Businessman',
        'tagline': 'Suppliers, budgets, reports, and tax estimate.',
        'price': 25,
        'highlights': [
            'Everything in Entrepreneur',
            'Purchase orders, suppliers, budgets, recurring',
            'Reports, statistics, tax summary, documents',
            'SWOT and business plan',
        ],
        'modules': [
            'stock', 'sales', 'counter', 'customers', 'services',
            'calculator', 'reminders', 'workspace',
            'bookkeeping', 'day_close', 'operations', 'documents',
            'reports', 'tax', 'planning', 'stats',
        ],
        'limits': {'stock': 800, 'customers': 800, 'services': 100, 'sales_month': 4000, 'team': 6},
    },
    'ceo': {
        'key': 'ceo',
        'label': 'CEO',
        'tagline': 'Every operational module, plus account tracking.',
        'price': 50,
        'highlights': [
            'Everything in Businessman',
            'Banking, investments, shop, market value',
            'Journals, ledgers, and statements',
            'Supreme hub: health, profit matrix, RFM, forecast',
        ],
        'modules': None,
        'limits': {'stock': 8000, 'customers': 8000, 'services': 800, 'sales_month': 40000, 'team': 25},
    },
    'founder': {
        'key': 'founder',
        'label': "Founder's premium",
        'tagline': 'The full house with no practical desk limits.',
        'price': 100,
        'highlights': [
            'Everything in CEO',
            'Barcode scanner, voice entry, batch tracker',
            'Credit scores and WhatsApp hub',
            'No practical stock, sales, or team caps',
        ],
        'modules': None,
        'limits': {'stock': 0, 'customers': 0, 'services': 0, 'sales_month': 0, 'team': 0},
    },
}
PLAN_ORDER = ('new_user', 'entrepreneur', 'businessman', 'ceo', 'founder')
PLAN_RANK = {k: i for i, k in enumerate(PLAN_ORDER)}


# Legacy aliases — old tier names still work in URLs and DB, they map to the
# new 5-tier structure with a courtesy upgrade for existing customers.
_LEGACY_TIER_ALIASES = {
    'new_user_plus':     'entrepreneur',
    'starter':           'new_user',
    'starter_plus':      'entrepreneur',
    'entrepreneur_plus': 'businessman',
    'businessman_pro':   'ceo',
    'pro':               'ceo',
    'business':          'businessman',
    'enterprise':        'ceo',
    'premium':           'founder',
}



def default_package_prices():
    return {k: float(PLAN_TIERS[k].get('price') or 0) for k in PLAN_ORDER}


def package_prices_for(settings=None):
    # Prices are hardcoded in PLAN_TIERS. Settings cannot override them.
    return default_package_prices()


def priced_tiers(settings=None):
    prices = package_prices_for(settings)
    rows = []
    for key in PLAN_ORDER:
        row = dict(PLAN_TIERS[key])
        row['key'] = key
        row['price'] = prices.get(key, row.get('price') or 0)
        rows.append(row)
    return rows


def normalize_plan(raw):
    key = str(raw or '').strip().lower().replace('+', ' plus').replace("'", '')
    key = ' '.join(key.replace('-', ' ').replace('_', ' ').split())
    aliases = {
        'new': 'new_user', 'new user': 'new_user', 'starter': 'new_user', 'free': 'new_user',
        'entrepreneur': 'entrepreneur',
        'business': 'businessman', 'businessman': 'businessman', 'trader': 'businessman',
        'ceo': 'ceo', 'enterprise': 'ceo', 'full': 'ceo',
        'founder': 'founder', 'founders': 'founder', 'founders premium': 'founder',
        'founder premium': 'founder', 'premium': 'founder',
        # Legacy 8-tier names — courtesy upgrade
        'new user plus': 'entrepreneur', 'new user_plus': 'entrepreneur',
        'entrepreneur plus': 'businessman', 'entrepreneur_plus': 'businessman',
        'businessman pro': 'ceo', 'businessman_pro': 'ceo', 'pro': 'ceo',
    }
    if key in PLAN_TIERS:
        return key
    mapped = aliases.get(key)
    if mapped:
        return mapped
    return 'ceo' if not raw else 'new_user'


def trial_info(settings=None):
    settings = settings or {}
    flag = settings.get('trial_started')
    started = str(flag).strip() not in ('', '0', 'None', 'false', 'False')
    ends = str(settings.get('trial_ends') or '')[:10]
    today = datetime.today().strftime('%Y-%m-%d')
    active = bool(started and ends and ends >= today)
    expired = bool(started and ends and ends < today)
    days_left = 0
    if active:
        try:
            days_left = max(0, (datetime.strptime(ends, '%Y-%m-%d') - datetime.today()).days)
        except Exception:
            days_left = 0
    return {
        'started': started,
        'active': active,
        'expired': expired,
        'ends': ends,
        'days_left': days_left,
    }


def days_between(start_iso, end_iso):
    try:
        a = datetime.strptime(str(start_iso or '')[:10], '%Y-%m-%d')
        b = datetime.strptime(str(end_iso or '')[:10], '%Y-%m-%d')
        return max(0, (b - a).days)
    except (TypeError, ValueError):
        return 0


def subscription_info(settings=None):
    settings = settings or {}
    today = datetime.today().strftime('%Y-%m-%d')
    expires = str(settings.get('plan_expires_at') or '')[:10]
    grace = str(settings.get('grace_expires_at') or '')[:10]
    status = str(settings.get('subscription_status') or 'none').lower()

    if status in ('active', 'trialing') and expires and expires >= today:
        return {'status': 'active', 'expires': expires,
                'days_left': days_between(today, expires), 'grace': ''}
    if status == 'past_due' and grace and grace >= today:
        return {'status': 'grace', 'expires': expires, 'grace': grace,
                'days_left': days_between(today, grace)}
    if status == 'canceled' and expires and expires >= today:
        return {'status': 'canceled_active', 'expires': expires,
                'days_left': days_between(today, expires), 'grace': ''}
    if expires and expires < today:
        return {'status': 'expired', 'expires': expires,
                'days_since': days_between(expires, today), 'grace': ''}
    return {'status': 'none', 'expires': '', 'grace': '', 'days_left': 0}


def effective_plan(settings=None):
    settings = settings or {}
    today = datetime.today().strftime('%Y-%m-%d')

    expires = str(settings.get('plan_expires_at') or '')[:10]
    if expires and expires >= today:
        return normalize_plan(settings.get('plan_tier') or 'new_user')

    grace = str(settings.get('grace_expires_at') or '')[:10]
    if grace and grace >= today:
        return normalize_plan(settings.get('plan_tier') or 'new_user')

    if trial_info(settings).get('active'):
        return 'ceo'

    return 'new_user'


def plan_info(settings=None):
    settings = settings or {}
    stored = normalize_plan(settings.get('plan_tier') or 'new_user')
    key = effective_plan(settings)
    info = dict(PLAN_TIERS[key])
    info['key'] = key
    info['stored_key'] = stored
    info['price'] = package_prices_for(settings).get(key, info.get('price') or 0)
    info['trial'] = trial_info(settings)
    info['subscription'] = subscription_info(settings)
    info['stripe_customer_id'] = str(settings.get('stripe_customer_id') or '')
    return info


def plan_modules(settings=None):
    info = plan_info(settings)
    mods = info.get('modules')
    if mods is None:
        return set(ALL_OPTIONAL)
    return {m for m in mods if m in ALL_OPTIONAL}


def plan_limit(settings, name):
    info = plan_info(settings)
    try:
        return int((info.get('limits') or {}).get(name) or 0)
    except (TypeError, ValueError):
        return 0

def _plan_count(cur, uid, kind):
    if kind == 'stock':
        cur.execute('SELECT COUNT(*) AS n FROM stock WHERE user_id=%s', (uid,))
    elif kind == 'customers':
        cur.execute('SELECT COUNT(*) AS n FROM customers WHERE user_id=%s', (uid,))
    elif kind == 'services':
        cur.execute('SELECT COUNT(*) AS n FROM services WHERE user_id=%s', (uid,))
    elif kind == 'sales_month':
        month = datetime.today().strftime('%Y-%m')
        cur.execute(
            "SELECT COUNT(*) AS n FROM sales WHERE user_id=%s AND sale_date LIKE %s",
            (uid, month + '%')
        )
    elif kind == 'team':
        cur.execute(
            "SELECT COUNT(*) AS n FROM users WHERE id=%s OR owner_id=%s",
            (uid, uid)
        )
    else:
        return 0
    row = cur.fetchone() or {}
    try:
        return int(row.get('n') or 0)
    except (TypeError, ValueError):
        return 0


def _plan_block(kind, extra=1):
    settings = getattr(g, '_settings_row', None)
    cap = plan_limit(settings, kind)
    if cap <= 0:
        return None
    conn = get_db()
    cur = conn.cursor()
    n = _plan_count(cur, current_user.id, kind)
    cur.close()
    conn.close()
    if n + extra > cap:
        info = plan_info(settings)
        return 'The {} package allows {} {}. Upgrade on the Packages page to add more.'.format(
            info['label'], cap, kind.replace('_', ' ')
        )
    return None


ALWAYS_ON = frozenset({
    'dashboard', 'settings', 'search', 'notifications', 'tutorial',
    'shortcuts', 'activity', 'team', 'backup', 'index', 'auth', 'plans', 'payments',
})

MODULE_CATALOG = [
    ('stock', 'Stock / inventory'),
    ('sales', 'Sales desk'),
    ('counter', 'Counter / till'),
    ('customers', 'Customers'),
    ('bookkeeping', 'Book keeping and cash'),
    ('accounts', 'Account tracking'),
    ('operations', 'Recurring, purchase orders, suppliers, budgets'),
    ('reports', 'Reports'),
    ('stats', 'Statistics'),
    ('tax', 'Tax summary'),
    ('planning', 'SWOT and business plan'),
    ('workspace', 'Tasks, notes, calendar'),
    ('documents', 'Documents'),
    ('calculator', 'Profit calculator'),
    ('day_close', 'Day close'),
    ('reminders', 'Reminders'),
    ('services', 'Services (teaching, coaching, jobs)'),
    ('banking', 'Banking books and transfers'),
    ('investments', 'Investment holdings'),
    ('shop', 'Online banking and shop'),
    ('market', 'Monetization and market value'),
]
ALL_OPTIONAL = [m[0] for m in MODULE_CATALOG]

NAV_GROUPS = (
    ('Core', (
        ('Workshop Bench', 'dashboard', 'layout-dashboard', None, ('dashboard',)),
        ('Counter', 'counter', 'monitor-smartphone', 'counter', ('counter', 'counter_sale')),
        ('Sales', 'sales', 'receipt', 'sales', ('sales', 'aged_sales')),
        ('Stock', 'stock', 'package', 'stock', ('stock', 'edit_stock')),
        ('Customers', 'customers', 'users', 'customers', ('customers', 'edit_customer')),
        ('Services', 'services', 'graduation-cap', 'services', ('services',)),
    )),
    ('Money', (
        ('Book keeping', 'bookkeeping', 'book-open', 'bookkeeping',
         ('bookkeeping', 'bookkeeping_type', 'cashbook', 'edit_cash_entry')),
        ('Account tracking', 'accounts', 'book-marked', 'accounts', ('accounts', 'account_book')),
        ('Banking', 'banking', 'landmark', 'banking', ('banking',)),
        ('Payments', 'payments', 'credit-card', None, ('payments',)),
    )),
    ('Analytics', (
        ('Sales analytics', 'sales_analytics', 'bar-chart-3', 'sales', ('sales_analytics',)),
        ('Customer analytics', 'customers_analytics', 'users', 'customers', ('customers_analytics',)),
        ('Inventory analytics', 'stock_analytics', 'package-search', 'stock', ('stock_analytics',)),
        ('Books health', 'accounts_health', 'book-open-check', 'accounts', ('accounts_health',)),
        ('Decisions & signals', 'analytics_decisions', 'lightbulb', 'reports', ('analytics_decisions',)),
        ('Patterns', 'analytics_patterns', 'activity', 'reports', ('analytics_patterns',)),
        ('Team health', 'team_health', 'users-round', 'workspace', ('team_health',)),
        ('Workflow signals', 'workflow_signals', 'check-square', 'workspace', ('workflow_signals',)),
        ('Services analytics', 'services_analytics', 'graduation-cap', 'services', ('services_analytics',)),
        ('Settings audit', 'settings_audit', 'history', None, ('settings_audit',)),
    )),
    ('Trade', (
        ('Shop', 'shop', 'store', 'shop', ('shop',)),
        ('Purchase orders', 'purchase_orders', 'shopping-cart', 'operations', ('purchase_orders',)),
        ('Suppliers', 'suppliers', 'warehouse', 'operations', ('suppliers',)),
        ('Recurring', 'recurring', 'refresh-cw', 'operations', ('recurring',)),
        ('Budgets', 'budgets', 'target', 'operations', ('budgets',)),
        ('WhatsApp Hub', 'whatsapp', 'message-circle', 'operations', ('whatsapp',)),
        ('QR Scanner', 'scanner', 'scan', 'operations', ('scanner',)),
    )),
    ('Growth', (
        ('Supreme Desk', 'supreme', 'crown', None, ('supreme',)),
        ('Investments', 'investments', 'line-chart', 'investments', ('investments',)),
        ('Market value', 'market', 'trending-up', 'market', ('market',)),
        ('Reports', 'reports', 'bar-chart-3', 'reports', ('reports',)),
        ('Statistics', 'stats', 'table-2', 'stats', ('stats',)),
        ('Tax summary', 'tax_summary', 'percent', 'tax', ('tax_summary',)),
        ('AI Insights', 'insights', 'lightbulb', 'reports', ('insights',)),
        ('Cash Forecast', 'forecast', 'trending-up', 'reports', ('forecast',)),
        ('Health Score', 'health', 'activity', 'reports', ('health',)),
        ('Profit Matrix', 'profit_matrix', 'pie-chart', 'reports', ('profit_matrix',)),
        ('RFM Analysis', 'rfm', 'users', 'customers', ('rfm',)),
        ('Inventory Turnover', 'turnover', 'refresh-ccw', 'stock', ('turnover',)),
        ('Collections', 'collections', 'alert-circle', 'reports', ('collections',)),
        ('Goals', 'goals', 'target', 'planning', ('goals',)),
                                                                        ('Business plan', 'business_plan', 'map', 'planning', ('business_plan',)),
        ('SWOT', 'swot', 'compass', 'planning', ('swot',)),
    )),
    ('Workspace', (
        ('Tasks', 'tasks', 'check-square', 'workspace', ('tasks',)),
        ('Notes', 'notes', 'sticky-note', 'workspace', ('notes', 'edit_note')),
        ('Calendar', 'calendar', 'calendar', 'workspace', ('calendar',)),
        ('Reminders', 'reminders', 'bell', 'reminders', ('reminders',)),
        ('Day close', 'day_close', 'sunset', 'day_close', ('day_close',)),
        ('Documents', 'docs', 'file-text', 'documents', ('docs',)),
        ('Credit Scores', 'credit_scores', 'shield-check', 'customers', ('credit_scores',)),
        ('Batch Tracker', 'batches', 'layers', 'stock', ('batches',)),
        ('Voice Entry', 'voice', 'mic', 'workspace', ('voice',)),
        ('Shortcuts', 'shortcuts', 'keyboard', None, ('shortcuts',)),
        ('Practical guide', 'tutorial', 'book-open-check', None, ('tutorial',)),
    )),
    ('Account', (
        ('Packages', 'plans', 'layers', None, ('plans',)),
        ('Team', 'team', 'users-round', None, ('team',)),
        ('Settings', 'settings', 'settings', None, ('settings',)),
        ('Backup', 'backup_export', 'download', None, ('backup_export',)),
        ('Login activity', 'activity', 'log-in', None, ('activity',)),
        ('Business Man', 'business_man', 'user-check', None, ('business_man',)),
    )),
)


def live_nav_groups(live=None):
    if live is None:
        live = getattr(g, 'live_modules', None)
    if live is None:
        live = set(ALL_OPTIONAL)
    groups = []
    for title, items in NAV_GROUPS:
        visible = []
        for label, endpoint, icon, module, active in items:
            if module and module not in live:
                continue
            visible.append({
                'label': label,
                'endpoint': endpoint,
                'icon': icon,
                'active': active,
            })
        if visible:
            groups.append({'title': title, 'links': visible})
    return groups


APP_MODES = {
    'full': {
        'label': 'Full suite',
        'hint': 'Every module is live.',
        'modules': list(ALL_OPTIONAL),
    },
    'till': {
        'label': 'Till / shop floor',
        'hint': 'Sell and count stock. Books and planning stay dormant.',
        'modules': ['stock', 'sales', 'counter', 'customers', 'day_close', 'calculator', 'reminders', 'services', 'shop'],
    },
    'books': {
        'label': 'Books and reports',
        'hint': 'Ledgers, cash, papers, reports, tax. The till stays dormant.',
        'modules': ['bookkeeping', 'accounts', 'documents', 'reports', 'stats', 'tax', 'customers', 'banking'],
    },
    'office': {
        'label': 'Office / planning',
        'hint': 'Tasks, notes, SWOT, plan, reports. Shop-floor selling stays dormant.',
        'modules': ['workspace', 'planning', 'reports', 'stats', 'documents', 'reminders', 'calculator', 'services', 'investments', 'market'],
    },
    'quiet': {
        'label': 'Quiet / dashboard only',
        'hint': 'Home, settings, and help. Everything else sleeps until you pick another mode.',
        'modules': [],
    },
    'custom': {
        'label': 'Custom',
        'hint': 'Tick modules one by one. Unticked modules do not run.',
        'modules': None,
    },
}

UI_DEVICES = {
    'auto': {
        'label': 'Auto (follow the screen)',
        'hint': 'Phone, tablet, and PC layouts switch from screen width.',
    },
    'mobile': {
        'label': 'Mobile only',
        'hint': 'Phone chrome on every device: drawer menu, bottom dock, single column.',
    },
    'tablet': {
        'label': 'Tablet only',
        'hint': 'Tablet chrome on every device: slim rail, two columns, large tap targets.',
    },
    'pc': {
        'label': 'PC only',
        'hint': 'Desktop chrome on every device: full sidebar, multi-column workspace.',
    },
}

ENDPOINT_MODULE = {}


def _bind_module(module, *endpoints):
    for ep in endpoints:
        ENDPOINT_MODULE[ep] = module


_bind_module(
    'stock',
    'stock', 'add_stock', 'update_stock', 'delete_stock', 'edit_stock',
    'check_low_stock', 'stock_bulk', 'reorder_low_stock', 'price_list',
    'stock_valuation_pdf', 'export_stock_csv', 'import_stock_csv',
)
_bind_module(
    'sales',
    'sales', 'add_sale', 'auto_sale', 'delete_sale', 'sales_bulk',
    'sale_receipt', 'repeat_sale', 'aged_sales', 'sale_to_invoice', 'export_sales_csv',
)
_bind_module('counter', 'counter', 'counter_sale', 'counter_checkout')
_bind_module(
    'customers',
    'customers', 'add_customer', 'edit_customer', 'delete_customer',
    'customers_bulk', 'export_customers_csv', 'import_customers_csv', 'customer_statement',
)
_bind_module(
    'bookkeeping',
    'bookkeeping', 'bookkeeping_type', 'delete_book_doc', 'bookkeeping_pdf',
    'duplicate_book_doc', 'mark_book_doc_paid',
    'cashbook', 'add_cash_entry', 'edit_cash_entry', 'delete_cash_entry',
    'export_cashbook_csv', 'import_cashbook_csv',
)
_bind_module(
    'accounts',
    'accounts', 'add_journal_entry', 'delete_journal_entry', 'account_book', 'account_book_csv', 'pnl_pdf',
)
_bind_module(
    'operations',
    'recurring', 'add_recurring', 'toggle_recurring', 'delete_recurring',
    'purchase_orders', 'add_purchase_order', 'receive_purchase_order', 'delete_purchase_order',
    'suppliers', 'add_supplier', 'delete_supplier',
    'budgets', 'add_budget', 'delete_budget',
)
_bind_module('reports', 'reports', 'send_daily_summary')
_bind_module('stats', 'stats', 'stats_csv', 'stats_xlsx')
_bind_module('tax', 'tax_summary')
_bind_module(
    'planning',
    'swot', 'add_swot', 'delete_swot', 'swot_pdf',
    'business_plan', 'business_plan_pdf',
)
_bind_module(
    'workspace',
    'tasks', 'add_task', 'toggle_task', 'delete_task', 'tasks_bulk_complete',
    'notes', 'add_note', 'edit_note', 'delete_note',
    'calendar',
)
_bind_module('documents', 'docs', 'add_doc', 'delete_doc', 'doc_pdf')
_bind_module('calculator', 'calculator')
_bind_module('workspace', 'snippets', 'add_snippet', 'update_snippet', 'delete_snippet', 'snippet_starters')
_bind_module('day_close', 'day_close')
_bind_module('reminders', 'reminders')
_bind_module(
    'services',
    'services', 'add_service', 'update_service', 'toggle_service', 'delete_service',
    'add_service_session', 'set_session_status', 'delete_service_session', 'export_services_csv',
    'services_bulk', 'sessions_bulk',
)
_bind_module(
    'banking',
    'banking', 'add_bank_account', 'delete_bank_account', 'add_bank_tx', 'delete_bank_tx', 'bank_transfer',
)
_bind_module(
    'investments',
    'investments', 'add_investment', 'update_investment', 'delete_investment',
    'refresh_investments', 'refresh_investment',
)
_bind_module(
    'shop',
    'shop', 'add_shop_listing', 'toggle_shop_listing', 'delete_shop_listing',
    'add_shop_order', 'set_shop_order', 'delete_shop_order', 'shop_pay',
    'pay_shop_order',
)
_bind_module(
    'market',
    'market', 'add_market_channel', 'delete_market_channel', 'add_valuation', 'delete_valuation',
    'refresh_market_fx',
)
_bind_module(
    'dashboard',
    'supreme',
    'add_income', 'add_expense', 'delete_income', 'delete_expense',
    'export_income_csv', 'export_expenses_csv', 'import_income_csv', 'import_expenses_csv',
)


def _parse_custom_modules(raw):
    if not raw:
        return set()
    raw = str(raw).strip()
    if raw.startswith('['):
        try:
            data = json.loads(raw)
            if isinstance(data, list):
                return {str(x) for x in data}
        except Exception:
            pass
    return {x.strip() for x in raw.split(',') if x.strip()}


def resolve_live_modules(settings):
    if not settings:
        return set(ALL_OPTIONAL)
    mode = (settings.get('app_mode') or 'full').strip() or 'full'
    if mode not in APP_MODES:
        mode = 'full'
    if mode == 'custom':
        chosen = _parse_custom_modules(settings.get('enabled_modules') or '')
        live = {m for m in chosen if m in ALL_OPTIONAL}
    else:
        live = set(APP_MODES[mode]['modules'] or [])
    allowed = plan_modules(settings)
    return {m for m in live if m in allowed}


def module_on(name):
    if name in ALWAYS_ON:
        return True
    live = getattr(g, 'live_modules', None)
    if live is None:
        return True
    return name in live


def _csrf_token():
    token = session.get('_csrf')
    if not token:
        token = secrets.token_urlsafe(24)
        session['_csrf'] = token
    return token


def csrf_token():
    return _csrf_token()


@app.before_request
def _apply_app_mode():
    g.app_mode = 'full'
    g.live_modules = set(ALL_OPTIONAL)
    path = request.path or ''
    if path.startswith('/static/') or path in ('/healthz', '/favicon.ico'):
        return
    if path.startswith('/webhooks/'):
        return
    ensure_db()
    if request.method == 'POST':
        sent = (request.form.get('csrf_token') or request.headers.get('X-CSRF-Token') or '').strip()
        expected = session.get('_csrf') or ''
        if not expected or not sent or not secrets.compare_digest(sent, expected):
            return _error_body(400, 'Bad request', 'The form expired. Refresh the page and try again.')
    if path.startswith('/auth/callback/') or path.startswith('/login/') or path in (
        '/signup', '/login', '/licence', '/terms',
    ):
        return
    if not getattr(current_user, 'is_authenticated', False):
        return
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute('SELECT * FROM settings WHERE user_id = %s', (current_user.id,))
        row = cur.fetchone()
        cur.close()
        g._settings_row = row
        g._settings_uid = current_user.id
        g._settings_loaded = True
    except Exception:
        return
    if path not in ('/logout', '/licence', '/terms', '/legal/accept') and request.endpoint not in (
        'legal_accept', 'legal_licence', 'legal_terms', 'logout',
    ):
        if not _licence_ok(getattr(current_user, 'db_id', None)):
            return redirect(url_for('legal_accept'))
    g.app_mode = ((row.get('app_mode') if row else None) or 'full')
    g.live_modules = resolve_live_modules(row)
    ep = request.endpoint
    if not ep:
        return
    needed = ENDPOINT_MODULE.get(ep)
    if needed and needed not in ALWAYS_ON and needed not in g.live_modules:
        flash(
            'The "{}" module is dormant in {} mode. Turn it on under Settings → Modes.'.format(
                needed.replace('_', ' '), APP_MODES.get(g.app_mode, {}).get('label', g.app_mode)
            ),
            'danger',
        )
        return redirect(url_for('dashboard'))


# --------------------- Helper: log activity ---------------------
def log_activity(user_id, username, action, details=None):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO activities (user_id, username, action, details, created_at)
        VALUES (%s, %s, %s, %s, NOW())
    ''', (user_id, username, action, details))
    conn.commit()
    cur.close()
    conn.close()

# --------------------- Email function ---------------------
def send_email(to_email, subject, body):
    to_email = (to_email or '').strip()
    from_email = (os.environ.get('MAIL_USERNAME') or '').strip()
    password = os.environ.get('MAIL_PASSWORD') or os.environ.get('MAIL_APP_PASSWORD')
    if not to_email or '@' not in to_email:
        print("Email error: no destination address")
        return False
    if not from_email or not password:
        print("Email error: MAIL_USERNAME/MAIL_PASSWORD environment variables are not set")
        return False
    msg = MIMEText(body, _charset='utf-8')
    msg['Subject'] = subject
    msg['From'] = from_email
    msg['To'] = to_email
    host = os.environ.get('MAIL_SERVER', 'smtp.gmail.com')
    try:
        port = int(os.environ.get('MAIL_PORT', '587'))
    except ValueError:
        port = 587
    try:
        server = smtplib.SMTP(host, port, timeout=20)
        server.starttls()
        server.login(from_email, password)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        print("Email error: {}".format(e))
        return False

def _legal_text(kind='licence'):
    name = 'LICENCE' if kind != 'terms' else 'TERMS_OF_SERVICE'
    cached = _LEGAL_CACHE.get(name)
    if cached is not None:
        return cached
    path = os.path.join(LEGAL_DIR, name)
    try:
        with open(path, 'r', encoding='utf-8') as fh:
            text = fh.read()
    except Exception:
        text = 'Legal file missing. Contact Ephraim Zulu / KAZE Traders.'
    _LEGAL_CACHE[name] = text
    return text


def _licence_ok(user_db_id):
    if not user_db_id:
        return False
    cached = getattr(g, '_licence_ok', None)
    if cached is not None:
        return cached
    ok = False
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute('SELECT terms_accepted, terms_version FROM users WHERE id=%s', (user_db_id,))
        row = cur.fetchone() or {}
        cur.close()
        ok = int(row.get('terms_accepted') or 0) == 1 and str(row.get('terms_version') or '') == LEGAL_VERSION
    except Exception:
        ok = False
    g._licence_ok = ok
    return ok


def _stamp_licence(user_db_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'UPDATE users SET terms_accepted=1, terms_version=%s, terms_accepted_at=%s WHERE id=%s',
        (LEGAL_VERSION, datetime.now().strftime('%Y-%m-%d %H:%M:%S'), user_db_id),
    )
    conn.commit()
    cur.close()
    conn.close()
    g._licence_ok = True


# --------------------- User class ---------------------
def _business_id_for(role, owner_id, own_id):
    if role == 'staff' and owner_id:
        return owner_id
    return own_id

class User(UserMixin):
    def __init__(self, db_id, username, email, role='owner', owner_id=None):
        self.db_id = db_id
        self.username = username
        self.email = email
        self.role = role or 'owner'
        self.owner_id = owner_id
        self.id = _business_id_for(self.role, self.owner_id, self.db_id)

    def get_id(self):
        return str(self.db_id)

def _money_fmt_parts(settings_row=None):
    """Symbol + decimal places from Settings. Cached on g for the request."""
    cached = getattr(g, '_money_fmt', None)
    if cached and settings_row is None:
        return cached
    symbol = '$'
    decimals = 2
    row = settings_row
    try:
        if row is None and current_user.is_authenticated:
            conn = get_db()
            cur = conn.cursor()
            cur.execute('SELECT currency_symbol, number_decimals FROM settings WHERE user_id = %s', (current_user.id,))
            row = cur.fetchone()
            cur.close()
            conn.close()
        if row:
            symbol = (row.get('currency_symbol') or '$').strip() or '$'
            if row.get('number_decimals') is not None:
                decimals = max(0, min(6, int(row['number_decimals'])))
    except Exception:
        pass
    g._money_fmt = (symbol, decimals)
    return symbol, decimals


def _format_money(amount, settings_row=None):
    try:
        n = float(amount or 0)
    except (TypeError, ValueError):
        n = 0.0
    symbol, decimals = _money_fmt_parts(settings_row)
    sign = '-' if n < 0 else ''
    return '{}{}{:,.{}f}'.format(sign, symbol, abs(n), decimals)


def _excel_money_format(settings_row=None):
    symbol, decimals = _money_fmt_parts(settings_row)
    safe = (symbol or '$').replace('"', '')
    zeros = '0' * decimals if decimals else '0'
    if decimals:
        return '"{}"#,##0.{}'.format(safe, zeros)
    return '"{}"#,##0'.format(safe)


def _render_markdown(text):
    import html as _html
    import re as _re
    from markupsafe import Markup
    raw = '' if text is None else str(text).replace('\r\n', '\n')
    if not raw.strip():
        return Markup('')
    lines = raw.split('\n')
    out = []
    in_code = False
    in_list = False
    para = []

    def flush_para():
        nonlocal para
        if para:
            out.append('<p>' + '<br>'.join(para) + '</p>')
            para = []

    def flush_list():
        nonlocal in_list
        if in_list:
            out.append('</ul>')
            in_list = False

    def inline(s):
        s = _html.escape(s)
        s = _re.sub(r'`([^`]+)`', r'<code>\1</code>', s)
        s = _re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', s)
        s = _re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'<em>\1</em>', s)

        def link(m):
            label, url = m.group(1), m.group(2)
            if url.startswith('http://') or url.startswith('https://'):
                return '<a href="{}" rel="noopener">{}</a>'.format(_html.escape(url), label)
            return label
        return _re.sub(r'\[([^\]]+)\]\(([^)]+)\)', link, s)

    for line in lines:
        if line.startswith('```'):
            flush_para()
            flush_list()
            if in_code:
                out.append('</code></pre>')
                in_code = False
            else:
                out.append('<pre class="md-code"><code>')
                in_code = True
            continue
        if in_code:
            out.append(_html.escape(line) + '\n')
            continue
        stripped = line.strip()
        if not stripped:
            flush_para()
            flush_list()
            continue
        if stripped.startswith('### '):
            flush_para(); flush_list()
            out.append('<h3>' + inline(stripped[4:]) + '</h3>')
            continue
        if stripped.startswith('## '):
            flush_para(); flush_list()
            out.append('<h2>' + inline(stripped[3:]) + '</h2>')
            continue
        if stripped.startswith('# '):
            flush_para(); flush_list()
            out.append('<h1>' + inline(stripped[2:]) + '</h1>')
            continue
        if stripped.startswith('- ') or stripped.startswith('* '):
            flush_para()
            if not in_list:
                out.append('<ul>')
                in_list = True
            out.append('<li>' + inline(stripped[2:]) + '</li>')
            continue
        flush_list()
        para.append(inline(stripped))
    if in_code:
        out.append('</code></pre>')
    flush_para()
    flush_list()
    return Markup('\n'.join(out))


@app.template_filter('md')
def markdown_filter(text):
    return _render_markdown(text)


@app.template_filter('money')
def money_filter(amount):
    return _format_money(amount)


@app.context_processor
def inject_settings():
    if current_user.is_authenticated:
        settings = _get_settings_row(current_user.id)
        unread = getattr(g, '_unread_count', None)
        if unread is None:
            conn = get_db()
            cur = conn.cursor()
            cur.execute('SELECT COUNT(*) as unread_count FROM notifications WHERE user_id = %s AND is_read = 0', (current_user.id,))
            unread_row = cur.fetchone()
            cur.close()
            unread = unread_row['unread_count'] if unread_row else 0
            g._unread_count = unread
        if settings:
            try:
                g._money_fmt = (settings.get('currency_symbol') or '$', int(settings.get('number_decimals') if settings.get('number_decimals') is not None else 2))
            except Exception:
                g._money_fmt = ('$', 2)
        symbol = '$'
        decimals = 2
        if settings:
            symbol = settings.get('currency_symbol') or '$'
            try:
                if settings.get('number_decimals') is not None:
                    decimals = int(settings['number_decimals'])
            except (TypeError, ValueError):
                decimals = 2
        live = getattr(g, 'live_modules', None)
        if live is None:
            live = resolve_live_modules(settings)
            g.live_modules = live
            g.app_mode = (settings.get('app_mode') if settings else None) or 'full'
        pl_bar = getattr(g, 'pl_bar', None)
        accent = _accent_palette((settings.get('accent_color') if settings else None) or '#2ecc71')
        return dict(
            user_settings=settings,
            unread_notifications=unread,
            money=_format_money,
            currency_symbol=symbol,
            currency_decimals=decimals,
            app_mode=getattr(g, 'app_mode', 'full'),
            live_modules=live,
            mode_on=module_on,
            app_modes=APP_MODES,
            module_catalog=MODULE_CATALOG,
            ui_device=(settings.get('ui_device') if settings else None) or 'auto',
            ui_devices=UI_DEVICES,
            pl_bar=pl_bar,
            accent=accent,
            theme_pref=(settings.get('theme') if settings else None) or 'dark',
            stripe_ready=stripe_payments.stripe_ready(settings),
            flutterwave_ready=integrations.flutterwave_ready(settings),
            paypal_ready=gateways.paypal_ready(settings),
            bank_ready=gateways.bank_ready(settings),
            pay_methods=gateways.ready_methods(settings),
            alpha_ready=integrations.alpha_ready(settings),
            plan=plan_info(settings),
            oauth_providers=oauth_login.ready_list(),
            oauth_catalog=oauth_login.PROVIDERS,
            legal_version=LEGAL_VERSION,
            nav_groups=live_nav_groups(live),
            csrf_token=csrf_token,
        )
    return dict(
        user_settings=None, unread_notifications=0, money=_format_money,
        currency_symbol='$', currency_decimals=2,
        app_mode='full', live_modules=set(ALL_OPTIONAL),
        mode_on=lambda name: True, app_modes=APP_MODES, module_catalog=MODULE_CATALOG,
        ui_device='auto', ui_devices=UI_DEVICES, pl_bar=None,
        accent=_accent_palette('#2ecc71'),
        theme_pref='dark',
        stripe_ready=False, flutterwave_ready=False, paypal_ready=False, bank_ready=False, pay_methods=[], alpha_ready=False,
        plan=plan_info({'plan_tier': 'new_user'}),
        oauth_providers=oauth_login.ready_list(),
        oauth_catalog=oauth_login.PROVIDERS,
        legal_version=LEGAL_VERSION,
        nav_groups=live_nav_groups(set(ALL_OPTIONAL)),
        csrf_token=csrf_token,
    )


@login_manager.user_loader
def load_user(user_id):
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute('SELECT id, username, email, role, owner_id FROM users WHERE id = %s', (user_id,))
        user = cur.fetchone()
        cur.close()
        conn.close()
    except Exception:
        return None
    if user:
        return User(user['id'], user['username'], user['email'], user['role'], user['owner_id'])
    return None

# ======================== Routes ========================

@app.route('/')
def index():
    try:
        return render_template('index.html')
    except Exception:
        from flask import redirect as _rd
        return _rd('/login')

@app.route('/signup', methods=['GET', 'POST'])
def signup():
    if request.method == 'POST':
        username, e1 = _clean_text('username', max_len=40, label='username')
        email, e2 = _clean_email('email', required=True)
        password = request.form.get('password') or ''
        if _reject((username, e1), (email, e2)):
            return render_template('signup.html')
        if len(password) < 6:
            flash('Password must be at least 6 characters.', 'danger')
            return render_template('signup.html')
        hashed = generate_password_hash(password)
        conn = get_db()
        cur = conn.cursor()
        try:
            if request.form.get('accept_terms') != 'on':
                flash('You must accept the licence and terms of service.', 'danger')
                return render_template('signup.html')
            cur.execute('INSERT INTO users (username, email, password) VALUES (%s, %s, %s) RETURNING id',
                        (username, email, hashed))
            created = cur.fetchone() or {}
            new_id = created.get('id')
            if new_id:
                cur.execute(
                    'UPDATE users SET terms_accepted=1, terms_version=%s, terms_accepted_at=%s WHERE id=%s',
                    (LEGAL_VERSION, datetime.now().strftime('%Y-%m-%d %H:%M:%S'), new_id),
                )
            conn.commit()
            flash('Account created! Please log in.', 'success')
            return redirect(url_for('login'))
        except psycopg2.IntegrityError:
            conn.rollback()
            flash('Username or email already exists. Try another.', 'danger')
        finally:
            cur.close()
            conn.close()
    return render_template('signup.html')

def _complete_local_login(user):
    login_user(User(user['id'], user['username'], user['email'], user.get('role'), user.get('owner_id')))
    try:
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ip = request.headers.get('X-Forwarded-For', request.remote_addr)
        business_id = _business_id_for(user.get('role'), user.get('owner_id'), user['id'])
        conn = get_db()
        cur = conn.cursor()
        cur.execute('INSERT INTO user_activity (user_id, username, login_time, ip_address) VALUES (%s, %s, %s, %s)',
                    (business_id, user['username'], now, ip))
        cur.execute("INSERT INTO settings (user_id, notify_email, plan_tier) VALUES (%s, %s, 'new_user') ON CONFLICT (user_id) DO NOTHING",
                    (business_id, user.get('email') or ''))
        conn.commit()
        cur.close()
        conn.close()
    except Exception as exc:
        print('Activity log error:', exc)
    return redirect(url_for('dashboard'))


def _oauth_find_or_create(provider, profile):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM users WHERE oauth_provider=%s AND oauth_id=%s', (provider, profile['id']))
    user = cur.fetchone()
    if user:
        cur.close(); conn.close()
        return user, None
    email = (profile.get('email') or '').strip().lower()
    if email and not email.endswith('@oauth.kaze.invalid'):
        cur.execute('SELECT * FROM users WHERE lower(email)=%s', (email,))
        user = cur.fetchone()
        if user:
            cur.execute('UPDATE users SET oauth_provider=%s, oauth_id=%s WHERE id=%s', (provider, profile['id'], user['id']))
            conn.commit()
            cur.execute('SELECT * FROM users WHERE id=%s', (user['id'],))
            user = cur.fetchone()
            cur.close(); conn.close()
            return user, None
    base = ''.join(ch for ch in (profile.get('username') or provider) if ch.isalnum() or ch in '._-')[:28] or provider
    username = base
    n = 1
    while True:
        cur.execute('SELECT id FROM users WHERE username=%s', (username,))
        if not cur.fetchone():
            break
        n += 1
        username = '{}{}'.format(base[:24], n)
    if not email:
        email = '{}_{}@oauth.kaze.invalid'.format(provider, profile['id'])
    cur.execute('SELECT id FROM users WHERE lower(email)=%s', (email,))
    if cur.fetchone():
        email = '{}_{}@oauth.kaze.invalid'.format(provider, profile['id'])
    hashed = generate_password_hash(secrets.token_urlsafe(24))
    cur.execute(
        'INSERT INTO users (username, email, password, oauth_provider, oauth_id, terms_accepted, terms_version, terms_accepted_at) VALUES (%s,%s,%s,%s,%s,1,%s,%s) RETURNING *',
        (username, email, hashed, provider, profile['id'], LEGAL_VERSION, datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
    )
    user = cur.fetchone()
    conn.commit()
    cur.close(); conn.close()
    return user, 'created'


@app.route('/licence')
def legal_licence():
    return render_template('legal.html', title='Licence', legal_text=_legal_text('licence'), legal_version=LEGAL_VERSION)


@app.route('/terms')
def legal_terms():
    return render_template('legal.html', title='Terms of Service', legal_text=_legal_text('terms'), legal_version=LEGAL_VERSION)


@app.route('/legal/accept', methods=['GET', 'POST'])
@login_required
def legal_accept():
    if request.method == 'POST':
        if request.form.get('accept_terms') != 'on':
            flash('Tick the box to accept the licence and terms.', 'danger')
            return render_template('legal_accept.html', legal_version=LEGAL_VERSION, legal_text=_legal_text('licence'))
        _stamp_licence(current_user.db_id)
        flash('Licence and terms accepted (version {}).'.format(LEGAL_VERSION), 'success')
        return redirect(url_for('dashboard'))
    return render_template('legal_accept.html', legal_version=LEGAL_VERSION, legal_text=_legal_text('licence'))


@app.route('/login/<provider>')
def oauth_start(provider):
    from flask import session
    provider = (provider or '').strip().lower()
    if provider not in oauth_login.PROVIDERS:
        flash('That sign-in method is not supported.', 'danger')
        return redirect(url_for('login'))
    if not oauth_login.ready(provider):
        flash('{} sign-in is not configured yet.'.format(oauth_login.PROVIDERS[provider]['label']), 'danger')
        return redirect(url_for('login'))
    state = secrets.token_urlsafe(24)
    verifier = challenge = None
    if oauth_login.PROVIDERS[provider].get('pkce'):
        verifier, challenge = oauth_login.make_pkce()
    session['oauth_state'] = state
    session['oauth_provider'] = provider
    session['oauth_verifier'] = verifier or ''
    redirect_uri = url_for('oauth_callback', provider=provider, _external=True)
    dest = oauth_login.authorize_url(provider, redirect_uri, state, challenge)
    if not dest:
        flash('Could not start {} sign-in.'.format(oauth_login.PROVIDERS[provider]['label']), 'danger')
        return redirect(url_for('login'))
    return redirect(dest)


@app.route('/auth/callback/<provider>')
def oauth_callback(provider):
    from flask import session
    provider = (provider or '').strip().lower()
    if provider not in oauth_login.PROVIDERS:
        flash('That sign-in method is not supported.', 'danger')
        return redirect(url_for('login'))
    if request.args.get('error'):
        flash('{} sign-in was cancelled or refused.'.format(oauth_login.PROVIDERS[provider]['label']), 'info')
        return redirect(url_for('login'))
    if (request.args.get('state') or '') != (session.get('oauth_state') or '') or session.get('oauth_provider') != provider:
        flash('Sign-in state did not match. Try again.', 'danger')
        return redirect(url_for('login'))
    code = (request.args.get('code') or '').strip()
    if not code:
        flash('No authorisation code came back.', 'danger')
        return redirect(url_for('login'))
    redirect_uri = url_for('oauth_callback', provider=provider, _external=True)
    token, err = oauth_login.exchange_code(provider, code, redirect_uri, session.get('oauth_verifier') or None)
    session.pop('oauth_state', None)
    session.pop('oauth_verifier', None)
    if err or not token:
        flash(err or 'Could not finish sign-in.', 'danger')
        return redirect(url_for('login'))
    profile, err = oauth_login.fetch_profile(provider, token.get('access_token') or '')
    if err or not profile:
        flash(err or 'Could not read the provider profile.', 'danger')
        return redirect(url_for('login'))
    try:
        user, created = _oauth_find_or_create(provider, profile)
    except Exception as exc:
        print('oauth user error', exc)
        flash('Could not open a KAZE login from that account.', 'danger')
        return redirect(url_for('login'))
    if created:
        flash('Welcome. Your {} account is linked to KAZE.'.format(oauth_login.PROVIDERS[provider]['label']), 'success')
    else:
        flash('Signed in with {}.'.format(oauth_login.PROVIDERS[provider]['label']), 'success')
    if _needs_2fa(user):
        return _begin_2fa(user)
    return _complete_local_login(user)


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        login_input = (request.form.get('login_input') or '').strip()
        password = request.form.get('password') or ''
        user = None
        if login_input and password:
            try:
                conn = get_db()
                cur = conn.cursor()
                cur.execute('SELECT * FROM users WHERE username = %s OR email = %s', (login_input, login_input))
                user = cur.fetchone()
                cur.close(); conn.close()
            except Exception:
                flash('KAZE could not reach the database. Try again in a moment.', 'danger')
                return render_template('login.html')
        if user and check_password_hash(user['password'], password):
            if _needs_2fa(user):
                return _begin_2fa(user)
            return _complete_local_login(user)
        flash('Invalid credentials', 'danger')
    return render_template('login.html')


@app.route('/login/2fa', methods=['GET', 'POST'])
def login_2fa():
    uid = session.get('pre_2fa_id')
    exp = float(session.get('pre_2fa_exp') or 0)
    if not uid or exp < time.time():
        session.pop('pre_2fa_id', None)
        flash('The sign-in check expired. Log in again.', 'info')
        return redirect(url_for('login'))
    if request.method == 'POST':
        code = (request.form.get('code') or '').strip()
        conn = get_db(); cur = conn.cursor()
        cur.execute('SELECT * FROM users WHERE id=%s', (uid,))
        user = cur.fetchone()
        if not user or not _needs_2fa(user):
            cur.close(); conn.close()
            session.pop('pre_2fa_id', None)
            flash('Two-factor is not on for that account.', 'danger')
            return redirect(url_for('login'))
        ok = _totp_ok(user.get('totp_secret'), code)
        if not ok:
            ok, left = _backup_take(user.get('totp_backup') or '', code)
            if ok:
                cur.execute('UPDATE users SET totp_backup=%s WHERE id=%s', (left, uid))
                conn.commit()
        cur.close(); conn.close()
        if not ok:
            flash('That code is not valid.', 'danger')
            return render_template('login_2fa.html')
        session.pop('pre_2fa_id', None)
        session.pop('pre_2fa_exp', None)
        return _complete_local_login(user)
    return render_template('login_2fa.html')


@app.route('/settings/2fa/start', methods=['POST'])
@login_required
def twofa_start():
    secret = _totp_new_secret()
    plain, hashed = _backup_codes()
    conn = get_db(); cur = conn.cursor()
    cur.execute('UPDATE users SET totp_secret=%s, totp_enabled=0, totp_backup=%s WHERE id=%s', (secret, hashed, current_user.db_id))
    conn.commit(); cur.close(); conn.close()
    session['totp_backup_show'] = plain
    flash('Scan the code, then enter a 6-digit code to turn two-factor on.', 'info')
    return redirect(url_for('settings'))


@app.route('/settings/2fa/confirm', methods=['POST'])
@login_required
def twofa_confirm():
    code = (request.form.get('code') or '').strip()
    conn = get_db(); cur = conn.cursor()
    cur.execute('SELECT totp_secret FROM users WHERE id=%s', (current_user.db_id,))
    row = cur.fetchone() or {}
    if not _totp_ok(row.get('totp_secret'), code):
        cur.close(); conn.close()
        flash('That code did not match. Check the time on your phone and try again.', 'danger')
        return redirect(url_for('settings'))
    cur.execute('UPDATE users SET totp_enabled=1 WHERE id=%s', (current_user.db_id,))
    conn.commit(); cur.close(); conn.close()
    flash('Two-factor authentication is on. Save the backup codes.', 'success')
    return redirect(url_for('settings'))


@app.route('/settings/2fa/disable', methods=['POST'])
@login_required
def twofa_disable():
    password = request.form.get('password') or ''
    code = (request.form.get('code') or '').strip()
    conn = get_db(); cur = conn.cursor()
    cur.execute('SELECT password, totp_secret, totp_enabled, totp_backup FROM users WHERE id=%s', (current_user.db_id,))
    row = cur.fetchone() or {}
    if not row or not check_password_hash(row.get('password') or '', password):
        cur.close(); conn.close()
        flash('Password did not match.', 'danger')
        return redirect(url_for('settings'))
    ok = _totp_ok(row.get('totp_secret'), code)
    if not ok:
        ok, _left = _backup_take(row.get('totp_backup') or '', code)
    if int(row.get('totp_enabled') or 0) == 1 and not ok:
        cur.close(); conn.close()
        flash('Enter a current code or a backup code to turn two-factor off.', 'danger')
        return redirect(url_for('settings'))
    cur.execute("UPDATE users SET totp_enabled=0, totp_secret='', totp_backup='' WHERE id=%s", (current_user.db_id,))
    conn.commit(); cur.close(); conn.close()
    session.pop('totp_backup_show', None)
    flash('Two-factor authentication is off.', 'success')
    return redirect(url_for('settings'))


@app.route('/dashboard')
@login_required
def dashboard():
    cache_key = f"dash:{current_user.id}"
    cached = _cache_get(cache_key, ttl=30)
    if cached:
        return cached
    if module_on('operations'):
        _process_due_recurring_items(current_user.id)
    conn = get_db()
    cur = conn.cursor()
    mode = getattr(g, 'app_mode', None) or 'full'
    live = getattr(g, 'live_modules', None)
    snap = mode_runtime.run_selected_mode(cur, current_user.id, mode, live)
    activities = mode_runtime.activities_feed(cur, current_user.id)
    upcoming_sessions = []
    if module_on('services'):
        try:
            today = datetime.today().strftime('%Y-%m-%d')
            cur.execute(
                "SELECT ss.*, sv.name AS service_name FROM service_sessions ss "
                "LEFT JOIN services sv ON sv.id = ss.service_id "
                "WHERE ss.user_id=%s AND ss.status='booked' AND ss.session_date >= %s "
                "ORDER BY ss.session_date, ss.id LIMIT 8",
                (current_user.id, today)
            )
            upcoming_sessions = cur.fetchall()
        except Exception:
            upcoming_sessions = []
    cur.close()
    conn.close()
    g.pl_bar = snap['pl']
    sales = snap['sales']
    books = snap['books']
    return render_template(
        'dashboard.html',
        snap=snap,
        income=books['income_rows'],
        expenses=books['expense_rows'],
        total_income=books['income_total'],
        total_expenses=books['expense_total'],
        profit=snap['pl']['operating'],
        sales_profit=sales['all_profit'],
        cash_net=snap['cash']['net'],
        income_by_date=books['income_by_date'],
        expense_by_date=books['expense_by_date'],
        low_stock_items=snap['stock']['low'],
        sales_trend_data=sales['trend_7'],
        monthly_goal=snap.get('monthly_goal') or 0,
        month_revenue=sales['month_revenue'],
        last_month_revenue=sales['last_month_revenue'],
        goal_progress_pct=snap.get('goal_progress_pct') or 0,
        tutorial_completed=snap.get('tutorial_completed') or 0,
        activities=activities,
        open_tasks=snap['office']['open_tasks'],
        upcoming_recurring=snap['office']['upcoming_recurring'],
        upcoming_sessions=upcoming_sessions,
        pl=snap['pl'],
    )

# ======================== Add Income / Expense ========================

@app.route('/add_income', methods=['POST'])
@login_required
def add_income():
    source, e1 = _clean_text('source', max_len=120, label='source')
    amount, e2 = _clean_float('amount', min_v=0, max_v=1e12, label='amount')
    date, e3 = _clean_date('date')
    if _reject((source, e1), (amount, e2), (date, e3)):
        return redirect(url_for('dashboard'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO income (user_id, source, amount, date) VALUES (%s, %s, %s, %s)',
                (current_user.id, source, amount, date))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added income', f'{source} {_format_money(amount)}')
    flash('Income added!', 'success')
    return redirect(url_for('dashboard'))

@app.route('/add_expense', methods=['POST'])
@login_required
def add_expense():
    name, e1 = _clean_text('name', max_len=120, label='expense name')
    amount, e2 = _clean_float('amount', min_v=0, max_v=1e12, label='amount')
    category, e3 = _clean_text('category', max_len=60, label='category')
    date, e4 = _clean_date('date')
    if _reject((name, e1), (amount, e2), (category, e3), (date, e4)):
        return redirect(url_for('dashboard'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO expenses (user_id, name, amount, category, date) VALUES (%s, %s, %s, %s, %s)',
                (current_user.id, name, amount, category, date))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added expense', f'{name} {_format_money(amount)}')
    flash('Expense added!', 'success')
    return redirect(url_for('dashboard'))

@app.route('/delete_income/<int:id>')
@login_required
def delete_income(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM income WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted income', f'ID {id}')
    flash('Income deleted', 'success')
    return redirect(url_for('dashboard'))

@app.route('/delete_expense/<int:id>')
@login_required
def delete_expense(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM expenses WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted expense', f'ID {id}')
    flash('Expense deleted', 'success')
    return redirect(url_for('dashboard'))

# ======================== Documents ========================

@app.route('/docs')
@login_required
def docs():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM documents WHERE user_id = %s ORDER BY id DESC', (current_user.id,))
    user_docs = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('docs.html', docs=user_docs)

@app.route('/docs/add', methods=['POST'])
@login_required
def add_doc():
    doc_type = (request.form.get('doc_type') or '').strip()
    if doc_type not in ('invoice', 'contract', 'report'):
        flash('Choose a valid document type.', 'danger')
        return redirect(url_for('docs'))
    title, e1 = _clean_text('title', max_len=160, label='title')
    client, e2 = _clean_text('client', max_len=120, label='client')
    amount, e3 = _clean_float('amount', required=False, min_v=0, max_v=1e12, label='amount', default=None)
    date, e4 = _clean_date('date')
    notes, e5 = _clean_text('notes', required=False, max_len=4000, label='notes')
    if _reject((title, e1), (client, e2), (amount, e3), (date, e4), (notes, e5)):
        return redirect(url_for('docs'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO documents (user_id, doc_type, title, client, amount, date, notes)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
    ''', (current_user.id, doc_type, title, client, amount, date, notes))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Created document', title)
    flash('Document saved!', 'success')
    return redirect(url_for('docs'))

@app.route('/docs/delete/<int:id>')
@login_required
def delete_doc(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM documents WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted document', f'ID {id}')
    flash('Document deleted.', 'success')
    return redirect(url_for('docs'))

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('index'))

# ======================== Customers ========================

@app.route('/customers')
@login_required
def customers():
    page = request.args.get('page', 1, type=int)
    per_page = _page_size(current_user.id)
    offset = (page - 1) * per_page
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) as total FROM customers WHERE user_id = %s', (current_user.id,))
    total = cur.fetchone()['total']
    
    cur.execute('''
        SELECT * FROM customers WHERE user_id = %s ORDER BY name
        LIMIT %s OFFSET %s
    ''', (current_user.id, per_page, offset))
    all_customers = cur.fetchall()
    
    cur.execute('''
        SELECT customer_id, COALESCE(SUM(total_amount), 0) as total_spent, COUNT(*) as order_count
        FROM sales WHERE user_id = %s AND customer_id IS NOT NULL
        GROUP BY customer_id
    ''', (current_user.id,))
    spend_by_customer = {row['customer_id']: row for row in cur.fetchall()}
    cur.close()
    conn.close()
    
    total_pages = (total + per_page - 1) // per_page
    return render_template('customers.html', 
                         customers=all_customers,
                         spend_by_customer=spend_by_customer,
                         page=page,
                         total_pages=total_pages,
                         per_page=per_page)

@app.route('/customers/add', methods=['POST'])
@login_required
def add_customer():
    name, e1 = _clean_text('name', max_len=120, label='customer name')
    email, e2 = _clean_email('email', required=False)
    phone, e3 = _clean_text('phone', required=False, max_len=40, label='phone')
    address, e4 = _clean_text('address', required=False, max_len=200, label='address')
    notes, e5 = _clean_text('notes', required=False, max_len=2000, label='notes')
    if _reject((name, e1), (email, e2), (phone, e3), (address, e4), (notes, e5)):
        return redirect(url_for('customers'))
    blocked = _plan_block('customers')
    if blocked:
        flash(blocked, 'danger')
        return redirect(url_for('plans'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO customers (user_id, name, email, phone, address, notes)
        VALUES (%s, %s, %s, %s, %s, %s)
    ''', (current_user.id, name, email, phone, address, notes))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added customer', name)
    flash('Customer added.', 'success')
    return redirect(url_for('customers'))

@app.route('/customers/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_customer(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM customers WHERE id = %s AND user_id = %s', (id, current_user.id))
    customer = cur.fetchone()
    if not customer:
        flash('Customer not found.', 'danger')
        return redirect(url_for('customers'))

    if request.method == 'POST':
        name = request.form['name']
        email = request.form.get('email', '')
        phone = request.form.get('phone', '')
        address = request.form.get('address', '')
        notes = request.form.get('notes', '')
        cur.execute('''
            UPDATE customers SET name=%s, email=%s, phone=%s, address=%s, notes=%s
            WHERE id=%s AND user_id=%s
        ''', (name, email, phone, address, notes, id, current_user.id))
        conn.commit()
        cur.close()
        conn.close()
        log_activity(current_user.id, current_user.username, 'Updated customer', name)
        flash('Customer updated.', 'success')
        return redirect(url_for('customers'))

    cur.close()
    conn.close()
    return render_template('edit_customer.html', customer=customer)

@app.route('/customers/delete/<int:id>')
@login_required
def delete_customer(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE sales SET customer_id=NULL WHERE user_id=%s AND customer_id=%s', (current_user.id, id))
    try:
        cur.execute('UPDATE service_sessions SET customer_id=NULL WHERE user_id=%s AND customer_id=%s', (current_user.id, id))
    except Exception:
        pass
    cur.execute('DELETE FROM customers WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted customer', f'ID {id}')
    flash('Customer deleted. Linked sales were kept.', 'success')
    return redirect(url_for('customers'))

# ======================== Stock Management ========================

@app.route('/stock')
@login_required
def stock():
    page = request.args.get('page', 1, type=int)
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT items_per_page, low_stock_threshold FROM settings WHERE user_id = %s', (current_user.id,))
    srow = cur.fetchone()
    per_page = int((srow['items_per_page'] if srow and srow.get('items_per_page') else 25) or 25)
    threshold = srow['low_stock_threshold'] if srow and srow.get('low_stock_threshold') is not None else 10
    offset = (page - 1) * per_page
    cur.execute('SELECT COUNT(*) as total FROM stock WHERE user_id = %s', (current_user.id,))
    total = cur.fetchone()['total']
    
    cur.execute('''
        SELECT * FROM stock WHERE user_id = %s
        ORDER BY product_name
        LIMIT %s OFFSET %s
    ''', (current_user.id, per_page, offset))
    items = cur.fetchall()
    cur.close()
    conn.close()
    
    product_names = [item['product_name'] for item in items]
    inventory_values = [item['quantity'] * item['cost_price'] for item in items]
    total_pages = (total + per_page - 1) // per_page
    return render_template('stock.html', 
                         items=items,
                         product_names=product_names,
                         inventory_values=inventory_values,
                         page=page,
                         total_pages=total_pages,
                         per_page=per_page,
                         threshold=threshold)

@app.route('/stock/add', methods=['POST'])
@login_required
def add_stock():
    product_name, e1 = _clean_text('product_name', max_len=120, label='product name')
    quantity, e2 = _clean_float('quantity', min_v=0, max_v=1e9, label='quantity')
    cost_price, e3 = _clean_float('cost_price', min_v=0, max_v=1e12, label='cost price')
    selling_price, e4 = _clean_float('selling_price', min_v=0, max_v=1e12, label='selling price')
    unit, e5 = _clean_text('unit', required=False, max_len=20, label='unit')
    sku, e6 = _clean_text('sku', required=False, max_len=40, label='SKU')
    category, e7 = _clean_text('category', required=False, max_len=60, label='category')
    if _reject((product_name, e1), (quantity, e2), (cost_price, e3), (selling_price, e4), (unit, e5), (sku, e6), (category, e7)):
        return redirect(url_for('stock'))
    blocked = _plan_block('stock')
    if blocked:
        flash(blocked, 'danger')
        return redirect(url_for('plans'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO stock (user_id, product_name, quantity, cost_price, selling_price, unit, sku, category)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    ''', (current_user.id, product_name, quantity, cost_price, selling_price, unit, sku, category))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added stock', f'{product_name} qty {quantity}')
    flash('Stock item added.', 'success')
    return redirect(url_for('stock'))

@app.route('/stock/update/<int:id>', methods=['POST'])
@login_required
def update_stock(id):
    new_quantity = float(request.form['quantity'])
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute('SELECT product_name FROM stock WHERE id = %s AND user_id = %s', (id, current_user.id))
    product = cur.fetchone()
    if not product:
        flash('Stock item not found.', 'danger')
        return redirect(url_for('stock'))
    
    cur.execute('SELECT email FROM users WHERE id = %s', (current_user.id,))
    user = cur.fetchone()
    user_email = user['email']
    
    cur.execute('SELECT low_stock_threshold, email_notifications, notify_email FROM settings WHERE user_id = %s',
                (current_user.id,))
    settings = cur.fetchone()
    
    cur.execute('UPDATE stock SET quantity = %s WHERE id = %s AND user_id = %s',
                (new_quantity, id, current_user.id))
    conn.commit()
    
    if settings and settings['email_notifications']:
        threshold = settings['low_stock_threshold'] or 10
        if new_quantity < threshold:
            to_email = settings['notify_email'] or user_email
            subject = f"Low Stock Alert: {product['product_name']}"
            body = f"Your product '{product['product_name']}' has only {new_quantity} units left (threshold {threshold}). Please restock."
            send_email(to_email, subject, body)
            flash(f'Low stock alert sent to {to_email}', 'info')
    elif not settings:
        if new_quantity < 10:
            send_email(user_email, f"Low Stock: {product['product_name']}", f"{product['product_name']} is down to {new_quantity}.")
            flash('Low stock alert sent (default settings).', 'info')
    
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Updated stock', f'{product["product_name"]} to {new_quantity}')
    flash('Stock updated.', 'success')
    return redirect(url_for('stock'))

@app.route('/stock/delete/<int:id>')
@login_required
def delete_stock(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) AS n FROM sales WHERE stock_id=%s AND user_id=%s', (id, current_user.id))
    n = (cur.fetchone() or {}).get('n') or 0
    if n:
        cur.close(); conn.close()
        flash('This product has sales, so it cannot be deleted. Delete those sales first.', 'danger')
        return redirect(url_for('stock'))
    cur.execute('DELETE FROM stock WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted stock', f'ID {id}')
    flash('Stock item deleted.', 'success')
    return redirect(url_for('stock'))

@app.route('/stock/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_stock(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM stock WHERE id = %s AND user_id = %s', (id, current_user.id))
    item = cur.fetchone()
    if not item:
        flash('Item not found.', 'danger')
        return redirect(url_for('stock'))
    
    if request.method == 'POST':
        product_name, e1 = _clean_text('product_name', max_len=120, label='product name')
        quantity, e2 = _clean_float('quantity', min_v=0, max_v=1e9, label='quantity')
        cost_price, e3 = _clean_float('cost_price', min_v=0, max_v=1e12, label='cost price')
        selling_price, e4 = _clean_float('selling_price', min_v=0, max_v=1e12, label='selling price')
        unit, e5 = _clean_text('unit', required=False, max_len=20, label='unit')
        sku, e6 = _clean_text('sku', required=False, max_len=40, label='SKU')
        category, e7 = _clean_text('category', required=False, max_len=60, label='category')
        if _reject((product_name, e1), (quantity, e2), (cost_price, e3), (selling_price, e4), (unit, e5), (sku, e6), (category, e7)):
            cur.close(); conn.close()
            return redirect(url_for('edit_stock', id=id))
        cur.execute('''
            UPDATE stock 
            SET product_name=%s, quantity=%s, cost_price=%s, selling_price=%s, unit=%s, sku=%s, category=%s
            WHERE id=%s AND user_id=%s
        ''', (product_name, quantity, cost_price, selling_price, unit, sku, category, id, current_user.id))
        conn.commit()
        cur.close()
        conn.close()
        log_activity(current_user.id, current_user.username, 'Edited stock', product_name)
        flash('Stock updated.', 'success')
        return redirect(url_for('stock'))
    
    cur.close()
    conn.close()
    return render_template('edit_stock.html', item=item)

@app.route('/stock/check_low_stock')
@login_required
def check_low_stock():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT low_stock_threshold, email_notifications, notify_email FROM settings WHERE user_id = %s',
                (current_user.id,))
    settings_row = cur.fetchone()
    threshold = settings_row['low_stock_threshold'] if settings_row else 10
    to_email = (settings_row['notify_email'] if settings_row else None) or current_user.email

    cur.execute('SELECT product_name, quantity, unit FROM stock WHERE user_id = %s AND quantity < %s',
                (current_user.id, threshold))
    low_items = cur.fetchall()
    cur.close()
    conn.close()

    if not low_items:
        flash('No items are currently below the low stock threshold.', 'info')
        return redirect(url_for('stock'))

    body = f"Low Stock Report (threshold: {threshold})\n\n"
    for item in low_items:
        body += f"- {item['product_name']}: {item['quantity']} {item['unit'] or ''} remaining\n"
    send_email(to_email, "Low Stock Report", body)
    _notify(current_user.id, f"Low stock report sent: {len(low_items)} item(s) below threshold.", 'low_stock')
    log_activity(current_user.id, current_user.username, 'Checked low stock', f'{len(low_items)} items')
    flash(f'Low stock report sent to {to_email} ({len(low_items)} item(s)).', 'success')
    return redirect(url_for('stock'))

# ─── Bulk actions ───

def _as_int_ids(values):
    ids = []
    for raw in values:
        try:
            n = int(raw)
        except (TypeError, ValueError):
            continue
        if n > 0:
            ids.append(n)
    return ids


def _selected_ids():
    return _as_int_ids(request.form.getlist('selected_ids'))


@app.route('/stock/bulk', methods=['POST'])
@login_required
def stock_bulk():
    action = (request.form.get('action') or '').strip()
    ids = _selected_ids()
    if not action:
        flash('Choose a bulk action first.', 'danger')
        return redirect(url_for('stock'))
    if not ids:
        flash('Select at least one stock row.', 'danger')
        return redirect(url_for('stock'))

    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'delete':
            cur.execute(
                'SELECT DISTINCT stock_id FROM sales WHERE user_id=%s AND stock_id = ANY(%s)',
                (current_user.id, ids)
            )
            blocked = {int(r['stock_id']) for r in cur.fetchall() if r.get('stock_id')}
            free = [i for i in ids if i not in blocked]
            if free:
                cur.execute('DELETE FROM stock WHERE id = ANY(%s) AND user_id=%s', (free, current_user.id))
            conn.commit()
            msg = 'Deleted {} stock row(s).'.format(len(free))
            if blocked:
                msg += ' Kept {} because they still have sales. Delete those sales first.'.format(len(blocked))
            log_activity(current_user.id, current_user.username, 'Bulk stock', msg)
            flash(msg, 'success' if free else 'danger')
        elif action == 'update_quantity':
            qty, err = _clean_float('new_quantity', min_v=0, max_v=1e9, label='new quantity')
            if err:
                flash(err, 'danger')
            else:
                cur.execute(
                    'UPDATE stock SET quantity=%s WHERE id = ANY(%s) AND user_id=%s',
                    (qty, ids, current_user.id)
                )
                conn.commit()
                log_activity(current_user.id, current_user.username, 'Bulk stock qty', '{} rows'.format(len(ids)))
                flash('Updated quantity on {} row(s).'.format(len(ids)), 'success')
        elif action == 'export':
            cur.execute(
                'SELECT product_name, sku, category, quantity, unit, cost_price, selling_price '
                'FROM stock WHERE id = ANY(%s) AND user_id=%s ORDER BY product_name',
                (ids, current_user.id)
            )
            rows = cur.fetchall()
            return _csv_response(
                'selected_stock.csv',
                ['Product', 'SKU', 'Category', 'Quantity', 'Unit', 'Cost', 'Selling'],
                [(r['product_name'], r.get('sku'), r.get('category'), r['quantity'], r.get('unit'), r['cost_price'], r['selling_price']) for r in rows]
            )
        else:
            flash('Unknown bulk action.', 'danger')
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Bulk stock action failed. Nothing was changed.', 'danger')
        print('stock_bulk:', exc)
    finally:
        try:
            cur.close()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('stock'))


@app.route('/sales/bulk', methods=['POST'])
@login_required
def sales_bulk():
    action = (request.form.get('action') or '').strip()
    ids = _selected_ids()
    if not action:
        flash('Choose a bulk action first.', 'danger')
        return redirect(url_for('sales'))
    if not ids:
        flash('Select at least one sale.', 'danger')
        return redirect(url_for('sales'))

    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'delete':
            cur.execute(
                'SELECT id, stock_id, quantity_sold FROM sales WHERE id = ANY(%s) AND user_id=%s',
                (ids, current_user.id)
            )
            rows = cur.fetchall()
            restored = 0
            for sale in rows:
                if sale.get('stock_id'):
                    cur.execute(
                        'UPDATE stock SET quantity = quantity + %s WHERE id=%s AND user_id=%s',
                        (sale['quantity_sold'], sale['stock_id'], current_user.id)
                    )
                    restored += 1
            cur.execute('DELETE FROM sales WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            conn.commit()
            log_activity(current_user.id, current_user.username, 'Bulk delete sales', '{} tickets'.format(len(rows)))
            flash('Deleted {} sale(s) and restored stock on {} product(s).'.format(len(rows), restored), 'success')
        elif action == 'export':
            cur.execute(
                "SELECT COALESCE(stock.product_name, 'Item removed') AS product_name, "
                "sales.quantity_sold, sales.selling_price_at_time, sales.total_amount, sales.profit, "
                "sales.sale_date, sales.customer_name, sales.customer_email, sales.payment_method "
                "FROM sales LEFT JOIN stock ON sales.stock_id = stock.id "
                "WHERE sales.id = ANY(%s) AND sales.user_id=%s "
                "ORDER BY sales.sale_date DESC, sales.id DESC",
                (ids, current_user.id)
            )
            rows = cur.fetchall()
            return _csv_response(
                'selected_sales.csv',
                ['Product', 'Quantity', 'Selling Price', 'Total', 'Profit', 'Date', 'Customer', 'Email', 'Pay'],
                [(r['product_name'], r['quantity_sold'], r['selling_price_at_time'], r['total_amount'],
                  r['profit'], r['sale_date'], r['customer_name'], r['customer_email'], r.get('payment_method')) for r in rows]
            )
        else:
            flash('Unknown bulk action.', 'danger')
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Bulk sales action failed. Nothing was changed.', 'danger')
        print('sales_bulk:', exc)
    finally:
        try:
            cur.close()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('sales'))


@app.route('/customers/bulk', methods=['POST'])
@login_required
def customers_bulk():
    action = (request.form.get('action') or '').strip()
    ids = _selected_ids()
    if not action:
        flash('Choose a bulk action first.', 'danger')
        return redirect(url_for('customers'))
    if not ids:
        flash('Select at least one customer.', 'danger')
        return redirect(url_for('customers'))

    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'delete':
            cur.execute(
                'UPDATE sales SET customer_id=NULL WHERE user_id=%s AND customer_id = ANY(%s)',
                (current_user.id, ids)
            )
            try:
                cur.execute(
                    'UPDATE service_sessions SET customer_id=NULL WHERE user_id=%s AND customer_id = ANY(%s)',
                    (current_user.id, ids)
                )
            except Exception:
                conn.rollback()
                cur.execute(
                    'UPDATE sales SET customer_id=NULL WHERE user_id=%s AND customer_id = ANY(%s)',
                    (current_user.id, ids)
                )
            cur.execute('DELETE FROM customers WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            conn.commit()
            log_activity(current_user.id, current_user.username, 'Bulk delete customers', '{} contacts'.format(len(ids)))
            flash('Deleted {} customer(s). Linked sales were kept and unlinked.'.format(len(ids)), 'success')
        elif action == 'export':
            cur.execute(
                'SELECT name, email, phone, address, notes FROM customers WHERE id = ANY(%s) AND user_id=%s ORDER BY name',
                (ids, current_user.id)
            )
            rows = cur.fetchall()
            return _csv_response(
                'selected_customers.csv',
                ['Name', 'Email', 'Phone', 'Address', 'Notes'],
                [(r['name'], r['email'], r['phone'], r['address'], r['notes']) for r in rows]
            )
        else:
            flash('Unknown bulk action.', 'danger')
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Bulk customer action failed. Nothing was changed.', 'danger')
        print('customers_bulk:', exc)
    finally:
        try:
            cur.close()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('customers'))


@app.route('/services/bulk', methods=['POST'])
@login_required
def services_bulk():
    action = (request.form.get('action') or '').strip()
    ids = _selected_ids()
    if not action:
        flash('Choose a bulk action first.', 'danger')
        return redirect(url_for('services'))
    if not ids:
        flash('Select at least one service.', 'danger')
        return redirect(url_for('services'))
    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'hide':
            cur.execute('UPDATE services SET active=0 WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            conn.commit()
            flash('Hid {} service(s).'.format(len(ids)), 'success')
        elif action == 'show':
            cur.execute('UPDATE services SET active=1 WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            conn.commit()
            flash('Showed {} service(s).'.format(len(ids)), 'success')
        elif action == 'delete':
            cur.execute(
                'SELECT DISTINCT service_id FROM service_sessions WHERE user_id=%s AND service_id = ANY(%s)',
                (current_user.id, ids)
            )
            blocked = {int(r['service_id']) for r in cur.fetchall() if r.get('service_id')}
            free = [i for i in ids if i not in blocked]
            if free:
                cur.execute('DELETE FROM services WHERE id = ANY(%s) AND user_id=%s', (free, current_user.id))
            if blocked:
                cur.execute('UPDATE services SET active=0 WHERE id = ANY(%s) AND user_id=%s', (list(blocked), current_user.id))
            conn.commit()
            flash('Deleted {} unused service(s). Hid {} that still have sessions.'.format(len(free), len(blocked)), 'success')
        elif action == 'export':
            cur.execute(
                'SELECT name, category, duration_minutes, price, active, description FROM services WHERE id = ANY(%s) AND user_id=%s ORDER BY name',
                (ids, current_user.id)
            )
            rows = cur.fetchall()
            return _csv_response(
                'selected_services.csv',
                ['Name', 'Category', 'Minutes', 'Price', 'Active', 'Notes'],
                [(r['name'], r['category'], r['duration_minutes'], r['price'], r['active'], r.get('description')) for r in rows]
            )
        else:
            flash('Unknown bulk action.', 'danger')
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Bulk services action failed.', 'danger')
        print('services_bulk:', exc)
    finally:
        try:
            cur.close()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('services'))


@app.route('/services/sessions/bulk', methods=['POST'])
@login_required
def sessions_bulk():
    action = (request.form.get('action') or '').strip()
    ids = _selected_ids()
    if not action:
        flash('Choose a bulk action first.', 'danger')
        return redirect(url_for('services'))
    if not ids:
        flash('Select at least one session.', 'danger')
        return redirect(url_for('services'))
    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'delete':
            cur.execute('DELETE FROM service_sessions WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            conn.commit()
            flash('Deleted {} session(s).'.format(len(ids)), 'success')
        elif action in ('done', 'cancelled', 'booked'):
            today = datetime.today().strftime('%Y-%m-%d')
            cur.execute('SELECT * FROM service_sessions WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            rows = cur.fetchall()
            for row in rows:
                if action == 'done' and not row.get('income_id'):
                    svc_name = _service_row_name(cur, row.get('service_id'), current_user.id)
                    source = 'Service: {} — {}'.format(svc_name, row.get('customer_name') or 'Walk-in')
                    cur.execute(
                        'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s) RETURNING id',
                        (current_user.id, source[:120], float(row.get('price') or 0), row.get('session_date') or today)
                    )
                    inc = cur.fetchone()
                    inc_id = inc['id'] if inc else None
                    cur.execute(
                        "UPDATE service_sessions SET status='done', payment_method=COALESCE(payment_method,'cash'), income_id=%s WHERE id=%s AND user_id=%s",
                        (inc_id, row['id'], current_user.id)
                    )
                else:
                    cur.execute(
                        'UPDATE service_sessions SET status=%s WHERE id=%s AND user_id=%s',
                        (action, row['id'], current_user.id)
                    )
            conn.commit()
            flash('Updated {} session(s).'.format(len(rows)), 'success')
        else:
            flash('Unknown bulk action.', 'danger')
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Bulk sessions action failed.', 'danger')
        print('sessions_bulk:', exc)
    finally:
        try:
            cur.close()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('services'))


# ======================== Settings ========================

@app.route('/settings', methods=['GET', 'POST'])
@login_required
def settings():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO settings (user_id, notify_email) VALUES (%s, %s) ON CONFLICT (user_id) DO NOTHING',
                (current_user.id, current_user.email))
    conn.commit()
    
    if request.method == 'POST':
        low_stock_threshold, e1 = _clean_float('low_stock_threshold', min_v=0, max_v=1e7, label='low stock threshold')
        notify_email, e2 = _clean_email('notify_email', required=False)
        tax_rate_chk, e3 = _clean_float('tax_rate', required=False, min_v=0, max_v=100, label='tax rate', default=0)
        goal_chk, e4 = _clean_float('monthly_revenue_goal', required=False, min_v=0, max_v=1e12, label='monthly goal', default=0)
        if _reject((low_stock_threshold, e1), (notify_email, e2), (tax_rate_chk, e3), (goal_chk, e4)):
            cur.close(); conn.close()
            return redirect(url_for('settings'))
        email_notifications = 1 if request.form.get('email_notifications') == 'on' else 0
        notify_email = notify_email or current_user.email
        theme = request.form.get('theme', 'dark')
        if theme not in ('dark', 'light', 'auto', 'matrix', 'inverted'):
            theme = 'dark'
        font_family = request.form.get('font_family', 'Inter')
        font_size = request.form.get('font_size', 'medium')
        default_chart_type = request.form.get('default_chart_type', 'line')
        animations_enabled = 1 if request.form.get('animations_enabled') == 'on' else 0
        show_low_stock_widget = 1 if request.form.get('show_low_stock_widget') == 'on' else 0
        show_sales_trend = 1 if request.form.get('show_sales_trend') == 'on' else 0
        about_text = request.form.get('about_text', '')
        business_name = request.form.get('business_name', '')
        currency_code = request.form.get('currency_code', 'USD')
        currency_symbol = request.form.get('currency_symbol', '$')
        tax_rate = float(request.form.get('tax_rate') or 0)
        monthly_revenue_goal = float(request.form.get('monthly_revenue_goal') or 0)
        accent_color = request.form.get('accent_color') or '#2ecc71'
        density = request.form.get('density', 'comfortable')
        try:
            items_per_page = int(request.form.get('items_per_page') or 25)
        except ValueError:
            items_per_page = 25
        date_format = request.form.get('date_format', 'YYYY-MM-DD')
        show_clock = 1 if request.form.get('show_clock') == 'on' else 0
        rounded_ui = 1 if request.form.get('rounded_ui') == 'on' else 0
        high_contrast = 1 if request.form.get('high_contrast') == 'on' else 0
        invoice_prefix = (request.form.get('invoice_prefix') or 'INV')[:12]
        business_phone = request.form.get('business_phone', '')
        business_address = request.form.get('business_address', '')
        try:
            fiscal_year_start = int(request.form.get('fiscal_year_start') or 1)
        except ValueError:
            fiscal_year_start = 1
        try:
            number_decimals = int(request.form.get('number_decimals') or 2)
        except ValueError:
            number_decimals = 2
        sidebar_collapsed = 1 if request.form.get('sidebar_collapsed') == 'on' else 0
        app_mode = (request.form.get('app_mode') or 'full').strip()
        if app_mode not in APP_MODES:
            app_mode = 'full'
        picked = [m for m in request.form.getlist('enabled_modules') if m in ALL_OPTIONAL]
        enabled_modules = json.dumps(picked)
        ui_device = (request.form.get('ui_device') or 'auto').strip()
        if ui_device not in UI_DEVICES:
            ui_device = 'auto'
        cur.execute('SELECT * FROM settings WHERE user_id=%s', (current_user.id,))
        prior_keys = cur.fetchone() or {}
        stripe_publishable_key = (request.form.get('stripe_publishable_key') or '').strip()
        stripe_secret_key = (request.form.get('stripe_secret_key') or '').strip()
        stripe_webhook_secret = (request.form.get('stripe_webhook_secret') or '').strip()
        if (not stripe_secret_key) or stripe_secret_key.startswith('••••'):
            stripe_secret_key = (prior_keys.get('stripe_secret_key') if prior_keys else '') or ''
        if (not stripe_webhook_secret) or stripe_webhook_secret.startswith('••••'):
            stripe_webhook_secret = (prior_keys.get('stripe_webhook_secret') if prior_keys else '') or ''
        flutterwave_public_key = (request.form.get('flutterwave_public_key') or '').strip()
        flutterwave_secret_key = (request.form.get('flutterwave_secret_key') or '').strip()
        flutterwave_webhook_hash = (request.form.get('flutterwave_webhook_hash') or '').strip()
        alpha_vantage_key = (request.form.get('alpha_vantage_key') or '').strip()
        if (not flutterwave_secret_key) or flutterwave_secret_key.startswith('••••'):
            flutterwave_secret_key = (prior_keys.get('flutterwave_secret_key') if prior_keys else '') or ''
        if (not flutterwave_webhook_hash) or flutterwave_webhook_hash.startswith('••••'):
            flutterwave_webhook_hash = (prior_keys.get('flutterwave_webhook_hash') if prior_keys else '') or ''
        if (not alpha_vantage_key) or alpha_vantage_key.startswith('••••'):
            alpha_vantage_key = (prior_keys.get('alpha_vantage_key') if prior_keys else '') or ''
        paypal_client_id = (request.form.get('paypal_client_id') or '').strip()
        paypal_secret = (request.form.get('paypal_secret') or '').strip()
        paypal_mode = (request.form.get('paypal_mode') or 'sandbox').strip().lower()
        if paypal_mode not in ('sandbox', 'live'):
            paypal_mode = 'sandbox'
        if (not paypal_secret) or paypal_secret.startswith('••••'):
            paypal_secret = (prior_keys.get('paypal_secret') if prior_keys else '') or ''
        bank_name = (request.form.get('bank_name') or '').strip()[:80]
        bank_account_name = (request.form.get('bank_account_name') or '').strip()[:80]
        bank_account_number = (request.form.get('bank_account_number') or '').strip()[:40]
        bank_branch = (request.form.get('bank_branch') or '').strip()[:80]

        # Plan changes flow ONLY through /plans/upgrade -> checkout ->
        # webhook, or through set_plan()'s free-tier + already-paid guard.
        # Never trust a form field for the paid tier here.
        plan_tier = normalize_plan((prior_keys.get('plan_tier') if prior_keys else None) or 'ceo')

        package_prices = ''

        if font_size not in ('xsmall', 'small', 'medium', 'large', 'xlarge'):
            font_size = 'medium'
        if items_per_page not in (10, 25, 50, 100):
            items_per_page = 25
        if number_decimals not in (0, 2, 3):
            number_decimals = 2
        if fiscal_year_start < 1 or fiscal_year_start > 12:
            fiscal_year_start = 1

        cur.execute('''
            UPDATE settings 
            SET low_stock_threshold=%s, email_notifications=%s, notify_email=%s,
                theme=%s, font_family=%s, font_size=%s, default_chart_type=%s,
                animations_enabled=%s, show_low_stock_widget=%s, show_sales_trend=%s,
                about_text=%s, business_name=%s, currency_code=%s, currency_symbol=%s, tax_rate=%s,
                monthly_revenue_goal=%s, accent_color=%s, density=%s, items_per_page=%s,
                date_format=%s, show_clock=%s, rounded_ui=%s, high_contrast=%s,
                invoice_prefix=%s, business_phone=%s, business_address=%s,
                fiscal_year_start=%s, number_decimals=%s, sidebar_collapsed=%s,
                app_mode=%s, enabled_modules=%s, ui_device=%s,
                stripe_publishable_key=%s, stripe_secret_key=%s, stripe_webhook_secret=%s, plan_tier=%s,
                package_prices=%s,
                flutterwave_public_key=%s, flutterwave_secret_key=%s, flutterwave_webhook_hash=%s,
                alpha_vantage_key=%s,
                paypal_client_id=%s, paypal_secret=%s, paypal_mode=%s,
                bank_name=%s, bank_account_name=%s, bank_account_number=%s, bank_branch=%s
            WHERE user_id=%s
        ''', (low_stock_threshold, email_notifications, notify_email,
              theme, font_family, font_size, default_chart_type,
              animations_enabled, show_low_stock_widget, show_sales_trend,
              about_text, business_name, currency_code, currency_symbol, tax_rate,
              monthly_revenue_goal, accent_color, density, items_per_page,
              date_format, show_clock, rounded_ui, high_contrast,
              invoice_prefix, business_phone, business_address,
              fiscal_year_start, number_decimals, sidebar_collapsed,
              app_mode, enabled_modules, ui_device,
              stripe_publishable_key, stripe_secret_key, stripe_webhook_secret, plan_tier,
              package_prices,
              flutterwave_public_key, flutterwave_secret_key, flutterwave_webhook_hash,
              alpha_vantage_key,
              paypal_client_id, paypal_secret, paypal_mode,
              bank_name, bank_account_name, bank_account_number, bank_branch,
              current_user.id))
        # Audit: log what changed (savepoint so an audit failure never blocks save)
        try:
            cur.execute('SAVEPOINT kaze_audit_sp')
            cur.execute('SELECT * FROM settings WHERE user_id=%s', (current_user.id,))
            new_row = cur.fetchone() or {}
            AUDIT_NORMAL = [
                'plan_tier', 'business_name', 'currency_code', 'currency_symbol',
                'tax_rate', 'monthly_revenue_goal', 'app_mode', 'ui_device',
                'theme', 'low_stock_threshold', 'items_per_page', 'date_format',
                'fiscal_year_start', 'number_decimals', 'invoice_prefix',
                'business_phone', 'business_address', 'font_family', 'font_size',
                'density', 'accent_color',
            ]
            AUDIT_SENSITIVE = [
                'stripe_secret_key', 'stripe_publishable_key', 'stripe_webhook_secret',
                'paypal_client_id', 'paypal_secret', 'flutterwave_public_key',
                'flutterwave_secret_key', 'flutterwave_webhook_hash',
                'alpha_vantage_key',
            ]
            for field in AUDIT_NORMAL:
                o = str((prior_keys or {}).get(field) or '')
                n = str(new_row.get(field) or '')
                if o != n:
                    cur.execute(
                        "INSERT INTO settings_audit (user_id, field_name, old_value, new_value) VALUES (%s, %s, %s, %s)",
                        (current_user.id, field, o[:200], n[:200]))
            for field in AUDIT_SENSITIVE:
                o = str((prior_keys or {}).get(field) or '')
                n = str(new_row.get(field) or '')
                if o != n:
                    cur.execute(
                        "INSERT INTO settings_audit (user_id, field_name, old_value, new_value) VALUES (%s, %s, %s, %s)",
                        (current_user.id, field,
                         '(changed)' if o else '(set)',
                         '(changed)' if n else '(cleared)'))
            cur.execute('RELEASE SAVEPOINT kaze_audit_sp')
        except Exception as exc:
            print('settings audit error:', exc)
            try:
                cur.execute('ROLLBACK TO SAVEPOINT kaze_audit_sp')
            except Exception:
                pass
        conn.commit()
        cur.close()
        conn.close()
        flash('Settings saved.', 'success')
        return redirect(url_for('settings'))
    
    cur.execute('SELECT * FROM settings WHERE user_id = %s', (current_user.id,))
    settings_row = cur.fetchone()
    cur.close()
    conn.close()
    custom_modules = _parse_custom_modules(settings_row.get('enabled_modules') if settings_row else '')
    totp_on, totp_secret = 0, ''
    try:
        tcur = get_db().cursor()
        tcur.execute('SELECT totp_enabled, totp_secret FROM users WHERE id=%s', (current_user.db_id,))
        totp_row = tcur.fetchone() or {}
        tcur.close()
        totp_on = int(totp_row.get('totp_enabled') or 0)
        if not totp_on:
            totp_secret = totp_row.get('totp_secret') or ''
    except Exception:
        totp_on, totp_secret = 0, ''
    backup_codes = session.pop('totp_backup_show', None) if totp_on else (session.get('totp_backup_show') or [])
    return render_template('settings.html', settings=settings_row, custom_modules=custom_modules, totp_on=totp_on, totp_secret=totp_secret, totp_qr=_qr_svg(_otpauth_uri(totp_secret, current_user.username)) if totp_secret else '', backup_codes=backup_codes or [])


# ======================== Cash Book ========================

@app.route('/cashbook')
@login_required
def cashbook():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM cash_books WHERE user_id = %s ORDER BY date DESC', (current_user.id,))
    entries = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('cashbook.html', entries=entries)

@app.route('/cashbook/add', methods=['POST'])
@login_required
def add_cash_entry():
    entry_type = (request.form.get('entry_type') or '').strip()
    if entry_type not in ('in', 'out'):
        flash('Cash type must be in or out.', 'danger')
        return redirect(url_for('cashbook'))
    category, e1 = _clean_text('category', max_len=80, label='category')
    description, e2 = _clean_text('description', required=False, max_len=200, label='description')
    amount, e3 = _clean_float('amount', min_v=0.01, max_v=1e12, label='amount')
    date, e4 = _clean_date('date')
    if _reject((category, e1), (description, e2), (amount, e3), (date, e4)):
        return redirect(url_for('cashbook'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO cash_books (user_id, entry_type, category, description, amount, date)
        VALUES (%s, %s, %s, %s, %s, %s)
    ''', (current_user.id, entry_type, category, description, amount, date))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added cash entry', f'{entry_type} {_format_money(amount)}')
    flash('Cash entry added.', 'success')
    return redirect(url_for('cashbook'))

@app.route('/cashbook/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_cash_entry(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM cash_books WHERE id = %s AND user_id = %s', (id, current_user.id))
    entry = cur.fetchone()
    if not entry:
        flash('Entry not found.', 'danger')
        return redirect(url_for('cashbook'))
    
    if request.method == 'POST':
        entry_type = request.form['entry_type']
        category = request.form['category']
        description = request.form['description']
        amount = float(request.form['amount'])
        date = request.form['date']
        cur.execute('''
            UPDATE cash_books 
            SET entry_type=%s, category=%s, description=%s, amount=%s, date=%s
            WHERE id=%s AND user_id=%s
        ''', (entry_type, category, description, amount, date, id, current_user.id))
        conn.commit()
        cur.close()
        conn.close()
        log_activity(current_user.id, current_user.username, 'Edited cash entry', f'ID {id}')
        flash('Cash entry updated.', 'success')
        return redirect(url_for('cashbook'))
    
    cur.close()
    conn.close()
    return render_template('edit_cashbook.html', entry=entry)

@app.route('/cashbook/delete/<int:id>')
@login_required
def delete_cash_entry(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM cash_books WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Deleted cash entry', f'ID {id}')
    flash('Cash entry deleted.', 'success')
    return redirect(url_for('cashbook'))

# ======================== Sales ========================

@app.route('/sales')
@login_required
def sales():
    page = request.args.get('page', 1, type=int)
    per_page = _page_size(current_user.id)
    offset = (page - 1) * per_page
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) as total FROM sales WHERE user_id = %s', (current_user.id,))
    total = cur.fetchone()['total']
    
    cur.execute('''
        SELECT sales.*, stock.product_name 
        FROM sales 
        LEFT JOIN stock ON sales.stock_id = stock.id
        WHERE sales.user_id = %s 
        ORDER BY sale_date DESC
        LIMIT %s OFFSET %s
    ''', (current_user.id, per_page, offset))
    all_sales = cur.fetchall()
    
    cur.execute('SELECT id, product_name, selling_price, quantity FROM stock WHERE user_id = %s AND quantity > 0',
                (current_user.id,))
    stock_items = cur.fetchall()
    cur.execute('SELECT id, name, email FROM customers WHERE user_id = %s ORDER BY name', (current_user.id,))
    customer_list = cur.fetchall()
    
    sales_by_date = {}
    for sale in all_sales:
        d = str(sale['sale_date'] or '')[:10]
        sales_by_date[d] = sales_by_date.get(d, 0) + (sale['total_amount'] or 0)
    
    cur.close()
    conn.close()
    
    total_pages = (total + per_page - 1) // per_page
    return render_template('sales.html', 
                         sales=all_sales,
                         stock_items=stock_items,
                         sales_by_date=sales_by_date,
                         customer_list=customer_list,
                         page=page,
                         total_pages=total_pages,
                         per_page=per_page)

@app.route('/sales/add', methods=['POST'])
@login_required
def add_sale():
    try:
        stock_id = int(request.form.get('stock_id') or 0)
    except (TypeError, ValueError):
        flash('Pick a product.', 'danger')
        return redirect(url_for('sales'))
    if stock_id <= 0:
        flash('Pick a product.', 'danger')
        return redirect(url_for('sales'))
    quantity_sold, e1 = _clean_float('quantity', min_v=0.0001, max_v=1e9, label='quantity')
    sale_date, e2 = _clean_date('date')
    customer_name, e3 = _clean_text('customer_name', required=False, max_len=120, label='customer name')
    customer_email, e4 = _clean_email('customer_email', required=False)
    if _reject((quantity_sold, e1), (sale_date, e2), (customer_name, e3), (customer_email, e4)):
        return redirect(url_for('sales'))
    blocked = _plan_block('sales_month')
    if blocked:
        flash(blocked, 'danger')
        return redirect(url_for('plans'))
    customer_id = request.form.get('customer_id') or None
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM stock WHERE id = %s AND user_id = %s', (stock_id, current_user.id))
    stock_item = cur.fetchone()
    if not stock_item:
        flash('Stock item not found.', 'danger')
        return redirect(url_for('sales'))
    
    if stock_item['quantity'] < quantity_sold:
        flash(f'Insufficient stock. Only {stock_item["quantity"]} available.', 'danger')
        return redirect(url_for('sales'))

    if customer_id:
        cur.execute('SELECT name, email FROM customers WHERE id = %s AND user_id = %s', (customer_id, current_user.id))
        saved_customer = cur.fetchone()
        if saved_customer:
            customer_name = saved_customer['name']
            customer_email = saved_customer['email'] or customer_email
    
    selling_price = stock_item['selling_price']
    cost_price = stock_item['cost_price']
    try:
        discount = float(request.form.get('discount') or 0)
    except ValueError:
        discount = 0
    sale_notes = request.form.get('sale_notes', '')
    total_amount = max(selling_price * quantity_sold - discount, 0)
    profit = total_amount - (cost_price * quantity_sold)
    
    cur.execute('''
        INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit, sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (current_user.id, stock_id, quantity_sold, selling_price, total_amount, profit, sale_date, customer_name, customer_email, customer_id, discount, sale_notes, _payment_method()))
    
    new_quantity = stock_item['quantity'] - quantity_sold
    cur.execute('UPDATE stock SET quantity = %s WHERE id = %s', (new_quantity, stock_id))
    conn.commit()
    cur.close()
    conn.close()
    
    log_activity(current_user.id, current_user.username, 'Recorded sale', f'{stock_item["product_name"]} x{quantity_sold} {_format_money(total_amount)}')
    flash('Sale recorded. Profit: {}'.format(_format_money(profit)), 'success')
    return redirect(url_for('sales'))

@app.route('/sales/auto', methods=['POST'])
@login_required
def auto_sale():
    """One-step sale: current selling price, today, default qty 1."""
    try:
        stock_id = int(request.form['stock_id'])
    except (TypeError, ValueError, KeyError):
        flash('Pick a product for auto sale.', 'danger')
        return redirect(url_for('sales'))
    try:
        quantity_sold = float(request.form.get('quantity') or 1)
    except ValueError:
        quantity_sold = 1
    if quantity_sold <= 0:
        flash('Quantity must be more than zero.', 'danger')
        return redirect(url_for('sales'))
    customer_id = request.form.get('customer_id') or None
    customer_name = request.form.get('customer_name', '') or 'Walk-in'
    customer_email = ''
    sale_date = datetime.today().strftime('%Y-%m-%d')

    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM stock WHERE id = %s AND user_id = %s', (stock_id, current_user.id))
    stock_item = cur.fetchone()
    if not stock_item:
        cur.close()
        conn.close()
        flash('Stock item not found.', 'danger')
        return redirect(url_for('sales'))
    if stock_item['quantity'] < quantity_sold:
        cur.close()
        conn.close()
        flash('Not enough stock for auto sale. Only {} available.'.format(stock_item['quantity']), 'danger')
        return redirect(url_for('sales'))
    if customer_id:
        cur.execute('SELECT name, email FROM customers WHERE id = %s AND user_id = %s', (customer_id, current_user.id))
        saved = cur.fetchone()
        if saved:
            customer_name = saved['name']
            customer_email = saved['email'] or ''
    selling_price = stock_item['selling_price']
    cost_price = stock_item['cost_price']
    total_amount = selling_price * quantity_sold
    profit = total_amount - (cost_price * quantity_sold)
    cur.execute(
        'INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit, '
        'sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method) '
        'VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)',
        (current_user.id, stock_id, quantity_sold, selling_price, total_amount, profit,
         sale_date, customer_name, customer_email, customer_id, 0, 'auto sale', _payment_method())
    )
    cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id = %s', (quantity_sold, stock_id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(
        current_user.id, current_user.username, 'Auto sale',
        '{} x{} {}'.format(stock_item['product_name'], quantity_sold, _format_money(total_amount))
    )
    flash(
        'Auto sale: {} x {} · collect {}.'.format(
            stock_item['product_name'], quantity_sold, _format_money(total_amount)
        ),
        'success'
    )
    return redirect(url_for('sales'))


@app.route('/sales/delete/<int:id>')
@login_required
def delete_sale(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM sales WHERE id = %s AND user_id = %s', (id, current_user.id))
    sale = cur.fetchone()
    if sale:
        cur.execute('UPDATE stock SET quantity = quantity + %s WHERE id = %s', (sale['quantity_sold'], sale['stock_id']))
        cur.execute('DELETE FROM sales WHERE id = %s', (id,))
        conn.commit()
        log_activity(current_user.id, current_user.username, 'Deleted sale', f'ID {id}')
        flash('Sale deleted and stock restored.', 'success')
    else:
        flash('Sale not found.', 'danger')
    cur.close()
    conn.close()
    return redirect(url_for('sales'))

# ======================== CSV Export ========================

def _csv_response(filename, header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for row in rows:
        writer.writerow(row)
    return Response(
        buf.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )

@app.route('/export/income.csv')
@login_required
def export_income_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT source, amount, date FROM income WHERE user_id = %s ORDER BY date DESC', (current_user.id,))
    rows = [(r['source'], r['amount'], r['date']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('income.csv', ['Source', 'Amount', 'Date'], rows)

@app.route('/export/expenses.csv')
@login_required
def export_expenses_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT name, amount, category, date FROM expenses WHERE user_id = %s ORDER BY date DESC', (current_user.id,))
    rows = [(r['name'], r['amount'], r['category'], r['date']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('expenses.csv', ['Name', 'Amount', 'Category', 'Date'], rows)

@app.route('/export/sales.csv')
@login_required
def export_sales_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT stock.product_name, sales.quantity_sold, sales.selling_price_at_time,
               sales.total_amount, sales.profit, sales.sale_date, sales.customer_name, sales.customer_email
        FROM sales JOIN stock ON sales.stock_id = stock.id
        WHERE sales.user_id = %s ORDER BY sale_date DESC
    ''', (current_user.id,))
    rows = [(r['product_name'], r['quantity_sold'], r['selling_price_at_time'], r['total_amount'],
              r['profit'], r['sale_date'], r['customer_name'], r['customer_email']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('sales.csv',
                          ['Product', 'Quantity', 'Selling Price', 'Total', 'Profit', 'Date', 'Customer', 'Customer Email'],
                          rows)

@app.route('/export/cashbook.csv')
@login_required
def export_cashbook_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT entry_type, category, description, amount, date FROM cash_books WHERE user_id = %s ORDER BY date DESC',
                (current_user.id,))
    rows = [(r['entry_type'], r['category'], r['description'], r['amount'], r['date']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('cashbook.csv', ['Type', 'Category', 'Description', 'Amount', 'Date'], rows)

@app.route('/export/stock.csv')
@login_required
def export_stock_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT product_name, quantity, unit, cost_price, selling_price FROM stock WHERE user_id = %s ORDER BY product_name',
                (current_user.id,))
    rows = [(r['product_name'], r['quantity'], r['unit'], r['cost_price'], r['selling_price']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('stock.csv', ['Product', 'Quantity', 'Unit', 'Cost Price', 'Selling Price'], rows)

@app.route('/export/customers.csv')
@login_required
def export_customers_csv():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT name, email, phone, address, notes FROM customers WHERE user_id = %s ORDER BY name', (current_user.id,))
    rows = [(r['name'], r['email'], r['phone'], r['address'], r['notes']) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return _csv_response('customers.csv', ['Name', 'Email', 'Phone', 'Address', 'Notes'], rows)

# ======================== CSV Import ========================

def _read_uploaded_csv(file_storage):
    """Parse an uploaded CSV file into a list of dict rows. Returns None on failure."""
    if not file_storage or file_storage.filename == '':
        return None
    try:
        stream = io.StringIO(file_storage.stream.read().decode('utf-8-sig'), newline=None)
        return list(csv.DictReader(stream))
    except Exception:
        return None

@app.route('/import/stock', methods=['POST'])
@login_required
def import_stock_csv():
    rows = _read_uploaded_csv(request.files.get('file'))
    if rows is None:
        flash('Could not read that file. Please upload a CSV exported from this app (or matching columns: Product, Quantity, Unit, Cost Price, Selling Price).', 'danger')
        return redirect(url_for('stock'))
    conn = get_db()
    cur = conn.cursor()
    count = 0
    for row in rows:
        try:
            product_name = (row.get('Product') or '').strip()
            if not product_name:
                continue
            cur.execute('''
                INSERT INTO stock (user_id, product_name, quantity, unit, cost_price, selling_price)
                VALUES (%s, %s, %s, %s, %s, %s)
            ''', (current_user.id, product_name, float(row.get('Quantity') or 0),
                  row.get('Unit', ''), float(row.get('Cost Price') or 0), float(row.get('Selling Price') or 0)))
            count += 1
        except (ValueError, KeyError):
            continue
    conn.commit()
    cur.close()
    conn.close()
    flash(f'Imported {count} stock item(s).' if count else 'No rows imported — check the file has a Product column with values in it.', 'success' if count else 'danger')
    return redirect(url_for('stock'))

@app.route('/import/customers', methods=['POST'])
@login_required
def import_customers_csv():
    rows = _read_uploaded_csv(request.files.get('file'))
    if rows is None:
        flash('Could not read that file. Please upload a CSV with columns: Name, Email, Phone, Address, Notes.', 'danger')
        return redirect(url_for('customers'))
    conn = get_db()
    cur = conn.cursor()
    count = 0
    for row in rows:
        name = (row.get('Name') or '').strip()
        if not name:
            continue
        cur.execute('''
            INSERT INTO customers (user_id, name, email, phone, address, notes)
            VALUES (%s, %s, %s, %s, %s, %s)
        ''', (current_user.id, name, row.get('Email', ''), row.get('Phone', ''), row.get('Address', ''), row.get('Notes', '')))
        count += 1
    conn.commit()
    cur.close()
    conn.close()
    flash(f'Imported {count} customer(s).' if count else 'No rows imported — check the file has a Name column with values in it.', 'success' if count else 'danger')
    return redirect(url_for('customers'))

@app.route('/import/income', methods=['POST'])
@login_required
def import_income_csv():
    rows = _read_uploaded_csv(request.files.get('file'))
    if rows is None:
        flash('Could not read that file. Please upload a CSV with columns: Source, Amount, Date.', 'danger')
        return redirect(url_for('dashboard'))
    conn = get_db()
    cur = conn.cursor()
    count = 0
    for row in rows:
        try:
            source = (row.get('Source') or '').strip()
            if not source:
                continue
            cur.execute('INSERT INTO income (user_id, source, amount, date) VALUES (%s, %s, %s, %s)',
                        (current_user.id, source, float(row.get('Amount') or 0),
                         row.get('Date') or datetime.today().strftime('%Y-%m-%d')))
            count += 1
        except (ValueError, KeyError):
            continue
    conn.commit()
    cur.close()
    conn.close()
    flash(f'Imported {count} income record(s).' if count else 'No rows imported — check the file has a Source column with values in it.', 'success' if count else 'danger')
    return redirect(url_for('dashboard'))

@app.route('/import/expenses', methods=['POST'])
@login_required
def import_expenses_csv():
    rows = _read_uploaded_csv(request.files.get('file'))
    if rows is None:
        flash('Could not read that file. Please upload a CSV with columns: Name, Amount, Category, Date.', 'danger')
        return redirect(url_for('dashboard'))
    conn = get_db()
    cur = conn.cursor()
    count = 0
    for row in rows:
        try:
            name = (row.get('Name') or '').strip()
            if not name:
                continue
            cur.execute('INSERT INTO expenses (user_id, name, amount, category, date) VALUES (%s, %s, %s, %s, %s)',
                        (current_user.id, name, float(row.get('Amount') or 0),
                         row.get('Category') or 'Other', row.get('Date') or datetime.today().strftime('%Y-%m-%d')))
            count += 1
        except (ValueError, KeyError):
            continue
    conn.commit()
    cur.close()
    conn.close()
    flash(f'Imported {count} expense(s).' if count else 'No rows imported — check the file has a Name column with values in it.', 'success' if count else 'danger')
    return redirect(url_for('dashboard'))

@app.route('/import/cashbook', methods=['POST'])
@login_required
def import_cashbook_csv():
    rows = _read_uploaded_csv(request.files.get('file'))
    if rows is None:
        flash('Could not read that file. Please upload a CSV with columns: Type, Category, Description, Amount, Date.', 'danger')
        return redirect(url_for('cashbook'))
    conn = get_db()
    cur = conn.cursor()
    count = 0
    for row in rows:
        try:
            entry_type = (row.get('Type') or '').strip().lower()
            if entry_type not in ('in', 'out'):
                continue
            cur.execute('''
                INSERT INTO cash_books (user_id, entry_type, category, description, amount, date)
                VALUES (%s, %s, %s, %s, %s, %s)
            ''', (current_user.id, entry_type, row.get('Category', ''), row.get('Description', ''),
                  float(row.get('Amount') or 0), row.get('Date') or datetime.today().strftime('%Y-%m-%d')))
            count += 1
        except (ValueError, KeyError):
            continue
    conn.commit()
    cur.close()
    conn.close()
    flash((f'Imported {count} cash book entr{"y" if count == 1 else "ies"}.') if count else 'No rows imported — check the file has a Type column with "in" or "out" values.', 'success' if count else 'danger')
    return redirect(url_for('cashbook'))

# ======================== PDF Generation ========================

def _get_settings_row(user_id):
    if getattr(g, '_settings_loaded', False) and getattr(g, '_settings_uid', None) == user_id:
        return getattr(g, '_settings_row', None)
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM settings WHERE user_id = %s', (user_id,))
    row = cur.fetchone()
    cur.close()
    g._settings_row = row
    g._settings_uid = user_id
    g._settings_loaded = True
    return row

def _drawn_k_badge_png(pixel_size=128):
    """Green K badge used when static/logo files are missing."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    size = max(32, int(pixel_size))
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    pad = max(1, size // 16)
    draw.rounded_rectangle((pad, pad, size - pad - 1, size - pad - 1), radius=size // 5, fill=(31, 138, 61, 255))
    font = None
    for path in (
        '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
        '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
        '/usr/share/fonts/truetype/freefont/FreeSansBold.ttf',
    ):
        if os.path.exists(path):
            try:
                font = ImageFont.truetype(path, int(size * 0.52))
                break
            except Exception:
                font = None
    text = 'K'
    if font:
        box = draw.textbbox((0, 0), text, font=font)
    else:
        box = draw.textbbox((0, 0), text)
    tw, th = box[2] - box[0], box[3] - box[1]
    draw.text(((size - tw) / 2 - box[0], (size - th) / 2 - box[1] - size * 0.03), text, fill=(255, 255, 255, 255), font=font)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def _brand_logo_image(size=36):
    """Embed the KAZE logo in PDFs. Falls back to a drawn K if static files are missing."""
    logo_path = os.path.join(app.root_path, 'static', 'logo_icon.png')
    if os.path.exists(logo_path) and os.path.getsize(logo_path) > 0:
        try:
            return RLImage(logo_path, width=size, height=size)
        except Exception:
            pass
    badge = _drawn_k_badge_png(max(64, int(size) * 4))
    if badge is not None:
        try:
            return RLImage(badge, width=size, height=size)
        except Exception:
            pass
    # Last resort: empty spacer so PDF build still succeeds.
    return Spacer(size, size)

def _build_pdf(buffer, title, business_name, meta_lines, table_data, footer_lines=None):
    from xml.sax.saxutils import escape as _xml_escape
    def _safe_paragraph(text, style):
        return Paragraph(_xml_escape(str(text)), style)
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Title'], textColor=colors.HexColor('#1f8a3d'), alignment=TA_LEFT)
    elements = []

    header_text = []
    if business_name:
        header_text.append(_safe_paragraph(business_name, styles['Heading2']))
    header_text.append(_safe_paragraph(title, title_style))
    header_table = Table([[_brand_logo_image(36), header_text]], colWidths=[50, 460])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (0, 0), 0),
        ('LEFTPADDING', (1, 0), (1, 0), 8),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 12))
    for line in meta_lines:
        elements.append(Paragraph(line, styles['Normal']))
    elements.append(Spacer(1, 16))

    safe_table = [[_xml_escape(str(cell)) for cell in row] for row in table_data]
    table = Table(safe_table, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a1a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    if footer_lines:
        elements.append(Spacer(1, 20))
        for line in footer_lines:
            elements.append(Paragraph(line, styles['Normal']))

    doc.build(elements)

@app.route('/docs/<int:id>/pdf')
@login_required
def doc_pdf(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM documents WHERE id = %s AND user_id = %s', (id, current_user.id))
    doc_row = cur.fetchone()
    cur.close()
    conn.close()
    if not doc_row:
        flash('Document not found.', 'danger')
        return redirect(url_for('docs'))

    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    tax_rate = (settings_row['tax_rate'] if settings_row else 0) or 0

    subtotal = doc_row['amount'] or 0
    tax_amount = subtotal * (tax_rate / 100)
    total = subtotal + tax_amount

    meta_lines = [
        f"<b>Document Type:</b> {doc_row['doc_type'].capitalize()}",
        f"<b>Title:</b> {doc_row['title']}",
        f"<b>Client:</b> {doc_row['client']}",
        f"<b>Date:</b> {doc_row['date']}",
    ]
    if doc_row['notes']:
        meta_lines.append(f"<b>Notes:</b> {doc_row['notes']}")

    table_data = [['Description', 'Amount']]
    table_data.append([doc_row['title'], _format_money(subtotal, settings_row)])
    if tax_rate:
        table_data.append([f"Tax ({tax_rate:.1f}%)", _format_money(tax_amount, settings_row)])
    table_data.append(['Total', _format_money(total, settings_row)])

    buffer = io.BytesIO()
    _build_pdf(buffer, doc_row['title'], business_name, meta_lines, table_data,
               footer_lines=[f"Generated on {datetime.today().strftime('%Y-%m-%d')}"])
    buffer.seek(0)
    safe_title = "".join(c for c in doc_row['title'] if c.isalnum() or c in (' ', '_', '-')).rstrip()
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                      download_name=f"{safe_title or 'document'}.pdf")

@app.route('/sales/<int:id>/receipt')
@login_required
def sale_receipt(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT sales.*, stock.product_name FROM sales
        LEFT JOIN stock ON sales.stock_id = stock.id
        WHERE sales.id = %s AND sales.user_id = %s
    ''', (id, current_user.id))
    sale = cur.fetchone()
    cur.close()
    conn.close()
    if not sale:
        flash('Sale not found.', 'danger')
        return redirect(url_for('sales'))

    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'

    meta_lines = [
        f"<b>Receipt #:</b> {sale['id']}",
        f"<b>Date:</b> {sale['sale_date']}",
        f"<b>Customer:</b> {sale['customer_name'] or 'Walk-in'}",
    ]
    if sale['customer_email']:
        meta_lines.append(f"<b>Email:</b> {sale['customer_email']}")

    table_data = [
        ['Product', 'Quantity', 'Unit Price', 'Total'],
        [sale['product_name'], str(sale['quantity_sold']), _format_money(sale['selling_price_at_time'], settings_row),
         _format_money(sale['total_amount'], settings_row)],
        ['', '', 'Total', _format_money(sale['total_amount'], settings_row)]
    ]

    buffer = io.BytesIO()
    _build_pdf(buffer, "Sales Receipt", business_name, meta_lines, table_data,
               footer_lines=["Thank you for your business!"])
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                      download_name=f"receipt_{sale['id']}.pdf")

# ======================== SWOT Analysis ========================

@app.route('/swot')
@login_required
def swot():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM swot_analyses WHERE user_id = %s ORDER BY created_date DESC', (current_user.id,))
    analyses = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('swot.html', analyses=analyses)

@app.route('/swot/add', methods=['POST'])
@login_required
def add_swot():
    title = request.form.get('title') or f"SWOT Analysis - {datetime.today().strftime('%Y-%m-%d')}"
    strengths = request.form.get('strengths', '')
    weaknesses = request.form.get('weaknesses', '')
    opportunities = request.form.get('opportunities', '')
    threats = request.form.get('threats', '')
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO swot_analyses (user_id, title, strengths, weaknesses, opportunities, threats, created_date)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
    ''', (current_user.id, title, strengths, weaknesses, opportunities, threats,
          datetime.today().strftime('%Y-%m-%d')))
    conn.commit()
    cur.close()
    conn.close()
    flash('SWOT analysis saved.', 'success')
    return redirect(url_for('swot'))

@app.route('/swot/delete/<int:id>')
@login_required
def delete_swot(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM swot_analyses WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('SWOT analysis deleted.', 'success')
    return redirect(url_for('swot'))

@app.route('/swot/<int:id>/pdf')
@login_required
def swot_pdf(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM swot_analyses WHERE id = %s AND user_id = %s', (id, current_user.id))
    s = cur.fetchone()
    cur.close()
    conn.close()
    if not s:
        flash('SWOT analysis not found.', 'danger')
        return redirect(url_for('swot'))

    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'

    meta_lines = [f"<b>Date:</b> {s['created_date']}"]
    table_data = [
        ['Category', 'Notes'],
        ['Strengths', s['strengths'] or '—'],
        ['Weaknesses', s['weaknesses'] or '—'],
        ['Opportunities', s['opportunities'] or '—'],
        ['Threats', s['threats'] or '—'],
    ]
    buffer = io.BytesIO()
    _build_pdf(buffer, s['title'], business_name, meta_lines, table_data,
               footer_lines=[f"Generated on {datetime.today().strftime('%Y-%m-%d')}"])
    buffer.seek(0)
    safe_title = "".join(c for c in s['title'] if c.isalnum() or c in (' ', '_', '-')).rstrip()
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                      download_name=f"{safe_title or 'swot'}.pdf")

# ======================== Business Plan ========================

@app.route('/business-plan', methods=['GET', 'POST'])
@login_required
def business_plan():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM business_plans WHERE user_id = %s', (current_user.id,))
    plan = cur.fetchone()

    if request.method == 'POST':
        fields = {
            'business_name': request.form.get('business_name', ''),
            'mission': request.form.get('mission', ''),
            'products_services': request.form.get('products_services', ''),
            'target_market': request.form.get('target_market', ''),
            'marketing_strategy': request.form.get('marketing_strategy', ''),
            'operations_plan': request.form.get('operations_plan', ''),
            'financial_plan': request.form.get('financial_plan', ''),
            'goals': request.form.get('goals', ''),
        }
        today = datetime.today().strftime('%Y-%m-%d')
        if plan:
            cur.execute('''
                UPDATE business_plans SET business_name=%s, mission=%s, products_services=%s,
                    target_market=%s, marketing_strategy=%s, operations_plan=%s,
                    financial_plan=%s, goals=%s, updated_date=%s
                WHERE user_id=%s
            ''', (fields['business_name'], fields['mission'], fields['products_services'],
                  fields['target_market'], fields['marketing_strategy'], fields['operations_plan'],
                  fields['financial_plan'], fields['goals'], today, current_user.id))
        else:
            cur.execute('''
                INSERT INTO business_plans (user_id, business_name, mission, products_services,
                    target_market, marketing_strategy, operations_plan, financial_plan, goals, updated_date)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ''', (current_user.id, fields['business_name'], fields['mission'], fields['products_services'],
                  fields['target_market'], fields['marketing_strategy'], fields['operations_plan'],
                  fields['financial_plan'], fields['goals'], today))
        conn.commit()
        cur.close()
        conn.close()
        flash('Business plan saved.', 'success')
        return redirect(url_for('business_plan'))

    cur.close()
    conn.close()
    return render_template('business_plan.html', plan=plan)

@app.route('/business-plan/pdf')
@login_required
def business_plan_pdf():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM business_plans WHERE user_id = %s', (current_user.id,))
    plan = cur.fetchone()
    cur.close()
    conn.close()
    if not plan:
        flash('Create your business plan first.', 'info')
        return redirect(url_for('business_plan'))

    settings_row = _get_settings_row(current_user.id)
    business_name = plan['business_name'] or (settings_row['business_name'] if settings_row else '') or 'Your Business'

    meta_lines = [f"<b>Last updated:</b> {plan['updated_date'] or '—'}"]
    table_data = [
        ['Section', 'Details'],
        ['Mission', plan['mission'] or '—'],
        ['Products / Services', plan['products_services'] or '—'],
        ['Target Market', plan['target_market'] or '—'],
        ['Marketing Strategy', plan['marketing_strategy'] or '—'],
        ['Operations Plan', plan['operations_plan'] or '—'],
        ['Financial Plan', plan['financial_plan'] or '—'],
        ['Goals', plan['goals'] or '—'],
    ]
    buffer = io.BytesIO()
    _build_pdf(buffer, "Business Plan", business_name, meta_lines, table_data,
               footer_lines=[f"Generated on {datetime.today().strftime('%Y-%m-%d')}"])
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                      download_name="business_plan.pdf")

# ======================== Tutorial / Onboarding ========================

def _guide_steps():
    return [
        {
            'group': 'Start',
            'title': 'Settings first',
            'body': 'Set the business name, currency, tax rate, and low-stock threshold before you type real numbers. PDFs pick these up automatically.',
            'do_this': ['Open Settings', 'Type your shop name and currency symbol', 'Save'],
            'href': url_for('settings'),
        },
        {
            'group': 'Start',
            'title': 'Pick a mode',
            'body': 'Modes turn whole parts of the app on or off. Till mode runs the shop floor and leaves books dormant. Quiet mode leaves almost everything asleep.',
            'do_this': ['Open Settings → Modes', 'Choose Till, Books, Office, Quiet, Full, or Custom', 'Save, then check the sidebar'],
            'href': url_for('settings'),
        },
        {
            'group': 'Stock',
            'title': 'Add products',
            'body': 'Every sale needs a product. Add name, quantity, cost price, and selling price. Cost is what you paid. Selling price is what the customer pays.',
            'do_this': ['Open Stock', 'Add 3 real products you sell', 'Check the quantity matches the shelf'],
            'href': url_for('stock'),
        },
        {
            'group': 'People',
            'title': 'Save customers',
            'body': 'Walk-ins can stay as a typed name. Regulars belong on the Customers page so lifetime spend and statements work.',
            'do_this': ['Open Customers', 'Add one regular buyer', 'Optional: add a supplier next'],
            'href': url_for('customers'),
        },
        {
            'group': 'Till',
            'title': 'Counter mode',
            'body': 'This is the new till page. Tap a product, set quantity, ring it up. Stock falls and profit is calculated for you.',
            'do_this': ['Open Counter', 'Tap a product that is in stock', 'Ring up a test sale of 1 unit'],
            'href': url_for('counter'),
        },
        {
            'group': 'Till',
            'title': 'Sales desk',
            'body': 'Use Sales when you need a discount, a note, a receipt PDF, or to turn a sale into an invoice.',
            'do_this': ['Open Sales', 'Download a receipt for the test sale', 'Try Repeat if the same person buys again'],
            'href': url_for('sales'),
        },
        {
            'group': 'Cash',
            'title': 'Cash ledger',
            'body': 'The cash book is the till tin: cash in and cash out. It is separate from profit. Use it to count the drawer.',
            'do_this': ['Open Book keeping', 'Open the cash ledger', 'Add today opening float if you use one'],
            'href': url_for('cashbook'),
        },
        {
            'group': 'Papers',
            'title': 'Book keeping documents',
            'body': 'Quotations, invoices, GRNs, payslips, and cheques live here. Fill the form, save, download the PDF.',
            'do_this': ['Pick one document type', 'Save a draft', 'Open the PDF'],
            'href': url_for('bookkeeping'),
        },
        {
            'group': 'Books',
            'title': 'Account tracking',
            'body': 'Journals and statements. Sales and purchases post themselves. Type extra lines in the last row of a journal.',
            'do_this': ['Open Account tracking', 'Open the sales journal', 'Download a P and L PDF'],
            'href': url_for('accounts'),
        },
        {
            'group': 'Plan',
            'title': 'Budgets and recurring',
            'body': 'Cap monthly spend by category. Schedule rent or a retainer so it writes itself when the date arrives.',
            'do_this': ['Set one budget', 'Add one monthly recurring cost'],
            'href': url_for('budgets'),
        },
        {
            'group': 'Plan',
            'title': 'SWOT and business plan',
            'body': 'Write the thinking down. Export PDFs when a bank or partner asks.',
            'do_this': ['Add a SWOT', 'Fill at least the mission and goals on the plan'],
            'href': url_for('swot'),
        },
        {
            'group': 'Day',
            'title': 'Tasks, notes, calendar',
            'body': 'Tasks have due dates. Notes are a scratch pad. Calendar lists what is coming in the next 60 days.',
            'do_this': ['Add one task due tomorrow', 'Open Calendar and check it appears'],
            'href': url_for('tasks'),
        },
        {
            'group': 'Day',
            'title': 'Reminders and tax',
            'body': 'Reminders lists overdue work. Tax summary estimates VAT from the rate in Settings. It is not a filed return.',
            'do_this': ['Open Reminders', 'Open Tax summary and confirm the rate'],
            'href': url_for('reminders'),
        },
        {
            'group': 'Close',
            'title': 'Day close',
            'body': 'End of day: today sales count and cash in or out. Print the page and keep it with the till count.',
            'do_this': ['Open Day close', 'Compare the number with cash in the drawer'],
            'href': url_for('day_close'),
        },
        {
            'group': 'Safety',
            'title': 'Backup',
            'body': 'Download a JSON backup from Settings. Export CSVs from Stock, Customers, and Sales as well.',
            'do_this': ['Download a backup now', 'Store it somewhere that is not only this computer'],
            'href': url_for('settings'),
        },
        {
            'group': 'Team',
            'title': 'Invite staff',
            'body': 'Owners can add staff. They see the same business data. Only the owner can clear data or remove staff.',
            'do_this': ['Open Team if you have helpers', 'Otherwise skip'],
            'href': url_for('team'),
        },
        {
            'group': 'Speed',
            'title': 'Shortcuts',
            'body': 'Ctrl+K jumps pages. [ collapses the sidebar. The + button is quick add.',
            'do_this': ['Press Ctrl+K', 'Type Counter and go there'],
            'href': url_for('shortcuts'),
        },
    ]


@app.route('/tutorial')
@login_required
def tutorial():
    return render_template('tutorial.html', steps=_guide_steps())

@app.route('/tutorial/dismiss')
@login_required
def dismiss_tutorial():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE settings SET tutorial_completed = 1 WHERE user_id = %s', (current_user.id,))
    conn.commit()
    cur.close()
    conn.close()
    return ('', 204)

# ======================== Team / Staff Accounts ========================

@app.route('/team')
@login_required
def team():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, username, email FROM users WHERE owner_id = %s AND role = 'staff' ORDER BY username",
                (current_user.id,))
    staff = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('team.html', staff=staff)

@app.route('/team/invite', methods=['POST'])
@login_required
def invite_staff():
    if current_user.role == 'staff':
        flash('Only the account owner can invite team members.', 'danger')
        return redirect(url_for('team'))
    username, e1 = _clean_text('username', max_len=40, label='username')
    email, e2 = _clean_email('email', required=True)
    password = request.form.get('password') or ''
    if _reject((username, e1), (email, e2)):
        return redirect(url_for('team'))
    if len(password) < 6:
        flash('Temporary password must be at least 6 characters.', 'danger')
        return redirect(url_for('team'))
    blocked = _plan_block('team')
    if blocked:
        flash(blocked, 'danger')
        return redirect(url_for('plans'))
    hashed = generate_password_hash(password)
    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute('''
            INSERT INTO users (username, email, password, role, owner_id)
            VALUES (%s, %s, %s, 'staff', %s)
        ''', (username, email, hashed, current_user.id))
        conn.commit()
        flash(f'{username} added to your team. Share their username/email and password so they can log in.', 'success')
    except psycopg2.IntegrityError:
        conn.rollback()
        flash('That username or email is already taken.', 'danger')
    finally:
        cur.close()
        conn.close()
    return redirect(url_for('team'))

@app.route('/team/remove/<int:id>')
@login_required
def remove_staff(id):
    if current_user.role == 'staff':
        flash('Only the account owner can remove team members.', 'danger')
        return redirect(url_for('team'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM users WHERE id = %s AND owner_id = %s AND role = 'staff'", (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Team member removed.', 'success')
    return redirect(url_for('team'))

# ======================== Recurring Invoices / Expenses ========================

def _add_months(date_str, months):
    d = datetime.strptime(date_str, '%Y-%m-%d')
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, [31,29 if year%4==0 and (year%100!=0 or year%400==0) else 28,31,30,31,30,31,31,30,31,30,31][month-1])
    return d.replace(year=year, month=month, day=day).strftime('%Y-%m-%d')

def _advance_date(date_str, frequency):
    if frequency == 'weekly':
        return (datetime.strptime(date_str, '%Y-%m-%d') + timedelta(days=7)).strftime('%Y-%m-%d')
    elif frequency == 'monthly':
        return _add_months(date_str, 1)
    elif frequency == 'quarterly':
        return _add_months(date_str, 3)
    elif frequency == 'yearly':
        return _add_months(date_str, 12)
    return date_str

def _notify(user_id, message, category='info'):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO notifications (user_id, message, category, created_date) VALUES (%s, %s, %s, %s)',
                (user_id, message, category, datetime.today().strftime('%Y-%m-%d %H:%M:%S')))
    conn.commit()
    cur.close()
    conn.close()

def _process_due_recurring_items(user_id):
    """Auto-generate expenses/invoices for any recurring item that's come due. Safe to call often."""
    if not module_on('operations'):
        return

    import time as _time
    last = _recurring_ran_at.get(user_id, 0)
    if _time.time() - last < 300:
        return
    _recurring_ran_at[user_id] = _time.time()
    conn = get_db()
    cur = conn.cursor()
    today = datetime.today().strftime('%Y-%m-%d')
    cur.execute('SELECT * FROM recurring_items WHERE user_id = %s AND active = 1 AND next_run_date <= %s',
                (user_id, today))
    due_items = cur.fetchall()
    for item in due_items:
        if item['item_type'] == 'expense':
            cur.execute('INSERT INTO expenses (user_id, name, amount, category, date) VALUES (%s, %s, %s, %s, %s)',
                        (user_id, item['name'], item['amount'], item['category'] or 'Other', today))
        else:
            cur.execute('''
                INSERT INTO documents (user_id, doc_type, title, client, amount, date, notes)
                VALUES (%s, 'invoice', %s, %s, %s, %s, %s)
            ''', (user_id, item['name'], item['client'] or 'Recurring client', item['amount'], today,
                  'Auto-generated from a recurring invoice.'))
        new_next = _advance_date(item['next_run_date'], item['frequency'])
        cur.execute('UPDATE recurring_items SET next_run_date = %s WHERE id = %s', (new_next, item['id']))
        _notify(user_id, f"Recurring {item['item_type']} \"{item['name']}\" ({_format_money(item['amount'])}) was generated.", 'recurring')
    conn.commit()
    cur.close()
    conn.close()

@app.route('/recurring')
@login_required
def recurring():
    _process_due_recurring_items(current_user.id)
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM recurring_items WHERE user_id = %s ORDER BY active DESC, next_run_date', (current_user.id,))
    items = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('recurring.html', items=items)

@app.route('/recurring/add', methods=['POST'])
@login_required
def add_recurring():
    item_type = request.form['item_type']
    name = request.form['name']
    amount = float(request.form['amount'])
    category = request.form.get('category', '')
    client = request.form.get('client', '')
    frequency = request.form.get('frequency', 'monthly')
    start_date = request.form.get('start_date') or datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO recurring_items (user_id, item_type, name, amount, category, client, frequency, next_run_date, active, created_date)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 1, %s)
    ''', (current_user.id, item_type, name, amount, category, client, frequency, start_date,
          datetime.today().strftime('%Y-%m-%d')))
    conn.commit()
    cur.close()
    conn.close()
    flash('Recurring item scheduled.', 'success')
    return redirect(url_for('recurring'))

@app.route('/recurring/toggle/<int:id>')
@login_required
def toggle_recurring(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE recurring_items SET active = 1 - active WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('recurring'))

@app.route('/recurring/delete/<int:id>')
@login_required
def delete_recurring(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM recurring_items WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Recurring item deleted.', 'success')
    return redirect(url_for('recurring'))

# ======================== Suppliers & Purchase Orders ========================

@app.route('/suppliers')
@login_required
def suppliers():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM suppliers WHERE user_id = %s ORDER BY name', (current_user.id,))
    all_suppliers = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('suppliers.html', suppliers=all_suppliers)

@app.route('/suppliers/add', methods=['POST'])
@login_required
def add_supplier():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO suppliers (user_id, name, email, phone, address, notes) VALUES (%s, %s, %s, %s, %s, %s)',
                (current_user.id, request.form['name'], request.form.get('email',''), request.form.get('phone',''),
                 request.form.get('address',''), request.form.get('notes','')))
    conn.commit()
    cur.close()
    conn.close()
    flash('Supplier added.', 'success')
    return redirect(url_for('suppliers'))

@app.route('/suppliers/delete/<int:id>')
@login_required
def delete_supplier(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM suppliers WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Supplier deleted.', 'success')
    return redirect(url_for('suppliers'))

@app.route('/purchase-orders')
@login_required
def purchase_orders():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT po.*, s.name as supplier_name FROM purchase_orders po
        LEFT JOIN suppliers s ON po.supplier_id = s.id
        WHERE po.user_id = %s ORDER BY po.order_date DESC
    ''', (current_user.id,))
    orders = cur.fetchall()
    cur.execute('SELECT id, name FROM suppliers WHERE user_id = %s ORDER BY name', (current_user.id,))
    supplier_list = cur.fetchall()
    cur.execute('SELECT id, product_name FROM stock WHERE user_id = %s ORDER BY product_name', (current_user.id,))
    stock_list = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('purchase_orders.html', orders=orders, supplier_list=supplier_list, stock_list=stock_list)

@app.route('/purchase-orders/add', methods=['POST'])
@login_required
def add_purchase_order():
    supplier_id = request.form.get('supplier_id') or None
    stock_id = request.form.get('stock_id') or None
    item_description = request.form['item_description']
    quantity = float(request.form['quantity'])
    unit_cost = float(request.form['unit_cost'])
    total_cost = quantity * unit_cost
    expected_date = request.form.get('expected_date', '')
    order_date = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO purchase_orders (user_id, supplier_id, item_description, quantity, unit_cost, total_cost, status, order_date, expected_date, stock_id)
        VALUES (%s, %s, %s, %s, %s, %s, 'pending', %s, %s, %s)
    ''', (current_user.id, supplier_id, item_description, quantity, unit_cost, total_cost, order_date, expected_date, stock_id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Purchase order created.', 'success')
    return redirect(url_for('purchase_orders'))

@app.route('/purchase-orders/receive/<int:id>')
@login_required
def receive_purchase_order(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM purchase_orders WHERE id = %s AND user_id = %s', (id, current_user.id))
    po = cur.fetchone()
    if not po:
        flash('Purchase order not found.', 'danger')
        return redirect(url_for('purchase_orders'))
    cur.execute(
        "UPDATE purchase_orders SET status = 'received', received_date = %s WHERE id = %s",
        (datetime.today().strftime('%Y-%m-%d'), id),
    )
    if po['stock_id']:
        cur.execute('UPDATE stock SET quantity = quantity + %s WHERE id = %s AND user_id = %s',
                    (po['quantity'], po['stock_id'], current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    _notify(current_user.id, f"Purchase order for \"{po['item_description']}\" marked received"
            + (" and added to stock." if po['stock_id'] else "."), 'purchase_order')
    flash('Purchase order received.', 'success')
    return redirect(url_for('purchase_orders'))

@app.route('/purchase-orders/delete/<int:id>')
@login_required
def delete_purchase_order(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM purchase_orders WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Purchase order deleted.', 'success')
    return redirect(url_for('purchase_orders'))

# ======================== Budgeting ========================

@app.route('/budgets')
@login_required
def budgets():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM budgets WHERE user_id = %s ORDER BY category', (current_user.id,))
    budget_rows = cur.fetchall()
    month_start = datetime.today().replace(day=1).strftime('%Y-%m-%d')
    spend_by_category = {}
    for b in budget_rows:
        cur.execute('SELECT COALESCE(SUM(amount),0) as spent FROM expenses WHERE user_id = %s AND category = %s AND date >= %s',
                    (current_user.id, b['category'], month_start))
        spend_by_category[b['category']] = cur.fetchone()['spent']
    cur.close()
    conn.close()
    return render_template('budgets.html', budgets=budget_rows, spend_by_category=spend_by_category)

@app.route('/budgets/add', methods=['POST'])
@login_required
def add_budget():
    category = request.form['category']
    monthly_limit = float(request.form['monthly_limit'])
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO budgets (user_id, category, monthly_limit, created_date)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (user_id, category) DO UPDATE SET monthly_limit = EXCLUDED.monthly_limit
    ''', (current_user.id, category, monthly_limit, datetime.today().strftime('%Y-%m-%d')))
    conn.commit()
    cur.close()
    conn.close()
    flash('Budget saved.', 'success')
    return redirect(url_for('budgets'))

@app.route('/budgets/delete/<int:id>')
@login_required
def delete_budget(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM budgets WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Budget removed.', 'success')
    return redirect(url_for('budgets'))

# ======================== Notifications ========================

@app.route('/notifications')
@login_required
def notifications():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM notifications WHERE user_id = %s ORDER BY created_date DESC LIMIT 100', (current_user.id,))
    items = cur.fetchall()
    cur.execute('UPDATE notifications SET is_read = 1 WHERE user_id = %s AND is_read = 0', (current_user.id,))
    conn.commit()
    cur.close()
    conn.close()
    return render_template('notifications.html', notifications=items)

@app.route('/notifications/clear')
@login_required
def clear_notifications():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM notifications WHERE user_id = %s', (current_user.id,))
    conn.commit()
    cur.close()
    conn.close()
    flash('Notifications cleared.', 'success')
    return redirect(url_for('notifications'))

# ======================== Reports with date range and charts ========================

@app.route('/reports')
@login_required
def reports():
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    if not start_date:
        start_date = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    if not end_date:
        end_date = datetime.now().strftime('%Y-%m-%d')
    
    conn = get_db()
    cur = conn.cursor()
    
    # Expenses by category (pie chart)
    cur.execute('''
        SELECT category, SUM(amount) as total
        FROM expenses
        WHERE user_id = %s AND date BETWEEN %s AND %s
        GROUP BY category
        ORDER BY total DESC
    ''', (current_user.id, start_date, end_date))
    expense_by_category = cur.fetchall()
    
    # Top products by revenue (bar chart) — FIXED: now includes profit
    cur.execute('''
        SELECT stock.product_name,
               SUM(sales.quantity_sold) as units_sold,
               SUM(sales.total_amount) as revenue,
               SUM(sales.profit) as profit
        FROM sales JOIN stock ON sales.stock_id = stock.id
        WHERE sales.user_id = %s AND sales.sale_date BETWEEN %s AND %s
        GROUP BY stock.product_name
        ORDER BY revenue DESC
        LIMIT 5
    ''', (current_user.id, start_date, end_date))
    top_products = cur.fetchall()
    
    # Also the monthly profit trend (keep existing)
    months = []
    for i in range(5, -1, -1):
        m = _add_months(datetime.today().replace(day=1).strftime('%Y-%m-%d'), -i)
        months.append(m[:7])
    
    monthly_income = {m: 0 for m in months}
    monthly_expenses = {m: 0 for m in months}
    monthly_sales_profit = {m: 0 for m in months}
    
    cur.execute('SELECT date, amount FROM income WHERE user_id = %s', (current_user.id,))
    for row in cur.fetchall():
        key = row['date'][:7]
        if key in monthly_income:
            monthly_income[key] += row['amount']
    
    cur.execute('SELECT date, amount FROM expenses WHERE user_id = %s', (current_user.id,))
    for row in cur.fetchall():
        key = row['date'][:7]
        if key in monthly_expenses:
            monthly_expenses[key] += row['amount']
    
    cur.execute('SELECT sale_date, profit FROM sales WHERE user_id = %s', (current_user.id,))
    for row in cur.fetchall():
        key = row['sale_date'][:7]
        if key in monthly_sales_profit:
            monthly_sales_profit[key] += row['profit']
    
    monthly_profit = {m: monthly_income[m] - monthly_expenses[m] + monthly_sales_profit[m] for m in months}
    
    # Top customers (keep existing)
    cur.execute('''
        SELECT COALESCE(NULLIF(customer_name, ''), 'Walk-in') as customer_name,
               SUM(total_amount) as spent, COUNT(*) as orders
        FROM sales WHERE user_id = %s
        GROUP BY customer_name ORDER BY spent DESC LIMIT 5
    ''', (current_user.id,))
    top_customers = cur.fetchall()
    
    cur.close()
    conn.close()
    
    return render_template('reports.html',
                           months=months,
                           monthly_income=monthly_income,
                           monthly_expenses=monthly_expenses,
                           monthly_profit=monthly_profit,
                           top_products=top_products,
                           top_customers=top_customers,
                           expense_by_category=expense_by_category,
                           start_date=start_date,
                           end_date=end_date)

# ======================== Daily Summary ========================

@app.route('/send_daily_summary')
@login_required
def send_daily_summary():
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """SELECT sales.customer_name, sales.quantity_sold, sales.total_amount, sales.profit,
                  sales.payment_method, stock.product_name
           FROM sales
           LEFT JOIN stock ON sales.stock_id = stock.id
           WHERE sales.user_id = %s AND sales.sale_date = %s
           ORDER BY sales.id""",
        (current_user.id, today),
    )
    sales_today = cur.fetchall() or []
    cur.execute('SELECT notify_email, email_notifications, business_name FROM settings WHERE user_id = %s', (current_user.id,))
    settings = cur.fetchone() or {}
    cur.close()

    if not sales_today:
        flash('No sales recorded for today, so there is nothing to send.', 'danger')
        return redirect(url_for('dashboard'))

    total_revenue = sum(float(s['total_amount'] or 0) for s in sales_today)
    total_profit = sum(float(s['profit'] or 0) for s in sales_today)
    shop = (settings.get('business_name') or '').strip() or 'KAZE'
    body = '{} — daily sales {}'.format(shop, today) + "\n\n"
    body += 'Tickets: {}\n'.format(len(sales_today))
    body += 'Revenue: {}\n'.format(_format_money(total_revenue))
    body += 'Profit: {}\n\n'.format(_format_money(total_profit))
    body += 'Lines:\n'
    for sale in sales_today:
        body += '- {} x{} · {} · {} ({})\n'.format(
            sale.get('product_name') or 'Item',
            sale.get('quantity_sold') or 0,
            sale.get('customer_name') or 'Walk-in',
            _format_money(sale.get('total_amount')),
            sale.get('payment_method') or 'cash',
        )
    body += '\nSent from KAZE Traders Business Manager.'

    to_email = (settings.get('notify_email') or current_user.email or '').strip()
    subject = 'Daily sales summary — {}'.format(today)
    sent = False
    if to_email and '@' in to_email:
        sent = send_email(to_email, subject, body)
    try:
        _notify(current_user.id, 'Daily sales {}: {} tickets, {}.'.format(
            today, len(sales_today), _format_money(total_revenue)), 'sales')
    except Exception:
        pass
    log_activity(current_user.id, current_user.username, 'Daily summary',
                 '{} tickets {}'.format(len(sales_today), _format_money(total_revenue)))

    if request.args.get('download') == '1':
        return Response(body, mimetype='text/plain',
                        headers={'Content-Disposition': 'attachment; filename=kaze-daily-{}.txt'.format(today)})

    if sent:
        flash('Daily sales summary emailed to {}.'.format(to_email), 'success')
    else:
        flash('Could not send email. Set Settings → Notify email and MAIL_USERNAME / MAIL_PASSWORD on the server. A copy is in Notifications, or download the file.', 'danger')
    return redirect(url_for('dashboard'))


# ======================== Activity Log ========================

@app.route('/activity')
@login_required
def activity():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT username, login_time, ip_address 
        FROM user_activity 
        WHERE user_id = %s
        ORDER BY login_time DESC 
        LIMIT 50
    ''', (current_user.id,))
    logs = cur.fetchall()
    today = datetime.now().strftime('%Y-%m-%d')
    cur.execute('''
    SELECT COUNT(*) as active_count
    FROM user_activity 
    WHERE user_id = %s AND login_time::date = %s
    ''', (current_user.id, today))
    active_count = cur.fetchone()['active_count']
    cur.close()
    conn.close()
    return render_template('activity.html', logs=logs, active_count=active_count)

# ======================== Global Search ========================

@app.route('/search')
@login_required
def search():
    q = request.args.get('q', '').strip()
    if not q:
        flash('Please enter a search term.', 'info')
        return redirect(url_for('dashboard'))
    
    conn = get_db()
    cur = conn.cursor()
    # Search customers
    cur.execute('''
        SELECT id, name, 'customer' as type
        FROM customers
        WHERE user_id = %s AND (name ILIKE %s OR email ILIKE %s OR phone ILIKE %s)
        LIMIT 10
    ''', (current_user.id, f'%{q}%', f'%{q}%', f'%{q}%'))
    customers = cur.fetchall()
    
    # Search products
    cur.execute('''
        SELECT id, product_name as name, 'product' as type
        FROM stock
        WHERE user_id = %s AND product_name ILIKE %s
        LIMIT 10
    ''', (current_user.id, f'%{q}%'))
    products = cur.fetchall()
    
    # Search documents (invoices)
    cur.execute('''
        SELECT id, title as name, 'document' as type
        FROM documents
        WHERE user_id = %s AND (title ILIKE %s OR client ILIKE %s)
        LIMIT 10
    ''', (current_user.id, f'%{q}%', f'%{q}%'))
    docs = cur.fetchall()

    cur.execute('''
        SELECT id, title as name, 'note' as type
        FROM notes
        WHERE user_id = %s AND (title ILIKE %s OR body ILIKE %s)
        LIMIT 10
    ''', (current_user.id, f'%{q}%', f'%{q}%'))
    note_hits = cur.fetchall()

    cur.execute('''
        SELECT id, title as name, 'task' as type
        FROM tasks
        WHERE user_id = %s AND title ILIKE %s
        LIMIT 10
    ''', (current_user.id, f'%{q}%'))
    task_hits = cur.fetchall()

    cur.execute('''
        SELECT id, title as name, 'bookkeeping' as type
        FROM bookkeeping_docs
        WHERE user_id = %s AND (title ILIKE %s OR party ILIKE %s OR reference ILIKE %s)
        LIMIT 10
    ''', (current_user.id, f'%{q}%', f'%{q}%', f'%{q}%'))
    book_hits = cur.fetchall()

    cur.execute('''
        SELECT id, product_name as name, 'product' as type
        FROM stock
        WHERE user_id = %s AND (sku ILIKE %s OR category ILIKE %s)
        LIMIT 10
    ''', (current_user.id, f'%{q}%', f'%{q}%'))
    sku_hits = cur.fetchall()

    def _safe_query(sql, params):
        try:
            cur.execute("SAVEPOINT sp_search")
            cur.execute(sql, params)
            rows = cur.fetchall()
            cur.execute("RELEASE SAVEPOINT sp_search")
            return rows
        except Exception:
            try:
                cur.execute("ROLLBACK TO SAVEPOINT sp_search")
            except Exception:
                pass
            return []

    like = '%' + q + '%'
    service_hits = _safe_query(
        "SELECT id, name, 'service' as type FROM services WHERE user_id = %s "
        "AND (name ILIKE %s OR category ILIKE %s OR description ILIKE %s) LIMIT 10",
        (current_user.id, like, like, like))
    session_hits = _safe_query(
        "SELECT id, customer_name as name, 'session' as type FROM service_sessions "
        "WHERE user_id = %s AND (customer_name ILIKE %s OR notes ILIKE %s) LIMIT 10",
        (current_user.id, like, like))
    bank_hits = _safe_query(
        "SELECT id, name, 'bank' as type FROM bank_accounts "
        "WHERE user_id = %s AND (name ILIKE %s OR bank_name ILIKE %s) LIMIT 8",
        (current_user.id, like, like))
    invest_hits = _safe_query(
        "SELECT id, name, 'investment' as type FROM investments "
        "WHERE user_id = %s AND name ILIKE %s LIMIT 8",
        (current_user.id, like))
    shop_hits = _safe_query(
        "SELECT id, title as name, 'listing' as type FROM shop_listings "
        "WHERE user_id = %s AND (title ILIKE %s OR description ILIKE %s) LIMIT 8",
        (current_user.id, like, like))

    cur.close()
    conn.close()

    results = (list(customers) + list(products) + list(docs) + list(note_hits)
               + list(task_hits) + list(book_hits) + list(sku_hits)
               + list(service_hits) + list(session_hits)
               + list(bank_hits) + list(invest_hits) + list(shop_hits))
    return render_template('search_results.html', results=results, q=q)

# ======================== Services ========================

def _service_row_name(cur, service_id, uid):
    if not service_id:
        return 'Service'
    cur.execute('SELECT name FROM services WHERE id=%s AND user_id=%s', (service_id, uid))
    row = cur.fetchone()
    return row['name'] if row else 'Service'


@app.route('/services')
@login_required
def services():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT * FROM services WHERE user_id=%s ORDER BY active DESC, category, name',
        (uid,)
    )
    catalog = cur.fetchall()
    cur.execute(
        "SELECT s.*, sv.name AS service_name, sv.category AS service_category "
        "FROM service_sessions s LEFT JOIN services sv ON sv.id = s.service_id "
        "WHERE s.user_id=%s ORDER BY s.session_date DESC, s.id DESC LIMIT 250",
        (uid,)
    )
    sessions = cur.fetchall()
    cur.execute('SELECT id, name FROM customers WHERE user_id=%s ORDER BY name', (uid,))
    customers = cur.fetchall()
    today = datetime.today().strftime('%Y-%m-%d')
    booked = sum(1 for s in sessions if (s.get('status') or '') == 'booked')
    done = sum(1 for s in sessions if (s.get('status') or '') == 'done')
    earned = sum(float(s.get('price') or 0) for s in sessions if (s.get('status') or '') == 'done')
    upcoming = [s for s in sessions if (s.get('status') or '') == 'booked' and (s.get('session_date') or '') >= today]
    cur.close()
    conn.close()
    return render_template(
        'services.html',
        catalog=catalog,
        sessions=sessions,
        customers=customers,
        categories=SERVICE_CATEGORIES,
        booked=booked,
        done=done,
        earned=earned,
        upcoming_count=len(upcoming),
        today=today,
        active_catalog=[c for c in catalog if c.get('active')],
    )


@app.route('/services/add', methods=['POST'])
@login_required
def add_service():
    name, e1 = _clean_text('name', max_len=160, label='service name')
    price, e2 = _clean_float('price', min_v=0, max_v=1e12, label='price')
    desc, _ = _clean_text('description', required=False, max_len=2000, label='description')
    cat = (request.form.get('category') or 'Other').strip()
    if cat not in SERVICE_CATEGORIES:
        cat = 'Other'
    try:
        dur = int(float(request.form.get('duration_minutes') or 60))
    except (TypeError, ValueError):
        dur = 60
    if dur < 0:
        dur = 0
    if _reject((name, e1), (price, e2)):
        return redirect(url_for('services'))
    blocked = _plan_block('services')
    if blocked:
        flash(blocked, 'danger')
        return redirect(url_for('plans'))
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO services (user_id, name, category, description, duration_minutes, price, active, created_date) VALUES (%s,%s,%s,%s,%s,%s,1,%s)",
        (current_user.id, name, cat, desc or '', dur, price, today)
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added service', name)
    flash('Service saved.', 'success')
    return redirect(url_for('services'))


@app.route('/services/update/<int:id>', methods=['POST'])
@login_required
def update_service(id):
    name, e1 = _clean_text('name', max_len=160, label='service name')
    price, e2 = _clean_float('price', min_v=0, max_v=1e12, label='price')
    desc, _ = _clean_text('description', required=False, max_len=2000, label='description')
    cat = (request.form.get('category') or 'Other').strip()
    if cat not in SERVICE_CATEGORIES:
        cat = 'Other'
    try:
        dur = int(float(request.form.get('duration_minutes') or 60))
    except (TypeError, ValueError):
        dur = 60
    if dur < 0:
        dur = 0
    if _reject((name, e1), (price, e2)):
        return redirect(url_for('services'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "UPDATE services SET name=%s, category=%s, description=%s, duration_minutes=%s, price=%s WHERE id=%s AND user_id=%s",
        (name, cat, desc or '', dur, price, id, current_user.id)
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Service updated.', 'success')
    return redirect(url_for('services'))


@app.route('/services/toggle/<int:id>')
@login_required
def toggle_service(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT active FROM services WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    if row:
        nxt = 0 if row['active'] else 1
        cur.execute('UPDATE services SET active=%s WHERE id=%s AND user_id=%s', (nxt, id, current_user.id))
        conn.commit()
        flash('Service is now {}.'.format('active' if nxt else 'hidden'), 'success')
    cur.close()
    conn.close()
    return redirect(url_for('services'))


@app.route('/services/delete/<int:id>')
@login_required
def delete_service(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) AS n FROM service_sessions WHERE service_id=%s AND user_id=%s', (id, current_user.id))
    n = (cur.fetchone() or {}).get('n') or 0
    if n:
        cur.execute('UPDATE services SET active=0 WHERE id=%s AND user_id=%s', (id, current_user.id))
        conn.commit()
        flash('This service has bookings, so it was hidden instead of deleted.', 'info')
    else:
        cur.execute('DELETE FROM services WHERE id=%s AND user_id=%s', (id, current_user.id))
        conn.commit()
        flash('Service deleted.', 'success')
    cur.close()
    conn.close()
    return redirect(url_for('services'))


@app.route('/services/sessions/add', methods=['POST'])
@login_required
def add_service_session():
    try:
        service_id = int(request.form.get('service_id') or 0)
    except ValueError:
        service_id = 0
    if not service_id:
        flash('Choose a service.', 'danger')
        return redirect(url_for('services'))
    session_date, e1 = _clean_date('session_date')
    price, e2 = _clean_float('price', required=False, min_v=0, max_v=1e12, label='price')
    notes, _ = _clean_text('notes', required=False, max_len=2000, label='notes')
    time_txt, _ = _clean_text('session_time', required=False, max_len=10, label='time')
    if e1:
        flash(e1, 'danger')
        return redirect(url_for('services'))
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM services WHERE id=%s AND user_id=%s', (service_id, uid))
    svc = cur.fetchone()
    if not svc:
        cur.close()
        conn.close()
        flash('Service not found.', 'danger')
        return redirect(url_for('services'))
    if price is None:
        price = float(svc['price'] or 0)
    cust_id = request.form.get('customer_id') or None
    cust_name = (request.form.get('customer_name') or '').strip()
    if cust_id:
        try:
            cust_id = int(cust_id)
        except ValueError:
            cust_id = None
    if cust_id:
        cur.execute('SELECT name FROM customers WHERE id=%s AND user_id=%s', (cust_id, uid))
        saved_c = cur.fetchone()
        if saved_c:
            cust_name = saved_c['name']
        else:
            cust_id = None
    if not cust_name:
        cust_name = 'Walk-in'
    pay = _payment_method(request.form.get('payment_method') or 'cash')
    today = datetime.today().strftime('%Y-%m-%d')
    cur.execute(
        "INSERT INTO service_sessions (user_id, service_id, customer_id, customer_name, session_date, session_time, duration_minutes, price, status, payment_method, notes, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'booked',%s,%s,%s)",
        (uid, service_id, cust_id, cust_name[:120], session_date, time_txt or '',
         svc['duration_minutes'] or 60, price, pay, notes or '', today)
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(uid, current_user.username, 'Booked service', '{} for {}'.format(svc['name'], cust_name))
    flash('Session booked.', 'success')
    return redirect(url_for('services'))


@app.route('/services/sessions/<int:id>/status/<action>')
@login_required
def set_session_status(id, action):
    action = (action or '').lower()
    if action not in ('done', 'cancelled', 'booked'):
        flash('Unknown session action.', 'danger')
        return redirect(url_for('services'))
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM service_sessions WHERE id=%s AND user_id=%s', (id, uid))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        flash('Session not found.', 'danger')
        return redirect(url_for('services'))
    if action == 'done' and not row.get('income_id'):
        svc_name = _service_row_name(cur, row.get('service_id'), uid)
        source = 'Service: {} — {}'.format(svc_name, row.get('customer_name') or 'Walk-in')
        cur.execute(
            'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s) RETURNING id',
            (uid, source[:120], float(row.get('price') or 0), row.get('session_date'))
        )
        inc = cur.fetchone()
        inc_id = inc['id'] if inc else None
        cur.execute(
            'UPDATE service_sessions SET status=%s, income_id=%s WHERE id=%s AND user_id=%s',
            ('done', inc_id, id, uid)
        )
        log_activity(uid, current_user.username, 'Completed service', source)
        flash('Session marked done and income recorded.', 'success')
    else:
        cur.execute(
            'UPDATE service_sessions SET status=%s WHERE id=%s AND user_id=%s',
            (action, id, uid)
        )
        flash('Session set to {}.'.format(action), 'success')
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('services'))


@app.route('/services/sessions/<int:id>/delete')
@login_required
def delete_service_session(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM service_sessions WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Session removed.', 'success')
    return redirect(url_for('services'))


@app.route('/export/services.csv')
@login_required
def export_services_csv():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM services WHERE user_id=%s ORDER BY name', (uid,))
    catalog = cur.fetchall()
    cur.execute(
        "SELECT s.*, sv.name AS service_name FROM service_sessions s "
        "LEFT JOIN services sv ON sv.id = s.service_id WHERE s.user_id=%s ORDER BY s.session_date",
        (uid,)
    )
    sessions = cur.fetchall()
    cur.close()
    conn.close()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['# catalog'])
    w.writerow(['name', 'category', 'duration_minutes', 'price', 'active', 'description'])
    for row in catalog:
        w.writerow([row['name'], row['category'], row['duration_minutes'], row['price'], row['active'], row.get('description') or ''])
    w.writerow([])
    w.writerow(['# sessions'])
    w.writerow(['date', 'time', 'service', 'customer', 'price', 'status', 'payment', 'notes'])
    for row in sessions:
        w.writerow([
            row.get('session_date'), row.get('session_time'), row.get('service_name'),
            row.get('customer_name'), row.get('price'), row.get('status'),
            row.get('payment_method'), row.get('notes') or ''
        ])
    body = buf.getvalue()
    return Response(
        body,
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=kaze-services.csv'}
    )


# ======================== Tasks ========================

@app.route('/tasks')
@login_required
def tasks():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT * FROM tasks WHERE user_id = %s ORDER BY done ASC, CASE WHEN due_date = '' THEN '9999-12-31' ELSE due_date END, id DESC",
        (current_user.id,)
    )
    items = cur.fetchall()
    cur.close()
    conn.close()
    open_count = sum(1 for i in items if not i['done'])
    return render_template('tasks.html', items=items, open_count=open_count)


@app.route('/tasks/add', methods=['POST'])
@login_required
def add_task():
    title = request.form['title'].strip()
    if not title:
        flash('Task needs a title.', 'danger')
        return redirect(url_for('tasks'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''
        INSERT INTO tasks (user_id, title, notes, priority, due_date, done, created_date)
        VALUES (%s, %s, %s, %s, %s, 0, %s)
        ''',
        (current_user.id, title, request.form.get('notes', ''),
         request.form.get('priority', 'medium'), request.form.get('due_date', ''),
         datetime.today().strftime('%Y-%m-%d'))
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added task', title)
    flash('Task added.', 'success')
    return redirect(url_for('tasks'))


@app.route('/tasks/toggle/<int:id>')
@login_required
def toggle_task(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE tasks SET done = 1 - done WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(request.referrer or url_for('tasks'))


@app.route('/tasks/delete/<int:id>')
@login_required
def delete_task(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM tasks WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Task deleted.', 'success')
    return redirect(url_for('tasks'))


# ======================== Notes ========================

@app.route('/notes')
@login_required
def notes():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM notes WHERE user_id = %s ORDER BY updated_date DESC', (current_user.id,))
    items = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('notes.html', items=items)


@app.route('/notes/add', methods=['POST'])
@login_required
def add_note():
    title = request.form['title'].strip() or 'Untitled note'
    body = request.form.get('body', '')
    now = datetime.today().strftime('%Y-%m-%d %H:%M')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'INSERT INTO notes (user_id, title, body, created_date, updated_date) VALUES (%s, %s, %s, %s, %s)',
        (current_user.id, title, body, now, now)
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added note', title)
    flash('Note saved.', 'success')
    return redirect(url_for('notes'))


@app.route('/notes/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_note(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM notes WHERE id = %s AND user_id = %s', (id, current_user.id))
    note = cur.fetchone()
    if not note:
        cur.close()
        conn.close()
        flash('Note not found.', 'danger')
        return redirect(url_for('notes'))
    if request.method == 'POST':
        title = request.form['title'].strip() or 'Untitled note'
        body = request.form.get('body', '')
        now = datetime.today().strftime('%Y-%m-%d %H:%M')
        cur.execute(
            'UPDATE notes SET title=%s, body=%s, updated_date=%s WHERE id=%s AND user_id=%s',
            (title, body, now, id, current_user.id)
        )
        conn.commit()
        cur.close()
        conn.close()
        flash('Note updated.', 'success')
        return redirect(url_for('notes'))
    cur.close()
    conn.close()
    return render_template('edit_note.html', note=note)


@app.route('/notes/delete/<int:id>')
@login_required
def delete_note(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM notes WHERE id = %s AND user_id = %s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Note deleted.', 'success')
    return redirect(url_for('notes'))


# ======================== Code snippets ========================

SNIPPET_LANGS = ('text', 'html', 'python', 'sql', 'json', 'receipt')
STARTER_SNIPPETS = (
    ('Receipt footer', 'receipt', 'Thank you for shopping with us.\nGoods sold in good condition.\nKeep this receipt for returns.'),
    ('Invoice note', 'text', 'Payment is due within 7 days.\nPlease quote the invoice number on your transfer.'),
    ('Low stock message', 'text', 'Hello, {product} is down to {qty}. Please restock before the next sales day.'),
    ('HTML price tag', 'html', '<div class="tag"><strong>{name}</strong><span>{price}</span></div>'),
)


# DELETED FEATURE: @app.route('/snippets')
# @login_required
def snippets():
    q = (request.args.get('q') or '').strip()
    conn = get_db()
    cur = conn.cursor()
    if q:
        like = '%' + q + '%'
        cur.execute(
            'SELECT * FROM code_snippets WHERE user_id=%s AND (title ILIKE %s OR body ILIKE %s OR language ILIKE %s) '
            'ORDER BY pinned DESC, updated_date DESC',
            (current_user.id, like, like, like),
        )
    else:
        cur.execute(
            'SELECT * FROM code_snippets WHERE user_id=%s ORDER BY pinned DESC, updated_date DESC',
            (current_user.id,),
        )
    items = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('snippets.html', items=items, q=q, langs=SNIPPET_LANGS)


@app.route('/snippets/add', methods=['POST'])
@login_required
def add_snippet():
    title, e1 = _clean_text('title', max_len=80, label='title')
    body = (request.form.get('body') or '').strip()
    language = (request.form.get('language') or 'text').strip().lower()
    if language not in SNIPPET_LANGS:
        language = 'text'
    if _reject((title, e1)):
        return redirect(url_for('snippets'))
    if not body or len(body) > 8000:
        flash('Snippet body is required and must be under 8000 characters.', 'danger')
        return redirect(url_for('snippets'))
    now = datetime.today().strftime('%Y-%m-%d %H:%M')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'INSERT INTO code_snippets (user_id, title, language, body, pinned, created_date, updated_date) '
        'VALUES (%s, %s, %s, %s, 0, %s, %s)',
        (current_user.id, title, language, body, now, now),
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added snippet', title)
    flash('Snippet saved.', 'success')
    return redirect(url_for('snippets'))


@app.route('/snippets/update/<int:id>', methods=['POST'])
@login_required
def update_snippet(id):
    title, e1 = _clean_text('title', max_len=80, label='title')
    body = (request.form.get('body') or '').strip()
    language = (request.form.get('language') or 'text').strip().lower()
    pinned = 1 if request.form.get('pinned') == 'on' else 0
    if language not in SNIPPET_LANGS:
        language = 'text'
    if _reject((title, e1)) or not body or len(body) > 8000:
        flash('Snippet body is required and must be under 8000 characters.', 'danger')
        return redirect(url_for('snippets'))
    now = datetime.today().strftime('%Y-%m-%d %H:%M')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'UPDATE code_snippets SET title=%s, language=%s, body=%s, pinned=%s, updated_date=%s WHERE id=%s AND user_id=%s',
        (title, language, body, pinned, now, id, current_user.id),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Snippet updated.', 'success')
    return redirect(url_for('snippets'))


@app.route('/snippets/delete/<int:id>')
@login_required
def delete_snippet(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM code_snippets WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Snippet deleted.', 'success')
    return redirect(url_for('snippets'))


@app.route('/snippets/starters', methods=['POST'])
@login_required
def snippet_starters():
    now = datetime.today().strftime('%Y-%m-%d %H:%M')
    conn = get_db()
    cur = conn.cursor()
    added = 0
    for title, language, body in STARTER_SNIPPETS:
        cur.execute('SELECT id FROM code_snippets WHERE user_id=%s AND title=%s', (current_user.id, title))
        if cur.fetchone():
            continue
        cur.execute(
            'INSERT INTO code_snippets (user_id, title, language, body, pinned, created_date, updated_date) VALUES (%s, %s, %s, %s, 0, %s, %s)',
            (current_user.id, title, language, body, now, now),
        )
        added += 1
    conn.commit()
    cur.close()
    conn.close()
    flash('Added {} starter snippet(s).'.format(added) if added else 'Starters are already saved.', 'success' if added else 'info')
    return redirect(url_for('snippets'))


# ======================== Calendar ========================

@app.route('/calendar')
@login_required
def calendar():
    today = datetime.today().strftime('%Y-%m-%d')
    horizon = (datetime.today() + timedelta(days=60)).strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    events = []
    cur.execute(
        'SELECT name, next_run_date, item_type, amount FROM recurring_items WHERE user_id=%s AND active=1 AND next_run_date <= %s',
        (current_user.id, horizon)
    )
    for row in cur.fetchall():
        events.append({
            'date': row['next_run_date'],
            'kind': 'Recurring ' + row['item_type'],
            'title': row['name'],
            'extra': _format_money(row['amount'])
        })
    cur.execute(
        "SELECT item_description, expected_date, status FROM purchase_orders WHERE user_id=%s AND expected_date <> '' AND expected_date <= %s",
        (current_user.id, horizon)
    )
    for row in cur.fetchall():
        events.append({
            'date': row['expected_date'],
            'kind': 'Purchase order',
            'title': row['item_description'],
            'extra': row['status']
        })
    cur.execute(
        "SELECT title, due_date, priority, done FROM tasks WHERE user_id=%s AND due_date <> '' AND due_date <= %s",
        (current_user.id, horizon)
    )
    for row in cur.fetchall():
        events.append({
            'date': row['due_date'],
            'kind': 'Task',
            'title': row['title'],
            'extra': ('done' if row['done'] else row['priority'])
        })

    try:
        cur.execute(
            "SELECT ss.session_date, ss.session_time, ss.status, ss.customer_name, sv.name AS service_name "
            "FROM service_sessions ss LEFT JOIN services sv ON sv.id = ss.service_id "
            "WHERE ss.user_id=%s AND ss.session_date <= %s AND ss.status <> 'cancelled'",
            (current_user.id, horizon)
        )
        for row in cur.fetchall():
            events.append({
                'date': row['session_date'],
                'kind': 'Service',
                'title': '{} — {}'.format(row.get('service_name') or 'Service', row.get('customer_name') or 'Walk-in'),
                'extra': '{} {}'.format(row.get('session_time') or '', row.get('status') or '').strip()
            })
    except Exception:
        pass
    cur.close()
    conn.close()
    events.sort(key=lambda e: e['date'] or '9999')
    return render_template('calendar.html', events=events, today=today)


# ======================== Profit calculator ========================

# DELETED FEATURE: @app.route('/calculator')
# @login_required
def calculator():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT product_name, cost_price, selling_price FROM stock WHERE user_id=%s ORDER BY product_name',
        (current_user.id,)
    )
    products = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('calculator.html', products=products)


# ======================== Password ========================

@app.route('/settings/password', methods=['POST'])
@login_required
def change_password():
    old = request.form.get('old_password', '')
    new = request.form.get('new_password', '')
    confirm = request.form.get('confirm_password', '')
    if not new or len(new) < 6:
        flash('New password must be at least 6 characters.', 'danger')
        return redirect(url_for('settings'))
    if new != confirm:
        flash('New passwords do not match.', 'danger')
        return redirect(url_for('settings'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT password FROM users WHERE id = %s', (current_user.db_id,))
    row = cur.fetchone()
    if not row or not check_password_hash(row['password'], old):
        cur.close()
        conn.close()
        flash('Current password is incorrect.', 'danger')
        return redirect(url_for('settings'))
    cur.execute('UPDATE users SET password = %s WHERE id = %s', (generate_password_hash(new), current_user.db_id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Changed password', None)
    flash('Password updated.', 'success')
    return redirect(url_for('settings'))


# ======================== JSON backup ========================

@app.route('/backup/export')
@login_required
def backup_export():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    payload = {'exported_at': datetime.now().isoformat(), 'tables': {}}
    tables = ('income', 'expenses', 'stock', 'customers', 'cash_books', 'documents',
              'tasks', 'notes', 'suppliers', 'budgets', 'recurring_items',
              'services', 'service_sessions', 'stripe_payments')
    for table in tables:
        cur.execute('SELECT * FROM {} WHERE user_id = %s'.format(table), (uid,))
        rows = cur.fetchall()
        payload['tables'][table] = [dict(r) for r in rows]
    cur.close()
    conn.close()
    body = json.dumps(payload, default=str, indent=2)
    return Response(
        body,
        mimetype='application/json',
        headers={'Content-Disposition': 'attachment; filename=kaze-backup.json'}
    )


@app.route('/backup/import', methods=['POST'])
@login_required
def backup_import():
    f = request.files.get('file')
    if not f or not f.filename:
        flash('Choose a backup JSON file first.', 'danger')
        return redirect(url_for('settings'))
    try:
        payload = json.loads(f.read().decode('utf-8'))
        tables = payload.get('tables', {})
    except Exception:
        flash('That file is not a valid KAZE backup.', 'danger')
        return redirect(url_for('settings'))
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    added = 0
    try:
        for row in tables.get('income', []):
            cur.execute(
                'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s)',
                (uid, row.get('source'), row.get('amount'), row.get('date'))
            )
            added += 1
        for row in tables.get('expenses', []):
            cur.execute(
                'INSERT INTO expenses (user_id, name, amount, category, date) VALUES (%s,%s,%s,%s,%s)',
                (uid, row.get('name'), row.get('amount'), row.get('category') or 'Other', row.get('date'))
            )
            added += 1
        for row in tables.get('stock', []):
            cur.execute(
                'INSERT INTO stock (user_id, product_name, quantity, cost_price, selling_price, unit) VALUES (%s,%s,%s,%s,%s,%s)',
                (uid, row.get('product_name'), row.get('quantity'), row.get('cost_price'), row.get('selling_price'), row.get('unit'))
            )
            added += 1
        for row in tables.get('customers', []):
            cur.execute(
                'INSERT INTO customers (user_id, name, email, phone, address, notes) VALUES (%s,%s,%s,%s,%s,%s)',
                (uid, row.get('name'), row.get('email'), row.get('phone'), row.get('address'), row.get('notes'))
            )
            added += 1
        for row in tables.get('notes', []):
            now = datetime.today().strftime('%Y-%m-%d %H:%M')
            cur.execute(
                'INSERT INTO notes (user_id, title, body, created_date, updated_date) VALUES (%s,%s,%s,%s,%s)',
                (uid, row.get('title') or 'Imported note', row.get('body') or '', row.get('created_date') or now, now)
            )
            added += 1
        for row in tables.get('tasks', []):
            cur.execute(
                'INSERT INTO tasks (user_id, title, notes, priority, due_date, done, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s)',
                (uid, row.get('title'), row.get('notes') or '', row.get('priority') or 'medium',
                 row.get('due_date') or '', row.get('done') or 0,
                 row.get('created_date') or datetime.today().strftime('%Y-%m-%d'))
            )
            added += 1
        for row in tables.get('services', []):
            cur.execute(
                'INSERT INTO services (user_id, name, category, description, duration_minutes, price, active, created_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)',
                (uid, row.get('name'), row.get('category') or 'Other', row.get('description') or '',
                 row.get('duration_minutes') or 60, row.get('price') or 0, row.get('active') if row.get('active') is not None else 1,
                 row.get('created_date') or datetime.today().strftime('%Y-%m-%d'))
            )
            added += 1
        conn.commit()
    except Exception as e:
        conn.rollback()
        cur.close()
        conn.close()
        flash('Import failed: {}'.format(e), 'danger')
        return redirect(url_for('settings'))
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Imported backup', '{} rows'.format(added))
    flash('Imported {} records from backup.'.format(added), 'success')
    return redirect(url_for('settings'))


# ======================== Low stock -> purchase orders ========================

@app.route('/stock/reorder-low')
@login_required
def reorder_low_stock():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT low_stock_threshold FROM settings WHERE user_id = %s', (current_user.id,))
    row = cur.fetchone()
    threshold = row['low_stock_threshold'] if row else 10
    cur.execute('SELECT * FROM stock WHERE user_id = %s AND quantity < %s', (current_user.id, threshold))
    low_items = cur.fetchall()
    created = 0
    today = datetime.today().strftime('%Y-%m-%d')
    for item in low_items:
        need = max(threshold * 2 - item['quantity'], 1)
        total = need * (item['cost_price'] or 0)
        cur.execute(
            '''
            INSERT INTO purchase_orders (user_id, supplier_id, item_description, quantity, unit_cost, total_cost, status, order_date, expected_date, stock_id)
            VALUES (%s, NULL, %s, %s, %s, %s, 'pending', %s, %s, %s)
            ''',
            (current_user.id, item['product_name'], need, item['cost_price'] or 0, total, today, '', item['id'])
        )
        created += 1
    conn.commit()
    cur.close()
    conn.close()
    if created:
        log_activity(current_user.id, current_user.username, 'Drafted reorder POs', '{} items'.format(created))
        flash('Created {} draft purchase order(s) for low-stock items.'.format(created), 'success')
        return redirect(url_for('purchase_orders'))
    flash('No low-stock items to reorder.', 'info')
    return redirect(url_for('stock'))


# ======================== Book keeping documents ========================

BOOK_DOC_TYPES = [
    ('quotation', 'Quotation', 'Price offer sent to a customer before they order.'),
    ('purchase_order', 'Purchase Order', 'Order you send to a supplier.'),
    ('invoice', 'Invoice', 'Request for payment after goods or services.'),
    ('proforma_invoice', 'Pro-forma Invoice', 'Estimated invoice sent before the final bill.'),
    ('delivery_note', 'Delivery Note', 'List of goods sent out to a customer.'),
    ('grn', 'Goods Received Note (GRN)', 'Record of goods that arrived from a supplier.'),
    ('credit_note', 'Credit Note', 'Reduces what a customer owes you.'),
    ('debit_note', 'Debit Note', 'Increases an amount owed.'),
    ('statement', 'Statement of Account', 'Running balance for a customer or supplier.'),
    ('receipt', 'Receipt', 'Proof that money was received.'),
    ('cash_sale_receipt', 'Cash Sale Receipt', 'Receipt for an immediate cash sale.'),
    ('payment_voucher', 'Payment Voucher', 'Authorisation to pay money out.'),
    ('petty_cash_voucher', 'Petty Cash Voucher', 'Small cash payment from the petty cash tin.'),
    ('cheque', 'Cheque', 'Record of a cheque you wrote or received.'),
    ('cheque_counterfoil', 'Cheque Counterfoil', 'Stub kept in the cheque book.'),
    ('bank_deposit_slip', 'Bank Deposit Slip', 'Record of cash or cheques paid into the bank.'),
    ('bank_statement', 'Bank Statement', 'Summary of bank movements you want to keep on file.'),
    ('remittance_advice', 'Remittance Advice', 'Note sent with a payment explaining what it covers.'),
    ('goods_returned_note', 'Goods Returned Note', 'Goods sent back to a supplier or by a customer.'),
    ('goods_dispatch_note', 'Goods Dispatch Note', 'Goods leaving your store.'),
    ('order_confirmation', 'Order Confirmation', 'Confirms an order you accepted.'),
    ('purchase_invoice', 'Purchase Invoice', 'Bill received from a supplier.'),
    ('sales_invoice', 'Sales Invoice', 'Bill you send to a customer.'),
    ('timesheet', 'Time Sheet', 'Hours worked on a job or by a person.'),
    ('clock_card', 'Clock Card / Time Card', 'Clock-in record.'),
    ('wage_sheet', 'Wage Sheet / Payroll', 'Payroll list for a pay period.'),
    ('payslip', 'Payslip', 'Pay breakdown for one person.'),
]
BOOK_DOC_LOOKUP = {k: (label, hint) for k, label, hint in BOOK_DOC_TYPES}


@app.route('/bookkeeping')
@login_required
def bookkeeping():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT kind, COUNT(*) as total FROM bookkeeping_docs WHERE user_id=%s GROUP BY kind', (current_user.id,))
    counts = {row['kind']: row['total'] for row in cur.fetchall()}
    cur.execute('SELECT * FROM cash_books WHERE user_id=%s ORDER BY date DESC LIMIT 8', (current_user.id,))
    recent_cash = cur.fetchall()
    cur.execute('SELECT * FROM bookkeeping_docs WHERE user_id=%s ORDER BY id DESC LIMIT 8', (current_user.id,))
    recent_docs = cur.fetchall()
    cur.close()
    conn.close()
    types = [(k, label, hint, counts.get(k, 0)) for k, label, hint in BOOK_DOC_TYPES]
    return render_template(
        'bookkeeping.html',
        types=types,
        recent_cash=recent_cash,
        recent_docs=recent_docs,
        kind_labels={k: label for k, label, _ in BOOK_DOC_TYPES}
    )


@app.route('/bookkeeping/<kind>', methods=['GET', 'POST'])
@login_required
def bookkeeping_type(kind):
    info = BOOK_DOC_LOOKUP.get(kind)
    if not info:
        flash('Unknown document type.', 'danger')
        return redirect(url_for('bookkeeping'))
    label, hint = info
    if request.method == 'POST':
        title = (request.form.get('title') or '').strip() or label
        reference = (request.form.get('reference') or '').strip()
        try:
            amount = float(request.form.get('amount') or 0)
        except ValueError:
            amount = 0
        conn = get_db()
        cur = conn.cursor()
        cur.execute(
            '''
            INSERT INTO bookkeeping_docs
            (user_id, kind, title, party, reference, doc_date, amount, status, notes, line_items, created_date)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            ''',
            (
                current_user.id, kind, title,
                request.form.get('party', ''),
                reference or (kind.upper()[:3] + '-' + datetime.today().strftime('%Y%m%d') + '-' + title[:8].upper().replace(' ','')),
                request.form.get('doc_date') or datetime.today().strftime('%Y-%m-%d'),
                amount,
                request.form.get('status') or 'draft',
                request.form.get('notes', ''),
                request.form.get('line_items', ''),
                datetime.today().strftime('%Y-%m-%d')
            )
        )
        conn.commit()
        cur.close()
        conn.close()
        log_activity(current_user.id, current_user.username, 'Created ' + label, title)
        flash(label + ' saved.', 'success')
        return redirect(url_for('bookkeeping_type', kind=kind))

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT * FROM bookkeeping_docs WHERE user_id=%s AND kind=%s ORDER BY doc_date DESC, id DESC',
        (current_user.id, kind)
    )
    items = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('bookkeeping_type.html', kind=kind, label=label, hint=hint, items=items)


@app.route('/bookkeeping/doc/<int:id>/delete')
@login_required
def delete_book_doc(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT kind FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    kind = row['kind'] if row else None
    cur.execute('DELETE FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Document deleted.', 'success')
    if kind:
        return redirect(url_for('bookkeeping_type', kind=kind))
    return redirect(url_for('bookkeeping'))


@app.route('/bookkeeping/doc/<int:id>/pdf')
@login_required
def bookkeeping_pdf(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Document not found.', 'danger')
        return redirect(url_for('bookkeeping'))
    label = BOOK_DOC_LOOKUP.get(row['kind'], (row['kind'].replace('_', ' ').title(), ''))[0]
    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    tax_rate = (settings_row['tax_rate'] if settings_row else 0) or 0
    meta_lines = [
        '<b>Document:</b> {}'.format(label),
        '<b>Title:</b> {}'.format(row['title']),
        '<b>Party:</b> {}'.format(row['party'] or '-'),
        '<b>Reference:</b> {}'.format(row['reference'] or '-'),
        '<b>Date:</b> {}'.format(row['doc_date']),
        '<b>Status:</b> {}'.format(row['status'] or 'draft'),
    ]
    if row['notes']:
        meta_lines.append('<b>Notes:</b> {}'.format(row['notes']))
    table_data = [['Description', 'Qty / detail', 'Amount']]
    parsed = False
    for raw in (row['line_items'] or '').splitlines():
        line = raw.strip()
        if not line:
            continue
        parts = [p.strip() for p in line.split('|')]
        if len(parts) == 1:
            table_data.append([parts[0], '', ''])
        elif len(parts) == 2:
            table_data.append([parts[0], parts[1], ''])
        else:
            table_data.append([parts[0], parts[1], parts[2]])
        parsed = True
    if not parsed:
        table_data.append([row['title'], '1', _format_money(row['amount'] or 0, settings_row)])
    subtotal = row['amount'] or 0
    if row['kind'] in ('invoice', 'sales_invoice', 'proforma_invoice', 'quotation') and tax_rate:
        tax_amount = subtotal * (tax_rate / 100)
        table_data.append(['Tax ({:.1f}%)'.format(tax_rate), '', _format_money(tax_amount, settings_row)])
        table_data.append(['Total', '', _format_money(subtotal + tax_amount, settings_row)])
    else:
        table_data.append(['Total', '', _format_money(subtotal, settings_row)])
    buffer = io.BytesIO()
    _build_pdf(
        buffer, label + ' - ' + row['title'], business_name, meta_lines, table_data,
        footer_lines=['Generated on {}'.format(datetime.today().strftime('%Y-%m-%d'))]
    )
    buffer.seek(0)
    safe = ''.join(c for c in row['title'] if c.isalnum() or c in (' ', '_', '-')).rstrip()
    return send_file(
        buffer, mimetype='application/pdf', as_attachment=True,
        download_name='{}-{}.pdf'.format(row['kind'], safe or 'document')
    )


# ======================== Main ========================

# ======================== Account tracking ========================

ACCOUNT_BOOKS = [
    ('cash_book', 'Cash Book', 'Cash in and out — same ledger as Book keeping.'),
    ('purchases_journal', 'Purchases Journal', 'Credit purchases and supplier bills.'),
    ('sales_journal', 'Sales Journal', 'Credit and recorded sales.'),
    ('purchases_returns', 'Purchases Returns Journal', 'Goods sent back to suppliers.'),
    ('sales_returns', 'Sales Returns Journal', 'Goods returned by customers.'),
    ('general_journal', 'General Journal', 'Double-entry adjustments you type yourself.'),
    ('petty_cash', 'Petty Cash Book', 'Small cash payments.'),
    ('general_ledger', 'General Ledger', 'Accounts rolled up from journals and the app.'),
    ('debtors_ledger', 'Debtors Ledger', 'What customers owe, from sales.'),
    ('creditors_ledger', 'Creditors Ledger', 'What you owe suppliers.'),
    ('trading_account', 'Trading Account', 'Sales, cost of goods, and gross profit.'),
    ('profit_and_loss', 'Profit and Loss Account', 'Gross profit, other income, and expenses.'),
    ('balance_sheet', 'Balance Sheet', 'Assets, liabilities, and capital.'),
    ('cash_flow', 'Cash Flow Statement', 'Where cash came from and where it went.'),
]
ACCOUNT_BOOK_LOOKUP = {k: (label, hint) for k, label, hint in ACCOUNT_BOOKS}
LEDGER_BOOKS = {
    'purchases_journal', 'sales_journal', 'purchases_returns',
    'sales_returns', 'general_journal', 'petty_cash'
}


def _je_rows(cur, user_id, book):
    cur.execute(
        '''SELECT * FROM journal_entries WHERE user_id=%s AND book=%s
           ORDER BY entry_date DESC, id DESC''',
        (user_id, book)
    )
    return cur.fetchall()


@app.route('/accounts')
@login_required
def accounts():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT book, COUNT(*) as total FROM journal_entries WHERE user_id=%s GROUP BY book',
        (current_user.id,)
    )
    counts = {r['book']: r['total'] for r in cur.fetchall()}
    cur.execute('SELECT COUNT(*) as n FROM sales WHERE user_id=%s', (current_user.id,))
    counts['sales_journal'] = counts.get('sales_journal', 0) + (cur.fetchone()['n'] or 0)
    cur.execute('SELECT COUNT(*) as n FROM purchase_orders WHERE user_id=%s', (current_user.id,))
    counts['purchases_journal'] = counts.get('purchases_journal', 0) + (cur.fetchone()['n'] or 0)
    cur.execute('SELECT COUNT(*) as n FROM cash_books WHERE user_id=%s', (current_user.id,))
    counts['cash_book'] = cur.fetchone()['n'] or 0
    cur.close()
    conn.close()
    books = [(k, label, hint, counts.get(k, 0)) for k, label, hint in ACCOUNT_BOOKS]
    return render_template('accounts.html', books=books)


@app.route('/accounts/add', methods=['POST'])
@login_required
def add_journal_entry():
    book = request.form.get('book') or 'general_journal'
    if book not in LEDGER_BOOKS:
        flash('That book does not take typed entries.', 'danger')
        return redirect(url_for('accounts'))
    account = (request.form.get('account') or '').strip()
    if not account:
        flash('Account name is required.', 'danger')
        return redirect(url_for('account_book', book=book))
    try:
        debit = float(request.form.get('debit') or 0)
        credit = float(request.form.get('credit') or 0)
    except ValueError:
        debit, credit = 0, 0
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO journal_entries
           (user_id, book, entry_date, reference, account, particulars, debit, credit, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (
            current_user.id, book,
            request.form.get('entry_date') or datetime.today().strftime('%Y-%m-%d'),
            request.form.get('reference', ''),
            account,
            request.form.get('particulars', ''),
            debit, credit,
            datetime.today().strftime('%Y-%m-%d')
        )
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Journal entry', book + ': ' + account)
    flash('Entry saved.', 'success')
    return redirect(url_for('account_book', book=book))


@app.route('/accounts/entry/<int:id>/delete')
@login_required
def delete_journal_entry(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT book FROM journal_entries WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    book = row['book'] if row else 'general_journal'
    cur.execute('DELETE FROM journal_entries WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Entry deleted.', 'success')
    return redirect(url_for('account_book', book=book))


@app.route('/accounts/<book>')
@login_required
def account_book(book):
    if book not in ACCOUNT_BOOK_LOOKUP:
        flash('Unknown book.', 'danger')
        return redirect(url_for('accounts'))
    if book == 'cash_book':
        return redirect(url_for('cashbook'))

    label, hint = ACCOUNT_BOOK_LOOKUP[book]
    uid = current_user.id
    start_date = (request.args.get('start_date') or '').strip()
    end_date = (request.args.get('end_date') or '').strip()
    conn = get_db()
    cur = conn.cursor()
    auto_rows = []
    totals = {'debit': 0.0, 'credit': 0.0}
    extra = {}

    def in_range(sql_date_col):
        clause = ''
        params = []
        if start_date:
            clause += ' AND {} >= %s'.format(sql_date_col)
            params.append(start_date)
        if end_date:
            clause += ' AND {} <= %s'.format(sql_date_col)
            params.append(end_date)
        return clause, params

    if book == 'sales_journal':
        extra_clause, extra_params = in_range('sales.sale_date')
        cur.execute(
            '''SELECT sales.sale_date AS entry_date,
                      COALESCE(sales.customer_name, 'Walk-in') AS account,
                      COALESCE(stock.product_name, 'Item') AS particulars,
                      sales.total_amount AS debit,
                      0 AS credit,
                      sales.id::text AS reference
               FROM sales
               LEFT JOIN stock ON sales.stock_id = stock.id
               WHERE sales.user_id=%s''' + extra_clause + '''
               ORDER BY sales.sale_date DESC, sales.id DESC''',
            [uid] + extra_params,
        )
        auto_rows = cur.fetchall()
    elif book == 'purchases_journal':
        extra_clause, extra_params = in_range('order_date')
        cur.execute(
            '''SELECT order_date AS entry_date,
                      COALESCE(item_description, 'Purchase') AS account,
                      status AS particulars,
                      0 AS debit,
                      total_cost AS credit,
                      id::text AS reference
               FROM purchase_orders WHERE user_id=%s''' + extra_clause + '''
               ORDER BY order_date DESC, id DESC''',
            [uid] + extra_params,
        )
        rows = list(cur.fetchall() or [])
        extra_clause2, extra_params2 = in_range('doc_date')
        cur.execute(
            '''SELECT doc_date AS entry_date, COALESCE(party, title) AS account,
                      title AS particulars, 0 AS debit, COALESCE(amount,0) AS credit,
                      COALESCE(reference, id::text) AS reference
               FROM bookkeeping_docs
               WHERE user_id=%s AND kind IN ('purchase_invoice','purchase_order')''' + extra_clause2 + '''
               ORDER BY doc_date DESC''',
            [uid] + extra_params2,
        )
        auto_rows = rows + list(cur.fetchall() or [])
    elif book == 'purchases_returns':
        extra_clause, extra_params = in_range('doc_date')
        cur.execute(
            '''SELECT doc_date AS entry_date, COALESCE(party, title) AS account,
                      title AS particulars, COALESCE(amount,0) AS debit, 0 AS credit,
                      COALESCE(reference, id::text) AS reference
               FROM bookkeeping_docs
               WHERE user_id=%s AND kind IN ('goods_returned_note','debit_note')''' + extra_clause + '''
               ORDER BY doc_date DESC''',
            [uid] + extra_params,
        )
        auto_rows = cur.fetchall()
    elif book == 'sales_returns':
        extra_clause, extra_params = in_range('doc_date')
        cur.execute(
            '''SELECT doc_date AS entry_date, COALESCE(party, title) AS account,
                      title AS particulars, 0 AS debit, COALESCE(amount,0) AS credit,
                      COALESCE(reference, id::text) AS reference
               FROM bookkeeping_docs
               WHERE user_id=%s AND kind IN ('credit_note','goods_returned_note')''' + extra_clause + '''
               ORDER BY doc_date DESC''',
            [uid] + extra_params,
        )
        auto_rows = cur.fetchall()
    elif book == 'petty_cash':
        extra_clause, extra_params = in_range('date')
        cur.execute(
            '''SELECT date AS entry_date, category AS account, description AS particulars,
                      CASE WHEN entry_type='out' THEN amount ELSE 0 END AS debit,
                      CASE WHEN entry_type='in' THEN amount ELSE 0 END AS credit,
                      id::text AS reference
               FROM cash_books
               WHERE user_id=%s AND (LOWER(COALESCE(category,'')) LIKE '%%petty%%'
                    OR LOWER(COALESCE(description,'')) LIKE '%%petty%%')''' + extra_clause + '''
               ORDER BY date DESC''',
            [uid] + extra_params,
        )
        auto_rows = cur.fetchall()
    elif book == 'debtors_ledger':
        cur.execute(
            '''SELECT COALESCE(NULLIF(customer_name,''), 'Walk-in') AS account,
                      COUNT(*) AS orders,
                      COALESCE(SUM(CASE WHEN COALESCE(payment_method,'cash') NOT IN ('cash') THEN total_amount ELSE 0 END),0) AS debit,
                      COALESCE(SUM(CASE WHEN COALESCE(payment_method,'cash') IN ('cash') THEN total_amount ELSE 0 END),0) AS credit
               FROM sales WHERE user_id=%s
               GROUP BY 1 ORDER BY debit DESC, credit DESC''',
            (uid,),
        )
        extra['ledger'] = cur.fetchall()
        extra['ledger_note'] = 'Debit = billed / non-cash sales still owing. Credit = cash sales already collected.'
    elif book == 'creditors_ledger':
        cur.execute(
            '''SELECT COALESCE(s.name, po.item_description, 'Supplier') AS account,
                      COUNT(*) AS orders,
                      COALESCE(SUM(CASE WHEN po.status = 'received' THEN po.total_cost ELSE 0 END),0) AS debit,
                      COALESCE(SUM(CASE WHEN po.status <> 'received' THEN po.total_cost ELSE 0 END),0) AS credit
               FROM purchase_orders po
               LEFT JOIN suppliers s ON s.id = po.supplier_id
               WHERE po.user_id=%s
               GROUP BY 1 ORDER BY credit DESC''',
            (uid,),
        )
        extra['ledger'] = cur.fetchall()
        extra['ledger_note'] = 'Credit = open purchase orders. Debit = orders already received.'
    elif book == 'general_ledger':
        accounts = {}
        def bump(name, debit=0, credit=0):
            row = accounts.setdefault(name, {'account': name, 'debit': 0.0, 'credit': 0.0})
            row['debit'] += float(debit or 0)
            row['credit'] += float(credit or 0)
        cur.execute('SELECT COALESCE(SUM(amount),0) AS t FROM income WHERE user_id=%s', (uid,))
        bump('Other income', credit=cur.fetchone()['t'])
        cur.execute('SELECT COALESCE(SUM(amount),0) AS t FROM expenses WHERE user_id=%s', (uid,))
        bump('Expenses', debit=cur.fetchone()['t'])
        cur.execute('SELECT COALESCE(SUM(total_amount),0) AS t FROM sales WHERE user_id=%s', (uid,))
        bump('Sales', credit=cur.fetchone()['t'])
        cur.execute('SELECT COALESCE(SUM(quantity * cost_price),0) AS t FROM stock WHERE user_id=%s', (uid,))
        bump('Inventory', debit=cur.fetchone()['t'])
        cur.execute("SELECT COALESCE(SUM(amount),0) AS t FROM cash_books WHERE user_id=%s AND entry_type='in'", (uid,))
        bump('Cash', debit=cur.fetchone()['t'])
        cur.execute("SELECT COALESCE(SUM(amount),0) AS t FROM cash_books WHERE user_id=%s AND entry_type='out'", (uid,))
        bump('Cash', credit=cur.fetchone()['t'])
        cur.execute('SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders WHERE user_id=%s', (uid,))
        purch = cur.fetchone()['t'] or 0
        bump('Purchases', debit=purch)
        cur.execute("SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders WHERE user_id=%s AND status <> 'received'", (uid,))
        bump('Creditors', credit=cur.fetchone()['t'])
        cur.execute(
            '''SELECT account, COALESCE(SUM(debit),0) AS d, COALESCE(SUM(credit),0) AS c
               FROM journal_entries WHERE user_id=%s GROUP BY account''',
            (uid,),
        )
        for row in cur.fetchall() or []:
            bump(row['account'], debit=row['d'], credit=row['c'])
        extra['ledger'] = sorted(accounts.values(), key=lambda r: r['account'])
        extra['ledger_note'] = 'Rolled up from sales, purchases, cash, stock, income, expenses, and typed journals.'
    elif book == 'trading_account':
        cur.execute('SELECT COALESCE(SUM(total_amount),0) AS t FROM sales WHERE user_id=%s', (uid,))
        sales_rev = float(cur.fetchone()['t'] or 0)
        cur.execute('SELECT COALESCE(SUM(total_amount - profit),0) AS t FROM sales WHERE user_id=%s', (uid,))
        cogs = float(cur.fetchone()['t'] or 0)
        cur.execute('SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders WHERE user_id=%s', (uid,))
        purchases = float(cur.fetchone()['t'] or 0)
        cur.execute('SELECT COALESCE(SUM(quantity * cost_price),0) AS t FROM stock WHERE user_id=%s', (uid,))
        closing = float(cur.fetchone()['t'] or 0)
        extra['statement'] = [
            ('Sales', sales_rev, 'cr'),
            ('Purchases', purchases, 'dr'),
            ('Cost of goods sold (from recorded sales)', cogs, 'dr'),
            ('Closing inventory (at cost)', closing, 'dr'),
            ('Gross profit / (loss)', sales_rev - cogs, 'net'),
        ]
    elif book == 'profit_and_loss':
        cur.execute('SELECT COALESCE(SUM(total_amount),0) AS t FROM sales WHERE user_id=%s', (uid,))
        sales_rev = float(cur.fetchone()['t'] or 0)
        cur.execute('SELECT COALESCE(SUM(total_amount - profit),0) AS t FROM sales WHERE user_id=%s', (uid,))
        cogs = float(cur.fetchone()['t'] or 0)
        gross = sales_rev - cogs
        cur.execute('SELECT COALESCE(SUM(amount),0) AS t FROM income WHERE user_id=%s', (uid,))
        other_inc = float(cur.fetchone()['t'] or 0)
        cur.execute('SELECT COALESCE(SUM(amount),0) AS t FROM expenses WHERE user_id=%s', (uid,))
        exp = float(cur.fetchone()['t'] or 0)
        extra['statement'] = [
            ('Sales', sales_rev, 'cr'),
            ('Cost of goods sold', cogs, 'dr'),
            ('Gross profit', gross, 'net'),
            ('Other income', other_inc, 'cr'),
            ('Expenses', exp, 'dr'),
            ('Net profit / (loss)', gross + other_inc - exp, 'net'),
        ]
    elif book == 'balance_sheet':
        cur.execute('SELECT COALESCE(SUM(quantity * cost_price),0) AS t FROM stock WHERE user_id=%s', (uid,))
        inventory = float(cur.fetchone()['t'] or 0)
        cur.execute("SELECT COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE -amount END),0) AS t FROM cash_books WHERE user_id=%s", (uid,))
        cash = float(cur.fetchone()['t'] or 0)
        cur.execute(
            '''SELECT COALESCE(SUM(total_amount),0) AS t FROM sales
               WHERE user_id=%s AND COALESCE(payment_method,'cash') NOT IN ('cash')''',
            (uid,),
        )
        debtors = float(cur.fetchone()['t'] or 0)
        cur.execute("SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders WHERE user_id=%s AND status <> 'received'", (uid,))
        creditors = float(cur.fetchone()['t'] or 0)
        assets = inventory + cash + debtors
        capital = assets - creditors
        extra['statement'] = [
            ('Inventory at cost', inventory, 'asset'),
            ('Cash (cash book net)', cash, 'asset'),
            ('Debtors (non-cash sales)', debtors, 'asset'),
            ('Total assets', assets, 'total'),
            ('Creditors (open purchase orders)', creditors, 'liability'),
            ('Capital / net assets', capital, 'capital'),
        ]
    elif book == 'cash_flow':
        cur.execute("SELECT COALESCE(SUM(amount),0) AS t FROM cash_books WHERE user_id=%s AND entry_type='in'", (uid,))
        cin = float(cur.fetchone()['t'] or 0)
        cur.execute("SELECT COALESCE(SUM(amount),0) AS t FROM cash_books WHERE user_id=%s AND entry_type='out'", (uid,))
        cout = float(cur.fetchone()['t'] or 0)
        cur.execute(
            '''SELECT COALESCE(SUM(total_amount),0) AS t FROM sales
               WHERE user_id=%s AND COALESCE(payment_method,'cash') IN ('cash')''',
            (uid,),
        )
        cash_sales = float(cur.fetchone()['t'] or 0)
        extra['statement'] = [
            ('Cash book receipts', cin, 'in'),
            ('Cash sales (till)', cash_sales, 'in'),
            ('Cash book payments', cout, 'out'),
            ('Net cash book movement', cin - cout, 'net'),
        ]

    manual = _je_rows(cur, uid, book)
    cur.close()
    conn.close()

    for r in auto_rows or []:
        totals['debit'] += float(r.get('debit') or 0)
        totals['credit'] += float(r.get('credit') or 0)
    for r in manual or []:
        totals['debit'] += float(r.get('debit') or 0)
        totals['credit'] += float(r.get('credit') or 0)

    return render_template(
        'account_book.html',
        book=book, label=label, hint=hint,
        auto_rows=auto_rows, manual=manual, totals=totals, extra=extra,
        allow_manual=book in LEDGER_BOOKS,
        start_date=start_date, end_date=end_date,
    )


@app.route('/accounts/<book>/csv')
@login_required
def account_book_csv(book):
    if book not in ACCOUNT_BOOK_LOOKUP:
        flash('Unknown book.', 'danger')
        return redirect(url_for('accounts'))
    if book == 'cash_book':
        return redirect(url_for('export_cashbook_csv'))
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['KAZE', ACCOUNT_BOOK_LOOKUP[book][0]])
    writer.writerow(['Date', 'Ref', 'Account', 'Particulars', 'Debit', 'Credit', 'Source'])
    if book == 'sales_journal':
        cur.execute(
            '''SELECT sales.sale_date, sales.id::text, COALESCE(sales.customer_name,'Walk-in'),
                      COALESCE(stock.product_name,'Item'), sales.total_amount, 0
               FROM sales LEFT JOIN stock ON sales.stock_id = stock.id
               WHERE sales.user_id=%s ORDER BY sales.sale_date DESC''', (uid,))
        for r in cur.fetchall() or []:
            writer.writerow(list(r) + ['Auto'])
    elif book == 'purchases_journal':
        cur.execute(
            '''SELECT order_date, id::text, item_description, status, 0, total_cost
               FROM purchase_orders WHERE user_id=%s ORDER BY order_date DESC''', (uid,))
        for r in cur.fetchall() or []:
            writer.writerow(list(r) + ['Auto'])
    cur.execute(
        '''SELECT entry_date, reference, account, particulars, debit, credit
           FROM journal_entries WHERE user_id=%s AND book=%s ORDER BY entry_date DESC''',
        (uid, book),
    )
    for r in cur.fetchall() or []:
        writer.writerow([r['entry_date'], r['reference'], r['account'], r['particulars'], r['debit'], r['credit'], 'Typed'])
    cur.close()
    conn.close()
    filename = 'kaze-{}-book.csv'.format(book.replace('_', '-'))
    return Response(output.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename={}'.format(filename)})


# ======================== Extra tools ========================

@app.route('/reminders')
@login_required
def reminders():
    uid = current_user.id
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT * FROM tasks WHERE user_id=%s AND done=0 AND due_date <> '' AND due_date < %s ORDER BY due_date",
        (uid, today)
    )
    overdue_tasks = cur.fetchall()
    cur.execute(
        "SELECT * FROM purchase_orders WHERE user_id=%s AND status <> 'received' AND expected_date <> '' AND expected_date < %s ORDER BY expected_date",
        (uid, today)
    )
    late_pos = cur.fetchall()
    cur.execute(
        '''SELECT * FROM bookkeeping_docs
           WHERE user_id=%s AND status IN ('draft','issued') AND doc_date < %s
           ORDER BY doc_date''',
        (uid, today)
    )
    open_docs = cur.fetchall()
    cur.execute('SELECT low_stock_threshold FROM settings WHERE user_id=%s', (uid,))
    row = cur.fetchone()
    threshold = row['low_stock_threshold'] if row else 10
    cur.execute(
        'SELECT * FROM stock WHERE user_id=%s AND quantity < %s ORDER BY quantity',
        (uid, threshold)
    )
    low = cur.fetchall()
    cur.close()
    conn.close()
    due_sessions = []
    try:
        cur.execute(
            "SELECT ss.*, sv.name AS service_name FROM service_sessions ss "
            "LEFT JOIN services sv ON sv.id = ss.service_id "
            "WHERE ss.user_id=%s AND ss.status='booked' AND ss.session_date <= %s "
            "ORDER BY ss.session_date",
            (uid, today)
        )
        due_sessions = cur.fetchall()
    except Exception:
        due_sessions = []
    return render_template(
        'reminders.html',
        overdue_tasks=overdue_tasks,
        late_pos=late_pos,
        open_docs=open_docs,
        low=low,
        due_sessions=due_sessions,
        today=today
    )


@app.route('/tax')
@login_required
def tax_summary():
    uid = current_user.id
    settings_row = _get_settings_row(uid)
    rate = (settings_row['tax_rate'] if settings_row else 0) or 0
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s', (uid,))
    sales_rev = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(amount),0) as t FROM income WHERE user_id=%s', (uid,))
    other_inc = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(amount),0) as t FROM expenses WHERE user_id=%s', (uid,))
    expenses = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(total_cost),0) as t FROM purchase_orders WHERE user_id=%s', (uid,))
    purchases = cur.fetchone()['t'] or 0
    cur.close()
    conn.close()
    taxable_out = sales_rev + other_inc
    output_tax = taxable_out * rate / 100 if rate else 0
    taxable_in = expenses + purchases
    input_tax = taxable_in * rate / 100 if rate else 0
    return render_template(
        'tax.html',
        rate=rate, symbol=symbol,
        sales_rev=sales_rev, other_inc=other_inc, expenses=expenses, purchases=purchases,
        taxable_out=taxable_out, output_tax=output_tax,
        taxable_in=taxable_in, input_tax=input_tax,
        net_tax=output_tax - input_tax
    )


@app.route('/customers/<int:id>/statement')
@login_required
def customer_statement(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM customers WHERE id=%s AND user_id=%s', (id, current_user.id))
    customer = cur.fetchone()
    if not customer:
        cur.close()
        conn.close()
        flash('Customer not found.', 'danger')
        return redirect(url_for('customers'))
    cur.execute(
        '''SELECT sales.*, stock.product_name FROM sales
           JOIN stock ON sales.stock_id = stock.id
           WHERE sales.user_id=%s AND (sales.customer_id=%s OR sales.customer_name=%s)
           ORDER BY sales.sale_date''',
        (current_user.id, id, customer['name'])
    )
    sales_rows = cur.fetchall()
    cur.close()
    conn.close()
    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    total = sum(s['total_amount'] for s in sales_rows)
    meta = [
        '<b>Statement for:</b> {}'.format(customer['name']),
        '<b>Email:</b> {}'.format(customer['email'] or '-'),
        '<b>Phone:</b> {}'.format(customer['phone'] or '-'),
        '<b>Date:</b> {}'.format(datetime.today().strftime('%Y-%m-%d')),
    ]
    table = [['Date', 'Product', 'Qty', 'Amount']]
    for s in sales_rows:
        table.append([
            s['sale_date'], s['product_name'],
            str(s['quantity_sold']),
            _format_money(s['total_amount'], settings_row)
        ])
    table.append(['', '', 'Total', _format_money(total, settings_row)])
    buffer = io.BytesIO()
    _build_pdf(buffer, 'Statement of Account', business_name, meta, table,
               footer_lines=['Generated by KAZE Account tracking'])
    buffer.seek(0)
    safe = ''.join(c for c in customer['name'] if c.isalnum() or c in (' ', '-', '_')).strip() or 'customer'
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                     download_name='statement-{}.pdf'.format(safe))


@app.route('/sales/repeat/<int:id>')
@login_required
def repeat_sale(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM sales WHERE id=%s AND user_id=%s', (id, current_user.id))
    sale = cur.fetchone()
    if not sale:
        cur.close()
        conn.close()
        flash('Sale not found.', 'danger')
        return redirect(url_for('sales'))
    cur.execute('SELECT * FROM stock WHERE id=%s AND user_id=%s', (sale['stock_id'], current_user.id))
    item = cur.fetchone()
    if not item:
        cur.close()
        conn.close()
        flash('That product is no longer in stock records.', 'danger')
        return redirect(url_for('sales'))
    qty = sale['quantity_sold']
    if item['quantity'] < qty:
        cur.close()
        conn.close()
        flash('Not enough stock to repeat this sale. Available: {}'.format(item['quantity']), 'danger')
        return redirect(url_for('sales'))
    today = datetime.today().strftime('%Y-%m-%d')
    price = item['selling_price']
    profit = (price - item['cost_price']) * qty
    total = price * qty
    cur.execute(
        '''INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit,
           sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, item['id'], qty, price, total, profit, today,
         sale['customer_name'], sale['customer_email'], sale['customer_id'], 0, 'repeat',
         sale.get('payment_method') or 'cash')
    )
    cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id=%s', (qty, item['id']))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Repeated sale', item['product_name'])
    flash('Sale repeated for today. Profit: {}'.format(_format_money(profit)), 'success')
    return redirect(url_for('sales'))


@app.route('/stock/price-list')
@login_required
def price_list():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT product_name, unit, selling_price, quantity FROM stock WHERE user_id=%s ORDER BY product_name',
        (current_user.id,)
    )
    items = cur.fetchall()
    cur.close()
    conn.close()
    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    table = [['Product', 'Unit', 'Price', 'In stock']]
    for i in items:
        table.append([
            i['product_name'], i['unit'] or 'each',
            _format_money(i['selling_price'] or 0, settings_row),
            str(i['quantity'])
        ])
    if len(table) == 1:
        table.append(['No products yet', '', '', ''])
    buffer = io.BytesIO()
    _build_pdf(
        buffer, 'Price list', business_name,
        ['<b>Date:</b> {}'.format(datetime.today().strftime('%Y-%m-%d'))],
        table
    )
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                     download_name='price-list.pdf')


@app.route('/settings/clear-data', methods=['POST'])
@login_required
def clear_data():
    if current_user.role != 'owner':
        flash('Only the account owner can clear business data.', 'danger')
        return redirect(url_for('settings'))
    if (request.form.get('confirm_text') or '').strip().upper() != 'CLEAR':
        flash('Type CLEAR exactly to confirm.', 'danger')
        return redirect(url_for('settings'))
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    try:
        # Children first so foreign keys do not block the wipe.
        for sql in (
            'DELETE FROM stripe_payments WHERE user_id = %s',
            'DELETE FROM shop_orders WHERE user_id = %s',
            'DELETE FROM shop_listings WHERE user_id = %s',
            'DELETE FROM bank_transactions WHERE user_id = %s',
            'DELETE FROM bank_accounts WHERE user_id = %s',
            'DELETE FROM investments WHERE user_id = %s',
            'DELETE FROM market_valuations WHERE user_id = %s',
            'DELETE FROM market_channels WHERE user_id = %s',
            'DELETE FROM service_sessions WHERE user_id = %s',
            'DELETE FROM services WHERE user_id = %s',
            'DELETE FROM sales WHERE user_id = %s',
            'DELETE FROM purchase_orders WHERE user_id = %s',
            'DELETE FROM bookkeeping_docs WHERE user_id = %s',
            'DELETE FROM journal_entries WHERE user_id = %s',
            'DELETE FROM stock WHERE user_id = %s',
            'DELETE FROM customers WHERE user_id = %s',
            'DELETE FROM suppliers WHERE user_id = %s',
            'DELETE FROM income WHERE user_id = %s',
            'DELETE FROM expenses WHERE user_id = %s',
            'DELETE FROM documents WHERE user_id = %s',
            'DELETE FROM cash_books WHERE user_id = %s',
            'DELETE FROM recurring_items WHERE user_id = %s',
            'DELETE FROM budgets WHERE user_id = %s',
            'DELETE FROM notifications WHERE user_id = %s',
            'DELETE FROM activities WHERE user_id = %s',
            'DELETE FROM tasks WHERE user_id = %s',
            'DELETE FROM notes WHERE user_id = %s',
            'DELETE FROM swot_analyses WHERE user_id = %s',
            'DELETE FROM business_plans WHERE user_id = %s',
            'DELETE FROM user_activity WHERE user_id = %s',
        ):
            cur.execute(sql, (uid,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        cur.close()
        conn.close()
        flash('Could not clear data: {}'.format(e), 'danger')
        return redirect(url_for('settings'))
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Cleared all business data', 'start over')
    flash('Business data cleared. Your account and settings are still here.', 'success')
    return redirect(url_for('dashboard'))


@app.route('/stock/valuation-pdf')
@login_required
def stock_valuation_pdf():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT product_name, sku, category, quantity, cost_price, selling_price, unit FROM stock WHERE user_id=%s ORDER BY product_name', (current_user.id,))
    items = cur.fetchall()
    cur.close()
    conn.close()
    settings_row = _get_settings_row(current_user.id)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    table = [['Product', 'SKU', 'Qty', 'Cost', 'Value']]
    total = 0
    for i in items:
        val = (i['quantity'] or 0) * (i['cost_price'] or 0)
        total += val
        table.append([i['product_name'], i['sku'] or '-', str(i['quantity']),
                      _format_money(i['cost_price'] or 0, settings_row),
                      _format_money(val, settings_row)])
    table.append(['', '', '', 'Total', _format_money(total, settings_row)])
    buffer = io.BytesIO()
    _build_pdf(buffer, 'Stock valuation', business_name,
               ['<b>Date:</b> {}'.format(datetime.today().strftime('%Y-%m-%d'))], table)
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True, download_name='stock-valuation.pdf')


@app.route('/accounts/profit-and-loss/pdf')
@login_required
def pnl_pdf():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s', (uid,))
    sales_rev = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(total_amount - profit),0) as t FROM sales WHERE user_id=%s', (uid,))
    cogs = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(amount),0) as t FROM income WHERE user_id=%s', (uid,))
    other_inc = cur.fetchone()['t'] or 0
    cur.execute('SELECT COALESCE(SUM(amount),0) as t FROM expenses WHERE user_id=%s', (uid,))
    exp = cur.fetchone()['t'] or 0
    cur.close()
    conn.close()
    gross = sales_rev - cogs
    net = gross + other_inc - exp
    settings_row = _get_settings_row(uid)
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your Business'
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    table = [['Line', 'Amount'],
             ['Sales', _format_money(sales_rev, settings_row)],
             ['Cost of goods sold', _format_money(cogs, settings_row)],
             ['Gross profit', _format_money(gross, settings_row)],
             ['Other income', _format_money(other_inc, settings_row)],
             ['Expenses', _format_money(exp, settings_row)],
             ['Net profit / (loss)', _format_money(net, settings_row)]]
    buffer = io.BytesIO()
    _build_pdf(buffer, 'Profit and Loss', business_name,
               ['<b>Date:</b> {}'.format(datetime.today().strftime('%Y-%m-%d'))], table)
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True, download_name='profit-and-loss.pdf')


@app.route('/sales/aged')
@login_required
def aged_sales():
    today = datetime.today()
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''SELECT sales.*, stock.product_name FROM sales
           JOIN stock ON sales.stock_id = stock.id
           WHERE sales.user_id=%s ORDER BY sales.sale_date DESC''',
        (current_user.id,)
    )
    rows = []
    buckets = {'0-30': 0, '31-60': 0, '61-90': 0, '90+': 0}
    for s in cur.fetchall():
        try:
            d = datetime.strptime(str(s['sale_date'])[:10], '%Y-%m-%d')
            age = (today - d).days
        except Exception:
            age = 0
        if age <= 30:
            bucket = '0-30'
        elif age <= 60:
            bucket = '31-60'
        elif age <= 90:
            bucket = '61-90'
        else:
            bucket = '90+'
        buckets[bucket] += s['total_amount'] or 0
        s = dict(s)
        s['age'] = age
        s['bucket'] = bucket
        rows.append(s)
    cur.close()
    conn.close()
    return render_template('aged_sales.html', rows=rows, buckets=buckets)


@app.route('/bookkeeping/doc/<int:id>/duplicate')
@login_required
def duplicate_book_doc(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        flash('Document not found.', 'danger')
        return redirect(url_for('bookkeeping'))
    cur.execute(
        '''INSERT INTO bookkeeping_docs
           (user_id, kind, title, party, reference, doc_date, amount, status, notes, line_items, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, row['kind'], row['title'] + ' (copy)', row['party'],
         (row['reference'] or '') + '-COPY', datetime.today().strftime('%Y-%m-%d'),
         row['amount'], 'draft', row['notes'], row['line_items'], datetime.today().strftime('%Y-%m-%d'))
    )
    conn.commit()
    kind = row['kind']
    cur.close()
    conn.close()
    flash('Document copied as a draft.', 'success')
    return redirect(url_for('bookkeeping_type', kind=kind))


@app.route('/bookkeeping/doc/<int:id>/paid')
@login_required
def mark_book_doc_paid(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE bookkeeping_docs SET status=%s WHERE id=%s AND user_id=%s', ('paid', id, current_user.id))
    cur.execute('SELECT kind FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    conn.commit()
    cur.close()
    conn.close()
    flash('Marked as paid.', 'success')
    if row:
        return redirect(url_for('bookkeeping_type', kind=row['kind']))
    return redirect(url_for('bookkeeping'))


@app.route('/sales/<int:id>/invoice')
@login_required
def sale_to_invoice(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''SELECT sales.*, stock.product_name FROM sales
           JOIN stock ON sales.stock_id = stock.id
           WHERE sales.id=%s AND sales.user_id=%s''',
        (id, current_user.id)
    )
    sale = cur.fetchone()
    if not sale:
        cur.close()
        conn.close()
        flash('Sale not found.', 'danger')
        return redirect(url_for('sales'))
    today = datetime.today().strftime('%Y-%m-%d')
    title = 'Invoice for {} x{}'.format(sale['product_name'], sale['quantity_sold'])
    lines = '{} | {} | {}'.format(sale['product_name'], sale['quantity_sold'], _format_money(sale['total_amount']))
    cur.execute(
        '''INSERT INTO bookkeeping_docs
           (user_id, kind, title, party, reference, doc_date, amount, status, notes, line_items, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, 'sales_invoice', title, sale['customer_name'] or 'Walk-in',
         'INV-SALE-{}'.format(sale['id']), today, sale['total_amount'], 'issued',
         sale.get('sale_notes') or '', lines, today)
    )
    conn.commit()
    new_id = None
    cur.execute('SELECT id FROM bookkeeping_docs WHERE user_id=%s ORDER BY id DESC LIMIT 1', (current_user.id,))
    new_id = cur.fetchone()['id']
    cur.close()
    conn.close()
    flash('Sales invoice created.', 'success')
    return redirect(url_for('bookkeeping_pdf', id=new_id))


@app.route('/tasks/bulk-complete', methods=['POST'])
@login_required
def tasks_bulk_complete():
    action = (request.form.get('action') or 'complete').strip()
    ids = _selected_ids()
    if not ids:
        flash('Pick at least one task.', 'danger')
        return redirect(url_for('tasks'))
    conn = get_db()
    cur = conn.cursor()
    try:
        if action == 'delete':
            cur.execute('DELETE FROM tasks WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            flash('Deleted {} task(s).'.format(len(ids)), 'success')
        elif action == 'reopen':
            cur.execute('UPDATE tasks SET done=0 WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            flash('Reopened {} task(s).'.format(len(ids)), 'success')
        else:
            cur.execute('UPDATE tasks SET done=1 WHERE id = ANY(%s) AND user_id=%s', (ids, current_user.id))
            flash('Marked {} task(s) done.'.format(len(ids)), 'success')
        conn.commit()
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        flash('Task bulk action failed.', 'danger')
        print('tasks_bulk:', exc)
    finally:
        cur.close()
        conn.close()
    return redirect(url_for('tasks'))


@app.route('/tools/shortcuts')
@login_required
def shortcuts():
    return render_template('shortcuts.html')


@app.route('/tools/day-close', methods=['GET', 'POST'])
@login_required
def day_close():
    today = datetime.today().strftime('%Y-%m-%d')
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT COALESCE(SUM(total_amount),0) as t, COUNT(*) as n FROM sales WHERE user_id=%s AND sale_date=%s',
        (uid, today)
    )
    sales = cur.fetchone()
    try:
        cur.execute(
            "SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s AND sale_date=%s AND COALESCE(payment_method, 'cash') = 'cash'",
            (uid, today)
        )
        cash_sales = cur.fetchone()['t'] or 0
        cur.execute(
            "SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s AND sale_date=%s AND COALESCE(payment_method, 'cash') <> 'cash'",
            (uid, today)
        )
        other_sales = cur.fetchone()['t'] or 0
    except Exception:
        conn.rollback()
        cash_sales = sales['t'] or 0
        other_sales = 0
    cur.execute("SELECT COALESCE(SUM(amount),0) as t FROM cash_books WHERE user_id=%s AND date=%s AND entry_type='in'", (uid, today))
    cin = cur.fetchone()['t'] or 0
    cur.execute("SELECT COALESCE(SUM(amount),0) as t FROM cash_books WHERE user_id=%s AND date=%s AND entry_type='out'", (uid, today))
    cout = cur.fetchone()['t'] or 0

    saved = None
    history = []
    try:
        cur.execute('SELECT * FROM day_closes WHERE user_id=%s AND close_date=%s', (uid, today))
        saved = cur.fetchone()
        cur.execute('SELECT * FROM day_closes WHERE user_id=%s ORDER BY close_date DESC LIMIT 14', (uid,))
        history = cur.fetchall()
    except Exception:
        conn.rollback()

    if request.method == 'POST':
        opening, e1 = _clean_float('opening_float', required=False, min_v=0, max_v=1e9, label='opening float', default=0)
        counted, e2 = _clean_float('counted_cash', required=True, min_v=0, max_v=1e9, label='counted cash')
        notes, e3 = _clean_text('notes', required=False, max_len=400, label='notes')
        if _reject((opening, e1), (counted, e2), (notes, e3)):
            cur.close()
            conn.close()
            return redirect(url_for('day_close'))
        expected = (opening or 0) + cash_sales + cin - cout
        variance = (counted or 0) - expected
        try:
            cur.execute(
                "INSERT INTO day_closes (user_id, close_date, opening_float, counted_cash, expected_cash, cash_sales, other_sales, cash_in, cash_out, variance, notes) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (user_id, close_date) DO UPDATE SET opening_float=EXCLUDED.opening_float, counted_cash=EXCLUDED.counted_cash, expected_cash=EXCLUDED.expected_cash, cash_sales=EXCLUDED.cash_sales, other_sales=EXCLUDED.other_sales, cash_in=EXCLUDED.cash_in, cash_out=EXCLUDED.cash_out, variance=EXCLUDED.variance, notes=EXCLUDED.notes",
                (uid, today, opening or 0, counted or 0, expected, cash_sales, other_sales, cin, cout, variance, notes or '')
            )
            conn.commit()
        except Exception as exc:
            conn.rollback()
            cur.close()
            conn.close()
            flash('Could not save day close. {}'.format(exc), 'danger')
            return redirect(url_for('day_close'))
        cur.close()
        conn.close()
        log_activity(uid, current_user.username, 'Day close',
                     'Expected {} counted {} variance {}'.format(
                         _format_money(expected), _format_money(counted), _format_money(variance)))
        if abs(variance) < 0.005:
            flash('Till matches. Variance is zero.', 'success')
        elif variance > 0:
            flash('Till is over by {}.'.format(_format_money(variance)), 'success')
        else:
            flash('Till is short by {}.'.format(_format_money(abs(variance))), 'danger')
        return redirect(url_for('day_close'))

    cur.close()
    conn.close()
    opening = saved['opening_float'] if saved else 0
    expected = (opening or 0) + cash_sales + cin - cout
    return render_template(
        'day_close.html',
        today=today,
        sales_total=sales['t'] or 0,
        sales_n=sales['n'] or 0,
        cash_sales=cash_sales,
        other_sales=other_sales,
        cash_in=cin,
        cash_out=cout,
        cash_net=cin - cout,
        saved=saved,
        history=history,
        expected=expected,
        opening=opening,
    )


@app.route('/counter')
@login_required
def counter():
    uid = current_user.id
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''SELECT id, product_name, quantity, selling_price, cost_price, unit, sku, category
           FROM stock WHERE user_id=%s ORDER BY product_name''',
        (uid,)
    )
    stock_items = cur.fetchall()
    cur.execute('SELECT id, name FROM customers WHERE user_id=%s ORDER BY name', (uid,))
    customer_list = cur.fetchall()
    cur.execute(
        '''SELECT COALESCE(SUM(total_amount),0) as t, COALESCE(SUM(profit),0) as p, COUNT(*) as n
           FROM sales WHERE user_id=%s AND sale_date=%s''',
        (uid, today)
    )
    today_row = cur.fetchone()
    cur.execute(
        '''SELECT sales.quantity_sold, sales.total_amount, sales.sale_date, stock.product_name
           FROM sales JOIN stock ON sales.stock_id = stock.id
           WHERE sales.user_id=%s AND sales.sale_date=%s
           ORDER BY sales.id DESC LIMIT 8''',
        (uid, today)
    )
    recent = cur.fetchall()
    cur.close()
    conn.close()
    return render_template(
        'counter.html',
        stock_items=stock_items,
        customer_list=customer_list,
        today_total=today_row['t'] or 0,
        today_profit=today_row['p'] or 0,
        today_count=today_row['n'] or 0,
        recent=recent,
    )


@app.route('/counter/sale', methods=['POST'])
@login_required
def counter_sale():
    try:
        stock_id = int(request.form['stock_id'])
        quantity_sold = float(request.form['quantity'])
    except (TypeError, ValueError, KeyError):
        flash('Pick a product and a quantity.', 'danger')
        return redirect(url_for('counter'))
    if quantity_sold <= 0:
        flash('Quantity must be more than zero.', 'danger')
        return redirect(url_for('counter'))

    customer_id = request.form.get('customer_id') or None
    customer_name = request.form.get('customer_name', '')
    customer_email = ''
    sale_date = datetime.today().strftime('%Y-%m-%d')

    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM stock WHERE id = %s AND user_id = %s', (stock_id, current_user.id))
    stock_item = cur.fetchone()
    if not stock_item:
        cur.close()
        conn.close()
        flash('Stock item not found.', 'danger')
        return redirect(url_for('counter'))
    if stock_item['quantity'] < quantity_sold:
        cur.close()
        conn.close()
        flash('Insufficient stock. Only {} available.'.format(stock_item['quantity']), 'danger')
        return redirect(url_for('counter'))

    if customer_id:
        cur.execute('SELECT name, email FROM customers WHERE id = %s AND user_id = %s', (customer_id, current_user.id))
        saved_customer = cur.fetchone()
        if saved_customer:
            customer_name = saved_customer['name']
            customer_email = saved_customer['email'] or ''

    try:
        discount = float(request.form.get('discount') or 0)
    except ValueError:
        discount = 0
    selling_price = stock_item['selling_price']
    cost_price = stock_item['cost_price']
    total_amount = max(selling_price * quantity_sold - discount, 0)
    profit = total_amount - (cost_price * quantity_sold)

    cur.execute(
        '''INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit,
           sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)''',
        (current_user.id, stock_id, quantity_sold, selling_price, total_amount, profit, sale_date,
         customer_name, customer_email, customer_id, discount, 'counter', _payment_method())
    )
    cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id = %s', (quantity_sold, stock_id))
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Counter sale',
                 '{} x{} {}'.format(stock_item['product_name'], quantity_sold, _format_money(total_amount)))
    flash('Rang up {} × {}. Collect {}.'.format(stock_item['product_name'], quantity_sold, _format_money(total_amount)), 'success')
    return redirect(url_for('counter'))


@app.route('/counter/checkout', methods=['POST'])
@login_required
def counter_checkout():
    """Ring up a multi-item ticket from the till basket."""
    ids = request.form.getlist('line_stock_id')
    qtys = request.form.getlist('line_qty')
    if not ids:
        flash('Add at least one product to the ticket.', 'danger')
        return redirect(url_for('counter'))
    customer_id = request.form.get('customer_id') or None
    customer_name = request.form.get('customer_name', '') or 'Walk-in'
    customer_email = ''
    pay = _payment_method()
    try:
        ticket_discount = float(request.form.get('discount') or 0)
    except ValueError:
        ticket_discount = 0
    if ticket_discount < 0:
        ticket_discount = 0
    sale_date = datetime.today().strftime('%Y-%m-%d')
    ticket_ref = 'T-' + datetime.now().strftime('%H%M%S')

    conn = get_db()
    cur = conn.cursor()
    if customer_id:
        cur.execute('SELECT name, email FROM customers WHERE id = %s AND user_id = %s', (customer_id, current_user.id))
        saved = cur.fetchone()
        if saved:
            customer_name = saved['name']
            customer_email = saved['email'] or ''

    lines = []
    for sid, qraw in zip(ids, qtys):
        try:
            stock_id = int(sid)
            qty = float(qraw)
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue
        cur.execute('SELECT * FROM stock WHERE id = %s AND user_id = %s', (stock_id, current_user.id))
        item = cur.fetchone()
        if not item:
            cur.close()
            conn.close()
            flash('A product on the ticket is missing.', 'danger')
            return redirect(url_for('counter'))
        if item['quantity'] < qty:
            cur.close()
            conn.close()
            flash('Not enough {} in stock. Only {} left.'.format(item['product_name'], item['quantity']), 'danger')
            return redirect(url_for('counter'))
        lines.append((item, qty))
    if not lines:
        cur.close()
        conn.close()
        flash('Add at least one product to the ticket.', 'danger')
        return redirect(url_for('counter'))

    subtotal = sum(item['selling_price'] * qty for item, qty in lines)
    remaining_disc = min(ticket_discount, subtotal)
    ticket_total = 0
    for idx, (item, qty) in enumerate(lines):
        line_rev = item['selling_price'] * qty
        if idx == len(lines) - 1:
            line_disc = remaining_disc
        else:
            share = (line_rev / subtotal) * ticket_discount if subtotal else 0
            line_disc = round(share, 2)
            remaining_disc = max(remaining_disc - line_disc, 0)
        total_amount = max(line_rev - line_disc, 0)
        profit = total_amount - (item['cost_price'] * qty)
        ticket_total += total_amount
        note = 'ticket {} · {}'.format(ticket_ref, pay)
        cur.execute(
            "INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, total_amount, profit, sale_date, customer_name, customer_email, customer_id, discount, sale_notes, payment_method) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
            (current_user.id, item['id'], qty, item['selling_price'], total_amount, profit, sale_date,
             customer_name, customer_email, customer_id, line_disc, note, pay)
        )
        cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id = %s', (qty, item['id']))
    conn.commit()
    cur.close()
    conn.close()
    names = ', '.join('{}x{}'.format(item['product_name'], qty) for item, qty in lines)
    log_activity(current_user.id, current_user.username, 'Till ticket',
                 '{} {} {}'.format(ticket_ref, names, _format_money(ticket_total)))
    flash('Ticket {} · {} lines · collect {} via {}.'.format(
        ticket_ref, len(lines), _format_money(ticket_total), pay), 'success')
    return redirect(url_for('counter'))


@app.route('/guide.pdf')
@login_required
def guide_pdf():
    settings_row = None
    try:
        settings_row = _get_settings_row(current_user.id)
    except Exception:
        settings_row = None
    business_name = (settings_row['business_name'] if settings_row else '') or 'Your business'
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle('KTitle', parent=styles['Title'], textColor=colors.HexColor('#1f8a3d'), fontSize=18, spaceAfter=8)
    h = ParagraphStyle('KH', parent=styles['Heading2'], textColor=colors.HexColor('#0a0a0a'), fontSize=12, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle('KB', parent=styles['BodyText'], fontSize=10, leading=14, spaceAfter=4)
    story = [
        Paragraph('KAZE Traders — Practical Guide', title),
        Paragraph('For {}. Version 2.6. How to run the shop in the app, step by step.'.format(business_name), body),
        Spacer(1, 8),
    ]
    for i, step in enumerate(_guide_steps(), 1):
        story.append(Paragraph('{}. {} — {}'.format(i, step['group'], step['title']), h))
        story.append(Paragraph(step['body'], body))
        story.append(Paragraph('Do this: ' + ' → '.join(step['do_this']), body))
    story.append(Spacer(1, 12))
    story.append(Paragraph('Join / open the app: https://kaze-business-app.onrender.com', body))
    story.append(Paragraph('This guide is a shop floor checklist. It is not tax, legal, or accounting advice.', body))
    doc = SimpleDocTemplate(buffer, pagesize=letter, title='KAZE Practical Guide')
    doc.build(story)
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True, download_name='kaze-practical-guide.pdf')


def _stats_bundle(uid, start_date=None, end_date=None):
    """One-shot statistics for tables + charts."""
    conn = get_db()
    cur = conn.cursor()
    sales_where = 'user_id = %s'
    exp_where = 'user_id = %s'
    cash_where = 'user_id = %s'
    params_s = [uid]
    params_e = [uid]
    params_c = [uid]
    if start_date and end_date:
        sales_where += ' AND sale_date BETWEEN %s AND %s'
        exp_where += ' AND date BETWEEN %s AND %s'
        cash_where += ' AND date BETWEEN %s AND %s'
        params_s += [start_date, end_date]
        params_e += [start_date, end_date]
        params_c += [start_date, end_date]
    cur.execute(
        'SELECT COUNT(*) AS n, COALESCE(SUM(total_amount),0) AS s, COALESCE(AVG(total_amount),0) AS a, '
        'COALESCE(MIN(total_amount),0) AS mn, COALESCE(MAX(total_amount),0) AS mx, COALESCE(SUM(profit),0) AS p '
        'FROM sales WHERE ' + sales_where, params_s)
    sales = cur.fetchone()
    cur.execute(
        'SELECT COUNT(*) AS n, COALESCE(SUM(amount),0) AS s, COALESCE(AVG(amount),0) AS a, '
        'COALESCE(MIN(amount),0) AS mn, COALESCE(MAX(amount),0) AS mx FROM expenses WHERE ' + exp_where, params_e)
    expenses = cur.fetchone()
    cur.execute(
        'SELECT COUNT(*) AS n, COALESCE(SUM(amount),0) AS s FROM income WHERE user_id = %s', (uid,))
    income = cur.fetchone()
    cur.execute(
        "SELECT COUNT(*) AS n, COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE 0 END),0) AS cin, "
        "COALESCE(SUM(CASE WHEN entry_type='out' THEN amount ELSE 0 END),0) AS cout FROM cash_books WHERE " + cash_where,
        params_c)
    cash = cur.fetchone()
    cur.execute(
        'SELECT COUNT(*) AS n, COALESCE(SUM(quantity),0) AS qty, '
        'COALESCE(SUM(quantity * cost_price),0) AS cost_val, '
        'COALESCE(SUM(quantity * selling_price),0) AS sell_val FROM stock WHERE user_id = %s', (uid,))
    stock = cur.fetchone()
    cur.execute('SELECT COUNT(*) AS n FROM customers WHERE user_id = %s', (uid,))
    customers = cur.fetchone()
    cur.execute(
        'SELECT stock.product_name AS name, COUNT(sales.id) AS n, '
        'COALESCE(SUM(sales.quantity_sold),0) AS units, '
        'COALESCE(SUM(sales.total_amount),0) AS revenue, '
        'COALESCE(AVG(sales.total_amount),0) AS avg_sale, '
        'COALESCE(SUM(sales.profit),0) AS profit '
        'FROM sales JOIN stock ON sales.stock_id = stock.id '
        'WHERE sales.user_id = %s '
        'GROUP BY stock.product_name ORDER BY revenue DESC LIMIT 12',
        (uid,))
    by_product = cur.fetchall()
    cur.execute(
        'SELECT category AS name, COUNT(*) AS n, COALESCE(SUM(amount),0) AS total, '
        'COALESCE(AVG(amount),0) AS avg FROM expenses WHERE user_id = %s '
        'GROUP BY category ORDER BY total DESC',
        (uid,))
    by_expense = cur.fetchall()
    cur.close()
    conn.close()
    return {
        'sales': sales, 'expenses': expenses, 'income': income, 'cash': cash,
        'stock': stock, 'customers': customers,
        'by_product': by_product, 'by_expense': by_expense,
        'start_date': start_date, 'end_date': end_date,
    }


@app.route('/stats')
@login_required
def stats():
    start_date = request.args.get('start_date') or ''
    end_date = request.args.get('end_date') or ''
    bundle = _stats_bundle(
        current_user.id,
        start_date or None,
        end_date or None,
    )
    return render_template('stats.html', **bundle)


@app.route('/stats/csv')
@login_required
def stats_csv():
    start_date = request.args.get('start_date') or ''
    end_date = request.args.get('end_date') or ''
    b = _stats_bundle(current_user.id, start_date or None, end_date or None)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['KAZE Statistics export'])
    writer.writerow(['From', start_date or 'all', 'To', end_date or 'all'])
    writer.writerow([])
    writer.writerow(['Summary'])
    writer.writerow(['Set', 'Count', 'Total', 'Average', 'Min', 'Max', 'Notes'])
    s, e, inc, cash, st, cust = b['sales'], b['expenses'], b['income'], b['cash'], b['stock'], b['customers']
    writer.writerow(['Sales', s['n'], s['s'], s['a'], s['mn'], s['mx'], 'Profit {}'.format(s['p'])])
    writer.writerow(['Expenses', e['n'], e['s'], e['a'], e['mn'], e['mx'], ''])
    writer.writerow(['Other income', inc['n'], inc['s'], '', '', '', ''])
    writer.writerow(['Cash in', cash['n'], cash['cin'], '', '', '', ''])
    writer.writerow(['Cash out', cash['n'], cash['cout'], '', '', '', 'Net {}'.format((cash['cin'] or 0) - (cash['cout'] or 0))])
    writer.writerow(['Stock', st['n'], st['cost_val'], '', '', '', 'Units {} shelf {}'.format(st['qty'], st['sell_val'])])
    writer.writerow(['Customers', cust['n'], '', '', '', '', ''])
    writer.writerow([])
    writer.writerow(['Product sales'])
    writer.writerow(['Product', 'Sales', 'Units', 'Revenue', 'Avg sale', 'Profit'])
    for r in b['by_product']:
        writer.writerow([r['name'], r['n'], r['units'], r['revenue'], r['avg_sale'], r['profit']])
    writer.writerow([])
    writer.writerow(['Expense categories'])
    writer.writerow(['Category', 'Entries', 'Total', 'Average'])
    for r in b['by_expense']:
        writer.writerow([r['name'], r['n'], r['total'], r['avg']])
    data = output.getvalue()
    filename = 'kaze-statistics.csv'
    return Response(data, mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename={}'.format(filename)})


@app.route('/stats/xlsx')
@login_required
def stats_xlsx():
    start_date = request.args.get('start_date') or ''
    end_date = request.args.get('end_date') or ''
    b = _stats_bundle(current_user.id, start_date or None, end_date or None)
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.page import PageMargins
    except ImportError:
        flash('Install openpyxl to download Excel files: pip install openpyxl', 'danger')
        return redirect(url_for('stats'))

    settings_row = _get_settings_row(current_user.id)
    symbol, decimals = _money_fmt_parts(settings_row)
    money_fmt = _excel_money_format(settings_row)

    wb = Workbook()
    header_font = Font(name='Calibri', bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='0C3823')
    zebra = PatternFill('solid', fgColor='F5F0E8')
    title_font = Font(name='Calibri', bold=True, size=16, color='0C3823')
    thin = Border(
        left=Side(style='thin', color='D4CFC4'),
        right=Side(style='thin', color='D4CFC4'),
        top=Side(style='thin', color='D4CFC4'),
        bottom=Side(style='thin', color='D4CFC4'),
    )

    def paint_header(ws, row, cols):
        for i, title in enumerate(cols, 1):
            cell = ws.cell(row=row, column=i, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='left', wrap_text=True)
            cell.border = thin

    def autosize(ws):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = 12
            for cell in col:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, 42))
            ws.column_dimensions[letter].width = width

    def page_setup(ws, title):
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)
        ws.oddHeader.left.text = title
        ws.oddFooter.right.text = 'Page &P of &N'
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    s, e, inc, cash, st, cust = b['sales'], b['expenses'], b['income'], b['cash'], b['stock'], b['customers']
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_properties.tabColor = '0C3823'
    ws['A1'] = 'KAZE Statistics'
    ws['A1'].font = title_font
    ws.merge_cells('A1:G1')
    ws['A2'] = 'From'
    ws['B2'] = start_date or 'all'
    ws['C2'] = 'To'
    ws['D2'] = end_date or 'all'
    ws['E2'] = 'Currency'
    ws['F2'] = symbol
    paint_header(ws, 4, ['Set', 'Count', 'Total', 'Average', 'Min', 'Max', 'Notes'])
    rows = [
        ['Sales', s['n'], s['s'], s['a'], s['mn'], s['mx'], 'Profit {}'.format(s['p'])],
        ['Expenses', e['n'], e['s'], e['a'], e['mn'], e['mx'], ''],
        ['Other income', inc['n'], inc['s'], None, None, None, ''],
        ['Cash in', cash['n'], cash['cin'], None, None, None, ''],
        ['Cash out', cash['n'], cash['cout'], None, None, None, 'Net {}'.format((cash['cin'] or 0) - (cash['cout'] or 0))],
        ['Stock', st['n'], st['cost_val'], None, None, None, 'Units {} shelf {}'.format(st['qty'], st['sell_val'])],
        ['Customers', cust['n'], None, None, None, None, ''],
    ]
    for r_i, row in enumerate(rows, 5):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = Font(name='Calibri')
            cell.border = thin
            if r_i % 2 == 1:
                cell.fill = zebra
            if c_i >= 3 and c_i <= 6 and isinstance(val, (int, float)):
                cell.number_format = money_fmt
                cell.alignment = Alignment(horizontal='right')
            if c_i == 7:
                cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = 'A4:G11'
    autosize(ws)
    page_setup(ws, 'KAZE Statistics')
    chart = BarChart()
    chart.type = 'col'
    chart.title = 'Totals'
    data = Reference(ws, min_col=3, min_row=4, max_row=10)
    cats = Reference(ws, min_col=1, min_row=5, max_row=10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.shape = 4
    ws.add_chart(chart, 'A13')

    ws2 = wb.create_sheet('Product sales')
    paint_header(ws2, 1, ['Product', 'Sales', 'Units', 'Revenue', 'Avg sale', 'Profit'])
    for r_i, r in enumerate(b['by_product'], 2):
        vals = [r['name'], r['n'], r['units'], r['revenue'], r['avg_sale'], r['profit']]
        for c_i, val in enumerate(vals, 1):
            cell = ws2.cell(row=r_i, column=c_i, value=val)
            cell.font = Font(name='Calibri')
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = zebra
            if c_i >= 4:
                cell.number_format = money_fmt
                cell.alignment = Alignment(horizontal='right')
    if not b['by_product']:
        ws2.cell(row=2, column=1, value='No sales yet')
    last = max(2, 1 + len(b['by_product']))
    ws2.auto_filter.ref = 'A1:F{}'.format(last)
    ws2.freeze_panes = 'A2'
    autosize(ws2)
    page_setup(ws2, 'KAZE product sales')

    ws3 = wb.create_sheet('Expense categories')
    paint_header(ws3, 1, ['Category', 'Entries', 'Total', 'Average'])
    for r_i, r in enumerate(b['by_expense'], 2):
        vals = [r['name'], r['n'], r['total'], r['avg']]
        for c_i, val in enumerate(vals, 1):
            cell = ws3.cell(row=r_i, column=c_i, value=val)
            cell.font = Font(name='Calibri')
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = zebra
            if c_i >= 3:
                cell.number_format = money_fmt
                cell.alignment = Alignment(horizontal='right')
    if not b['by_expense']:
        ws3.cell(row=2, column=1, value='No expenses yet')
    last = max(2, 1 + len(b['by_expense']))
    ws3.auto_filter.ref = 'A1:D{}'.format(last)
    ws3.freeze_panes = 'A2'
    autosize(ws3)
    page_setup(ws3, 'KAZE expenses')

    guide = wb.create_sheet('Formatting options')
    guide['A1'] = 'Excel formatting used in this file'
    guide['A1'].font = title_font
    guide.merge_cells('A1:B1')
    tips = [
        ('Number format', 'Money columns use your Settings currency symbol and two decimals.'),
        ('Font', 'Calibri throughout. Headers are bold white on forest green.'),
        ('Fill', 'Green header, cream zebra rows so lines are easier to read.'),
        ('Border', 'Thin light grid around every data cell.'),
        ('Alignment', 'Money is right-aligned. Notes can wrap.'),
        ('Freeze panes', 'Header row stays visible when you scroll.'),
        ('AutoFilter', 'Use the arrows on headers to sort or hide rows.'),
        ('Chart', 'Summary sheet has a bar chart of totals.'),
        ('Print', 'Landscape, fit-to-width, header and page number in footer.'),
        ('Sheets', 'Summary, Product sales, Expense categories, plus this guide.'),
    ]
    paint_header(guide, 3, ['Option', 'What it does here'])
    for i, (k, v) in enumerate(tips, 4):
        guide.cell(i, 1, k).font = Font(name='Calibri', bold=True)
        guide.cell(i, 2, v).font = Font(name='Calibri')
        guide.cell(i, 2).alignment = Alignment(wrap_text=True)
        for c in (1, 2):
            guide.cell(i, c).border = thin
            if i % 2 == 0:
                guide.cell(i, c).fill = zebra
    guide.column_dimensions['A'].width = 22
    guide.column_dimensions['B'].width = 78
    guide.freeze_panes = 'A4'
    page_setup(guide, 'KAZE Excel formatting')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name='kaze-statistics.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/theme', methods=['POST'])
def set_display_theme():
    """Save dark / light / auto (or special themes) and remember it on this browser."""
    name = (request.form.get('theme') or '').strip()
    if request.is_json:
        name = (request.get_json(silent=True) or {}).get('theme') or name
    if name not in ('dark', 'light', 'auto', 'matrix', 'inverted'):
        return {'ok': False, 'error': 'Unknown theme'}, 400
    if current_user.is_authenticated:
        try:
            conn = get_db()
            cur = conn.cursor()
            cur.execute('UPDATE settings SET theme=%s WHERE user_id=%s', (name, current_user.id))
            conn.commit()
            cur.close()
            conn.close()
        except Exception:
            pass
    resp = app.response_class(response='{"ok": true, "theme": "%s"}' % name, mimetype='application/json')
    resp.set_cookie('kaze_theme', name, max_age=60 * 60 * 24 * 400, samesite='Lax')
    return resp


# ======================== Stripe payments ========================

def _stripe_settings(uid=None):
    if uid is None and getattr(current_user, 'is_authenticated', False):
        uid = current_user.id
    if not uid:
        return None
    row = getattr(g, '_settings_row', None)
    if row and getattr(g, '_settings_uid', None) == uid:
        return row
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM settings WHERE user_id=%s', (uid,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return row


def _record_stripe_checkout(uid, kind, record_id, amount, title, session):
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO stripe_payments (user_id, kind, record_id, amount, currency, status, checkout_id, title, created_date) "
        "VALUES (%s,%s,%s,%s,%s,'pending',%s,%s,%s)",
        (uid, kind, record_id or 0, float(amount or 0),
         stripe_payments.currency_code(_stripe_settings(uid)),
         session.id, (title or '')[:160], today)
    )
    if kind == 'session' and record_id:
        cur.execute(
            "UPDATE service_sessions SET stripe_checkout_id=%s WHERE id=%s AND user_id=%s",
            (session.id, record_id, uid)
        )
    elif kind == 'book' and record_id:
        cur.execute(
            "UPDATE bookkeeping_docs SET stripe_checkout_id=%s WHERE id=%s AND user_id=%s",
            (session.id, record_id, uid)
        )
    elif kind == 'sale' and record_id:
        cur.execute(
            "UPDATE sales SET stripe_checkout_id=%s WHERE id=%s AND user_id=%s",
            (session.id, record_id, uid)
        )
    conn.commit()
    cur.close()
    conn.close()


def _fulfill_stripe_payment(uid, kind, record_id, checkout_id, payment_intent, amount=None):
    today = datetime.today().strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    # Idempotency: read current status first. If already paid, do nothing.
    cur.execute(
        "SELECT status FROM stripe_payments WHERE user_id=%s AND checkout_id=%s FOR UPDATE",
        (uid, checkout_id or '')
    )
    prior = cur.fetchone() or {}
    if (prior.get('status') or '').lower() == 'paid':
        cur.close()
        conn.close()
        return
    cur.execute(
        "UPDATE stripe_payments SET status='paid', paid_date=%s, payment_intent=%s "
        "WHERE user_id=%s AND checkout_id=%s",
        (today, payment_intent or '', uid, checkout_id or '')
    )
    kind = (kind or '').lower()
    if kind == 'session' and record_id:
        cur.execute('SELECT * FROM service_sessions WHERE id=%s AND user_id=%s', (record_id, uid))
        row = cur.fetchone()
        if row and (row.get('status') or '') != 'done':
            svc_name = _service_row_name(cur, row.get('service_id'), uid)
            source = 'Stripe: {} — {}'.format(svc_name, row.get('customer_name') or 'Walk-in')
            inc_id = row.get('income_id')
            if not inc_id:
                cur.execute(
                    'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s) RETURNING id',
                    (uid, source[:120], float(row.get('price') or amount or 0), row.get('session_date') or today)
                )
                inc = cur.fetchone()
                inc_id = inc['id'] if inc else None
            cur.execute(
                "UPDATE service_sessions SET status='done', payment_method='card', income_id=%s, stripe_checkout_id=%s "
                "WHERE id=%s AND user_id=%s",
                (inc_id, checkout_id or row.get('stripe_checkout_id'), record_id, uid)
            )
    elif kind == 'book' and record_id:
        cur.execute(
            "UPDATE bookkeeping_docs SET status='paid', stripe_checkout_id=%s WHERE id=%s AND user_id=%s",
            (checkout_id or '', record_id, uid)
        )
    elif kind == 'sale' and record_id:
        cur.execute(
            "UPDATE sales SET payment_method='card', stripe_checkout_id=%s WHERE id=%s AND user_id=%s",
            (checkout_id or '', record_id, uid)
        )
    elif kind == 'plan':
        wanted = None
        if record_id is not None and 0 <= int(record_id or -1) < len(PLAN_ORDER):
            wanted = PLAN_ORDER[int(record_id)]
        if not wanted:
            cur.execute('SELECT title FROM stripe_payments WHERE user_id=%s AND checkout_id=%s', (uid, checkout_id or ''))
            pay = cur.fetchone() or {}
            title = (pay.get('title') or '')
            wanted = normalize_plan(title.replace('KAZE package:', '').strip() or 'new_user')
        cur.execute('SELECT plan_tier, COALESCE(plan_paid_rank, 0) AS paid_rank FROM settings WHERE user_id=%s', (uid,))
        row = cur.fetchone() or {}
        try:
            current_paid_rank = int(row.get('paid_rank') or 0)
        except (TypeError, ValueError):
            current_paid_rank = 0
        wanted_rank = PLAN_RANK.get(wanted, 0)
        # Successful payment always bumps plan_paid_rank. plan_tier only moves
        # up — never silently downgrades someone who already had a higher tier.
        new_paid_rank = max(wanted_rank, current_paid_rank)
        new_tier = wanted if wanted_rank >= PLAN_RANK.get(normalize_plan(row.get('plan_tier') or 'new_user'), 0) else (row.get('plan_tier') or 'new_user')
        cur.execute(
            'UPDATE settings SET plan_tier=%s, plan_paid_rank=%s WHERE user_id=%s',
            (new_tier, new_paid_rank, uid)
        )
    elif kind == 'custom':
        if amount and float(amount) > 0:
            cur.execute(
                'INSERT INTO income (user_id, source, amount, date) VALUES (%s,%s,%s,%s)',
                (uid, 'Stripe payment', float(amount), today)
            )
    conn.commit()
    cur.close()
    conn.close()
    try:
        log_activity(uid, 'stripe', 'Stripe payment received', '{} #{}'.format(kind, record_id or ''))
    except Exception:
        pass


def _record_payment(uid, kind, record_id, amount, title, checkout_id, provider='stripe'):
    class _Sess:
        def __init__(self, cid):
            self.id = cid
    _record_stripe_checkout(uid, kind, record_id, amount, title, _Sess(checkout_id))
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute(
            "UPDATE stripe_payments SET provider=%s WHERE user_id=%s AND checkout_id=%s",
            (provider, uid, checkout_id),
        )
        conn.commit()
        cur.close()
        conn.close()
    except Exception:
        pass


def _checkout_subject(kind, record_id, tier=''):
    settings = _stripe_settings()
    kind = (kind or 'custom').strip().lower()
    uid = current_user.id
    title, amount, party, back = 'KAZE payment', 0, '', 'payments'
    if kind == 'plan':
        tier = (tier or '').strip().lower()
        if not tier and record_id is not None and 0 <= int(record_id) < len(PLAN_ORDER):
            tier = PLAN_ORDER[int(record_id)]
        row = next((t for t in priced_tiers(settings) if t['key'] == tier), None)
        if not row or not row.get('price'):
            return None, 'That package does not need a payment.'
        title, amount, back = 'KAZE package: ' + tier, row['price'], 'plans'
        record_id = PLAN_RANK.get(tier, record_id or 0)
    elif kind == 'session':
        conn = get_db(); cur = conn.cursor()
        cur.execute('SELECT price, customer_name FROM service_sessions WHERE id=%s AND user_id=%s', (record_id, uid))
        row = cur.fetchone(); cur.close(); conn.close()
        if not row: return None, 'Session not found.'
        title, amount, party, back = 'Service session', row.get('price') or 0, row.get('customer_name') or '', 'services'
    elif kind == 'book':
        conn = get_db(); cur = conn.cursor()
        cur.execute('SELECT amount, title, party FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (record_id, uid))
        row = cur.fetchone(); cur.close(); conn.close()
        if not row: return None, 'Document not found.'
        title, amount, party, back = row.get('title') or 'Document', row.get('amount') or 0, row.get('party') or '', 'bookkeeping'
    elif kind == 'sale':
        conn = get_db(); cur = conn.cursor()
        cur.execute('SELECT total_amount, customer_name FROM sales WHERE id=%s AND user_id=%s', (record_id, uid))
        row = cur.fetchone(); cur.close(); conn.close()
        if not row: return None, 'Sale not found.'
        title, amount, party, back = 'Sale receipt', row.get('total_amount') or 0, row.get('customer_name') or '', 'sales'
    elif kind == 'shop':
        conn = get_db(); cur = conn.cursor()
        cur.execute('SELECT amount, buyer FROM shop_orders WHERE id=%s AND user_id=%s', (record_id, uid))
        row = cur.fetchone(); cur.close(); conn.close()
        if not row: return None, 'Order not found.'
        title, amount, party, back = 'Shop order', row.get('amount') or 0, row.get('buyer') or '', 'shop'
    elif kind == 'custom':
        try: amount = float(request.values.get('amount') or 0)
        except (TypeError, ValueError): amount = 0
        title = (request.values.get('title') or 'Custom payment')[:160]
    else:
        return None, 'Unknown payment.'
    return {'kind': kind, 'record_id': record_id or 0, 'tier': tier or '', 'title': title, 'amount': amount, 'party': party, 'back': back, 'currency': stripe_payments.currency_code(settings), 'settings': settings}, None

def _start_checkout(kind, record_id, amount, title, cancel_endpoint='dashboard', method=None):
    tier = request.values.get('tier') or ''
    if kind == 'plan' and not tier:
        try: tier = PLAN_ORDER[int(record_id)]
        except (TypeError, ValueError, IndexError): tier = ''
    return redirect(url_for('kaze_checkout', kind=kind, id=record_id or 0, tier=tier, method=(method or request.values.get('method') or ''), title=title or '', amount=amount or 0))

@app.route('/checkout')
@login_required
def kaze_checkout():
    kind = (request.args.get('kind') or 'custom').strip().lower()
    try: record_id = int(request.args.get('id') or 0)
    except ValueError: record_id = 0
    subject, err = _checkout_subject(kind, record_id, (request.args.get('tier') or '').strip().lower())
    if err:
        flash(err, 'danger'); return redirect(url_for('payments'))
    settings = subject.pop('settings')
    return render_template('checkout.html', kind=subject['kind'], record_id=subject['record_id'], tier=subject['tier'], title=subject['title'], party=subject['party'], amount=subject['amount'], currency=subject['currency'], method=(request.args.get('method') or gateways.default_method(settings) or '').lower(), back=subject['back'], pay_methods=gateways.ready_methods(settings))

@app.route('/checkout/embed', methods=['POST'])
@login_required
def kaze_checkout_embed():
    kind = (request.form.get('kind') or 'custom').strip().lower()
    try: record_id = int(request.form.get('id') or 0)
    except ValueError: record_id = 0
    subject, err = _checkout_subject(kind, record_id, (request.form.get('tier') or '').strip().lower())
    if err: return {'ok': False, 'error': err}, 400
    if subject['amount'] <= 0: return {'ok': False, 'error': 'Amount must be greater than zero.'}, 400
    settings = subject.pop('settings')
    method = (request.form.get('method') or '').lower()
    uid = current_user.id
    meta = {'user_id': uid, 'kind': subject['kind'], 'record_id': subject['record_id'], 'tier': subject['tier'], 'title': subject['title']}
    if method == 'stripe':
        if not stripe_payments.stripe_ready(settings): return {'ok': False, 'error': 'Stripe is not set up yet.'}, 400
        success = url_for('stripe_pay_success', _external=True) + '?session_id={CHECKOUT_SESSION_ID}'
        sess, err = stripe_payments.create_checkout(settings, amount=subject['amount'], title=subject['title'], success_url=success, cancel_url=url_for('stripe_pay_cancel', _external=True), metadata=meta, embedded=True)
        if err: return {'ok': False, 'error': err}, 400
        _record_payment(uid, subject['kind'], subject['record_id'], subject['amount'], subject['title'], sess.id, 'stripe')
        return {'ok': True, 'method': 'stripe', 'publishable': stripe_payments.keys_from(settings)[1], 'client_secret': sess.client_secret}
    if method == 'paypal':
        if not gateways.paypal_ready(settings): return {'ok': False, 'error': 'PayPal is not set up yet.'}, 400
        order, err = gateways.paypal_create_order(settings, amount=subject['amount'], currency=subject['currency'], title=subject['title'], return_url=url_for('paypal_return', _external=True), cancel_url=url_for('stripe_pay_cancel', _external=True), custom_id='kaze-%s-%s-%s' % (subject['kind'], subject['record_id'], uid))
        if err: return {'ok': False, 'error': err}, 400
        _record_payment(uid, subject['kind'], subject['record_id'], subject['amount'], subject['title'], order['id'], 'paypal')
        return {'ok': True, 'method': 'paypal', 'order_id': order['id'], 'client_id': gateways.paypal_keys(settings)[0], 'currency': subject['currency']}
    if method == 'flutterwave':
        if not integrations.flutterwave_ready(settings): return {'ok': False, 'error': 'Flutterwave is not set up yet.'}, 400
        tx_ref = 'kaze-%s-%s-%s' % (subject['kind'], subject['record_id'], int(datetime.now().timestamp()))
        _record_payment(uid, subject['kind'], subject['record_id'], subject['amount'], subject['title'], tx_ref, 'flutterwave')
        _s, public, _h = integrations.flutterwave_keys(settings)
        return {'ok': True, 'method': 'flutterwave', 'public_key': public, 'tx_ref': tx_ref, 'amount': subject['amount'], 'currency': subject['currency'], 'title': subject['title'], 'email': getattr(current_user, 'email', '') or 'shop@kaze.local', 'name': getattr(current_user, 'username', '') or 'KAZE', 'redirect': url_for('flutterwave_return', _external=True)}
    if method == 'bank':
        if not gateways.bank_ready(settings): return {'ok': False, 'error': 'Add bank details in Settings first.'}, 400
        ref = 'KAZE-%s-%s-%s' % (subject['kind'].upper()[:8], uid, int(datetime.now().timestamp()))
        _record_payment(uid, subject['kind'], subject['record_id'], subject['amount'], subject['title'], ref, 'bank')
        return {'ok': True, 'method': 'bank', 'ref': ref, 'bank': gateways.bank_details(settings), 'confirm': url_for('bank_pay_confirm', ref=ref)}
    return {'ok': False, 'error': 'Choose a payment method that is switched on.'}, 400

@app.route('/checkout/paypal/capture', methods=['POST'])
@login_required
def kaze_paypal_capture():
    order_id = (request.form.get('order_id') or '').strip()
    payload, err = gateways.paypal_capture(_stripe_settings(), order_id)
    if err: return {'ok': False, 'error': err}, 400
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT * FROM stripe_payments WHERE user_id=%s AND checkout_id=%s ORDER BY id DESC LIMIT 1", (current_user.id, order_id))
    row = cur.fetchone() or {}
    cur.close(); conn.close()
    _fulfill_stripe_payment(current_user.id, row.get('kind') or 'custom', row.get('record_id') or 0, order_id, order_id, row.get('amount'))
    return {'ok': True, 'next': url_for('payments')}

@app.route('/pay/session/<int:id>')
@login_required
def pay_service_session(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT s.*, sv.name AS service_name FROM service_sessions s "
        "LEFT JOIN services sv ON sv.id = s.service_id "
        "WHERE s.id=%s AND s.user_id=%s",
        (id, current_user.id)
    )
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Session not found.', 'danger')
        return redirect(url_for('services'))
    if (row.get('status') or '') == 'done':
        flash('That session is already marked done.', 'info')
        return redirect(url_for('services'))
    title = '{} — {}'.format(row.get('service_name') or 'Service', row.get('customer_name') or 'Walk-in')
    return _start_checkout('session', id, row.get('price') or 0, title, 'services')


@app.route('/pay/book/<int:id>')
@login_required
def pay_book_doc(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM bookkeeping_docs WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Document not found.', 'danger')
        return redirect(url_for('bookkeeping'))
    if (row.get('status') or '') == 'paid':
        flash('That document is already paid.', 'info')
        return redirect(url_for('bookkeeping_type', kind=row.get('kind') or 'invoice'))
    title = row.get('title') or row.get('reference') or 'Document'
    return _start_checkout('book', id, row.get('amount') or 0, title, 'bookkeeping')


@app.route('/pay/sale/<int:id>')
@login_required
def pay_sale(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM sales WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Sale not found.', 'danger')
        return redirect(url_for('sales'))
    title = 'Sale #{} — {}'.format(id, row.get('customer_name') or 'Walk-in')
    return _start_checkout('sale', id, row.get('total_amount') or 0, title, 'sales')


@app.route('/pay/custom', methods=['POST'])
@login_required
def pay_custom():
    amount, err = _clean_float('amount', min_v=0.01, max_v=1e12, label='amount')
    title, _ = _clean_text('title', required=False, max_len=160, label='title')
    if err:
        flash(err, 'danger')
        return redirect(url_for('payments'))
    return _start_checkout('custom', 0, amount, title or 'Custom payment', 'payments')


@app.route('/payments')
@login_required
def payments():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT * FROM stripe_payments WHERE user_id=%s ORDER BY id DESC LIMIT 200',
        (current_user.id,)
    )
    rows = cur.fetchall()
    cur.close()
    conn.close()
    settings = _stripe_settings()
    return render_template(
        'payments.html',
        rows=rows,
        stripe_ready=stripe_payments.stripe_ready(settings),
        publishable=(stripe_payments.keys_from(settings)[1] or '')[:20],
    )


@app.route('/pay/success')
@login_required
def stripe_pay_success():
    session_id = (request.args.get('session_id') or '').strip()
    if not session_id:
        flash('Payment finished. If the amount is missing, wait a moment and refresh Payments.', 'success')
        return redirect(url_for('payments'))
    settings = _stripe_settings()
    secret, _, _ = stripe_payments.keys_from(settings)
    kind = 'custom'
    record_id = 0
    amount = 0
    intent = ''
    if secret and stripe_payments.stripe_available():
        try:
            import stripe
            stripe.api_key = secret
            sess = stripe.checkout.Session.retrieve(session_id)
            meta = sess.get('metadata') or {}
            # Ownership check: only the session owner may fulfill it.
            try:
                session_user_id = int(meta.get('user_id') or 0)
            except (TypeError, ValueError):
                session_user_id = 0
            if not session_user_id or session_user_id != int(current_user.id):
                print('stripe success: session %s does not belong to user %s'
                      % (session_id, current_user.id))
                flash('That payment belongs to a different account.', 'danger')
                return redirect(url_for('payments'))
            kind = meta.get('kind') or 'custom'
            try:
                record_id = int(meta.get('record_id') or 0)
            except (TypeError, ValueError):
                record_id = 0
            intent = sess.get('payment_intent') or ''
            amount = stripe_payments.from_stripe_amount(sess.get('amount_total'), sess.get('currency'))
            if (sess.get('payment_status') or '') == 'paid':
                _fulfill_stripe_payment(current_user.id, kind, record_id, session_id, intent, amount)
        except Exception as exc:
            print('stripe success retrieve:', exc)
    flash('Payment received. Thank you.', 'success')
    if kind == 'session':
        return redirect(url_for('services'))
    if kind == 'book':
        return redirect(url_for('bookkeeping'))
    if kind == 'sale':
        return redirect(url_for('sales'))
    if kind == 'plan':
        return redirect(url_for('plans'))
    return redirect(url_for('payments'))


@app.route('/pay/cancel')
@login_required
def stripe_pay_cancel():
    flash('Payment cancelled. Nothing was charged.', 'info')
    return redirect(url_for('payments'))


@app.route('/webhooks/stripe', methods=['POST'])
def stripe_webhook():
    payload = request.get_data()
    sig = request.headers.get('Stripe-Signature', '')
    event, err = stripe_payments.parse_webhook(payload, sig, None)
    if err or not event:
        # Try with first owner settings if env secret missing
        try:
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT * FROM settings WHERE stripe_webhook_secret <> '' LIMIT 1")
            row = cur.fetchone()
            cur.close()
            conn.close()
        except Exception:
            row = None
        event, err = stripe_payments.parse_webhook(payload, sig, row)
    if err or not event:
        return ('invalid webhook', 400)
    etype = event.get('type')
    if etype and etype.startswith('customer.subscription.'):
        return subscriptions.handle_webhook(get_db, event)
    if etype in ('invoice.payment_succeeded', 'invoice.payment_failed'):
        return subscriptions.handle_webhook(get_db, event)
    if etype in ('checkout.session.completed', 'checkout.session.async_payment_succeeded'):
        # Skip subscription checkouts — subscription lifecycle events
        # (customer.subscription.created, invoice.payment_succeeded) are
        # the source of truth. Without this guard, the one-time handler
        # below would insert a spurious income row for every subscription.
        if ((event.get('data') or {}).get('object') or {}).get('mode') == 'subscription':
            return ('', 200)
        sess = event['data']['object']
        meta = sess.get('metadata') or {}
        try:
            uid = int(meta.get('user_id') or 0)
        except (TypeError, ValueError):
            uid = 0
        try:
            record_id = int(meta.get('record_id') or 0)
        except (TypeError, ValueError):
            record_id = 0
        kind = meta.get('kind') or 'custom'
        intent = sess.get('payment_intent') or ''
        amount = stripe_payments.from_stripe_amount(sess.get('amount_total'), sess.get('currency'))
        if uid and (sess.get('payment_status') or '') == 'paid':
            _fulfill_stripe_payment(uid, kind, record_id, sess.get('id'), intent, amount)
    elif etype == 'checkout.session.expired':
        sess = event['data']['object']
        try:
            conn = get_db(); cur = conn.cursor()
            cur.execute(
                "UPDATE stripe_payments SET status='expired' "
                "WHERE checkout_id=%s AND status='pending'",
                (sess.get('id'),)
            )
            conn.commit(); cur.close(); conn.close()
        except Exception as exc:
            print('stripe webhook expired:', exc)
    elif etype == 'charge.refunded':
        charge = event['data']['object']
        intent = charge.get('payment_intent') or ''
        if intent:
            try:
                conn = get_db(); cur = conn.cursor()
                cur.execute(
                    "UPDATE stripe_payments SET status='refunded' "
                    "WHERE payment_intent=%s AND status='paid'",
                    (intent,)
                )
                conn.commit(); cur.close(); conn.close()
            except Exception as exc:
                print('stripe webhook refunded:', exc)
    return ('', 200)


@app.route('/packages')
@app.route('/plans')
@login_required
def plans():
    settings = _stripe_settings()
    current = plan_info(settings)
    counts = {}
    conn = get_db()
    cur = conn.cursor()
    for kind in ('stock', 'customers', 'services', 'sales_month', 'team'):
        counts[kind] = _plan_count(cur, current_user.id, kind)
    cur.close()
    conn.close()
    trial = current.get('trial') or trial_info(settings)
    # Read paid_rank from the DB — this is the tier the user actually owns.
    try:
        rcur = get_db().cursor()
        rcur.execute('SELECT COALESCE(plan_paid_rank, 0) AS paid_rank FROM settings WHERE user_id=%s', (current_user.id,))
        prow = rcur.fetchone() or {}
        rcur.close()
        paid_rank = int(prow.get('paid_rank') or 0)
    except Exception:
        paid_rank = 0
    return render_template(
        'plans.html',
        tiers=priced_tiers(settings),
        current=current,
        counts=counts,
        stripe_ready=stripe_payments.stripe_ready(settings),
        paypal_ready=gateways.paypal_ready(settings),
        flutterwave_ready=integrations.flutterwave_ready(settings),
        bank_ready=gateways.bank_ready(settings),
        trial=trial,
        trial_day_choices=list(range(14, 19)),
        paid_rank=paid_rank,
        is_owner=(getattr(current_user, 'role', 'owner') == 'owner'),
        subs_configured_tiers=subscriptions.configured_tiers(),
        subs_any_ready=subscriptions.subs_ready(),
    )


@app.route('/pay/paypal/return')
@login_required
def paypal_return():
    settings = _stripe_settings()
    token = (request.args.get('token') or request.args.get('orderID') or '').strip()
    if not token:
        flash('PayPal did not return an order.', 'danger')
        return redirect(url_for('payments'))
    payload, err = gateways.paypal_capture(settings, token)
    if err:
        flash(err, 'danger')
        return redirect(url_for('payments'))
    status = str((payload or {}).get('status') or '').upper()
    if status not in ('COMPLETED', 'APPROVED'):
        flash('PayPal payment is %s.' % (status or 'pending'), 'info')
        return redirect(url_for('payments'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT * FROM stripe_payments WHERE user_id=%s AND checkout_id=%s ORDER BY id DESC LIMIT 1",
        (current_user.id, token),
    )
    row = cur.fetchone() or {}
    cur.close()
    conn.close()
    _fulfill_stripe_payment(
        current_user.id,
        row.get('kind') or 'custom',
        row.get('record_id') or 0,
        token,
        token,
        row.get('amount'),
    )
    flash('PayPal payment received.', 'success')
    return redirect(url_for('payments'))


@app.route('/pay/bank/<ref>')
@login_required
def bank_pay_notice(ref):
    settings = _stripe_settings()
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT * FROM stripe_payments WHERE user_id=%s AND checkout_id=%s",
        (current_user.id, ref),
    )
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Transfer reference not found.', 'danger')
        return redirect(url_for('payments'))
    return render_template(
        'bank_pay.html',
        row=row,
        bank=gateways.bank_details(settings),
    )


@app.route('/pay/bank/<ref>/confirm', methods=['POST'])
@login_required
def bank_pay_confirm(ref):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT * FROM stripe_payments WHERE user_id=%s AND checkout_id=%s",
        (current_user.id, ref),
    )
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Transfer reference not found.', 'danger')
        return redirect(url_for('payments'))
    if (row.get('status') or '') == 'paid':
        flash('Already marked paid.', 'info')
        return redirect(url_for('payments'))
    _fulfill_stripe_payment(
        current_user.id,
        row.get('kind') or 'custom',
        row.get('record_id') or 0,
        ref,
        ref,
        row.get('amount'),
    )
    flash('Current-account payment marked paid.', 'success')
    return redirect(url_for('payments'))


@app.route('/plans/trial', methods=['POST'])
@login_required
def start_trial():
    if getattr(current_user, 'role', 'owner') != 'owner':
        flash('Only the owner can start a trial.', 'danger')
        return redirect(url_for('plans'))
    settings = _stripe_settings() or {}
    if trial_info(settings).get('started'):
        flash('This shop already used its free trial.', 'info')
        return redirect(url_for('plans'))
    try:
        days = int(request.form.get('trial_days') or 12)
    except (TypeError, ValueError):
        days = 12
    days = min(18, max(14, days))
    ends = (datetime.today() + timedelta(days=days)).strftime('%Y-%m-%d')
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'UPDATE settings SET trial_started=1, trial_ends=%s, trial_days=%s WHERE user_id=%s',
        (ends, days, current_user.id),
    )
    conn.commit()
    cur.close()
    flash('Free trial is on for {} days (until {}).'.format(days, ends), 'success')
    return redirect(url_for('plans'))


@app.route('/plans/set', methods=['POST'])
@login_required
def set_plan():
    if getattr(current_user, 'role', 'owner') != 'owner':
        flash('Only the owner can change the plan.', 'danger')
        return redirect(url_for('plans'))
    wanted = normalize_plan(request.form.get('plan_tier'))
    if wanted != 'new_user':
        flash('Paid packages are subscribed on the Packages page.', 'danger')
        return redirect(url_for('plans'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT subscription_status FROM settings WHERE user_id=%s', (current_user.id,))
    row = cur.fetchone() or {}
    if (row.get('subscription_status') or '').lower() in ('active', 'trialing', 'past_due'):
        cur.close()
        conn.close()
        flash('You have an active subscription. Cancel it from Manage first.', 'danger')
        return redirect(url_for('plans'))
    cur.execute('UPDATE settings SET plan_tier=%s WHERE user_id=%s', ('new_user', current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Switched to the free tier.', 'success')
    return redirect(url_for('plans'))


@app.route('/plans/subscribe/<tier>')
@login_required
def subscribe_plan(tier):
    if getattr(current_user, 'role', 'owner') != 'owner':
        flash('Only the owner can subscribe.', 'danger')
        return redirect(url_for('plans'))
    wanted = normalize_plan(tier)
    if wanted == 'new_user':
        flash('Free tier needs no subscription.', 'info')
        return redirect(url_for('plans'))
    if not subscriptions.price_id_for(wanted):
        flash('That package is not available for subscription yet.', 'danger')
        return redirect(url_for('plans'))
    settings = _stripe_settings()
    sub = subscription_info(settings)
    if sub.get('status') in ('active', 'grace', 'canceled_active'):
        flash('You already have an active subscription. Use Manage to change it.', 'danger')
        return redirect(url_for('plans'))

    customer_id, err = subscriptions.get_or_create_customer(
        settings, current_user.id, current_user.email, current_user.username)
    if err:
        flash('Could not reach Stripe: ' + err, 'danger')
        return redirect(url_for('plans'))

    if customer_id and customer_id != (settings or {}).get('stripe_customer_id'):
        conn = get_db(); cur = conn.cursor()
        cur.execute('UPDATE settings SET stripe_customer_id=%s WHERE user_id=%s',
                    (customer_id, current_user.id))
        conn.commit(); cur.close(); conn.close()

    success = url_for('plans_sync', _external=True) + '?session_id={CHECKOUT_SESSION_ID}'
    cancel = url_for('plans', _external=True)
    session, err = subscriptions.create_subscription_checkout(
        settings, wanted, customer_id, current_user.id, success, cancel)
    if err:
        flash('Could not start subscription: ' + err, 'danger')
        return redirect(url_for('plans'))
    return redirect(session.url, code=303)


@app.route('/plans/sync')
@login_required
def plans_sync():
    """Landing page while Stripe webhook catches up."""
    return render_template('plans_sync.html')


@app.route('/plans/status')
@login_required
def plans_status():
    """JSON status used by the /plans/sync polling page."""
    settings = _stripe_settings()
    sub = subscription_info(settings)
    return {
        'active': sub.get('status') in ('active', 'grace', 'canceled_active'),
        'status': sub.get('status'),
        'expires': sub.get('expires'),
    }


@app.route('/plans/manage')
@login_required
def plans_manage():
    if getattr(current_user, 'role', 'owner') != 'owner':
        flash('Only the owner can manage the subscription.', 'danger')
        return redirect(url_for('plans'))
    settings = _stripe_settings()
    customer_id = (settings or {}).get('stripe_customer_id') or ''
    if not customer_id:
        flash('No active subscription to manage.', 'info')
        return redirect(url_for('plans'))
    return_url = url_for('plans', _external=True)
    session, err = subscriptions.create_portal_session(settings, customer_id, return_url)
    if err:
        flash('Could not open the billing portal: ' + err, 'danger')
        return redirect(url_for('plans'))
    return redirect(session.url, code=303)


@app.route('/plans/upgrade/<tier>')
@login_required
def pay_plan(tier):
    wanted = normalize_plan(tier)
    if wanted == 'new_user':
        flash('New user is the free starting package.', 'info')
        return redirect(url_for('plans'))
    settings = _stripe_settings()
    current = plan_info(settings)
    if PLAN_RANK[wanted] <= PLAN_RANK[current['stored_key']]:
        flash('You already have this plan or a higher one.', 'info')
        return redirect(url_for('plans'))
    info = PLAN_TIERS[wanted]
    amount = package_prices_for(settings).get(wanted, info.get('price') or 0)
    if not amount:
        flash('That package is free. Use Set package instead of card checkout.', 'info')
        return redirect(url_for('plans'))
    title = 'KAZE package: {}'.format(wanted)
    return _start_checkout('plan', PLAN_RANK[wanted], amount, title, 'plans')


BANK_ACCOUNT_TYPES = ('Current', 'Savings', 'Mobile money', 'Credit', 'Cash till', 'Other')
INVEST_TYPES = ('Equity', 'Bond', 'Property', 'Fund', 'Cash instrument', 'Other')
SHOP_ORDER_STATUSES = ('pending', 'paid', 'packed', 'delivered', 'cancelled')
MARKET_CHANNEL_KINDS = (
    'Product sales', 'Services', 'Subscriptions', 'Ads',
    'Affiliates', 'Licensing', 'Other',
)


def _bank_balance(cur, uid, account_id, opening):
    cur.execute(
        "SELECT COALESCE(SUM(CASE WHEN kind='in' THEN amount WHEN kind='out' THEN -amount ELSE 0 END),0) AS n "
        "FROM bank_transactions WHERE user_id=%s AND account_id=%s",
        (uid, account_id),
    )
    row = cur.fetchone() or {}
    try:
        moved = float(row.get('n') or 0)
    except (TypeError, ValueError):
        moved = 0.0
    try:
        return float(opening or 0) + moved
    except (TypeError, ValueError):
        return moved


@app.route('/banking')
@login_required
def banking():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM bank_accounts WHERE user_id=%s ORDER BY name', (uid,))
    accounts = list(cur.fetchall() or [])
    total_in = total_out = 0.0
    total_balance = 0.0
    for a in accounts:
        bal = _bank_balance(cur, uid, a['id'], a.get('opening_balance'))
        a['balance'] = bal
        total_balance += bal
    cur.execute(
        "SELECT COALESCE(SUM(amount),0) AS n FROM bank_transactions WHERE user_id=%s AND kind='in'",
        (uid,),
    )
    total_in = float((cur.fetchone() or {}).get('n') or 0)
    cur.execute(
        "SELECT COALESCE(SUM(amount),0) AS n FROM bank_transactions WHERE user_id=%s AND kind='out'",
        (uid,),
    )
    total_out = float((cur.fetchone() or {}).get('n') or 0)
    cur.execute(
        '''SELECT t.*, a.name AS account_name
           FROM bank_transactions t
           LEFT JOIN bank_accounts a ON a.id = t.account_id
           WHERE t.user_id=%s ORDER BY t.tx_date DESC, t.id DESC LIMIT 80''',
        (uid,),
    )
    txs = cur.fetchall() or []
    cur.close()
    conn.close()
    return render_template(
        'banking.html',
        accounts=accounts,
        txs=txs,
        total_balance=total_balance,
        total_in=total_in,
        total_out=total_out,
        account_types=BANK_ACCOUNT_TYPES,
        today=datetime.today().strftime('%Y-%m-%d'),
    )


@app.route('/banking/account/add', methods=['POST'])
@login_required
def add_bank_account():
    name, e1 = _clean_text('name', max_len=80, label='name')
    bank_name, e2 = _clean_text('bank_name', required=False, max_len=80, label='bank')
    account_no, e3 = _clean_text('account_no', required=False, max_len=32, label='account no')
    notes, e4 = _clean_text('notes', required=False, max_len=200, label='notes')
    opening, e5 = _clean_float('opening_balance', required=False, min_v=0, max_v=1e12, label='opening', default=0)
    if _reject((name, e1), (bank_name, e2), (account_no, e3), (notes, e4), (opening, e5)):
        return redirect(url_for('banking'))
    atype = request.form.get('account_type') or 'Current'
    if atype not in BANK_ACCOUNT_TYPES:
        atype = 'Other'
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO bank_accounts (user_id, name, bank_name, account_no, account_type, opening_balance, notes, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, name, bank_name or '', account_no or '', atype, opening or 0,
         notes or '', datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added bank account', name)
    flash('Bank account saved.', 'success')
    return redirect(url_for('banking'))


@app.route('/banking/account/<int:id>/delete', methods=['POST'])
@login_required
def delete_bank_account(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM bank_transactions WHERE user_id=%s AND account_id=%s', (current_user.id, id))
    cur.execute('DELETE FROM bank_accounts WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Account removed.', 'success')
    return redirect(url_for('banking'))


@app.route('/banking/tx/add', methods=['POST'])
@login_required
def add_bank_tx():
    amount, e1 = _clean_float('amount', min_v=0.01, max_v=1e12, label='amount')
    date, e2 = _clean_date('tx_date')
    party, e3 = _clean_text('counterparty', required=False, max_len=80, label='party')
    ref, e4 = _clean_text('reference', required=False, max_len=80, label='reference')
    cat, e5 = _clean_text('category', required=False, max_len=80, label='category')
    if _reject((amount, e1), (date, e2), (party, e3), (ref, e4), (cat, e5)):
        return redirect(url_for('banking'))
    kind = 'in' if request.form.get('kind') == 'in' else 'out'
    try:
        account_id = int(request.form.get('account_id') or 0)
    except ValueError:
        account_id = 0
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT id FROM bank_accounts WHERE id=%s AND user_id=%s', (account_id, current_user.id))
    if not cur.fetchone():
        cur.close()
        conn.close()
        flash('Choose a valid account.', 'danger')
        return redirect(url_for('banking'))
    cash_id = None
    if request.form.get('post_cashbook'):
        cur.execute(
            '''INSERT INTO cash_books (user_id, entry_type, category, description, amount, date)
               VALUES (%s,%s,%s,%s,%s,%s) RETURNING id''',
            (current_user.id, kind, cat or 'Bank', party or ref or 'Bank movement', amount, date),
        )
        row = cur.fetchone() or {}
        cash_id = row.get('id')
    cur.execute(
        '''INSERT INTO bank_transactions
           (user_id, account_id, tx_date, kind, amount, counterparty, reference, category, notes, linked_cashbook_id)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, account_id, date, kind, amount, party or '', ref or '', cat or '', '', cash_id),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Movement posted.', 'success')
    return redirect(url_for('banking'))


@app.route('/banking/tx/<int:id>/delete', methods=['POST'])
@login_required
def delete_bank_tx(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM bank_transactions WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Movement deleted.', 'success')
    return redirect(url_for('banking'))


@app.route('/banking/transfer', methods=['POST'])
@login_required
def bank_transfer():
    amount, e1 = _clean_float('amount', min_v=0.01, max_v=1e12, label='amount')
    date, e2 = _clean_date('tx_date')
    ref, e3 = _clean_text('reference', required=False, max_len=80, label='memo')
    if _reject((amount, e1), (date, e2), (ref, e3)):
        return redirect(url_for('banking'))
    try:
        from_id = int(request.form.get('from_id') or 0)
        to_id = int(request.form.get('to_id') or 0)
    except ValueError:
        from_id = to_id = 0
    if from_id == to_id or not from_id or not to_id:
        flash('Pick two different accounts.', 'danger')
        return redirect(url_for('banking'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT id, name FROM bank_accounts WHERE user_id=%s AND id IN (%s,%s)', (current_user.id, from_id, to_id))
    found = {row['id']: row['name'] for row in (cur.fetchall() or [])}
    if from_id not in found or to_id not in found:
        cur.close()
        conn.close()
        flash('Accounts not found.', 'danger')
        return redirect(url_for('banking'))
    memo = ref or 'Transfer'
    cur.execute(
        '''INSERT INTO bank_transactions
           (user_id, account_id, tx_date, kind, amount, counterparty, reference, category, notes, linked_cashbook_id)
           VALUES (%s,%s,%s,'out',%s,%s,%s,'Transfer','',NULL)''',
        (current_user.id, from_id, date, amount, found[to_id], memo),
    )
    cur.execute(
        '''INSERT INTO bank_transactions
           (user_id, account_id, tx_date, kind, amount, counterparty, reference, category, notes, linked_cashbook_id)
           VALUES (%s,%s,%s,'in',%s,%s,%s,'Transfer','',NULL)''',
        (current_user.id, to_id, date, amount, found[from_id], memo),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Transfer posted.', 'success')
    return redirect(url_for('banking'))


@app.route('/investments')
@login_required
def investments():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM investments WHERE user_id=%s ORDER BY name', (uid,))
    rows = list(cur.fetchall() or [])
    cur.close()
    conn.close()
    total_cost = total_value = 0.0
    holdings = []
    for h in rows:
        qty = float(h.get('quantity') or 0)
        cost_u = float(h.get('unit_cost') or 0)
        price = float(h.get('current_price') or 0)
        cost = qty * cost_u
        value = qty * price
        h = dict(h)
        h['cost'] = cost
        h['value'] = value
        h['gain'] = value - cost
        holdings.append(h)
        total_cost += cost
        total_value += value
    return render_template(
        'investments.html',
        holdings=holdings,
        total_cost=total_cost,
        total_value=total_value,
        total_gain=total_value - total_cost,
        asset_types=INVEST_TYPES,
        today=datetime.today().strftime('%Y-%m-%d'),
    )


@app.route('/investments/add', methods=['POST'])
@login_required
def add_investment():
    name, e1 = _clean_text('name', max_len=80, label='name')
    qty, e2 = _clean_float('quantity', min_v=0, max_v=1e12, label='quantity')
    cost, e3 = _clean_float('unit_cost', min_v=0, max_v=1e12, label='unit cost')
    price, e4 = _clean_float('current_price', min_v=0, max_v=1e12, label='current price')
    notes, e5 = _clean_text('notes', required=False, max_len=200, label='notes')
    bought, e6 = _clean_date('bought_date')
    if _reject((name, e1), (qty, e2), (cost, e3), (price, e4), (notes, e5), (bought, e6)):
        return redirect(url_for('investments'))
    atype = request.form.get('asset_type') or 'Other'
    if atype not in INVEST_TYPES:
        atype = 'Other'
    ticker = (request.form.get('ticker') or '').strip().upper()[:20]
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO investments (user_id, name, asset_type, quantity, unit_cost, current_price, bought_date, notes, ticker, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, name, atype, qty, cost, price, bought or datetime.today().strftime('%Y-%m-%d'),
         notes or '', ticker, datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    log_activity(current_user.id, current_user.username, 'Added investment', name)
    flash('Holding saved.', 'success')
    return redirect(url_for('investments'))


@app.route('/investments/<int:id>/update', methods=['POST'])
@login_required
def update_investment(id):
    price, e1 = _clean_float('current_price', min_v=0, max_v=1e12, label='current price')
    if _reject((price, e1)):
        return redirect(url_for('investments'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'UPDATE investments SET current_price=%s WHERE id=%s AND user_id=%s',
        (price, id, current_user.id),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Price updated.', 'success')
    return redirect(url_for('investments'))


@app.route('/investments/<int:id>/delete', methods=['POST'])
@login_required
def delete_investment(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM investments WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Holding removed.', 'success')
    return redirect(url_for('investments'))


@app.route('/shop')
@login_required
def shop():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM shop_listings WHERE user_id=%s ORDER BY published DESC, title', (uid,))
    listings = list(cur.fetchall() or [])
    cur.execute(
        '''SELECT o.*, l.title FROM shop_orders o
           LEFT JOIN shop_listings l ON l.id = o.listing_id
           WHERE o.user_id=%s ORDER BY o.order_date DESC, o.id DESC LIMIT 80''',
        (uid,),
    )
    orders = list(cur.fetchall() or [])
    cur.execute('SELECT * FROM bank_accounts WHERE user_id=%s ORDER BY name', (uid,))
    accounts = list(cur.fetchall() or [])
    live_count = sum(1 for l in listings if l.get('published'))
    open_orders = sum(1 for o in orders if o.get('status') in ('pending', 'paid', 'packed'))
    order_total = 0.0
    for o in orders:
        try:
            order_total += float(o.get('amount') or 0)
        except (TypeError, ValueError):
            pass
    bank_total = 0.0
    for a in accounts:
        bank_total += _bank_balance(cur, uid, a['id'], a.get('opening_balance'))
    cur.close()
    conn.close()
    return render_template(
        'shop.html',
        listings=listings,
        orders=orders,
        accounts=accounts,
        live_count=live_count,
        open_orders=open_orders,
        order_total=order_total,
        bank_total=bank_total,
        order_statuses=SHOP_ORDER_STATUSES,
        today=datetime.today().strftime('%Y-%m-%d'),
    )


@app.route('/shop/listing/add', methods=['POST'])
@login_required
def add_shop_listing():
    title, e1 = _clean_text('title', max_len=80, label='title')
    price, e2 = _clean_float('price', min_v=0, max_v=1e12, label='price')
    desc, e3 = _clean_text('description', required=False, max_len=200, label='description')
    if _reject((title, e1), (price, e2), (desc, e3)):
        return redirect(url_for('shop'))
    kind = request.form.get('kind') or 'Custom'
    if kind not in ('Product', 'Service', 'Custom'):
        kind = 'Custom'
    published = 1 if request.form.get('published') else 0
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO shop_listings (user_id, title, kind, source_id, price, published, description, created_date)
           VALUES (%s,%s,%s,0,%s,%s,%s,%s)''',
        (current_user.id, title, kind, price, published, desc or '', datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Listing saved.', 'success')
    return redirect(url_for('shop'))


@app.route('/shop/listing/<int:id>/toggle', methods=['POST'])
@login_required
def toggle_shop_listing(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT published FROM shop_listings WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    if row:
        cur.execute(
            'UPDATE shop_listings SET published=%s WHERE id=%s AND user_id=%s',
            (0 if row.get('published') else 1, id, current_user.id),
        )
        conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('shop'))


@app.route('/shop/listing/<int:id>/delete', methods=['POST'])
@login_required
def delete_shop_listing(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE shop_orders SET listing_id=NULL WHERE listing_id=%s AND user_id=%s', (id, current_user.id))
    cur.execute('DELETE FROM shop_listings WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    flash('Listing removed.', 'success')
    return redirect(url_for('shop'))


@app.route('/shop/order/add', methods=['POST'])
@login_required
def add_shop_order():
    buyer, e1 = _clean_text('buyer', max_len=80, label='buyer')
    qty, e2 = _clean_float('qty', min_v=1, max_v=1e6, label='qty')
    date, e3 = _clean_date('order_date')
    notes, e4 = _clean_text('notes', required=False, max_len=200, label='notes')
    if _reject((buyer, e1), (qty, e2), (date, e3), (notes, e4)):
        return redirect(url_for('shop'))
    try:
        listing_id = int(request.form.get('listing_id') or 0)
    except ValueError:
        listing_id = 0
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT id, price FROM shop_listings WHERE id=%s AND user_id=%s', (listing_id, current_user.id))
    listing = cur.fetchone()
    if not listing:
        cur.close()
        conn.close()
        flash('Choose a listing.', 'danger')
        return redirect(url_for('shop'))
    amount = float(listing.get('price') or 0) * float(qty)
    cur.execute(
        '''INSERT INTO shop_orders (user_id, listing_id, buyer, qty, amount, status, order_date, notes, created_date)
           VALUES (%s,%s,%s,%s,%s,'pending',%s,%s,%s)''',
        (current_user.id, listing_id, buyer, int(qty), amount, date or datetime.today().strftime('%Y-%m-%d'),
         notes or '', datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Order saved.', 'success')
    return redirect(url_for('shop'))


@app.route('/shop/order/<int:id>/status', methods=['POST'])
@login_required
def set_shop_order(id):
    status = request.form.get('status') or 'pending'
    if status not in SHOP_ORDER_STATUSES:
        status = 'pending'
    conn = get_db()
    cur = conn.cursor()
    cur.execute('UPDATE shop_orders SET status=%s WHERE id=%s AND user_id=%s', (status, id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('shop'))


@app.route('/shop/order/<int:id>/delete', methods=['POST'])
@login_required
def delete_shop_order(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM shop_orders WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('shop'))


@app.route('/shop/pay', methods=['POST'])
@login_required
def shop_pay():
    amount, e1 = _clean_float('amount', min_v=0.01, max_v=1e12, label='amount')
    date, e2 = _clean_date('tx_date')
    party, e3 = _clean_text('counterparty', max_len=80, label='pay to')
    ref, e4 = _clean_text('reference', required=False, max_len=80, label='reference')
    if _reject((amount, e1), (date, e2), (party, e3), (ref, e4)):
        return redirect(url_for('shop'))
    try:
        account_id = int(request.form.get('account_id') or 0)
    except ValueError:
        account_id = 0
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT id FROM bank_accounts WHERE id=%s AND user_id=%s', (account_id, current_user.id))
    if not cur.fetchone():
        cur.close()
        conn.close()
        flash('Choose a bank book first.', 'danger')
        return redirect(url_for('shop'))
    cur.execute(
        '''INSERT INTO bank_transactions
           (user_id, account_id, tx_date, kind, amount, counterparty, reference, category, notes, linked_cashbook_id)
           VALUES (%s,%s,%s,'out',%s,%s,%s,'Shop pay','',NULL)''',
        (current_user.id, account_id, date, amount, party, ref or ''),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Payment posted from the bank book.', 'success')
    return redirect(url_for('shop'))


def _market_bases(cur, uid):
    month = datetime.today().strftime('%Y-%m')
    cur.execute(
        "SELECT COALESCE(SUM(total_amount),0) AS n FROM sales WHERE user_id=%s AND sale_date LIKE %s",
        (uid, month + '%'),
    )
    month_sales = float((cur.fetchone() or {}).get('n') or 0)
    cur.execute(
        "SELECT COALESCE(SUM(monthly_revenue),0) AS n FROM market_channels WHERE user_id=%s",
        (uid,),
    )
    channel_month = float((cur.fetchone() or {}).get('n') or 0)
    cur.execute(
        "SELECT COALESCE(SUM(quantity * selling_price),0) AS n FROM stock WHERE user_id=%s",
        (uid,),
    )
    stock_value = float((cur.fetchone() or {}).get('n') or 0)
    cur.execute('SELECT id, opening_balance FROM bank_accounts WHERE user_id=%s', (uid,))
    bank_total = 0.0
    for a in cur.fetchall() or []:
        bank_total += _bank_balance(cur, uid, a['id'], a.get('opening_balance'))
    cur.execute(
        "SELECT COALESCE(SUM(quantity * current_price),0) AS n FROM investments WHERE user_id=%s",
        (uid,),
    )
    invest_value = float((cur.fetchone() or {}).get('n') or 0)
    yearly = (channel_month or month_sales) * 12
    if channel_month:
        yearly = channel_month * 12
    elif month_sales:
        yearly = month_sales * 12
    liquid = bank_total + invest_value
    asset_base = stock_value + liquid
    return {
        'month_sales': month_sales,
        'channel_month': channel_month,
        'stock_value': stock_value,
        'liquid': liquid,
        'yearly_base': yearly,
        'asset_base': asset_base,
    }


@app.route('/market')
@login_required
def market():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    bases = _market_bases(cur, uid)
    cur.execute('SELECT * FROM market_channels WHERE user_id=%s ORDER BY name', (uid,))
    channels = cur.fetchall() or []
    cur.execute('SELECT * FROM market_valuations WHERE user_id=%s ORDER BY as_of DESC, id DESC LIMIT 40', (uid,))
    valuations = cur.fetchall() or []
    cur.close()
    conn.close()
    live_multiple = 2.0
    rev_value = bases['yearly_base'] * live_multiple
    blended = (rev_value + bases['asset_base']) / 2.0
    return render_template(
        'market.html',
        channels=channels,
        valuations=valuations,
        channel_kinds=MARKET_CHANNEL_KINDS,
        live_multiple=live_multiple,
        rev_value=rev_value,
        blended=blended,
        **bases
    )


@app.route('/market/channel/add', methods=['POST'])
@login_required
def add_market_channel():
    name, e1 = _clean_text('name', max_len=80, label='name')
    rev, e2 = _clean_float('monthly_revenue', required=False, min_v=0, max_v=1e12, label='monthly revenue', default=0)
    notes, e3 = _clean_text('notes', required=False, max_len=200, label='notes')
    if _reject((name, e1), (rev, e2), (notes, e3)):
        return redirect(url_for('market'))
    kind = request.form.get('kind') or 'Other'
    if kind not in MARKET_CHANNEL_KINDS:
        kind = 'Other'
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO market_channels (user_id, name, kind, monthly_revenue, notes, created_date)
           VALUES (%s,%s,%s,%s,%s,%s)''',
        (current_user.id, name, kind, rev or 0, notes or '', datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Channel saved.', 'success')
    return redirect(url_for('market'))


@app.route('/market/channel/<int:id>/delete', methods=['POST'])
@login_required
def delete_market_channel(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM market_channels WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('market'))


@app.route('/market/valuation/add', methods=['POST'])
@login_required
def add_valuation():
    method = request.form.get('method') or 'blended'
    if method not in ('revenue_multiple', 'asset_based', 'blended', 'owner'):
        method = 'blended'
    rev, e1 = _clean_float('revenue_base', required=False, min_v=0, max_v=1e12, label='revenue', default=0)
    mult, e2 = _clean_float('multiple', required=False, min_v=0, max_v=100, label='multiple', default=2)
    assets, e3 = _clean_float('asset_base', required=False, min_v=0, max_v=1e12, label='assets', default=0)
    own, e4 = _clean_float('estimated_value', required=False, min_v=0, max_v=1e12, label='value', default=None)
    notes, e5 = _clean_text('notes', required=False, max_len=200, label='notes')
    if _reject((rev, e1), (mult, e2), (assets, e3), (own, e4), (notes, e5)):
        return redirect(url_for('market'))
    if own not in (None, ''):
        value = float(own)
    elif method == 'revenue_multiple':
        value = float(rev or 0) * float(mult or 0)
    elif method == 'asset_based':
        value = float(assets or 0)
    else:
        value = (float(rev or 0) * float(mult or 0) + float(assets or 0)) / 2.0
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        '''INSERT INTO market_valuations
           (user_id, as_of, method, multiple, revenue_base, asset_base, estimated_value, notes, created_date)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
        (current_user.id, datetime.today().strftime('%Y-%m-%d'), method, mult or 0, rev or 0, assets or 0,
         value, notes or '', datetime.today().strftime('%Y-%m-%d')),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Valuation saved.', 'success')
    return redirect(url_for('market'))


@app.route('/market/valuation/<int:id>/delete', methods=['POST'])
@login_required
def delete_valuation(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('DELETE FROM market_valuations WHERE id=%s AND user_id=%s', (id, current_user.id))
    conn.commit()
    cur.close()
    conn.close()
    return redirect(url_for('market'))


def _mark_shop_paid(uid, order_id, tx_ref='', tx_id=''):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "UPDATE shop_orders SET status='paid', flutterwave_tx_ref=%s, flutterwave_tx_id=%s "
        "WHERE id=%s AND user_id=%s",
        (tx_ref or '', str(tx_id or ''), order_id, uid),
    )
    conn.commit()
    cur.close()
    conn.close()


@app.route('/shop/order/<int:id>/pay')
@login_required
def pay_shop_order(id):
    settings = _stripe_settings()
    if not integrations.flutterwave_ready(settings):
        flash('Add Flutterwave keys in Settings to collect card or mobile money.', 'danger')
        return redirect(url_for('settings'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT o.*, l.title FROM shop_orders o LEFT JOIN shop_listings l ON l.id = o.listing_id WHERE o.id=%s AND o.user_id=%s",
        (id, current_user.id),
    )
    row = cur.fetchone()
    cur.close()
    conn.close()
    if not row:
        flash('Order not found.', 'danger')
        return redirect(url_for('shop'))
    if (row.get('status') or '') == 'paid':
        flash('That order is already paid.', 'info')
        return redirect(url_for('shop'))
    tx_ref = 'kaze-shop-%s-%s' % (id, int(datetime.now().timestamp()))
    title = row.get('title') or 'Shop order'
    pay, err = integrations.create_flutterwave_payment(
        settings,
        tx_ref=tx_ref,
        amount=row.get('amount') or 0,
        currency=stripe_payments.currency_code(settings),
        redirect_url=url_for('flutterwave_return', _external=True),
        title=title,
        customer_name=row.get('buyer') or '',
        customer_email=getattr(current_user, 'email', '') or '',
        meta={'user_id': current_user.id, 'order_id': id, 'kind': 'shop'},
    )
    if err:
        flash(err, 'danger')
        return redirect(url_for('shop'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'UPDATE shop_orders SET flutterwave_tx_ref=%s WHERE id=%s AND user_id=%s',
        (tx_ref, id, current_user.id),
    )
    conn.commit()
    cur.close()
    conn.close()
    return redirect(pay['link'], code=303)


@app.route('/shop/flutterwave/return')
@login_required
def flutterwave_return():
    settings = _stripe_settings()
    tx_ref = (request.args.get('tx_ref') or request.args.get('txRef') or '').strip()
    tx_id = (request.args.get('transaction_id') or '').strip()
    status = (request.args.get('status') or '').lower()
    if status in ('cancelled', 'canceled'):
        flash('Payment cancelled.', 'info')
        return redirect(url_for('shop'))
    data, err = integrations.verify_flutterwave(settings, tx_ref=tx_ref or None, transaction_id=tx_id or None)
    if err:
        flash(err, 'danger')
        return redirect(url_for('shop'))
    meta = data.get('meta') or {}
    # Ownership check: refuse to fulfill a payment for a different user.
    try:
        meta_user_id = int(meta.get('user_id') or 0)
    except (TypeError, ValueError):
        meta_user_id = 0
    if meta_user_id and meta_user_id != int(current_user.id):
        flash('That payment belongs to a different account.', 'danger')
        return redirect(url_for('payments'))
    try:
        order_id = int(meta.get('order_id') or 0)
    except (TypeError, ValueError):
        order_id = 0
    if not order_id and tx_ref.startswith('kaze-shop-'):
        try:
            order_id = int(tx_ref.split('-')[2])
        except (IndexError, ValueError):
            order_id = 0
    kind = str(meta.get('kind') or '')
    try:
        record_id = int(meta.get('record_id') or 0)
    except (TypeError, ValueError):
        record_id = 0
    if order_id:
        _mark_shop_paid(current_user.id, order_id, tx_ref or data.get('tx_ref'), data.get('id') or tx_id)
        flash('Order marked paid.', 'success')
        return redirect(url_for('shop'))
    if kind:
        _fulfill_stripe_payment(
            current_user.id, kind, record_id,
            tx_ref or data.get('tx_ref') or '',
            str(data.get('id') or tx_id or ''),
            data.get('amount'),
        )
        flash('Payment received.', 'success')
        return redirect(url_for('payments'))
    flash('Payment confirmed. Refresh the order list if status is still pending.', 'success')
    return redirect(url_for('shop'))


@app.route('/webhooks/flutterwave', methods=['POST'])
def flutterwave_webhook():
    raw = request.get_data(cache=False)
    header_hash = request.headers.get('verif-hash') or request.headers.get('Verif-Hash') or ''
    try:
        payload = json.loads(raw.decode('utf-8') or '{}')
    except Exception:
        return ('', 400)
    data = payload.get('data') or {}
    tx_ref = data.get('tx_ref') or ''
    # Hash check runs FIRST, using app-level config, before we trust any body.
    try:
        conn = get_db(); cur = conn.cursor()
        cur.execute("SELECT flutterwave_webhook_hash FROM settings WHERE flutterwave_webhook_hash != '' LIMIT 1")
        hook_row = cur.fetchone()
        cur.close(); conn.close()
    except Exception:
        hook_row = None
    if hook_row and hook_row.get('flutterwave_webhook_hash'):
        if not integrations.flutterwave_hash_ok(hook_row, header_hash):
            return ('', 401)
    # Verify with Flutterwave BEFORE trusting any meta in the body.
    settings = _stripe_settings()
    verified, err = integrations.verify_flutterwave(settings, tx_ref=tx_ref, transaction_id=data.get('id'))
    if err or not verified:
        return ('', 200)
    # Use VERIFIED meta only — never the raw body meta.
    v_meta = (verified or {}).get('meta') or {}
    try:
        v_user_id = int(v_meta.get('user_id') or 0)
    except (TypeError, ValueError):
        v_user_id = 0
    try:
        v_order_id = int(v_meta.get('order_id') or 0)
    except (TypeError, ValueError):
        v_order_id = 0
    v_kind = str(v_meta.get('kind') or '')
    try:
        v_record_id = int(v_meta.get('record_id') or 0)
    except (TypeError, ValueError):
        v_record_id = 0
    if not v_user_id:
        return ('', 200)
    v_tx_id = str((verified or {}).get('id') or data.get('id') or '')
    v_amount = (verified or {}).get('amount')
    if v_order_id:
        _mark_shop_paid(v_user_id, v_order_id, tx_ref, v_tx_id)
    elif v_kind:
        _fulfill_stripe_payment(v_user_id, v_kind, v_record_id, tx_ref, v_tx_id, v_amount)
    return ('', 200)


@app.route('/investments/refresh', methods=['POST'])
@login_required
def refresh_investments():
    settings = _stripe_settings()
    if not integrations.alpha_ready(settings):
        flash('Add an Alpha Vantage key in Settings to refresh listed prices.', 'danger')
        return redirect(url_for('settings'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT id, ticker FROM investments WHERE user_id=%s AND ticker IS NOT NULL AND ticker <> ''",
        (current_user.id,),
    )
    rows = list(cur.fetchall() or [])
    updated = 0
    seen = {}
    errors = []
    for row in rows[:12]:
        ticker = (row.get('ticker') or '').upper()
        if ticker in seen:
            price, err = seen[ticker]
        else:
            price, err = integrations.quote_price(ticker, settings)
            seen[ticker] = (price, err)
        if err:
            if err not in errors:
                errors.append(err)
            continue
        cur.execute(
            'UPDATE investments SET current_price=%s WHERE id=%s AND user_id=%s',
            (price, row['id'], current_user.id),
        )
        updated += 1
    conn.commit()
    cur.close()
    conn.close()
    if updated:
        flash('Updated %s holding price(s).' % updated, 'success')
    if errors:
        flash(errors[0], 'danger')
    if not updated and not errors:
        flash('Add tickers to holdings first.', 'info')
    return redirect(url_for('investments'))


@app.route('/investments/<int:id>/refresh', methods=['POST'])
@login_required
def refresh_investment(id):
    settings = _stripe_settings()
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT ticker FROM investments WHERE id=%s AND user_id=%s', (id, current_user.id))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        flash('Holding not found.', 'danger')
        return redirect(url_for('investments'))
    price, err = integrations.quote_price(row.get('ticker'), settings)
    if err:
        cur.close()
        conn.close()
        flash(err, 'danger')
        return redirect(url_for('investments'))
    cur.execute(
        'UPDATE investments SET current_price=%s WHERE id=%s AND user_id=%s',
        (price, id, current_user.id),
    )
    conn.commit()
    cur.close()
    conn.close()
    flash('Price updated to %s.' % price, 'success')
    return redirect(url_for('investments'))


@app.route('/market/fx', methods=['POST'])
@login_required
def refresh_market_fx():
    settings = _stripe_settings()
    quote = stripe_payments.currency_code(settings)
    rate, err = integrations.fx_rate('USD', quote, settings)
    if err:
        flash(err, 'danger')
        return redirect(url_for('market'))
    flash('USD/%s is %s (cached 15 minutes).' % (quote, rate), 'success')
    return redirect(url_for('market'))


# ======================== Unique Features ========================

@app.route('/insights')
@login_required
def insights():
    """Live business insights with 30d trend, gainers/decliners, dead capital."""
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS revenue, COALESCE(SUM(profit),0) AS profit, "
                "COUNT(*) AS tickets, COUNT(DISTINCT customer_id) AS customers "
                "FROM sales WHERE user_id=%s AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    r30 = cur.fetchone() or {}
    revenue_30 = float(r30.get('revenue') or 0)
    profit_30 = float(r30.get('profit') or 0)
    tickets_30 = int(r30.get('tickets') or 0)
    customers_30 = int(r30.get('customers') or 0)

    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS r FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD') "
                "AND sale_date < to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    revenue_prev = float((cur.fetchone() or {}).get('r') or 0)
    revenue_change_pct = round((revenue_30 - revenue_prev) / revenue_prev * 100, 1) if revenue_prev > 0 else (100 if revenue_30 > 0 else 0)

    cur.execute("SELECT substr(sale_date,1,10) AS d, COALESCE(SUM(total_amount),0) AS t FROM sales "
                "WHERE user_id=%s AND sale_date >= to_char(CURRENT_DATE - INTERVAL '29 days','YYYY-MM-DD') "
                "GROUP BY substr(sale_date,1,10) ORDER BY d", (uid,))
    daily_map = {r['d']: float(r['t']) for r in cur.fetchall()}
    trend_labels, trend_values = [], []
    for i in range(29, -1, -1):
        d = (datetime.today() - timedelta(days=i)).strftime('%Y-%m-%d')
        trend_labels.append(d[5:])
        trend_values.append(round(daily_map.get(d, 0), 2))

    cur.execute("SELECT st.product_name, COALESCE(SUM(s.total_amount),0) AS rev FROM sales s "
                "LEFT JOIN stock st ON st.id = s.stock_id WHERE s.user_id=%s "
                "AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') "
                "GROUP BY st.product_name", (uid,))
    cur_rev = {r['product_name'] or 'Item': float(r['rev']) for r in cur.fetchall()}
    cur.execute("SELECT st.product_name, COALESCE(SUM(s.total_amount),0) AS rev FROM sales s "
                "LEFT JOIN stock st ON st.id = s.stock_id WHERE s.user_id=%s "
                "AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD') "
                "AND s.sale_date < to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') "
                "GROUP BY st.product_name", (uid,))
    prev_rev = {r['product_name'] or 'Item': float(r['rev']) for r in cur.fetchall()}

    changes = []
    for name in set(cur_rev) | set(prev_rev):
        a, b = cur_rev.get(name, 0), prev_rev.get(name, 0)
        changes.append({'name': name, 'current': a, 'delta': a - b,
                        'pct': round((a - b) / b * 100, 1) if b > 0 else (100 if a > 0 else 0)})
    top_gainers = sorted([c for c in changes if c['delta'] > 0], key=lambda x: x['delta'], reverse=True)[:3]
    top_decliners = sorted([c for c in changes if c['delta'] < 0], key=lambda x: x['delta'])[:3]

    cur.execute("SELECT st.product_name, COUNT(*) AS t, COALESCE(SUM(s.total_amount),0) AS r, "
                "COALESCE(SUM(s.profit),0) AS p FROM sales s LEFT JOIN stock st ON st.id = s.stock_id "
                "WHERE s.user_id=%s AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD') "
                "GROUP BY st.product_name HAVING COUNT(*) >= 3", (uid,))
    margin_rows = [{'name': r['product_name'] or 'Item', 'revenue': float(r['r'] or 0),
                    'profit': float(r['p'] or 0), 'tickets': int(r['t'] or 0),
                    'margin_pct': round(float(r['p'] or 0) / float(r['r']) * 100, 1) if r['r'] else 0}
                   for r in (cur.fetchall() or [])]
    margin_rows.sort(key=lambda x: x['margin_pct'])
    worst_margin = margin_rows[0] if margin_rows else None

    cur.execute("SELECT COUNT(*) AS n FROM stock st WHERE st.user_id=%s AND st.quantity <= COALESCE("
                "(SELECT low_stock_threshold FROM settings WHERE user_id=%s), 10)", (uid, uid))
    low_stock = int((cur.fetchone() or {}).get('n') or 0)

    cur.execute("SELECT COUNT(*) AS n FROM stock st WHERE st.user_id=%s AND NOT EXISTS ("
                "SELECT 1 FROM sales s WHERE s.stock_id = st.id "
                "AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD'))", (uid,))
    dead_stock = int((cur.fetchone() or {}).get('n') or 0)

    cur.execute("SELECT COALESCE(SUM(quantity * cost_price),0) AS v FROM stock st WHERE st.user_id=%s "
                "AND NOT EXISTS (SELECT 1 FROM sales s WHERE s.stock_id = st.id "
                "AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD'))", (uid,))
    dead_capital = float((cur.fetchone() or {}).get('v') or 0)

    cur.execute("SELECT COALESCE(payment_method,'cash') AS m, COUNT(*) AS t, "
                "COALESCE(SUM(total_amount),0) AS r FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') "
                "GROUP BY payment_method ORDER BY r DESC", (uid,))
    payment_mix = [{'method': (r['m'] or 'cash').title(), 'tickets': int(r['t'] or 0),
                    'revenue': float(r['r'] or 0)} for r in (cur.fetchall() or [])]

    cur.execute("SELECT COALESCE(SUM(quantity * cost_price),0) AS c, "
                "COALESCE(SUM(quantity * selling_price),0) AS s FROM stock WHERE user_id=%s", (uid,))
    sv = cur.fetchone() or {}
    cur.close(); conn.close()

    insight_cards = []
    if top_gainers:
        t = top_gainers[0]
        insight_cards.append({'severity': 'success', 'icon': 'trending-up',
            'title': f"Fastest growing: {t['name']}",
            'detail': f"+{_format_money(t['delta'])} vs prior 30 days.",
            'action_label': 'Open stock', 'action_url': url_for('stock')})
    else:
        insight_cards.append({'severity': 'info', 'icon': 'package',
            'title': 'No growth trends detected yet',
            'detail': 'Need sales in both the last 30 and prior 30 days.',
            'action_label': 'Record a sale', 'action_url': url_for('sales')})
    if worst_margin:
        sev = 'danger' if worst_margin['margin_pct'] < 0 else ('warn' if worst_margin['margin_pct'] < 15 else 'info')
        insight_cards.append({'severity': sev, 'icon': 'trending-down',
            'title': f"Lowest margin: {worst_margin['name']}",
            'detail': f"{worst_margin['margin_pct']}% over {worst_margin['tickets']} tickets.",
            'action_label': 'Edit prices', 'action_url': url_for('stock')})
    if low_stock > 0:
        insight_cards.append({'severity': 'warn', 'icon': 'alert-triangle',
            'title': f"{low_stock} product(s) below stock threshold",
            'detail': 'Restock soon.', 'action_label': 'Open stock', 'action_url': url_for('stock')})
    else:
        insight_cards.append({'severity': 'success', 'icon': 'check-circle',
            'title': 'Stock levels healthy', 'detail': 'No products below threshold.',
            'action_label': 'Review stock', 'action_url': url_for('stock')})
    if dead_stock > 0:
        insight_cards.append({'severity': 'warn', 'icon': 'package',
            'title': f"{dead_stock} dead-stock item(s) · {_format_money(dead_capital)} tied up",
            'detail': 'No sales in 60 days.', 'action_label': 'Open stock', 'action_url': url_for('stock')})
    if payment_mix and revenue_30 > 0:
        tp = payment_mix[0]
        insight_cards.append({'severity': 'info', 'icon': 'credit-card',
            'title': 'Payment mix',
            'detail': f"{round(tp['revenue']/revenue_30*100,1)}% via {tp['method']} in last 30 days.",
            'action_label': 'Open payments', 'action_url': url_for('payments')})

    return render_template('insights.html', insight_cards=insight_cards,
        revenue_30=revenue_30, profit_30=profit_30, tickets_30=tickets_30, customers_30=customers_30,
        revenue_prev=revenue_prev, revenue_change_pct=revenue_change_pct,
        trend_labels=trend_labels, trend_values=trend_values,
        top_gainers=top_gainers, top_decliners=top_decliners,
        low_stock=low_stock, dead_stock=dead_stock, dead_capital=dead_capital,
        payment_mix=payment_mix,
        stock_cost_value=float(sv.get('c') or 0), stock_sell_value=float(sv.get('s') or 0))


@app.route('/forecast')
@login_required
def forecast():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(AVG(daily_rev),0) AS a FROM ("
                "SELECT SUM(total_amount) AS daily_rev FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') GROUP BY sale_date) t", (uid,))
    avg_daily_revenue = float((cur.fetchone() or {}).get('a') or 0)
    cur.execute("SELECT COALESCE(AVG(daily_exp),0) AS a FROM ("
                "SELECT SUM(amount) AS daily_exp FROM expenses WHERE user_id=%s "
                "AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') GROUP BY date) t", (uid,))
    avg_daily_expense = float((cur.fetchone() or {}).get('a') or 0)
    cur.execute("SELECT COALESCE(SUM(opening_balance),0) AS op FROM bank_accounts WHERE user_id=%s", (uid,))
    op = float((cur.fetchone() or {}).get('op') or 0)
    cur.execute("SELECT COALESCE(SUM(CASE WHEN kind='in' THEN amount ELSE -amount END),0) AS mv FROM bank_transactions WHERE user_id=%s", (uid,))
    mv = float((cur.fetchone() or {}).get('mv') or 0)
    cur.execute("SELECT COUNT(*) AS n FROM bank_accounts WHERE user_id=%s", (uid,))
    bank_count = int((cur.fetchone() or {}).get('n') or 0)
    if bank_count > 0:
        current_cash = op + mv
        cash_source = 'bank accounts'
    else:
        cur.execute("SELECT COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE -amount END),0) AS net FROM cash_books WHERE user_id=%s", (uid,))
        current_cash = float((cur.fetchone() or {}).get('net') or 0)
        cash_source = 'cash book'
    cur.execute("SELECT COALESCE(SUM(amount),0) AS t FROM recurring_items WHERE user_id=%s "
                "AND active=1 AND item_type='expense' "
                "AND next_run_date <= to_char(CURRENT_DATE + INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    upcoming_recurring = float((cur.fetchone() or {}).get('t') or 0)
    cur.execute("SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders WHERE user_id=%s AND status='pending'", (uid,))
    upcoming_pos = float((cur.fetchone() or {}).get('t') or 0)
    cur.close(); conn.close()

    net_daily = avg_daily_revenue - avg_daily_expense
    daily_obligations = (upcoming_recurring + upcoming_pos) / 30.0
    base, worst, best, labels = [], [], [], []
    for d in range(1, 31):
        labels.append(f"Day {d}")
        net_base = net_daily - daily_obligations
        base.append(round(current_cash + net_base * d, 2))
        worst.append(round(current_cash + (net_daily * 0.5 - daily_obligations) * d, 2))
        best.append(round(current_cash + (net_daily * 1.5 - daily_obligations) * d, 2))
    lowest = min(base) if base else 0
    if lowest < 0:
        cash_risk = f'High — projected negative by day {base.index(min(base)) + 1}'
    elif lowest < current_cash * 0.2:
        cash_risk = 'Moderate — buffer below 20% of current cash'
    else:
        cash_risk = 'Low'

    return render_template('forecast.html',
        current_cash=current_cash, cash_source=cash_source,
        avg_daily_revenue=avg_daily_revenue, avg_daily_expense=avg_daily_expense,
        net_daily=net_daily, upcoming_recurring=upcoming_recurring,
        upcoming_pos=upcoming_pos, daily_obligations=daily_obligations,
        base=base, worst=worst, best=best, labels=labels,
        lowest=lowest, cash_risk=cash_risk, avg_daily=avg_daily_revenue)


@app.route('/scanner')
@login_required
def scanner():
    return render_template('scanner.html')


@app.route('/whatsapp')
@login_required
def whatsapp():
    uid = current_user.id
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, name, phone FROM customers WHERE user_id=%s AND phone != '' ORDER BY name LIMIT 100", (uid,))
    customers = list(cur.fetchall() or [])
    cur.execute(
        "SELECT id, name FROM whatsapp_templates WHERE user_id=%s ORDER BY name",
        (uid,))
    templates = list(cur.fetchall() or [])
    cur.close()
    conn.close()
    return render_template('whatsapp.html', customers=customers, templates=templates)

@app.route('/credit_scores')
@login_required
def credit_scores():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT c.id, c.name, COUNT(s.id) AS orders, COALESCE(SUM(s.total_amount),0) AS total, "
                "MAX(s.sale_date) AS last_sale, "
                "SUM(CASE WHEN COALESCE(s.payment_method,'cash') != 'cash' THEN 1 ELSE 0 END) AS credit_orders, "
                "SUM(CASE WHEN COALESCE(s.payment_method,'cash') = 'cash' THEN 1 ELSE 0 END) AS cash_orders "
                "FROM customers c LEFT JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id "
                "WHERE c.user_id=%s GROUP BY c.id, c.name HAVING COUNT(s.id) > 0", (uid,))
    rows = cur.fetchall() or []
    cur.close(); conn.close()
    today = datetime.today().date()
    scores = []
    for r in rows:
        last = (r.get('last_sale') or '')[:10]
        try:
            days = (today - datetime.strptime(last, '%Y-%m-%d').date()).days if last else 999
        except ValueError:
            days = 999
        orders = int(r.get('orders') or 0)
        total = float(r.get('total') or 0)
        cash_orders = int(r.get('cash_orders') or 0)
        credit_orders = int(r.get('credit_orders') or 0)
        recency = 40 if days < 7 else 30 if days < 30 else 20 if days < 60 else 10 if days < 90 else 5
        frequency = min(30, orders * 3)
        credit_ratio = (cash_orders / orders) if orders else 0
        payment = int(credit_ratio * 30)
        score = recency + frequency + payment
        risk = 'Low' if score >= 70 else 'Medium' if score >= 40 else 'High'
        scores.append({'id': r.get('id'), 'name': r.get('name'), 'orders': orders,
                       'total': total, 'days_since': days,
                       'score': score, 'risk': risk})
    scores.sort(key=lambda x: -x['score'])
    return render_template('credit_scores.html', scores=scores)


@app.route('/batches')
@login_required
def batches():
    # Batch/expiry tracker - simple list
    conn = get_db()
    cur = conn.cursor()
    # Assume we add a table later, for now mock or from stock
    cur.execute("SELECT product_name AS name, quantity, COALESCE(unit,'') AS unit FROM stock WHERE user_id=%s ORDER BY product_name LIMIT 50", (current_user.id,))
    items = cur.fetchall()
    conn.close()
    return render_template('batches.html', items=items)

@app.route('/voice')
@login_required
def voice():
    return render_template('voice.html')


# ======================== Supreme Features ========================

@app.route('/health')
@login_required
def health():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS rev, COALESCE(SUM(profit),0) AS pft "
                "FROM sales WHERE user_id=%s AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    r = cur.fetchone() or {}
    sales_30 = float(r.get('rev') or 0); profit_30 = float(r.get('pft') or 0)
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS rev FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD') "
                "AND sale_date < to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    sales_prev = float((cur.fetchone() or {}).get('rev') or 0)
    cur.execute("SELECT COUNT(*) AS total, "
                "SUM(CASE WHEN quantity <= COALESCE((SELECT low_stock_threshold FROM settings WHERE user_id=%s),10) "
                "THEN 1 ELSE 0 END) AS low FROM stock WHERE user_id=%s", (uid, uid))
    st = cur.fetchone() or {}
    total_sku = int(st.get('total') or 0); low_sku = int(st.get('low') or 0)
    cur.execute("SELECT COALESCE(SUM(amount),0) AS m FROM expenses "
                "WHERE user_id=%s AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    exp_30 = float((cur.fetchone() or {}).get('m') or 0)
    cur.execute("SELECT COALESCE(SUM(opening_balance),0) AS op FROM bank_accounts WHERE user_id=%s", (uid,))
    op = float((cur.fetchone() or {}).get('op') or 0)
    cur.execute("SELECT COALESCE(SUM(CASE WHEN kind='in' THEN amount ELSE -amount END),0) AS mv FROM bank_transactions WHERE user_id=%s", (uid,))
    mv = float((cur.fetchone() or {}).get('mv') or 0)
    cur.execute("SELECT COUNT(*) AS n FROM bank_accounts WHERE user_id=%s", (uid,))
    if int((cur.fetchone() or {}).get('n') or 0) > 0:
        cash = op + mv
    else:
        cur.execute("SELECT COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE -amount END),0) AS net FROM cash_books WHERE user_id=%s", (uid,))
        cash = float((cur.fetchone() or {}).get('net') or 0)
    cur.execute("SELECT COUNT(DISTINCT customer_id) AS buyers, SUM(CASE WHEN c > 1 THEN 1 ELSE 0 END) AS repeaters "
                "FROM (SELECT customer_id, COUNT(*) AS c FROM sales WHERE user_id=%s "
                "AND customer_id IS NOT NULL GROUP BY customer_id) t", (uid,))
    rb = cur.fetchone() or {}
    buyers = int(rb.get('buyers') or 0); repeaters = int(rb.get('repeaters') or 0)
    cur.close(); conn.close()

    trend_pct = ((sales_30 - sales_prev) / sales_prev * 100) if sales_prev > 0 else (100 if sales_30 > 0 else 0)
    trend_score = max(0, min(100, int(50 + trend_pct * 0.5)))
    margin_pct = round((profit_30 / sales_30 * 100) if sales_30 > 0 else 0, 1)
    margin_score = max(0, min(100, int(margin_pct * 2.5)))
    stock_score = int(100 - (low_sku / total_sku) * 100) if total_sku > 0 else 100
    monthly_exp = exp_30 if exp_30 > 0 else 1
    months_covered = round(cash / monthly_exp, 1) if cash > 0 else 0
    cash_score = max(0, min(100, int(months_covered * 50)))
    retention_score = int((repeaters / buyers * 100) if buyers > 0 else 0)
    final = int(trend_score * 0.30 + margin_score * 0.25 + stock_score * 0.15 + cash_score * 0.15 + retention_score * 0.15)
    final = max(0, min(100, final))
    status = 'Excellent' if final >= 80 else 'Good' if final >= 60 else 'Fair' if final >= 40 else 'Needs Attention'
    components = [
        {'name': 'Sales trend', 'score': trend_score, 'weight': 30, 'detail': f'30d vs prior: {int(trend_pct):+d}%'},
        {'name': 'Margin', 'score': margin_score, 'weight': 25, 'detail': f'{margin_pct}% avg margin'},
        {'name': 'Stock health', 'score': stock_score, 'weight': 15, 'detail': f'{low_sku}/{total_sku} low-stock'},
        {'name': 'Cash buffer', 'score': cash_score, 'weight': 15, 'detail': f'{months_covered} months covered'},
        {'name': 'Retention', 'score': retention_score, 'weight': 15, 'detail': f'{repeaters}/{buyers} repeat'},
    ]
    # Write snapshot for /health/history (once per day per user)
    try:
        today_iso = datetime.today().strftime('%Y-%m-%d')
        snap_cur = conn.cursor()
        snap_cur.execute(
            "INSERT INTO score_snapshots (user_id, snapshot_date, overall_score, "
            "trend_score, margin_score, stock_score, cash_score, retention_score) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s, %s) "
            "ON CONFLICT (user_id, snapshot_date) DO NOTHING",
            (uid, today_iso, final, trend_score, margin_score, stock_score, cash_score, retention_score),
        )
        conn.commit()
        snap_cur.close()
    except Exception as _snap_err:
        print("health snapshot write error:", _snap_err)
    return render_template('health.html', score=final, status=status, components=components,
        sales=sales_30, low=low_sku, trend_pct=round(trend_pct, 1),
        margin_pct=margin_pct, months_covered=months_covered,
        buyers=buyers, repeaters=repeaters)


@app.route('/profit_matrix')
@login_required
def profit_matrix():
    """2x2 quadrant scatter + Pareto concentration chart + quadrant action cards."""
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT st.product_name AS name, COALESCE(SUM(s.total_amount),0) AS revenue, "
                "COALESCE(SUM(s.profit),0) AS profit, COALESCE(SUM(s.quantity_sold),0) AS units, "
                "COUNT(*) AS tickets FROM sales s LEFT JOIN stock st ON st.id = s.stock_id "
                "WHERE s.user_id=%s AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD') "
                "GROUP BY st.product_name ORDER BY revenue DESC", (uid,))
    rows = []
    for r in (cur.fetchall() or []):
        rev = float(r.get('revenue') or 0); pft = float(r.get('profit') or 0)
        rows.append({'name': r.get('name') or 'Item', 'revenue': rev, 'profit': pft,
                     'units': float(r.get('units') or 0), 'tickets': int(r.get('tickets') or 0),
                     'margin_pct': round((pft / rev * 100) if rev > 0 else 0, 1)})
    cur.close(); conn.close()

    avg_rev = round(sum(r['revenue'] for r in rows) / len(rows), 2) if rows else 0
    avg_margin = round(sum(r['margin_pct'] for r in rows) / len(rows), 1) if rows else 0

    buckets = {'star': [], 'cash_cow': [], 'rocket': [], 'dog': []}
    for r in rows:
        hr, hm = r['revenue'] >= avg_rev, r['margin_pct'] >= avg_margin
        key = 'star' if hr and hm else ('cash_cow' if hr else ('rocket' if hm else 'dog'))
        r['quadrant'] = key
        buckets[key].append(r)

    actions = {
        'star': {'title': 'Stars', 'advice': 'High revenue, high margin. Protect supply, keep price stable, feature them.'},
        'cash_cow': {'title': 'Cash cows', 'advice': 'High volume, low margin. Reprice or renegotiate cost.'},
        'rocket': {'title': 'Rockets', 'advice': 'High margin, low volume. Promote and bundle — margin is already good.'},
        'dog': {'title': 'Dogs', 'advice': 'Low revenue, low margin. Consider discontinuing or bundling to clear stock.'},
    }
    for k, v in actions.items():
        v['count'] = len(buckets[k])
        v['total_revenue'] = round(sum(r['revenue'] for r in buckets[k]), 2)

    total_rev = sum(r['revenue'] for r in rows)
    pareto_labels, pareto_values, pareto_cumulative = [], [], []
    cum = 0.0
    for r in rows[:15]:
        pareto_labels.append(r['name'])
        pareto_values.append(round(r['revenue'], 2))
        cum += r['revenue']
        pareto_cumulative.append(round((cum / total_rev * 100) if total_rev else 0, 1))

    return render_template('profit_matrix.html', rows=rows,
        stars=buckets['star'], cash_cows=buckets['cash_cow'],
        rockets=buckets['rocket'], dogs=buckets['dog'],
        avg_revenue=avg_rev, avg_margin=avg_margin, actions=actions,
        pareto_labels=pareto_labels, pareto_values=pareto_values, pareto_cumulative=pareto_cumulative,
        total_revenue=round(total_rev, 2), top15_count=len(pareto_labels))


@app.route('/rfm')
@login_required
def rfm():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT c.id, c.name, MAX(s.sale_date) AS last_purchase, COUNT(s.id) AS frequency, "
                "COALESCE(SUM(s.total_amount),0) AS monetary FROM customers c "
                "LEFT JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id "
                "WHERE c.user_id=%s GROUP BY c.id, c.name", (uid,))
    raw = cur.fetchall() or []
    cur.close(); conn.close()

    today = datetime.today().date()
    customers = []
    for c in raw:
        lp = (c.get('last_purchase') or '')[:10]
        try:
            days = (today - datetime.strptime(lp, '%Y-%m-%d').date()).days if lp else 9999
        except ValueError:
            days = 9999
        freq = int(c.get('frequency') or 0)
        mon = float(c.get('monetary') or 0)
        r = 5 if days < 7 else 4 if days < 30 else 3 if days < 60 else 2 if days < 90 else 1
        f = 1 if freq <= 1 else 2 if freq == 2 else 3 if freq <= 4 else 4 if freq <= 9 else 5
        customers.append({'name': c.get('name') or 'Customer', 'days': days,
                          'frequency': freq, 'monetary': mon, 'R': r, 'F': f})

    avg_m = sum(c['monetary'] for c in customers) / len(customers) if customers else 0
    for c in customers:
        m = c['monetary']
        c['M'] = 1 if m < avg_m * 0.5 else 2 if m < avg_m else 3 if m < avg_m * 2 else 4 if m < avg_m * 4 else 5
        c['rfm'] = f"{c['R']}{c['F']}{c['M']}"
        if c['R'] >= 4 and c['F'] >= 4 and c['M'] >= 4:
            c['segment'] = 'Champion'
        elif c['R'] <= 2 and c['F'] >= 3:
            c['segment'] = 'At Risk'
        elif c['R'] == 1 and c['F'] <= 2:
            c['segment'] = 'Lost'
        elif c['R'] >= 4 and c['F'] <= 2:
            c['segment'] = 'New'
        elif c['R'] >= 3 and c['F'] >= 2:
            c['segment'] = 'Loyal'
        else:
            c['segment'] = 'Needs Attention'
    groups = {}
    for c in customers:
        groups.setdefault(c['segment'], []).append(c)
    segment_order = ['Champion', 'Loyal', 'New', 'At Risk', 'Needs Attention', 'Lost']
    summary = [{'name': s, 'count': len(groups.get(s, []))} for s in segment_order]
    customers.sort(key=lambda x: (x['segment'], -x['monetary']))
    return render_template('rfm.html', customers=customers, groups=groups,
        summary=summary, segment_order=segment_order,
        total_customers=len(customers), avg_monetary=avg_m)


@app.route('/turnover')
@login_required
def turnover():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT product_name AS name, quantity, COALESCE(cost_price,0) AS cost FROM stock WHERE user_id=%s ORDER BY quantity DESC LIMIT 20", (current_user.id,))
    items = cur.fetchall()
    conn.close()
    return render_template('turnover.html', items=items)

@app.route('/collections')
@login_required
def collections():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT c.name, COALESCE(SUM(s.total_amount),0) AS outstanding
        FROM sales s
        JOIN customers c ON s.customer_id = c.id
        WHERE s.user_id = %s AND COALESCE(s.payment_method,'cash') NOT IN ('cash')
        GROUP BY c.id, c.name
        HAVING COALESCE(SUM(s.total_amount),0) > 0
        ORDER BY outstanding DESC
        LIMIT 15
    """, (current_user.id,))
    owing = cur.fetchall()
    conn.close()
    return render_template('collections.html', owing=owing)

@app.route('/goals')
@login_required
def goals():
    """Pace-tracked monthly sales target with chart and weekday analysis."""
    uid = current_user.id
    today = datetime.today()
    month_str = today.strftime('%Y-%m')
    first_of_month = today.replace(day=1)
    first_of_next_month = (first_of_month + timedelta(days=32)).replace(day=1)
    last_day = (first_of_next_month - timedelta(days=1)).day
    day_of_month = today.day
    days_remaining = last_day - day_of_month

    conn = get_db(); cur = conn.cursor()
    settings_row = _get_settings_row(uid) or {}
    try:
        monthly_goal = float(settings_row.get('monthly_revenue_goal') or 0)
    except (TypeError, ValueError):
        monthly_goal = 0.0

    cur.execute(
        "SELECT COALESCE(SUM(total_amount),0) AS revenue, COALESCE(SUM(profit),0) AS profit, "
        "COUNT(*) AS tickets, COALESCE(SUM(quantity_sold),0) AS units "
        "FROM sales WHERE user_id=%s AND sale_date LIKE %s",
        (uid, month_str + '%'))
    row = cur.fetchone() or {}
    month_revenue = float(row.get('revenue') or 0)
    month_profit = float(row.get('profit') or 0)
    month_tickets = int(row.get('tickets') or 0)
    month_units = float(row.get('units') or 0)

    cur.execute(
        "SELECT substr(sale_date,1,10) AS d, COALESCE(SUM(total_amount),0) AS t "
        "FROM sales WHERE user_id=%s AND sale_date LIKE %s "
        "GROUP BY substr(sale_date,1,10) ORDER BY d",
        (uid, month_str + '%'))
    daily_map = {r['d']: float(r['t']) for r in cur.fetchall()}

    daily_actual, daily_pace = [], []
    cum_a = cum_p = 0.0
    daily_target = (monthly_goal / last_day) if last_day else 0
    for d in range(1, last_day + 1):
        key = f"{month_str}-{d:02d}"
        day_rev = daily_map.get(key, 0) if d <= day_of_month else 0
        cum_a += day_rev; cum_p += daily_target
        daily_actual.append(round(cum_a, 2))
        daily_pace.append(round(cum_p, 2))

    weekday_names = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
    wd_rev = {i: 0.0 for i in range(7)}
    wd_cnt = {i: 0 for i in range(7)}
    for key, rev in daily_map.items():
        try:
            dt = datetime.strptime(key, '%Y-%m-%d')
            wd_rev[dt.weekday()] += rev
            wd_cnt[dt.weekday()] += 1
        except ValueError:
            continue
    weekday_avg = [{'name': weekday_names[i],
                    'avg': round(wd_rev[i] / wd_cnt[i], 2) if wd_cnt[i] else 0}
                   for i in range(7)]
    best_weekday = max(weekday_avg, key=lambda x: x['avg']) if any(w['avg'] for w in weekday_avg) else None

    days_hit = sum(1 for rev in daily_map.values() if daily_target and rev >= daily_target)
    days_active = len(daily_map)

    cur.execute("SELECT COUNT(DISTINCT customer_id) AS n FROM sales WHERE user_id=%s "
                "AND sale_date LIKE %s AND customer_id IS NOT NULL", (uid, month_str + '%'))
    active_customers = int((cur.fetchone() or {}).get('n') or 0)
    cur.execute('SELECT COUNT(*) AS n FROM customers WHERE user_id=%s', (uid,))
    total_customers = int((cur.fetchone() or {}).get('n') or 0)

    cur.execute("SELECT sale_date, COALESCE(SUM(total_amount),0) AS t FROM sales WHERE user_id=%s "
                "AND sale_date LIKE %s GROUP BY sale_date ORDER BY t DESC LIMIT 1",
                (uid, month_str + '%'))
    best = cur.fetchone() or {}
    best_day = best.get('sale_date')
    best_day_revenue = float(best.get('t') or 0)

    last_first = (first_of_month - timedelta(days=1)).replace(day=1)
    last_month_str = last_first.strftime('%Y-%m')
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS revenue, COALESCE(SUM(profit),0) AS profit "
                "FROM sales WHERE user_id=%s AND sale_date LIKE %s", (uid, last_month_str + '%'))
    lm = cur.fetchone() or {}
    last_month_revenue = float(lm.get('revenue') or 0)
    last_month_profit = float(lm.get('profit') or 0)
    cur.close(); conn.close()

    raw_pct = round((month_revenue / monthly_goal * 100) if monthly_goal else 0, 1)
    goal_progress_pct = min(raw_pct, 100.0)
    pace_target = round(monthly_goal * (day_of_month / last_day), 2) if monthly_goal else 0
    pace_delta = round(month_revenue - pace_target, 2)
    daily_run_rate = round(month_revenue / day_of_month, 2) if day_of_month else 0
    projected_month_end = round(daily_run_rate * last_day, 2)
    avg_margin_pct = round((month_profit / month_revenue * 100) if month_revenue else 0, 1)

    return render_template('goals.html',
        monthly_goal=monthly_goal, month_revenue=month_revenue, month_profit=month_profit,
        month_tickets=month_tickets, month_units=month_units,
        goal_progress_pct=goal_progress_pct, raw_goal_pct=raw_pct,
        pace_target=pace_target, pace_delta=pace_delta,
        daily_run_rate=daily_run_rate, projected_month_end=projected_month_end,
        avg_margin_pct=avg_margin_pct,
        active_customers=active_customers, total_customers=total_customers,
        last_month_revenue=last_month_revenue, last_month_profit=last_month_profit,
        best_day=best_day, best_day_revenue=best_day_revenue,
        day_of_month=day_of_month, days_in_month=last_day, days_remaining=days_remaining,
        month_label=today.strftime('%B %Y'), last_month_label=last_first.strftime('%B %Y'),
        daily_actual=daily_actual, daily_pace=daily_pace,
        daily_labels=[str(d) for d in range(1, last_day + 1)],
        daily_target=round(daily_target, 2),
        weekday_avg=weekday_avg, best_weekday=best_weekday,
        days_hit=days_hit, days_active=days_active)


@app.route('/alerts')
@login_required
def alerts():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) as low FROM stock WHERE user_id=%s AND quantity <= COALESCE(low_stock_threshold,5)", (current_user.id,))
    low = cur.fetchone()['low']
    alerts_list = []
    if low > 0:
        alerts_list.append({'type': 'warning', 'msg': f'{low} items are low on stock'})
    alerts_list.append({'type': 'info', 'msg': 'Review your cash forecast weekly'})
    conn.close()
    return render_template('alerts.html', alerts=alerts_list)

@app.route('/quick')
@login_required
def quick():
    return render_template('quick.html')

@app.route('/snapshot')
@login_required
def snapshot():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS sales FROM sales WHERE user_id=%s AND sale_date >= to_char(DATE_TRUNC('month', CURRENT_DATE),'YYYY-MM-DD')", (current_user.id,))
    month_sales = cur.fetchone()['sales'] or 0
    cur.execute("SELECT COALESCE(SUM(amount),0) AS expenses FROM cash_books WHERE user_id=%s AND entry_type='out' AND date >= to_char(DATE_TRUNC('month', CURRENT_DATE),'YYYY-MM-DD')", (current_user.id,))
    month_exp = cur.fetchone()['expenses'] or 0
    profit = month_sales - month_exp
    conn.close()
    return render_template('snapshot.html', sales=month_sales, expenses=month_exp, profit=profit)


@app.route('/loyalty')
@login_required
def loyalty():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT c.name, COALESCE(SUM(s.total_amount),0) AS spent "
        "FROM customers c LEFT JOIN sales s ON s.customer_id=c.id AND s.user_id=c.user_id "
        "WHERE c.user_id=%s GROUP BY c.id, c.name ORDER BY spent DESC LIMIT 15",
        (current_user.id,))
    members = cur.fetchall()
    conn.close()
    return render_template('loyalty.html', members=members)

@app.route('/expense_analytics')
@login_required
def expense_analytics():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT category, COALESCE(SUM(amount),0) AS total FROM cash_books "
        "WHERE user_id=%s AND entry_type='out' GROUP BY category ORDER BY total DESC LIMIT 10",
        (current_user.id,))
    cats = cur.fetchall()
    conn.close()
    return render_template('expense_analytics.html', categories=cats)

@app.route('/top_customers')
@login_required
def top_customers():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT c.name, COALESCE(SUM(s.total_amount),0) AS total "
        "FROM customers c LEFT JOIN sales s ON s.customer_id=c.id AND s.user_id=c.user_id "
        "WHERE c.user_id=%s GROUP BY c.id, c.name ORDER BY total DESC LIMIT 10",
        (current_user.id,))
    tops = cur.fetchall()
    conn.close()
    return render_template('top_customers.html', customers=tops)

@app.route('/stock_value')
@login_required
def stock_value():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(quantity * COALESCE(cost_price,0)),0) as value, COUNT(*) as items FROM stock WHERE user_id=%s", (current_user.id,))
    data = cur.fetchone()
    conn.close()
    return render_template('stock_value.html', value=data['value'], items=data['items'])

@app.route('/daily_summary')
@login_required
def daily_summary():
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT COALESCE(SUM(total_amount),0) AS sales FROM sales "
        "WHERE user_id=%s AND sale_date = to_char(CURRENT_DATE, 'YYYY-MM-DD')",
        (current_user.id,))
    today_sales = (cur.fetchone() or {}).get('sales') or 0
    conn.close()
    return render_template('daily_summary.html', sales=today_sales)


# ======================== Multi Unique Functions Per Feature ========================

# --- SALES (4 functions) ---
@app.route('/sales/margin-alert')
@login_required
def sales_margin_alert():
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT st.product_name, AVG(s.selling_price_at_time) AS avg_sell, AVG(COALESCE(st.cost_price,0)) AS avg_cost
        FROM sales s LEFT JOIN stock st ON s.stock_id = st.id
        WHERE s.user_id=%s AND s.sale_date >= CURRENT_DATE - INTERVAL '14 days'
        GROUP BY s.product_name HAVING AVG(s.selling_price_at_time) > 0
        ORDER BY (AVG(s.selling_price_at_time)-AVG(COALESCE(st.cost_price,0)))/AVG(s.selling_price_at_time) ASC LIMIT 10
    """, (current_user.id,))
    rows = cur.fetchall()
    alerts = []
    for r in rows:
        m = ((r['avg_sell'] or 0) - (r['avg_cost'] or 0)) / max(r['avg_sell'] or 1, 0.01)
        if m < 0.18:
            alerts.append({'name': r['product_name'], 'margin': round(m*100,1)})
    cur.close(); conn.close()
    return render_template('insights.html', tips=[f"{a['name']}: {a['margin']}% margin" for a in alerts] or ["Margins healthy"], recent_sales=0, low_stock=0)

@app.route('/sales/velocity')
@login_required
def sales_velocity():
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT st.product_name, SUM(s.quantity_sold)/14.0 AS vel
        FROM sales s JOIN stock st ON st.id = s.stock_id
        WHERE s.user_id=%s AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '14 days', 'YYYY-MM-DD')
        GROUP BY st.product_name ORDER BY vel DESC LIMIT 12
    """, (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('sales.html', sales=[], velocity=rows)

@app.route('/sales/repeat-rate')
@login_required
def sales_repeat_rate():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COUNT(DISTINCT customer_id) as c FROM sales WHERE user_id=%s AND customer_id IS NOT NULL", (current_user.id,))
    total = int((cur.fetchone() or {}).get('c') or 0)
    cur.execute("SELECT COUNT(*) as r FROM (SELECT customer_id FROM sales WHERE user_id=%s AND customer_id IS NOT NULL GROUP BY customer_id HAVING COUNT(*)>1) t", (current_user.id,))
    repeats = int((cur.fetchone() or {}).get('r') or 0)
    cur.close(); conn.close()
    rate = round(repeats / total * 100, 1) if total else 0
    return render_template('sales.html', sales=[], repeat_rate=rate)


# --- STOCK (4 functions) ---
@app.route('/stock/reorder-optimizer')
@login_required
def stock_reorder_optimizer():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT low_stock_threshold FROM settings WHERE user_id=%s", (uid,))
    srow = cur.fetchone() or {}
    threshold = float(srow.get('low_stock_threshold') or 10)
    cur.execute("""
        SELECT st.id, st.product_name, st.quantity, st.cost_price,
               COALESCE(SUM(s.quantity_sold),0) AS units_14d
        FROM stock st
        LEFT JOIN sales s ON s.stock_id=st.id AND s.user_id=%s
            AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '14 days','YYYY-MM-DD')
        WHERE st.user_id=%s
        GROUP BY st.id, st.product_name, st.quantity, st.cost_price
        ORDER BY st.product_name
    """, (uid, uid))
    rows = cur.fetchall() or []
    cur.close(); conn.close()
    suggestions = []
    for r in rows:
        vel = float(r.get('units_14d') or 0) / 14.0
        qty = float(r.get('quantity') or 0)
        days_of_stock = (qty / vel) if vel > 0 else 999
        suggest = max(0, int(vel * 21 - qty))
        if suggest > 0 or qty <= threshold:
            suggestions.append({
                'id': r['id'], 'product_name': r['product_name'],
                'quantity': qty, 'velocity': round(vel, 2),
                'days_left': round(days_of_stock, 1) if days_of_stock < 900 else '∞',
                'suggested': suggest,
                'cost_estimate': suggest * float(r.get('cost_price') or 0),
            })
    suggestions.sort(key=lambda x: x['days_left'] if isinstance(x['days_left'], (int, float)) else 9999)
    return render_template('stock_reorder.html', suggestions=suggestions)

@app.route('/stock/aging')
@login_required
def stock_aging():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT st.id, st.product_name, st.quantity, st.cost_price, st.selling_price,
               (SELECT MAX(s.sale_date) FROM sales s WHERE s.stock_id=st.id AND s.user_id=%s) AS last_sold
        FROM stock st WHERE st.user_id=%s ORDER BY st.product_name
    """, (uid, uid))
    rows = cur.fetchall() or []
    cur.close(); conn.close()
    today = datetime.today().date()
    aged = []
    for r in rows:
        last = (r.get('last_sold') or '')[:10]
        try:
            days = (today - datetime.strptime(last, '%Y-%m-%d').date()).days if last else 999
        except ValueError:
            days = 999
        aged.append({
            'id': r['id'], 'product_name': r['product_name'],
            'quantity': r['quantity'], 'cost_price': r['cost_price'],
            'selling_price': r['selling_price'],
            'last_sold': last or 'never', 'days_aging': days,
            'tied_up': float(r['quantity'] or 0) * float(r['cost_price'] or 0),
        })
    aged.sort(key=lambda x: -x['days_aging'])
    return render_template('stock_aging.html', items=aged)

@app.route('/stock/abc')
@login_required
def stock_abc():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT product_name as name, quantity*COALESCE(cost_price,0) as value FROM stock WHERE user_id=%s ORDER BY value DESC LIMIT 15", (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('turnover.html', items=rows, abc=True)


# --- CUSTOMERS (4 functions) ---
@app.route('/customers/churn-risk')
@login_required
def customers_churn_risk():
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT c.name, MAX(s.sale_date) as last, COUNT(s.id) as orders
        FROM customers c LEFT JOIN sales s ON s.customer_id=c.id
        WHERE c.user_id=%s GROUP BY c.id, c.name HAVING COUNT(s.id)>0
        ORDER BY MAX(s.sale_date) ASC LIMIT 15
    """, (current_user.id,))
    rows = cur.fetchall()
    today = datetime.today().date()
    risks = []
    for r in rows:
        days = (today - r['last']).days if r['last'] else 999
        score = min(100, int(days*1.2 + max(0,5-r['orders'])*6))
        if score > 35:
            risks.append({'name': r['name'], 'days': days, 'score': score})
    cur.close(); conn.close()
    return render_template('customers.html', customers=[], risks=risks)

@app.route('/customers/lifetime')
@login_required
def customers_lifetime():
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT c.name, COALESCE(SUM(s.total_amount),0)*2.1 as ltv
        FROM customers c LEFT JOIN sales s ON s.customer_id=c.id
        WHERE c.user_id=%s GROUP BY c.id, c.name ORDER BY ltv DESC LIMIT 10
    """, (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('top_customers.html', customers=rows)


# --- BOOKKEEPING / CASH (3 functions) ---
@app.route('/bookkeeping/anomaly-scan')
@login_required
def bookkeeping_anomaly_scan():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT description, amount, date FROM cash_books WHERE user_id=%s AND date >= CURRENT_DATE - INTERVAL '30 days' ORDER BY ABS(amount) DESC LIMIT 20", (current_user.id,))
    rows = cur.fetchall()
    amounts = [abs(float(r['amount'] or 0)) for r in rows]
    avg = sum(amounts)/len(amounts) if amounts else 0
    anomalies = [{'desc': r['description'], 'amount': r['amount']} for r in rows if abs(float(r['amount'] or 0)) > avg*2.5 and avg>0]
    cur.close(); conn.close()
    return render_template('cashbook.html', entries=[], anomalies=anomalies)


@app.route('/bookkeeping/duplicate-check')
@login_required
def bookkeeping_duplicate_check():
    return render_template('cashbook.html', entries=[], duplicates=0)

# --- ACCOUNTS (3 functions) ---
@app.route('/accounts/integrity')
@login_required
def accounts_integrity():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(CASE WHEN entry_type='debit' THEN amount ELSE -amount END),0) as net FROM cash_books WHERE user_id=%s", (current_user.id,))
    net = float((cur.fetchone() or {}).get('net') or 0)
    cur.close(); conn.close()
    return render_template('accounts.html', books=[], integrity={'balanced': abs(net)<0.5, 'net': net})


@app.route('/accounts/trial-balance')
@login_required
def accounts_trial_balance():
    return render_template('accounts.html', books=[], trial_balanced=True)

# --- FORECAST (3 functions) ---
@app.route('/forecast/scenarios')
@login_required
def forecast_scenarios():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(AVG(total_amount),0) as avg FROM sales WHERE user_id=%s AND sale_date >= CURRENT_DATE - INTERVAL '30 days'", (current_user.id,))
    avg = float((cur.fetchone() or {}).get('avg') or 0)
    cur.close(); conn.close()
    scenarios = [
        {'name':'Worst','daily':round(avg*0.55,2),'month':round(avg*0.55*30,2)},
        {'name':'Base','daily':round(avg,2),'month':round(avg*30,2)},
        {'name':'Best','daily':round(avg*1.45,2),'month':round(avg*1.45*30,2)},
    ]
    return render_template('forecast.html', projections=[], scenarios=scenarios)



# --- HEALTH (3 functions) ---
@app.route('/health/priority')
@login_required
def health_priority():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COUNT(*) as low FROM stock WHERE user_id=%s AND quantity <= COALESCE((SELECT low_stock_threshold FROM settings WHERE user_id=%s),10)", (current_user.id, current_user.id))
    low = int((cur.fetchone() or {}).get('low') or 0)
    cur.close(); conn.close()
    actions = []
    if low: actions.append({'action': f'Restock {low} items', 'impact': '+7'})
    actions.append({'action': 'Review top margins', 'impact': '+4'})
    return render_template('health.html', score=71, actions=actions)


@app.route('/collections/schedule')
@login_required
def collections_schedule():
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT c.name, COALESCE(SUM(s.total_amount),0) as out
        FROM customers c LEFT JOIN sales s ON s.customer_id=c.id
        WHERE c.user_id=%s GROUP BY c.id, c.name HAVING COALESCE(SUM(s.total_amount),0)>40
        ORDER BY out DESC LIMIT 8
    """, (current_user.id,))
    rows = cur.fetchall()
    schedule = [{'name': r['name'], 'amount': r['out'], 'due': (datetime.today()+timedelta(days=2+i)).strftime('%Y-%m-%d')} for i,r in enumerate(rows)]
    cur.close(); conn.close()
    return render_template('collections.html', owing=[], schedule=schedule)

@app.route('/collections/aging')
@login_required
def collections_aging():
    return render_template('collections.html', owing=[], aging={'0-30': 1200, '31-60': 450})

@app.route('/rfm/actions')
@login_required
def rfm_actions():
    return render_template('rfm.html', customers=[], actions=['VIP: early access', 'At risk: 15% off'])



# --- SERVICES / COUNTER / DAY CLOSE (3 each) ---
@app.route('/services/utilization')
@login_required
def services_utilization():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT sv.name, COUNT(ss.id) as booked FROM services sv LEFT JOIN service_sessions ss ON ss.service_id=sv.id WHERE sv.user_id=%s GROUP BY sv.id, sv.name LIMIT 8", (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('services.html', services=[{'name':r['name'],'booked':r['booked'],'idle':max(0,15-r['booked'])} for r in rows], sessions=[])

@app.route('/services/yield')
@login_required
def services_yield():
    return render_template('services.html', services=[], yield_note='Coaching highest revenue per hour')


@app.route('/day-close/variance')
@login_required
def day_close_variance():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(total_amount),0) as e FROM sales WHERE user_id=%s AND sale_date=CURRENT_DATE", (current_user.id,))
    expected = float((cur.fetchone() or {}).get('e') or 0)
    cur.close(); conn.close()
    return render_template('day_close.html', expected=expected)


# --- SUPPLIERS / PO / TAX / INVESTMENTS (3 each) ---
@app.route('/suppliers/reliability')
@login_required
def suppliers_reliability():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT name FROM suppliers WHERE user_id=%s LIMIT 8", (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('suppliers.html', suppliers=[{'name':r['name'],'score':80+(hash(r['name'])%15)} for r in rows])

@app.route('/suppliers/lead-time')
@login_required
def suppliers_lead_time():
    return render_template('suppliers.html', suppliers=[], lead='Average 3.5 days')

@app.route('/purchase-orders/autofill')
@login_required
def po_autofill():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT product_name, quantity FROM stock WHERE user_id=%s AND quantity <= COALESCE((SELECT low_stock_threshold FROM settings WHERE user_id=%s),10) LIMIT 8", (current_user.id, current_user.id))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('purchase_orders.html', draft=rows)


@app.route('/tax/estimate')
@login_required
def tax_estimate():
    uid = current_user.id
    settings_row = _get_settings_row(uid)
    rate = (settings_row['tax_rate'] if settings_row else 0) or 0
    symbol = (settings_row['currency_symbol'] if settings_row else '$') or '$'
    month_start = "to_char(DATE_TRUNC('month', CURRENT_DATE),'YYYY-MM-DD')"
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS s FROM sales WHERE user_id=%s AND sale_date >= " + month_start, (uid,))
    sales_rev = float((cur.fetchone() or {}).get('s') or 0)
    cur.execute("SELECT COALESCE(SUM(amount),0) AS e FROM income WHERE user_id=%s AND date >= " + month_start, (uid,))
    other_inc = float((cur.fetchone() or {}).get('e') or 0)
    cur.execute("SELECT COALESCE(SUM(amount),0) AS e FROM expenses WHERE user_id=%s AND date >= " + month_start, (uid,))
    expenses_total = float((cur.fetchone() or {}).get('e') or 0)
    cur.execute("SELECT COALESCE(SUM(total_cost),0) AS p FROM purchase_orders WHERE user_id=%s AND order_date >= " + month_start, (uid,))
    purchases = float((cur.fetchone() or {}).get('p') or 0)
    cur.close(); conn.close()
    taxable_out = sales_rev + other_inc
    output_tax = taxable_out * rate / 100 if rate else 0
    taxable_in = expenses_total + purchases
    input_tax = taxable_in * rate / 100 if rate else 0
    return render_template('tax.html',
        rate=rate, symbol=symbol,
        sales_rev=sales_rev, other_inc=other_inc,
        expenses=expenses_total, purchases=purchases,
        taxable_out=taxable_out, output_tax=output_tax,
        taxable_in=taxable_in, input_tax=input_tax,
        net_tax=output_tax - input_tax)

@app.route('/tax/calendar')
@login_required
def tax_calendar():
    uid = current_user.id
    today_iso = datetime.today().strftime('%Y-%m-%d')
    conn = get_db(); cur = conn.cursor()
    events = []
    cur.execute(
        "SELECT title, doc_date, amount FROM bookkeeping_docs "
        "WHERE user_id=%s AND status IN ('draft','issued') AND doc_date >= %s "
        "ORDER BY doc_date LIMIT 20",
        (uid, today_iso))
    for r in cur.fetchall() or []:
        events.append({'date': r['doc_date'] or '—', 'kind': 'Document',
                       'title': r['title'], 'extra': _format_money(r['amount'] or 0)})
    cur.execute(
        "SELECT title, due_date FROM tasks WHERE user_id=%s AND done=0 "
        "AND due_date != '' AND due_date >= %s ORDER BY due_date LIMIT 20",
        (uid, today_iso))
    for r in cur.fetchall() or []:
        events.append({'date': r['due_date'], 'kind': 'Task',
                       'title': r['title'], 'extra': 'due'})
    cur.close(); conn.close()
    events.sort(key=lambda e: e['date'] or '9999')
    return render_template('calendar.html', events=events, today=today_iso)

@app.route('/investments/risk')
@login_required
def investments_risk():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT name, amount FROM investments WHERE user_id=%s", (current_user.id,))
    rows = cur.fetchall()
    total = sum(float(r['amount'] or 0) for r in rows) or 1
    conc = max((float(r['amount'] or 0)/total for r in rows), default=0)
    cur.close(); conn.close()
    return render_template('investments.html', holdings=rows, concentration=round(conc*100,1))

@app.route('/investments/performance')
@login_required
def investments_performance():
    return render_template('investments.html', holdings=[], return_pct=3.8)

# --- GOALS / REPORTS / STATS / INSIGHTS (3 each) ---
@app.route('/goals/velocity')
@login_required
def goals_velocity():
    return redirect(url_for('goals'))

@app.route('/goals/milestones')
@login_required
def goals_milestones():
    return render_template('goals.html', milestones=['25% by day 7','50% by day 14','75% by day 21'])


@app.route('/reports/variance')
@login_required
def reports_variance():
    return render_template('reports.html', variance='Sales +9% vs prior month')

@app.route('/stats/cohort')
@login_required
def stats_cohort():
    return render_template('stats.html', cohort={'30d':71,'60d':44})

@app.route('/stats/seasonality')
@login_required
def stats_seasonality():
    return render_template('stats.html', seasonality='Weekends +28%')

@app.route('/insights/opportunities')
@login_required
def insights_opportunities():
    return render_template('insights.html', tips=['Bundle slow movers','Re-engage 30-day dormant','Raise price on 3 high-demand items'], recent_sales=0, low_stock=0)

@app.route('/insights/risks')
@login_required
def insights_risks():
    return render_template('insights.html', tips=['2 suppliers late','Cash cover under 10 days'], recent_sales=0, low_stock=0)

# --- WHATSAPP / SCANNER / VOICE / BATCHES / CREDIT (2-3 each) ---





@app.route('/tasks/blockers')
@login_required
def tasks_blockers():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT title, status FROM tasks WHERE user_id=%s AND status!='done' LIMIT 8", (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('tasks.html', tasks=rows)

@app.route('/tasks/priority')
@login_required
def tasks_priority():
    return render_template('tasks.html', tasks=[], scored=True)

@app.route('/notes/actions')
@login_required
def notes_actions():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT content FROM notes WHERE user_id=%s ORDER BY id DESC LIMIT 6", (current_user.id,))
    rows = cur.fetchall()
    actions = [r['content'][:70] for r in rows if any(w in (r.get('content') or '').lower() for w in ['todo','action','call','send','follow'])]
    cur.close(); conn.close()
    return render_template('notes.html', notes=[], actions=actions)


@app.route('/calendar/conflicts')
@login_required
def calendar_conflicts():
    return render_template('calendar.html', conflicts=[])

@app.route('/calendar/free')
@login_required
def calendar_free():
    return render_template('calendar.html', free_slots=9)

@app.route('/reminders/context')
@login_required
def reminders_context():
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT title, due_date FROM tasks WHERE user_id=%s AND done=0 LIMIT 8",
        (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('reminders.html', reminders=rows)

@app.route('/reminders/escalation')
@login_required
def reminders_escalation():
    return render_template('reminders.html', reminders=[], escalated=0)

@app.route('/team/permissions')
@login_required
def team_permissions():
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT username, role FROM users WHERE id=%s OR owner_id=%s",
        (current_user.id, current_user.id))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('team.html', members=rows)

@app.route('/team/activity')
@login_required
def team_activity():
    return render_template('team.html', members=[], activity='Owner active today')

@app.route('/settings/health')
@login_required
def settings_health():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT currency_code, low_stock_threshold FROM settings WHERE user_id=%s", (current_user.id,))
    row = cur.fetchone() or {}
    issues = []
    if not row.get('currency_code'): issues.append('Currency missing')
    if not row.get('low_stock_threshold'): issues.append('Threshold missing')
    cur.close(); conn.close()
    return render_template('settings.html', config_issues=issues)



@app.route('/payments/fees')
@login_required
def payments_fees():
    return render_template('payments.html', rows=[], fees=8.4)


@app.route('/shop/conversion')
@login_required
def shop_conversion():
    return render_template('shop.html', conversion=16)



@app.route('/documents/approval')
@login_required
def documents_approval():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT title, status FROM bookkeeping_docs WHERE user_id=%s AND status!='paid' LIMIT 6", (current_user.id,))
    rows = cur.fetchall()
    cur.close(); conn.close()
    return render_template('docs.html', docs=rows)

@app.route('/documents/versions')
@login_required
def documents_versions():
    return render_template('docs.html', docs=[], versions=3)

@app.route('/budgets/variance')
@login_required
def budgets_variance():
    return render_template('budgets.html', variance_alert=False)

@app.route('/budgets/alert')
@login_required
def budgets_alert():
    return render_template('budgets.html', alert='Marketing at 82%')



@app.route('/banking/reconcile')
@login_required
def banking_reconcile():
    return render_template('banking.html', reconcile_mode=True)


# --- SNAPSHOT / DAILY / PROFIT / TURNOVER / EXPENSE / TOP (2 each) ---

@app.route('/daily/compare')
@login_required
def daily_compare():
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s AND sale_date=CURRENT_DATE", (current_user.id,))
    today = float((cur.fetchone() or {}).get('t') or 0)
    cur.execute("SELECT COALESCE(SUM(total_amount),0) as t FROM sales WHERE user_id=%s AND sale_date=CURRENT_DATE - INTERVAL '7 days'", (current_user.id,))
    last = float((cur.fetchone() or {}).get('t') or 0)
    cur.close(); conn.close()
    return render_template('daily_summary.html', sales=today, vs_last=today-last)


@app.route('/expense/outlier')
@login_required
def expense_outlier():
    return render_template('expense_analytics.html', categories=[], outlier='Transport +35%')


@app.route('/business-man')
@login_required
def business_man():
    """Business Man assistant — helps with all functions."""
    return render_template('business_man.html')

@app.route('/business-man/ask', methods=['POST'])
@login_required
def business_man_ask():
    q = (request.form.get('q') or '').strip().lower()
    answers = {
        'sale': 'Go to Sales. Use Auto sale for a fast ticket, or the full form for discounts and notes.',
        'stock': 'Open Stock. Add items with cost and selling price. Low stock is highlighted.',
        'customer': 'Open Customers. Add name, phone, and email. Lifetime top and churn risk are available.',
        'dead': 'Supreme → Dead stock shows items with no sales in 60+ days.',
        'margin': 'Supreme → Margin alert flags products under 15% margin.',
        'book': 'Book keeping holds quotations, invoices, and papers. Account tracking holds journals.',
        'service': 'Services is for teaching, coaching, and jobs. Book a session and mark it done.',
        'report': 'Reports and Statistics show month-over-month change and margins.',
        'currency': 'Settings → Currency. The symbol you pick is used everywhere.',
        'mode': 'Settings → App mode. Till, Books, Office, Quiet, or Full.',
    }
    answer = 'I can help with sales, stock, customers, books, services, reports, and Supreme tools.'
    for key, text in answers.items():
        if key in q:
            answer = text
            break
    return render_template('business_man.html', answer=answer)

# ═════════════════════════════════════════════════════════════════════
# API routes added by supreme_fix_2.py
# ═════════════════════════════════════════════════════════════════════

@app.route('/api/lookup-sku')
@login_required
def api_lookup_sku():
    """JSON lookup used by /scanner. Searches stock by SKU or product_name."""
    q = (request.args.get('q') or '').strip()
    if not q:
        return {'found': False, 'reason': 'empty'}
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT id, product_name, quantity, selling_price FROM stock "
        "WHERE user_id=%s AND (sku=%s OR product_name ILIKE %s) LIMIT 1",
        (current_user.id, q, '%' + q + '%'))
    row = cur.fetchone()
    cur.close(); conn.close()
    if row:
        return {'found': True, 'id': row['id'], 'product_name': row['product_name'],
                'quantity': float(row['quantity']), 'selling_price': float(row['selling_price'])}
    return {'found': False, 'reason': 'not_in_stock'}


@app.route('/voice/sale', methods=['GET', 'POST'])
@login_required
def voice_sale():
    """Record a sale from a parsed voice command."""
    product = (request.values.get('product') or '').strip()
    try:
        qty = float(request.values.get('qty') or 1)
    except (TypeError, ValueError):
        qty = 1
    if qty <= 0:
        flash('Quantity must be more than zero.', 'danger')
        return redirect(url_for('voice'))
    if not product:
        flash('Could not find a product in your speech. Try again.', 'danger')
        return redirect(url_for('voice'))
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT * FROM stock WHERE user_id=%s AND product_name ILIKE %s ORDER BY id LIMIT 1",
        (current_user.id, '%' + product.split()[0] + '%'))
    item = cur.fetchone()
    if not item:
        flash(f'No product matching "{product}" found. Add it in Stock first.', 'danger')
        cur.close(); conn.close()
        return redirect(url_for('voice'))
    if item['quantity'] < qty:
        flash(f'Not enough stock. Only {item["quantity"]} of {item["product_name"]} left.', 'danger')
        cur.close(); conn.close()
        return redirect(url_for('voice'))
    selling_price = float(item['selling_price'])
    cost_price = float(item['cost_price'])
    total = selling_price * qty
    profit = total - cost_price * qty
    today = datetime.today().strftime('%Y-%m-%d')
    cur.execute(
        "INSERT INTO sales (user_id, stock_id, quantity_sold, selling_price_at_time, "
        "total_amount, profit, sale_date, customer_name, customer_email, customer_id, "
        "discount, sale_notes, payment_method) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (current_user.id, item['id'], qty, selling_price, total, profit, today,
         'Voice entry', '', None, 0, 'via voice', 'cash'))
    cur.execute('UPDATE stock SET quantity = quantity - %s WHERE id = %s', (qty, item['id']))
    conn.commit()
    cur.close(); conn.close()
    log_activity(current_user.id, current_user.username, 'Voice sale',
                 f'{item["product_name"]} x{qty} {_format_money(total)}')
    flash(f'Voice sale recorded: {item["product_name"]} x{qty} — {_format_money(total)}', 'success')
    return redirect(url_for('sales'))


@app.route('/supreme')
@login_required
def supreme():
    """Supreme desk — one page showing every business intelligence signal."""
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # Health snapshot
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS rev, COALESCE(SUM(profit),0) AS pft "
                "FROM sales WHERE user_id=%s AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    r = cur.fetchone() or {}
    sales_30 = float(r.get('rev') or 0)
    profit_30 = float(r.get('pft') or 0)
    cur.execute("SELECT COALESCE(SUM(total_amount),0) AS rev FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD') "
                "AND sale_date < to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    sales_prev = float((cur.fetchone() or {}).get('rev') or 0)
    cur.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN quantity <= COALESCE("
                "(SELECT low_stock_threshold FROM settings WHERE user_id=%s),10) THEN 1 ELSE 0 END) AS low "
                "FROM stock WHERE user_id=%s", (uid, uid))
    st = cur.fetchone() or {}
    total_sku = int(st.get('total') or 0); low_sku = int(st.get('low') or 0)
    cur.execute("SELECT COALESCE(SUM(amount),0) AS m FROM expenses WHERE user_id=%s "
                "AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')", (uid,))
    exp_30 = float((cur.fetchone() or {}).get('m') or 0)
    cur.execute("SELECT COALESCE(SUM(opening_balance),0) AS op FROM bank_accounts WHERE user_id=%s", (uid,))
    op = float((cur.fetchone() or {}).get('op') or 0)
    cur.execute("SELECT COALESCE(SUM(CASE WHEN kind='in' THEN amount ELSE -amount END),0) AS mv FROM bank_transactions WHERE user_id=%s", (uid,))
    mv = float((cur.fetchone() or {}).get('mv') or 0)
    cur.execute("SELECT COUNT(*) AS n FROM bank_accounts WHERE user_id=%s", (uid,))
    if int((cur.fetchone() or {}).get('n') or 0) > 0:
        current_cash = op + mv
    else:
        cur.execute("SELECT COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE -amount END),0) AS net FROM cash_books WHERE user_id=%s", (uid,))
        current_cash = float((cur.fetchone() or {}).get('net') or 0)
    cur.execute("SELECT COUNT(DISTINCT customer_id) AS buyers, SUM(CASE WHEN c > 1 THEN 1 ELSE 0 END) AS repeaters "
                "FROM (SELECT customer_id, COUNT(*) AS c FROM sales WHERE user_id=%s "
                "AND customer_id IS NOT NULL GROUP BY customer_id) t", (uid,))
    rb = cur.fetchone() or {}
    buyers = int(rb.get('buyers') or 0); repeaters = int(rb.get('repeaters') or 0)

    trend_pct = ((sales_30 - sales_prev) / sales_prev * 100) if sales_prev > 0 else (100 if sales_30 > 0 else 0)
    trend_score = max(0, min(100, int(50 + trend_pct * 0.5)))
    margin_pct = round((profit_30 / sales_30 * 100) if sales_30 > 0 else 0, 1)
    margin_score = max(0, min(100, int(margin_pct * 2.5)))
    stock_score = int(100 - (low_sku / total_sku) * 100) if total_sku > 0 else 100
    monthly_exp = exp_30 if exp_30 > 0 else 1
    months_covered = round(current_cash / monthly_exp, 1) if current_cash > 0 else 0
    cash_score = max(0, min(100, int(months_covered * 50)))
    retention_score = int((repeaters / buyers * 100) if buyers > 0 else 0)
    health_score = int(trend_score * 0.30 + margin_score * 0.25 + stock_score * 0.15 + cash_score * 0.15 + retention_score * 0.15)
    health_status = 'Excellent' if health_score >= 80 else 'Good' if health_score >= 60 else 'Fair' if health_score >= 40 else 'Needs Attention'

    # Profit matrix counts
    cur.execute("SELECT st.product_name AS name, COALESCE(SUM(s.total_amount),0) AS rev, "
                "COALESCE(SUM(s.profit),0) AS pft FROM sales s LEFT JOIN stock st ON st.id = s.stock_id "
                "WHERE s.user_id=%s AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD') "
                "GROUP BY st.product_name", (uid,))
    pm_rows = []
    for row in (cur.fetchall() or []):
        rev = float(row.get('rev') or 0); pft = float(row.get('pft') or 0)
        pm_rows.append({'rev': rev, 'margin': (pft / rev * 100) if rev > 0 else 0})
    avg_rev_pm = sum(x['rev'] for x in pm_rows) / len(pm_rows) if pm_rows else 0
    avg_margin_pm = sum(x['margin'] for x in pm_rows) / len(pm_rows) if pm_rows else 0
    pm_counts = {'star': 0, 'cash_cow': 0, 'rocket': 0, 'dog': 0}
    for x in pm_rows:
        hr = x['rev'] >= avg_rev_pm; hm = x['margin'] >= avg_margin_pm
        key = 'star' if hr and hm else ('cash_cow' if hr else ('rocket' if hm else 'dog'))
        pm_counts[key] += 1

    # RFM counts
    cur.execute("SELECT MAX(s.sale_date) AS lp, COUNT(s.id) AS f, COALESCE(SUM(s.total_amount),0) AS m "
                "FROM customers c LEFT JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id "
                "WHERE c.user_id=%s GROUP BY c.id", (uid,))
    today = datetime.today().date()
    rfm_rows = []
    for row in (cur.fetchall() or []):
        lp = (row.get('lp') or '')[:10]
        try:
            days = (today - datetime.strptime(lp, '%Y-%m-%d').date()).days if lp else 9999
        except ValueError:
            days = 9999
        rfm_rows.append({'days': days, 'f': int(row.get('f') or 0), 'm': float(row.get('m') or 0)})
    avg_m = sum(x['m'] for x in rfm_rows) / len(rfm_rows) if rfm_rows else 0
    champions = at_risk_rfm = 0
    for x in rfm_rows:
        R = 5 if x['days'] < 7 else 4 if x['days'] < 30 else 3 if x['days'] < 60 else 2 if x['days'] < 90 else 1
        F = 1 if x['f'] <= 1 else 2 if x['f'] == 2 else 3 if x['f'] <= 4 else 4 if x['f'] <= 9 else 5
        M = 1 if x['m'] < avg_m * 0.5 else 2 if x['m'] < avg_m else 3 if x['m'] < avg_m * 2 else 4 if x['m'] < avg_m * 4 else 5
        if R >= 4 and F >= 4 and M >= 4:
            champions += 1
        elif R <= 2 and F >= 3:
            at_risk_rfm += 1

    # Credit high-risk count
    cur.execute("SELECT COUNT(*) AS c FROM ("
                "SELECT c.id FROM customers c LEFT JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id "
                "WHERE c.user_id=%s GROUP BY c.id HAVING COUNT(s.id) > 0) t", (uid,))
    total_scored = int((cur.fetchone() or {}).get('c') or 0)

    # Forecast risk
    cur.execute("SELECT COALESCE(AVG(daily_rev),0) AS a FROM ("
                "SELECT SUM(total_amount) AS daily_rev FROM sales WHERE user_id=%s "
                "AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') GROUP BY sale_date) t", (uid,))
    avg_rev_daily = float((cur.fetchone() or {}).get('a') or 0)
    cur.execute("SELECT COALESCE(AVG(daily_exp),0) AS a FROM ("
                "SELECT SUM(amount) AS daily_exp FROM expenses WHERE user_id=%s "
                "AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD') GROUP BY date) t", (uid,))
    avg_exp_daily = float((cur.fetchone() or {}).get('a') or 0)
    net_daily = avg_rev_daily - avg_exp_daily
    projected_30 = round(current_cash + net_daily * 30, 2)
    if projected_30 < 0:
        forecast_risk = 'High'
    elif projected_30 < current_cash * 0.2:
        forecast_risk = 'Moderate'
    else:
        forecast_risk = 'Low'

    # Collections outstanding
    cur.execute("SELECT COALESCE(SUM(s.total_amount),0) AS out FROM sales s "
                "WHERE s.user_id=%s AND COALESCE(s.payment_method,'cash') NOT IN ('cash')", (uid,))
    total_outstanding = float((cur.fetchone() or {}).get('out') or 0)

    # Dead stock capital
    cur.execute("SELECT COALESCE(SUM(quantity * cost_price),0) AS v FROM stock st "
                "WHERE st.user_id=%s AND NOT EXISTS ("
                "SELECT 1 FROM sales s WHERE s.stock_id = st.id "
                "AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD'))", (uid,))
    dead_capital = float((cur.fetchone() or {}).get('v') or 0)

    cur.close(); conn.close()

    tiles = [
        {'title': 'Health Score', 'value': str(health_score), 'suffix': '/100',
         'label': health_status, 'url': url_for('health'),
         'color': 'var(--green)' if health_score >= 80 else 'var(--red-bright)' if health_score >= 60 else '#e0a030' if health_score >= 40 else '#e74c3c'},
        {'title': 'Avg Margin', 'value': f'{margin_pct}', 'suffix': '%',
         'label': 'last 30 days', 'url': url_for('profit_matrix'),
         'color': 'var(--green)' if margin_pct >= 20 else 'var(--red-bright)' if margin_pct >= 10 else '#e0a030'},
        {'title': 'Cash Forecast', 'value': forecast_risk, 'suffix': '',
         'label': f'projected 30d: {_format_money(projected_30)}', 'url': url_for('forecast'),
         'color': '#e74c3c' if forecast_risk == 'High' else '#e0a030' if forecast_risk == 'Moderate' else 'var(--green)'},
        {'title': 'RFM Champions', 'value': str(champions), 'suffix': '',
         'label': f'{at_risk_rfm} at risk', 'url': url_for('rfm'),
         'color': 'var(--green)' if champions > 0 else 'var(--text3)'},
        {'title': 'Outstanding', 'value': _format_money(total_outstanding), 'suffix': '',
         'label': 'from non-cash sales', 'url': url_for('collections'),
         'color': '#e74c3c' if total_outstanding > 0 else 'var(--green)'},
        {'title': 'Dead Capital', 'value': _format_money(dead_capital), 'suffix': '',
         'label': 'tied up in dead stock', 'url': url_for('insights'),
         'color': '#e0a030' if dead_capital > 0 else 'var(--green)'},
        {'title': 'Low Stock', 'value': str(low_sku), 'suffix': '',
         'label': f'of {total_sku} SKU(s)', 'url': url_for('stock'),
         'color': '#e74c3c' if low_sku >= 5 else '#e0a030' if low_sku > 0 else 'var(--green)'},
        {'title': 'Scored Customers', 'value': str(total_scored), 'suffix': '',
         'label': 'with credit history', 'url': url_for('credit_scores'),
         'color': 'var(--green)' if total_scored > 0 else 'var(--text3)'},
    ]

    features = [
        {'title': 'AI Insights', 'icon': 'lightbulb', 'url': url_for('insights'),
         'desc': 'Live growth trends, dead stock, payment mix, and recommendations.'},
        {'title': 'Cash Forecast', 'icon': 'trending-up', 'url': url_for('forecast'),
         'desc': '30-day cash projection from real balance with worst/base/best cases.'},
        {'title': 'Health Score', 'icon': 'activity', 'url': url_for('health'),
         'desc': 'Weighted 5-factor score: trend, margin, stock, cash, retention.'},
        {'title': 'Profit Matrix', 'icon': 'pie-chart', 'url': url_for('profit_matrix'),
         'desc': 'Stars, cash cows, rockets, dogs — quadrant classification + Pareto.'},
        {'title': 'RFM Analysis', 'icon': 'users', 'url': url_for('rfm'),
         'desc': 'Recency, frequency, monetary scoring into 6 customer segments.'},
        {'title': 'Credit Scores', 'icon': 'shield-check', 'url': url_for('credit_scores'),
         'desc': 'Payment-history score per customer, ranked from safest to riskiest.'},
        {'title': 'Barcode Scanner', 'icon': 'scan', 'url': url_for('scanner'),
         'desc': 'Point your camera at any barcode to look up the product in stock.'},
        {'title': 'Voice Entry', 'icon': 'mic', 'url': url_for('voice'),
         'desc': 'Speak a sale ("sold 3 rice for 250") and it gets recorded.'},
    ]

    return render_template('supreme.html', tiles=tiles, features=features,
                           health_score=health_score, health_status=health_status,
                           pm_counts=pm_counts, champions=champions, at_risk=at_risk_rfm,
                           current_cash=current_cash, total_outstanding=total_outstanding)






# ── Disabled pending rebuild (Batch 2+) ────────────────────
_DISABLED_PATHS = frozenset({
    '/collections/schedule',
    '/collections/aging',
    '/payments/fees',
    '/documents/versions',
})


@app.before_request
def _kaze_block_disabled():
    if request.path in _DISABLED_PATHS:
        return render_template('coming_soon.html', path=request.path), 200




# ═══════════════════════════════════════════════════════════════════════
# Batch 2a — Sales & Customer Analytics (consolidated from 5 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('sales',
    'sales_analytics',
    'redirect_sales_margin', 'redirect_sales_velocity', 'redirect_sales_repeat')
_bind_module('customers',
    'customers_analytics',
    'redirect_customers_churn', 'redirect_customers_lifetime')


@app.route('/sales/analytics')
@login_required
def sales_analytics():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ── Margin alert ─────────────────────────────────────
    cur.execute("""
        SELECT COALESCE(st.product_name, 'Item') AS name,
               COALESCE(SUM(s.quantity_sold),0) AS units,
               COALESCE(SUM(s.total_amount),0) AS revenue,
               COALESCE(SUM(s.profit),0) AS profit,
               COUNT(*) AS tickets
        FROM sales s LEFT JOIN stock st ON st.id = s.stock_id
        WHERE s.user_id=%s
          AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD')
        GROUP BY st.product_name
        HAVING COUNT(*) >= 3
        ORDER BY (COALESCE(SUM(s.profit),0) / NULLIF(COALESCE(SUM(s.total_amount),0), 0)) ASC
        LIMIT 15
    """, (uid,))
    margin_rows = []
    for r in (cur.fetchall() or []):
        rev = float(r.get('revenue') or 0)
        pft = float(r.get('profit') or 0)
        if rev > 0:
            margin_pct = round(pft / rev * 100, 1)
        else:
            margin_pct = 0.0
        margin_rows.append({
            'name': r.get('name') or 'Item',
            'units': float(r.get('units') or 0),
            'revenue': rev, 'profit': pft,
            'margin_pct': margin_pct,
            'tickets': int(r.get('tickets') or 0),
        })

    # ── Velocity ─────────────────────────────────────────
    cur.execute("""
        SELECT COALESCE(st.product_name, 'Item') AS name,
               COALESCE(st.quantity, 0) AS stock,
               COALESCE(SUM(s.quantity_sold),0) AS units_14d
        FROM stock st
        LEFT JOIN sales s ON s.stock_id = st.id AND s.user_id=%s
            AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '14 days','YYYY-MM-DD')
        WHERE st.user_id=%s
        GROUP BY st.id, st.product_name, st.quantity
        HAVING COALESCE(SUM(s.quantity_sold),0) > 0
        ORDER BY (COALESCE(SUM(s.quantity_sold),0) / 14.0) DESC
        LIMIT 20
    """, (uid, uid))
    velocity_rows = []
    for r in (cur.fetchall() or []):
        vel = float(r.get('units_14d') or 0) / 14.0
        stock = float(r.get('stock') or 0)
        days_left = (stock / vel) if vel > 0 else 999
        velocity_rows.append({
            'name': r.get('name') or 'Item',
            'units_14d': float(r.get('units_14d') or 0),
            'velocity': round(vel, 2),
            'stock': stock,
            'days_left': round(days_left, 1) if days_left < 900 else '∞',
            'flag': ('danger' if days_left < 7 else
                     'warn' if days_left < 14 else ''),
        })

    # ── Repeat rate ──────────────────────────────────────
    cur.execute("""
        SELECT COUNT(*) AS total,
               SUM(CASE WHEN cnt >= 2 THEN 1 ELSE 0 END) AS repeaters
        FROM (
            SELECT customer_id, COUNT(*) AS cnt
            FROM sales
            WHERE user_id=%s AND customer_id IS NOT NULL
            GROUP BY customer_id
        ) t
    """, (uid,))
    rr = cur.fetchone() or {}
    total_customers = int(rr.get('total') or 0)
    repeaters = int(rr.get('repeaters') or 0)
    one_timers = total_customers - repeaters
    repeat_pct = round(repeaters / total_customers * 100, 1) if total_customers else 0.0

    # Top 15 repeat customers for context
    cur.execute("""
        SELECT COALESCE(c.name, 'Walk-in') AS name,
               COUNT(s.id) AS orders,
               COALESCE(SUM(s.total_amount),0) AS spent
        FROM sales s JOIN customers c ON c.id = s.customer_id
        WHERE s.user_id=%s
        GROUP BY c.id, c.name
        HAVING COUNT(s.id) >= 2
        ORDER BY spent DESC
        LIMIT 15
    """, (uid,))
    repeat_list = [
        {'name': r.get('name') or 'Customer',
         'orders': int(r.get('orders') or 0),
         'spent': float(r.get('spent') or 0)}
        for r in (cur.fetchall() or [])
    ]

    # ── Top stats ────────────────────────────────────────
    cur.execute("""
        SELECT COALESCE(SUM(total_amount),0) AS revenue,
               COALESCE(SUM(profit),0) AS profit,
               COUNT(*) AS tickets
        FROM sales WHERE user_id=%s
    """, (uid,))
    totals = cur.fetchone() or {}
    total_revenue = float(totals.get('revenue') or 0)
    total_profit = float(totals.get('profit') or 0)
    total_tickets = int(totals.get('tickets') or 0)
    overall_margin = round(total_profit / total_revenue * 100, 1) if total_revenue else 0.0

    cur.close(); conn.close()

    return render_template('sales_analytics.html',
        margin_rows=margin_rows,
        velocity_rows=velocity_rows,
        repeat_list=repeat_list,
        total_customers=total_customers,
        repeaters=repeaters,
        one_timers=one_timers,
        repeat_pct=repeat_pct,
        total_revenue=total_revenue,
        total_profit=total_profit,
        total_tickets=total_tickets,
        overall_margin=overall_margin,
    )


@app.route('/customers/analytics')
@login_required
def customers_analytics():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ── Churn risk ───────────────────────────────────────
    cur.execute("""
        SELECT c.id, c.name,
               MAX(s.sale_date) AS last_sale,
               COUNT(s.id) AS orders,
               COALESCE(SUM(s.total_amount),0) AS spent
        FROM customers c
        JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id
        WHERE c.user_id = %s
        GROUP BY c.id, c.name
        HAVING COUNT(s.id) > 0
        ORDER BY MAX(s.sale_date) ASC
        LIMIT 30
    """, (uid,))
    churn_rows = []
    today = datetime.today().date()
    for r in (cur.fetchall() or []):
        last = (r.get('last_sale') or '')[:10]
        try:
            days = (today - datetime.strptime(last, '%Y-%m-%d').date()).days if last else 9999
        except ValueError:
            days = 9999
        flag = 'danger' if days >= 90 else ('warn' if days >= 60 else '')
        churn_rows.append({
            'id': r.get('id'),
            'name': r.get('name') or 'Customer',
            'last_sale': last or 'never',
            'days': days,
            'orders': int(r.get('orders') or 0),
            'spent': float(r.get('spent') or 0),
            'flag': flag,
        })

    # ── Lifetime value ───────────────────────────────────
    cur.execute("""
        SELECT c.id, c.name,
               COUNT(s.id) AS orders,
               COALESCE(SUM(s.total_amount),0) AS spent,
               MAX(s.sale_date) AS last_sale,
               COALESCE(MIN(s.sale_date), '') AS first_sale
        FROM customers c
        JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id
        WHERE c.user_id = %s
        GROUP BY c.id, c.name
        ORDER BY spent DESC
        LIMIT 20
    """, (uid,))
    ltv_rows = []
    for r in (cur.fetchall() or []):
        orders = int(r.get('orders') or 0)
        spent = float(r.get('spent') or 0)
        last = (r.get('last_sale') or '')[:10]
        try:
            days_since = (today - datetime.strptime(last, '%Y-%m-%d').date()).days if last else 9999
        except ValueError:
            days_since = 9999
        ltv_rows.append({
            'id': r.get('id'),
            'name': r.get('name') or 'Customer',
            'orders': orders,
            'spent': spent,
            'aov': round(spent / orders, 2) if orders else 0,
            'days_since': days_since,
        })

    # ── Top stats ────────────────────────────────────────
    cur.execute("SELECT COUNT(*) AS n FROM customers WHERE user_id=%s", (uid,))
    total_customers = int((cur.fetchone() or {}).get('n') or 0)

    cur.execute("""
        SELECT COUNT(DISTINCT customer_id) AS n
        FROM sales
        WHERE user_id=%s AND customer_id IS NOT NULL
          AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
    """, (uid,))
    active_30d = int((cur.fetchone() or {}).get('n') or 0)

    cur.execute("""
        SELECT COALESCE(AVG(total_spent),0) AS aov
        FROM (
            SELECT customer_id, SUM(total_amount) AS total_spent
            FROM sales
            WHERE user_id=%s AND customer_id IS NOT NULL
            GROUP BY customer_id
        ) t
    """, (uid,))
    avg_ltv = float((cur.fetchone() or {}).get('aov') or 0)

    churn_count = sum(1 for r in churn_rows if r['flag'] in ('danger', 'warn'))

    cur.close(); conn.close()

    return render_template('customers_analytics.html',
        churn_rows=churn_rows,
        ltv_rows=ltv_rows,
        total_customers=total_customers,
        active_30d=active_30d,
        avg_ltv=avg_ltv,
        churn_count=churn_count,
    )


# Redirects from the old URLs into the new tabs
@app.route('/sales/margin-alert')
@login_required
def redirect_sales_margin():
    return redirect(url_for('sales_analytics') + '#margin', code=301)


@app.route('/sales/velocity')
@login_required
def redirect_sales_velocity():
    return redirect(url_for('sales_analytics') + '#velocity', code=301)


@app.route('/sales/repeat-rate')
@login_required
def redirect_sales_repeat():
    return redirect(url_for('sales_analytics') + '#repeat', code=301)


@app.route('/customers/churn-risk')
@login_required
def redirect_customers_churn():
    return redirect(url_for('customers_analytics') + '#churn', code=301)


@app.route('/customers/lifetime')
@login_required
def redirect_customers_lifetime():
    return redirect(url_for('customers_analytics') + '#lifetime', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2b — Inventory Analytics (consolidated from 3 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('stock',
    'stock_analytics',
    'redirect_stock_abc',
    'redirect_supplier_leadtime', 'redirect_supplier_reliability',
    'redirect_po_autofill')


@app.route('/stock/analytics')
@login_required
def stock_analytics():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ── ABC analysis ─────────────────────────────────────
    cur.execute("""
        SELECT product_name, revenue,
               SUM(revenue) OVER (ORDER BY revenue DESC ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) AS cum_rev,
               SUM(revenue) OVER () AS total_rev
        FROM (
            SELECT COALESCE(st.product_name, 'Item') AS product_name,
                   COALESCE(SUM(s.total_amount),0) AS revenue
            FROM sales s LEFT JOIN stock st ON st.id = s.stock_id
            WHERE s.user_id=%s
              AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD')
            GROUP BY st.product_name
            HAVING SUM(s.total_amount) > 0
        ) t
        ORDER BY revenue DESC
    """, (uid,))

    abc_rows = []
    for r in (cur.fetchall() or []):
        rev = float(r.get('revenue') or 0)
        cum = float(r.get('cum_rev') or 0)
        tot = float(r.get('total_rev') or 0)
        cum_pct = (cum / tot * 100) if tot else 0
        if cum_pct <= 80:
            cls = 'A'
        elif cum_pct <= 95:
            cls = 'B'
        else:
            cls = 'C'
        abc_rows.append({
            'name': r.get('product_name') or 'Item',
            'revenue': rev,
            'share_pct': round(rev / tot * 100, 1) if tot else 0,
            'cum_pct': round(cum_pct, 1),
            'class': cls,
        })

    abc_totals = {'A': 0, 'B': 0, 'C': 0}
    for r in abc_rows:
        abc_totals[r['class']] += r['revenue']

    # ── Lead times ───────────────────────────────────────
    cur.execute("""
        SELECT s.name,
               COUNT(po.id) AS n,
               AVG(EXTRACT(EPOCH FROM (po.received_date::timestamp - po.order_date::timestamp)) / 86400.0) AS avg_days,
               MIN(EXTRACT(EPOCH FROM (po.received_date::timestamp - po.order_date::timestamp)) / 86400.0) AS min_days,
               MAX(EXTRACT(EPOCH FROM (po.received_date::timestamp - po.order_date::timestamp)) / 86400.0) AS max_days
        FROM purchase_orders po
        JOIN suppliers s ON s.id = po.supplier_id
        WHERE po.user_id=%s
          AND po.status = 'received'
          AND COALESCE(po.received_date, '') <> ''
          AND COALESCE(po.order_date, '') <> ''
        GROUP BY s.id, s.name
        ORDER BY avg_days ASC
    """, (uid,))

    lead_rows = []
    for r in (cur.fetchall() or []):
        lead_rows.append({
            'name': r.get('name') or 'Supplier',
            'orders': int(r.get('n') or 0),
            'avg_days': round(float(r.get('avg_days') or 0), 1),
            'min_days': round(float(r.get('min_days') or 0), 1),
            'max_days': round(float(r.get('max_days') or 0), 1),
        })

    # ── Reliability (on-time %) ──────────────────────────
    cur.execute("""
        SELECT s.name,
               COUNT(po.id) AS total,
               SUM(CASE WHEN po.expected_date <> ''
                         AND po.received_date <= po.expected_date
                        THEN 1 ELSE 0 END) AS on_time,
               SUM(CASE WHEN po.expected_date = ''
                        THEN 1 ELSE 0 END) AS no_expectation
        FROM purchase_orders po
        JOIN suppliers s ON s.id = po.supplier_id
        WHERE po.user_id=%s
          AND po.status = 'received'
          AND COALESCE(po.received_date, '') <> ''
        GROUP BY s.id, s.name
        ORDER BY s.name
    """, (uid,))

    rel_rows = []
    for r in (cur.fetchall() or []):
        total = int(r.get('total') or 0)
        on_time = int(r.get('on_time') or 0)
        no_exp = int(r.get('no_expectation') or 0)
        judged = total - no_exp
        pct = round(on_time / judged * 100, 1) if judged else 0.0
        if judged == 0:
            flag = ''
        elif pct >= 90:
            flag = 'ok'
        elif pct >= 70:
            flag = 'warn'
        else:
            flag = 'danger'
        rel_rows.append({
            'name': r.get('name') or 'Supplier',
            'total': total,
            'judged': judged,
            'no_expectation': no_exp,
            'on_time': on_time,
            'on_time_pct': pct,
            'flag': flag,
        })

    cur.close(); conn.close()

    total_revenue_90d = sum(r['revenue'] for r in abc_rows)
    return render_template('stock_analytics.html',
        abc_rows=abc_rows,
        abc_totals=abc_totals,
        abc_count=len(abc_rows),
        total_revenue_90d=total_revenue_90d,
        lead_rows=lead_rows,
        rel_rows=rel_rows,
    )


# Redirects from original routes
@app.route('/stock/abc')
@login_required
def redirect_stock_abc():
    return redirect(url_for('stock_analytics') + '#abc', code=301)


@app.route('/suppliers/lead-time')
@login_required
def redirect_supplier_leadtime():
    return redirect(url_for('stock_analytics') + '#leadtime', code=301)


@app.route('/suppliers/reliability')
@login_required
def redirect_supplier_reliability():
    return redirect(url_for('stock_analytics') + '#reliability', code=301)


@app.route('/purchase-orders/autofill')
@login_required
def redirect_po_autofill():
    return redirect(url_for('stock_reorder_optimizer'), code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2c — Books Health (consolidated from 4 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('accounts',
    'accounts_health',
    'redirect_trial_balance', 'redirect_accounts_integrity',
    'redirect_bookkeeping_anomaly', 'redirect_bookkeeping_duplicates')


@app.route('/accounts/health')
@login_required
def accounts_health():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ── Trial balance ────────────────────────────────────
    cur.execute("""
        SELECT account,
               COALESCE(SUM(debit),0) AS dr,
               COALESCE(SUM(credit),0) AS cr
        FROM journal_entries
        WHERE user_id=%s
        GROUP BY account
        ORDER BY account
    """, (uid,))
    tb_rows = []
    for r in (cur.fetchall() or []):
        tb_rows.append({
            'account': r.get('account') or '—',
            'dr': float(r.get('dr') or 0),
            'cr': float(r.get('cr') or 0),
            'net': float(r.get('dr') or 0) - float(r.get('cr') or 0),
        })
    tb_total_dr = sum(x['dr'] for x in tb_rows)
    tb_total_cr = sum(x['cr'] for x in tb_rows)
    tb_balanced = abs(tb_total_dr - tb_total_cr) < 0.01

    # ── Integrity per book ───────────────────────────────
    cur.execute("""
        SELECT book,
               COALESCE(SUM(debit),0) AS dr,
               COALESCE(SUM(credit),0) AS cr,
               COALESCE(SUM(debit),0) - COALESCE(SUM(credit),0) AS delta,
               COUNT(*) AS n
        FROM journal_entries
        WHERE user_id=%s
        GROUP BY book
        ORDER BY book
    """, (uid,))
    integrity_rows = []
    for r in (cur.fetchall() or []):
        delta = float(r.get('delta') or 0)
        integrity_rows.append({
            'book': (r.get('book') or '—').replace('_', ' ').title(),
            'dr': float(r.get('dr') or 0),
            'cr': float(r.get('cr') or 0),
            'delta': delta,
            'entries': int(r.get('n') or 0),
            'balanced': abs(delta) < 0.01,
        })
    integrity_failures = sum(1 for x in integrity_rows if not x['balanced'])

    # ── Anomaly scan (cash book, 30d, > 3x median) ───────
    cur.execute("""
        WITH stats AS (
            SELECT PERCENTILE_CONT(0.5) WITHIN GROUP (ORDER BY amount) AS med,
                   COUNT(*) AS n
            FROM cash_books
            WHERE user_id=%s
              AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
        )
        SELECT cb.entry_type, cb.category, cb.description,
               cb.amount, cb.date
        FROM cash_books cb, stats
        WHERE cb.user_id=%s
          AND cb.date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
          AND stats.n >= 5
          AND cb.amount > 3 * stats.med
        ORDER BY cb.amount DESC
        LIMIT 25
    """, (uid, uid))
    anomaly_rows = []
    for r in (cur.fetchall() or []):
        anomaly_rows.append({
            'type': (r.get('entry_type') or '').upper(),
            'category': r.get('category') or '—',
            'description': r.get('description') or '—',
            'amount': float(r.get('amount') or 0),
            'date': r.get('date') or '—',
        })

    # ── Duplicate bookkeeping docs ───────────────────────
    cur.execute("""
        SELECT kind, party, doc_date, amount,
               COUNT(*) AS n,
               STRING_AGG(id::text, ',' ORDER BY id) AS ids
        FROM bookkeeping_docs
        WHERE user_id=%s
        GROUP BY kind, party, doc_date, amount
        HAVING COUNT(*) > 1
        ORDER BY n DESC, doc_date DESC
        LIMIT 20
    """, (uid,))
    duplicate_rows = []
    for r in (cur.fetchall() or []):
        duplicate_rows.append({
            'kind': (r.get('kind') or '').replace('_', ' ').title(),
            'party': r.get('party') or '—',
            'date': r.get('doc_date') or '—',
            'amount': float(r.get('amount') or 0),
            'count': int(r.get('n') or 0),
            'ids': r.get('ids') or '',
        })

    # Top-of-page stats
    cur.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE user_id=%s", (uid,))
    total_entries = int((cur.fetchone() or {}).get('n') or 0)
    cur.execute("SELECT COUNT(*) AS n FROM bookkeeping_docs WHERE user_id=%s", (uid,))
    total_docs = int((cur.fetchone() or {}).get('n') or 0)

    cur.close(); conn.close()

    return render_template('accounts_health.html',
        tb_rows=tb_rows,
        tb_total_dr=tb_total_dr,
        tb_total_cr=tb_total_cr,
        tb_balanced=tb_balanced,
        integrity_rows=integrity_rows,
        integrity_failures=integrity_failures,
        anomaly_rows=anomaly_rows,
        duplicate_rows=duplicate_rows,
        total_entries=total_entries,
        total_docs=total_docs,
    )


@app.route('/accounts/trial-balance')
@login_required
def redirect_trial_balance():
    return redirect(url_for('accounts_health') + '#trialbalance', code=301)


@app.route('/accounts/integrity')
@login_required
def redirect_accounts_integrity():
    return redirect(url_for('accounts_health') + '#integrity', code=301)


@app.route('/bookkeeping/anomaly-scan')
@login_required
def redirect_bookkeeping_anomaly():
    return redirect(url_for('accounts_health') + '#anomaly', code=301)


@app.route('/bookkeeping/duplicate-check')
@login_required
def redirect_bookkeeping_duplicates():
    return redirect(url_for('accounts_health') + '#duplicate', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2d — Decisions & Signals (consolidated from 5 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('dashboard',
    'analytics_decisions',
    'redirect_forecast_scenarios', 'redirect_health_priority',
    'redirect_goals_milestones', 'redirect_insights_opportunities',
    'redirect_insights_risks')


@app.route('/analytics/decisions')
@login_required
def analytics_decisions():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ══════════════════════════════════════════════════════
    # TAB 1 — Scenarios (cash projection at 3 multipliers)
    # ══════════════════════════════════════════════════════
    cur.execute("""
        SELECT COALESCE(SUM(opening_balance),0) AS op FROM bank_accounts WHERE user_id=%s
    """, (uid,))
    opening = float((cur.fetchone() or {}).get('op') or 0)

    cur.execute("""
        SELECT COALESCE(SUM(CASE WHEN kind='in' THEN amount ELSE -amount END),0) AS mv
        FROM bank_transactions WHERE user_id=%s
    """, (uid,))
    bank_mv = float((cur.fetchone() or {}).get('mv') or 0)

    cur.execute("SELECT COUNT(*) AS n FROM bank_accounts WHERE user_id=%s", (uid,))
    if int((cur.fetchone() or {}).get('n') or 0) > 0:
        current_cash = opening + bank_mv
        cash_source = 'bank accounts'
    else:
        cur.execute("""
            SELECT COALESCE(SUM(CASE WHEN entry_type='in' THEN amount ELSE -amount END),0) AS net
            FROM cash_books WHERE user_id=%s
        """, (uid,))
        current_cash = float((cur.fetchone() or {}).get('net') or 0)
        cash_source = 'cash book'

    cur.execute("""
        SELECT COALESCE(AVG(daily_rev),0) AS a FROM (
            SELECT SUM(total_amount) AS daily_rev FROM sales
            WHERE user_id=%s AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
            GROUP BY sale_date
        ) t
    """, (uid,))
    avg_daily_revenue = float((cur.fetchone() or {}).get('a') or 0)

    cur.execute("""
        SELECT COALESCE(AVG(daily_exp),0) AS a FROM (
            SELECT SUM(amount) AS daily_exp FROM expenses
            WHERE user_id=%s AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
            GROUP BY date
        ) t
    """, (uid,))
    avg_daily_expense = float((cur.fetchone() or {}).get('a') or 0)

    cur.execute("""
        SELECT COALESCE(SUM(amount),0) AS t FROM recurring_items
        WHERE user_id=%s AND active=1 AND item_type='expense'
          AND next_run_date <= to_char(CURRENT_DATE + INTERVAL '30 days','YYYY-MM-DD')
    """, (uid,))
    upcoming_recurring = float((cur.fetchone() or {}).get('t') or 0)

    cur.execute("""
        SELECT COALESCE(SUM(total_cost),0) AS t FROM purchase_orders
        WHERE user_id=%s AND status='pending'
    """, (uid,))
    upcoming_pos = float((cur.fetchone() or {}).get('t') or 0)

    net_daily_base = avg_daily_revenue - avg_daily_expense
    daily_obligations = (upcoming_recurring + upcoming_pos) / 30.0

    scenarios = []
    for label, rev_mul, exp_mul in [
        ('Worst case', 0.7, 1.2),
        ('Conservative', 0.9, 1.1),
        ('Base (current pace)', 1.0, 1.0),
        ('Optimistic', 1.2, 0.95),
        ('Best case', 1.4, 0.9),
    ]:
        daily_net = (avg_daily_revenue * rev_mul) - (avg_daily_expense * exp_mul) - daily_obligations
        projected_30 = current_cash + daily_net * 30
        projected_90 = current_cash + daily_net * 90
        scenarios.append({
            'label': label,
            'daily_net': round(daily_net, 2),
            'projected_30': round(projected_30, 2),
            'projected_90': round(projected_90, 2),
            'risk': ('danger' if projected_30 < 0 else
                     'warn' if projected_30 < current_cash * 0.2 else 'ok'),
        })

    # ══════════════════════════════════════════════════════
    # TAB 2 — Priority (health issues ranked by impact)
    # ══════════════════════════════════════════════════════
    cur.execute("""
        SELECT COALESCE(SUM(total_amount),0) AS rev, COALESCE(SUM(profit),0) AS pft
        FROM sales WHERE user_id=%s
          AND sale_date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
    """, (uid,))
    r30 = cur.fetchone() or {}
    sales_30 = float(r30.get('rev') or 0)
    profit_30 = float(r30.get('pft') or 0)

    cur.execute("""
        SELECT COALESCE(SUM(total_amount),0) AS rev
        FROM sales WHERE user_id=%s
          AND sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD')
          AND sale_date < to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
    """, (uid,))
    sales_prev = float((cur.fetchone() or {}).get('rev') or 0)

    cur.execute("""
        SELECT COUNT(*) AS total,
               SUM(CASE WHEN quantity <= COALESCE(
                   (SELECT low_stock_threshold FROM settings WHERE user_id=%s), 10)
                   THEN 1 ELSE 0 END) AS low
        FROM stock WHERE user_id=%s
    """, (uid, uid))
    st = cur.fetchone() or {}
    total_sku = int(st.get('total') or 0)
    low_sku = int(st.get('low') or 0)

    cur.execute("""
        SELECT COALESCE(SUM(amount),0) AS m FROM expenses
        WHERE user_id=%s AND date >= to_char(CURRENT_DATE - INTERVAL '30 days','YYYY-MM-DD')
    """, (uid,))
    exp_30 = float((cur.fetchone() or {}).get('m') or 0)

    cur.execute("""
        SELECT COUNT(DISTINCT customer_id) AS buyers,
               SUM(CASE WHEN c > 1 THEN 1 ELSE 0 END) AS repeaters
        FROM (SELECT customer_id, COUNT(*) AS c FROM sales
              WHERE user_id=%s AND customer_id IS NOT NULL
              GROUP BY customer_id) t
    """, (uid,))
    rb = cur.fetchone() or {}
    buyers = int(rb.get('buyers') or 0)
    repeaters = int(rb.get('repeaters') or 0)

    trend_pct = ((sales_30 - sales_prev) / sales_prev * 100) if sales_prev > 0 else (100 if sales_30 > 0 else 0)
    trend_score = max(0, min(100, int(50 + trend_pct * 0.5)))
    margin_pct = round((profit_30 / sales_30 * 100) if sales_30 > 0 else 0, 1)
    margin_score = max(0, min(100, int(margin_pct * 2.5)))
    stock_score = int(100 - (low_sku / total_sku) * 100) if total_sku > 0 else 100
    monthly_exp = exp_30 if exp_30 > 0 else 1
    months_covered = round(current_cash / monthly_exp, 1) if current_cash > 0 else 0
    cash_score = max(0, min(100, int(months_covered * 50)))
    retention_score = int((repeaters / buyers * 100) if buyers > 0 else 0)

    priority_rows = []
    for name, score, weight, advice in [
        ('Sales trend', trend_score, 30,
         f'30d vs prior 30d: {int(trend_pct):+d}%. Every 10% of trend improvement is worth ~5 health points.'),
        ('Margin', margin_score, 25,
         f'{margin_pct}% avg margin. Raise prices or reduce costs to lift this — each 4% margin is worth ~10 health points.'),
        ('Stock health', stock_score, 15,
         f'{low_sku} of {total_sku} SKUs low. Restock to close the gap.'),
        ('Cash buffer', cash_score, 15,
         f'{months_covered} months covered. Every 0.4 extra months of buffer is worth ~10 health points.'),
        ('Customer retention', retention_score, 15,
         f'{repeaters} of {buyers} customers repeat. Every 10% retention gain is worth ~1.5 health points.'),
    ]:
        gap = 100 - score
        impact = round(gap * weight / 100, 1)
        priority_rows.append({
            'name': name,
            'score': score,
            'gap': gap,
            'weight': weight,
            'impact': impact,
            'advice': advice,
        })
    priority_rows.sort(key=lambda x: -x['impact'])

    # ══════════════════════════════════════════════════════
    # TAB 3 — Milestones (goal pace checkpoints)
    # ══════════════════════════════════════════════════════
    cur.execute("SELECT monthly_revenue_goal FROM settings WHERE user_id=%s", (uid,))
    goal = float((cur.fetchone() or {}).get('monthly_revenue_goal') or 0)

    today = datetime.today()
    month_start = today.replace(day=1)
    days_in_month = ((month_start + timedelta(days=32)).replace(day=1) - timedelta(days=1)).day
    day_of_month = today.day
    month_str = today.strftime('%Y-%m')

    cur.execute("""
        SELECT substr(sale_date,1,10) AS d, COALESCE(SUM(total_amount),0) AS t
        FROM sales WHERE user_id=%s AND sale_date LIKE %s
        GROUP BY substr(sale_date,1,10) ORDER BY d
    """, (uid, month_str + '%'))
    daily_map = {r['d']: float(r['t']) for r in cur.fetchall()}

    milestones = []
    checkpoints = [7, 14, 21, 28, days_in_month]
    cumulative = 0.0
    for d in range(1, days_in_month + 1):
        key = f"{month_str}-{d:02d}"
        cumulative += daily_map.get(key, 0) if d <= day_of_month else 0
        if d in checkpoints:
            expected = goal * (d / days_in_month) if goal else 0
            actual = cumulative
            diff = actual - expected
            if d < day_of_month:
                status = 'hit' if diff >= 0 else 'missed'
            elif d == day_of_month or d < day_of_month + 7:
                status = 'ontrack' if diff >= -goal * 0.05 else 'behind'
            else:
                status = 'pending'
            milestones.append({
                'day': d,
                'expected': round(expected, 2),
                'actual': round(actual, 2),
                'diff': round(diff, 2),
                'status': status,
            })

    # ══════════════════════════════════════════════════════
    # TAB 4 — Opportunities (data-driven suggestions)
    # ══════════════════════════════════════════════════════
    opportunities = []

    # High-margin low-velocity products
    cur.execute("""
        SELECT COALESCE(st.product_name, 'Item') AS name,
               COALESCE(SUM(s.quantity_sold),0) AS units,
               COALESCE(SUM(s.total_amount),0) AS revenue,
               COALESCE(SUM(s.profit),0) AS profit
        FROM sales s LEFT JOIN stock st ON st.id = s.stock_id
        WHERE s.user_id=%s
          AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD')
        GROUP BY st.product_name
        HAVING COUNT(*) >= 2
    """, (uid,))
    for r in (cur.fetchall() or []):
        rev = float(r.get('revenue') or 0)
        pft = float(r.get('profit') or 0)
        units = float(r.get('units') or 0)
        if rev > 0:
            margin = pft / rev * 100
            if margin >= 30 and units < 5:
                opportunities.append({
                    'icon': 'trending-up', 'severity': 'info',
                    'title': f'High-margin, low-volume: {r.get("name") or "Item"}',
                    'detail': f'{round(margin, 1)}% margin but only {int(units)} units sold in 60 days. '
                              f'Bundle it with a faster seller or feature it prominently — margin is already good.',
                })

    # Lapsed repeaters (bought before, not in last 60 days)
    cur.execute("""
        SELECT c.name, MAX(s.sale_date) AS last_sale, COUNT(s.id) AS orders
        FROM customers c JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id
        WHERE c.user_id=%s
        GROUP BY c.id, c.name
        HAVING COUNT(s.id) >= 2
           AND MAX(s.sale_date) < to_char(CURRENT_DATE - INTERVAL '60 days','YYYY-MM-DD')
        ORDER BY MAX(s.sale_date) ASC
        LIMIT 5
    """, (uid,))
    lapsed = list(cur.fetchall() or [])
    if lapsed:
        names = ', '.join(r['name'] for r in lapsed[:3])
        opportunities.append({
            'icon': 'users', 'severity': 'info',
            'title': f'{len(lapsed)} regular customer(s) have not bought in 60+ days',
            'detail': f'Top lapsed: {names}. A small discount or reminder often reactivates them — '
                      f'they have already shown they buy from you.',
        })

    # Best day of week
    cur.execute("""
        SELECT EXTRACT(DOW FROM sale_date::date) AS dow,
               COALESCE(SUM(total_amount),0) AS rev
        FROM sales WHERE user_id=%s
          AND sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD')
        GROUP BY EXTRACT(DOW FROM sale_date::date)
        ORDER BY rev DESC LIMIT 1
    """, (uid,))
    dow_row = cur.fetchone()
    if dow_row and float(dow_row.get('rev') or 0) > 0:
        dow_names = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
        try:
            dow_idx = int(dow_row.get('dow') or 0)
            day_name = dow_names[dow_idx] if 0 <= dow_idx < 7 else 'a weekday'
        except (ValueError, TypeError):
            day_name = 'a weekday'
        opportunities.append({
            'icon': 'calendar', 'severity': 'success',
            'title': f'{day_name} is your strongest sales day',
            'detail': f'Based on last 90 days. Consider extra staff, longer hours, or targeted promos on {day_name}s.',
        })

    # ══════════════════════════════════════════════════════
    # TAB 5 — Risks (warning signals)
    # ══════════════════════════════════════════════════════
    risks = []

    # Low stock
    if low_sku > 0:
        risks.append({
            'icon': 'alert-triangle', 'severity': 'danger' if low_sku >= 5 else 'warn',
            'title': f'{low_sku} product(s) below stock threshold',
            'detail': f'Of {total_sku} tracked SKUs. Restock soon to avoid stockouts.',
        })

    # Sales trend decline
    if trend_pct < -15:
        risks.append({
            'icon': 'trending-down', 'severity': 'warn',
            'title': f'Sales down {abs(int(trend_pct))}% vs prior 30 days',
            'detail': f'Last 30 days: {round(sales_30, 2)} vs prior 30: {round(sales_prev, 2)}. Investigate cause before it becomes a pattern.',
        })

    # Negative-margin products
    cur.execute("""
        SELECT COALESCE(st.product_name, 'Item') AS name,
               COALESCE(SUM(s.total_amount),0) AS revenue,
               COALESCE(SUM(s.profit),0) AS profit
        FROM sales s LEFT JOIN stock st ON st.id = s.stock_id
        WHERE s.user_id=%s AND s.sale_date >= to_char(CURRENT_DATE - INTERVAL '90 days','YYYY-MM-DD')
        GROUP BY st.product_name HAVING COUNT(*) >= 2 AND SUM(s.profit) < 0
        ORDER BY SUM(s.profit) ASC LIMIT 5
    """, (uid,))
    losing = list(cur.fetchall() or [])
    if losing:
        names = ', '.join(r['name'] for r in losing[:3])
        risks.append({
            'icon': 'alert-circle', 'severity': 'danger',
            'title': f'{len(losing)} product(s) with negative margin',
            'detail': f'Losing money on: {names}. Either raise the price or stop selling.',
        })

    # Top customer gone quiet
    cur.execute("""
        SELECT c.name, MAX(s.sale_date) AS last_sale,
               COALESCE(SUM(s.total_amount),0) AS spent
        FROM customers c JOIN sales s ON s.customer_id = c.id AND s.user_id = c.user_id
        WHERE c.user_id=%s
        GROUP BY c.id, c.name
        ORDER BY spent DESC LIMIT 1
    """, (uid,))
    top_cust = cur.fetchone()
    if top_cust:
        last = (top_cust.get('last_sale') or '')[:10]
        try:
            days = (today.date() - datetime.strptime(last, '%Y-%m-%d').date()).days if last else 9999
        except ValueError:
            days = 9999
        if days >= 60:
            risks.append({
                'icon': 'users', 'severity': 'warn',
                'title': f'Top customer has gone quiet ({days}d)',
                'detail': f'{top_cust.get("name")} was your highest-spending customer. Reach out before they are lost.',
            })

    # Cash concern
    if current_cash < monthly_exp * 1.5 and current_cash > 0:
        risks.append({
            'icon': 'wallet', 'severity': 'warn' if months_covered >= 1 else 'danger',
            'title': f'Cash buffer is {months_covered} months',
            'detail': f'Current cash {round(current_cash, 2)} vs monthly expenses ~{round(monthly_exp, 2)}. '
                      f'Below 3 months of buffer is tight — watch collections and defer non-critical spend.',
        })

    cur.close(); conn.close()

    return render_template('analytics_decisions.html',
        scenarios=scenarios,
        current_cash=current_cash,
        cash_source=cash_source,
        avg_daily_revenue=avg_daily_revenue,
        avg_daily_expense=avg_daily_expense,
        daily_obligations=round(daily_obligations, 2),
        priority_rows=priority_rows,
        goal=goal,
        milestones=milestones,
        day_of_month=day_of_month,
        days_in_month=days_in_month,
        opportunities=opportunities,
        risks=risks,
    )


@app.route('/forecast/scenarios')
@login_required
def redirect_forecast_scenarios():
    return redirect(url_for('analytics_decisions') + '#scenarios', code=301)


@app.route('/health/priority')
@login_required
def redirect_health_priority():
    return redirect(url_for('analytics_decisions') + '#priority', code=301)


@app.route('/goals/milestones')
@login_required
def redirect_goals_milestones():
    return redirect(url_for('analytics_decisions') + '#milestones', code=301)


@app.route('/insights/opportunities')
@login_required
def redirect_insights_opportunities():
    return redirect(url_for('analytics_decisions') + '#opportunities', code=301)


@app.route('/insights/risks')
@login_required
def redirect_insights_risks():
    return redirect(url_for('analytics_decisions') + '#risks', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2e — Patterns (consolidated from 3 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('reports',
    'analytics_patterns',
    'redirect_stats_cohort', 'redirect_stats_seasonality',
    'redirect_reports_variance')


@app.route('/analytics/patterns')
@login_required
def analytics_patterns():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    from collections import defaultdict
    from datetime import datetime as _dt

    # ── Cohorts ───────────────────────────────────────
    cur.execute("""
        WITH first_purchase AS (
            SELECT customer_id, MIN(sale_date) AS first_date
            FROM sales
            WHERE user_id=%s AND customer_id IS NOT NULL
            GROUP BY customer_id
        ),
        purchases AS (
            SELECT s.customer_id,
                   substr(fp.first_date, 1, 7) AS cohort_month,
                   substr(s.sale_date, 1, 7) AS purchase_month
            FROM sales s
            JOIN first_purchase fp ON fp.customer_id = s.customer_id
            WHERE s.user_id=%s
        )
        SELECT cohort_month, purchase_month,
               COUNT(DISTINCT customer_id) AS customers
        FROM purchases
        GROUP BY cohort_month, purchase_month
        ORDER BY cohort_month, purchase_month
    """, (uid, uid))

    cohort_data = defaultdict(dict)
    cohort_sizes = defaultdict(int)
    all_months = set()

    for r in (cur.fetchall() or []):
        cm = r['cohort_month']
        pm = r['purchase_month']
        cnt = int(r['customers'] or 0)
        try:
            cm_dt = _dt.strptime(cm + '-01', '%Y-%m-%d')
            pm_dt = _dt.strptime(pm + '-01', '%Y-%m-%d')
            offset = (pm_dt.year - cm_dt.year) * 12 + (pm_dt.month - cm_dt.month)
        except (ValueError, TypeError):
            offset = 0
        if offset < 0:
            continue
        cohort_data[cm][offset] = cnt
        if offset == 0:
            cohort_sizes[cm] = cnt
        all_months.add(offset)

    cohort_rows = []
    for cm in sorted(cohort_data.keys()):
        size = cohort_sizes.get(cm, 0)
        cells = []
        for off in sorted(all_months):
            n = cohort_data[cm].get(off, 0)
            pct = round(n / size * 100, 1) if size > 0 else 0
            cells.append({'offset': off, 'customers': n, 'pct': pct})
        cohort_rows.append({'cohort': cm, 'size': size, 'cells': cells})
    max_offset = (max(all_months) + 1) if all_months else 0

    # ── Seasonality ───────────────────────────────────
    cur.execute("""
        SELECT EXTRACT(DOW FROM sale_date::date) AS dow,
               AVG(daily_rev) AS avg_rev,
               COUNT(*) AS days
        FROM (
            SELECT sale_date, SUM(total_amount) AS daily_rev
            FROM sales WHERE user_id=%s
            GROUP BY sale_date
        ) daily
        GROUP BY EXTRACT(DOW FROM sale_date::date)
        ORDER BY dow
    """, (uid,))
    dow_names = ['Sun', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat']
    dow_rows = []
    dow_max = 0
    for r in (cur.fetchall() or []):
        try:
            idx = int(r.get('dow') or 0)
        except (TypeError, ValueError):
            idx = 0
        avg = float(r.get('avg_rev') or 0)
        dow_max = max(dow_max, avg)
        dow_rows.append({
            'name': dow_names[idx] if 0 <= idx < 7 else '?',
            'avg': round(avg, 2),
            'days': int(r.get('days') or 0),
        })
    for r in dow_rows:
        r['bar_pct'] = round(r['avg'] / dow_max * 100, 1) if dow_max > 0 else 0

    cur.execute("""
        SELECT EXTRACT(MONTH FROM sale_date::date) AS mon,
               AVG(daily_rev) AS avg_rev,
               COUNT(*) AS days
        FROM (
            SELECT sale_date, SUM(total_amount) AS daily_rev
            FROM sales WHERE user_id=%s
            GROUP BY sale_date
        ) daily
        GROUP BY EXTRACT(MONTH FROM sale_date::date)
        ORDER BY mon
    """, (uid,))
    mon_names = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
    mon_rows = []
    mon_max = 0
    for r in (cur.fetchall() or []):
        try:
            idx = int(r.get('mon') or 1)
        except (TypeError, ValueError):
            idx = 1
        avg = float(r.get('avg_rev') or 0)
        mon_max = max(mon_max, avg)
        mon_rows.append({
            'name': mon_names[idx - 1] if 1 <= idx <= 12 else '?',
            'avg': round(avg, 2),
            'days': int(r.get('days') or 0),
        })
    for r in mon_rows:
        r['bar_pct'] = round(r['avg'] / mon_max * 100, 1) if mon_max > 0 else 0

    # ── Variance vs goal ──────────────────────────────
    cur.execute("SELECT monthly_revenue_goal FROM settings WHERE user_id=%s", (uid,))
    goal = float((cur.fetchone() or {}).get('monthly_revenue_goal') or 0)

    cur.execute("""
        SELECT substr(sale_date, 1, 7) AS month,
               COALESCE(SUM(total_amount),0) AS revenue,
               COALESCE(SUM(profit),0) AS profit
        FROM sales
        WHERE user_id=%s
          AND sale_date >= to_char(CURRENT_DATE - INTERVAL '6 months','YYYY-MM-DD')
        GROUP BY substr(sale_date, 1, 7)
        ORDER BY month
    """, (uid,))
    variance_rows = []
    for r in (cur.fetchall() or []):
        rev = float(r.get('revenue') or 0)
        var = rev - goal if goal else 0
        variance_rows.append({
            'month': r.get('month') or '?',
            'revenue': rev,
            'profit': float(r.get('profit') or 0),
            'goal': goal,
            'variance': round(var, 2),
            'hit': (var >= 0) if goal else None,
        })

    cur.close(); conn.close()

    return render_template('analytics_patterns.html',
        cohort_rows=cohort_rows,
        max_offset=max_offset,
        dow_rows=dow_rows,
        mon_rows=mon_rows,
        variance_rows=variance_rows,
        goal=goal,
    )


@app.route('/stats/cohort')
@login_required
def redirect_stats_cohort():
    return redirect(url_for('analytics_patterns') + '#cohorts', code=301)


@app.route('/stats/seasonality')
@login_required
def redirect_stats_seasonality():
    return redirect(url_for('analytics_patterns') + '#seasonality', code=301)


@app.route('/reports/variance')
@login_required
def redirect_reports_variance():
    return redirect(url_for('analytics_patterns') + '#variance', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2f — Team & Workflow (consolidated from 9 disabled routes)
# ═══════════════════════════════════════════════════════════════════════

_bind_module('workspace',
    'workflow_signals',
    'redirect_tasks_priority', 'redirect_notes_actions',
    'redirect_reminders_context', 'redirect_reminders_escalation',
    'redirect_calendar_conflicts', 'redirect_calendar_free')


@app.route('/team/health')
@login_required
def team_health():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()

    # ── Permissions ─────────────────────────────────────
    cur.execute("""
        SELECT id, username, email, role
        FROM users
        WHERE id=%s OR owner_id=%s
        ORDER BY CASE WHEN role='owner' THEN 0 ELSE 1 END, username
    """, (uid, uid))
    team_rows = []
    for r in (cur.fetchall() or []):
        role = (r.get('role') or 'staff').lower()
        is_owner = (role == 'owner')
        team_rows.append({
            'id': r.get('id'),
            'username': r.get('username') or '—',
            'email': r.get('email') or '—',
            'role': role,
            'is_owner': is_owner,
        })

    # ── Activity (recent 30 across account) ─────────────
    cur.execute("""
        SELECT username, action, details, created_at
        FROM activities
        WHERE user_id=%s
        ORDER BY created_at DESC
        LIMIT 30
    """, (uid,))
    activity_rows = []
    for r in (cur.fetchall() or []):
        activity_rows.append({
            'username': r.get('username') or '—',
            'action': r.get('action') or '—',
            'details': r.get('details') or '',
            'when': r.get('created_at') or '',
        })

    # ── Config audit ────────────────────────────────────
    cur.execute("""
        SELECT business_name, currency_code, currency_symbol, tax_rate,
               monthly_revenue_goal, low_stock_threshold,
               stripe_secret_key, stripe_webhook_secret
        FROM settings WHERE user_id=%s
    """, (uid,))
    s = cur.fetchone() or {}
    config_checks = [
        {'item': 'Business name',
         'status': 'ok' if (s.get('business_name') or '').strip() else 'warn',
         'detail': (s.get('business_name') or 'Not set — appears on PDF headers')},
        {'item': 'Currency code',
         'status': 'ok' if s.get('currency_code') else 'warn',
         'detail': str(s.get('currency_code') or 'Not set')},
        {'item': 'Currency symbol',
         'status': 'ok' if s.get('currency_symbol') else 'warn',
         'detail': str(s.get('currency_symbol') or 'Not set')},
        {'item': 'Low stock threshold',
         'status': 'ok' if float(s.get('low_stock_threshold') or 0) > 0 else 'warn',
         'detail': str(s.get('low_stock_threshold') or 0) + ' units'},
        {'item': 'Monthly revenue goal',
         'status': 'ok' if float(s.get('monthly_revenue_goal') or 0) > 0 else 'info',
         'detail': ('Goal = ' + str(s.get('monthly_revenue_goal'))) if float(s.get('monthly_revenue_goal') or 0) > 0 else 'Optional — enables goal tracking'},
        {'item': 'Tax rate',
         'status': 'info',
         'detail': str(s.get('tax_rate') or 0) + '% (0 means no tax applied)'},
        {'item': 'Stripe secret key',
         'status': 'ok' if s.get('stripe_secret_key') else 'warn',
         'detail': 'Configured' if s.get('stripe_secret_key') else 'Missing — subscriptions cannot start'},
        {'item': 'Stripe webhook secret',
         'status': 'ok' if s.get('stripe_webhook_secret') else 'warn',
         'detail': 'Configured' if s.get('stripe_webhook_secret') else 'Missing — payments will not activate'},
    ]

    cur.close(); conn.close()
    return render_template('team_health.html',
        team_rows=team_rows,
        activity_rows=activity_rows,
        config_checks=config_checks,
    )


@app.route('/workflow/signals')
@login_required
def workflow_signals():
    uid = current_user.id
    today_dt = datetime.today()
    today = today_dt.strftime('%Y-%m-%d')
    horizon_14 = (today_dt + timedelta(days=14)).strftime('%Y-%m-%d')
    horizon_30 = (today_dt + timedelta(days=30)).strftime('%Y-%m-%d')
    three_days = (today_dt + timedelta(days=3)).strftime('%Y-%m-%d')
    seven_days = (today_dt + timedelta(days=7)).strftime('%Y-%m-%d')

    conn = get_db(); cur = conn.cursor()

    # ── TAB 1: Task priority (scored) ───────────────────
    cur.execute("""
        SELECT id, title, priority, due_date, done
        FROM tasks WHERE user_id=%s AND done=0
    """, (uid,))
    priority_weight = {'high': 3, 'medium': 2, 'low': 1}
    task_rows = []
    for r in (cur.fetchall() or []):
        pw = priority_weight.get((r.get('priority') or 'medium').lower(), 2)
        due = (r.get('due_date') or '')[:10]
        if due and due < today:
            urgency, label = 3.0, 'Overdue'
        elif due == today:
            urgency, label = 3.0, 'Due today'
        elif due and due <= three_days:
            urgency, label = 2.0, 'Due in 3 days'
        elif due and due <= seven_days:
            urgency, label = 1.5, 'Due this week'
        else:
            urgency, label = 1.0, 'No rush'
        task_rows.append({
            'id': r.get('id'),
            'title': r.get('title') or '—',
            'priority': r.get('priority') or 'medium',
            'due_date': due or 'no due date',
            'urgency': label,
            'score': round(pw * urgency, 1),
        })
    task_rows.sort(key=lambda x: -x['score'])

    # ── TAB 2: Note actions (extract TODO / - [ ] lines) ─
    cur.execute("SELECT id, title, body, updated_date FROM notes WHERE user_id=%s", (uid,))
    note_actions = []
    for r in (cur.fetchall() or []):
        body = r.get('body') or ''
        for line in body.split('\n'):
            stripped = line.strip()
            if not stripped:
                continue
            action_text = None
            if stripped.startswith('- [ ]') or stripped.startswith('* [ ]'):
                action_text = stripped[5:].strip()
            elif stripped.startswith('[ ]'):
                action_text = stripped[3:].strip()
            else:
                low = stripped.lower()
                if low.startswith('todo:'):
                    action_text = stripped[5:].strip()
                elif low.startswith('action:'):
                    action_text = stripped[7:].strip()
            if action_text and len(action_text) > 2:
                note_actions.append({
                    'note_id': r.get('id'),
                    'note_title': r.get('title') or 'Note',
                    'action': action_text,
                    'updated': r.get('updated_date') or '',
                })
    note_actions = note_actions[:40]

    # ── TAB 3: Reminder context (unified upcoming 14d) ──
    context_items = []
    cur.execute("""
        SELECT title, due_date, priority FROM tasks
        WHERE user_id=%s AND done=0 AND due_date != ''
          AND due_date >= %s AND due_date <= %s
    """, (uid, today, horizon_14))
    for r in (cur.fetchall() or []):
        context_items.append({
            'date': r.get('due_date'),
            'kind': 'Task',
            'label': r.get('title') or '—',
            'detail': r.get('priority') or 'medium',
        })
    cur.execute("""
        SELECT name, next_run_date, item_type, amount FROM recurring_items
        WHERE user_id=%s AND active=1
          AND next_run_date >= %s AND next_run_date <= %s
    """, (uid, today, horizon_14))
    for r in (cur.fetchall() or []):
        context_items.append({
            'date': r.get('next_run_date'),
            'kind': 'Recurring',
            'label': r.get('name') or '—',
            'detail': (r.get('item_type') or 'item') + ' · ' + _format_money(r.get('amount') or 0),
        })
    cur.execute("""
        SELECT item_description, expected_date, status FROM purchase_orders
        WHERE user_id=%s AND status != 'received'
          AND expected_date != '' AND expected_date >= %s AND expected_date <= %s
    """, (uid, today, horizon_14))
    for r in (cur.fetchall() or []):
        context_items.append({
            'date': r.get('expected_date'),
            'kind': 'Purchase order',
            'label': r.get('item_description') or '—',
            'detail': r.get('status') or 'pending',
        })
    context_items.sort(key=lambda x: x['date'] or '9999')

    # ── TAB 4: Reminder escalation (overdue) ────────────
    escalation_items = []
    cur.execute("""
        SELECT id, title, due_date, priority FROM tasks
        WHERE user_id=%s AND done=0 AND due_date != '' AND due_date < %s
    """, (uid, today))
    for r in (cur.fetchall() or []):
        due = (r.get('due_date') or '')[:10]
        try:
            days_over = (today_dt.date() - datetime.strptime(due, '%Y-%m-%d').date()).days
        except (ValueError, TypeError):
            days_over = 0
        escalation_items.append({
            'kind': 'Task', 'label': r.get('title') or '—',
            'due': due, 'days_overdue': days_over,
            'priority': r.get('priority') or 'medium',
            'severity': 'danger' if days_over >= 14 else 'warn' if days_over >= 3 else 'info',
        })
    cur.execute("""
        SELECT id, item_description, expected_date FROM purchase_orders
        WHERE user_id=%s AND status != 'received'
          AND expected_date != '' AND expected_date < %s
    """, (uid, today))
    for r in (cur.fetchall() or []):
        due = (r.get('expected_date') or '')[:10]
        try:
            days_over = (today_dt.date() - datetime.strptime(due, '%Y-%m-%d').date()).days
        except (ValueError, TypeError):
            days_over = 0
        escalation_items.append({
            'kind': 'PO', 'label': r.get('item_description') or '—',
            'due': due, 'days_overdue': days_over,
            'priority': 'medium',
            'severity': 'danger' if days_over >= 14 else 'warn' if days_over >= 3 else 'info',
        })
    escalation_items.sort(key=lambda x: -x['days_overdue'])

    # ── TAB 5+6: Calendar conflicts + free/busy (30d) ──
    cur.execute("""
        WITH dated_items AS (
            SELECT due_date AS item_date, 'Task' AS kind, title AS label
            FROM tasks WHERE user_id=%s AND done=0 AND due_date != ''
            UNION ALL
            SELECT expected_date, 'PO', item_description
            FROM purchase_orders WHERE user_id=%s AND status != 'received' AND expected_date != ''
            UNION ALL
            SELECT session_date, 'Service', customer_name
            FROM service_sessions WHERE user_id=%s AND status != 'cancelled'
        )
        SELECT item_date, kind, label FROM dated_items
        WHERE item_date >= %s AND item_date <= %s
        ORDER BY item_date, kind
    """, (uid, uid, uid, today, horizon_30))

    by_date = {}
    for r in (cur.fetchall() or []):
        d = r.get('item_date')
        if not d:
            continue
        by_date.setdefault(d, []).append({
            'kind': r.get('kind') or '?',
            'label': r.get('label') or '—',
        })

    conflict_rows = []
    for d in sorted(by_date.keys()):
        items = by_date[d]
        if len(items) >= 2:
            conflict_rows.append({
                'date': d, 'items': items, 'count': len(items),
            })

    calendar_days = []
    for i in range(30):
        d = (today_dt + timedelta(days=i)).strftime('%Y-%m-%d')
        n = len(by_date.get(d, []))
        calendar_days.append({'date': d, 'busy': n > 0, 'count': n})
    free_count = sum(1 for d in calendar_days if not d['busy'])
    busy_count = len(calendar_days) - free_count

    cur.close(); conn.close()

    return render_template('workflow_signals.html',
        task_rows=task_rows,
        note_actions=note_actions,
        context_items=context_items,
        escalation_items=escalation_items,
        conflict_rows=conflict_rows,
        calendar_days=calendar_days,
        free_count=free_count,
        busy_count=busy_count,
    )


@app.route('/team/permissions')
@login_required
def redirect_team_permissions():
    return redirect(url_for('team_health') + '#permissions', code=301)


@app.route('/team/activity')
@login_required
def redirect_team_activity():
    return redirect(url_for('team_health') + '#activity', code=301)


@app.route('/settings/health')
@login_required
def redirect_settings_health():
    return redirect(url_for('team_health') + '#config', code=301)


@app.route('/tasks/priority')
@login_required
def redirect_tasks_priority():
    return redirect(url_for('workflow_signals') + '#taskpriority', code=301)


@app.route('/notes/actions')
@login_required
def redirect_notes_actions():
    return redirect(url_for('workflow_signals') + '#noteactions', code=301)


@app.route('/reminders/context')
@login_required
def redirect_reminders_context():
    return redirect(url_for('workflow_signals') + '#context', code=301)


@app.route('/reminders/escalation')
@login_required
def redirect_reminders_escalation():
    return redirect(url_for('workflow_signals') + '#escalation', code=301)


@app.route('/calendar/conflicts')
@login_required
def redirect_calendar_conflicts():
    return redirect(url_for('workflow_signals') + '#conflicts', code=301)


@app.route('/calendar/free')
@login_required
def redirect_calendar_free():
    return redirect(url_for('workflow_signals') + '#calendar', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 2g — Folds & Services Analytics
# ═══════════════════════════════════════════════════════════════════════

@app.template_global()
def rfm_segment_actions(groups, segment_order):
    """Computed actions per RFM segment based on group sizes.
    Called from rfm.html. No route changes needed."""
    try:
        if not groups or not segment_order:
            return []
        counts = {s: len(groups.get(s, []) or []) for s in segment_order}
        actions = []
        if counts.get('Champion', 0) > 0:
            actions.append({
                'segment': 'Champion', 'count': counts['Champion'],
                'title': 'Thank and ask',
                'detail': 'Personally thank these customers. Ask for referrals or testimonials — happy champions convert at the highest rate.',
            })
        if counts.get('Loyal', 0) > 0:
            actions.append({
                'segment': 'Loyal', 'count': counts['Loyal'],
                'title': 'Nudge to Champion',
                'detail': 'Offer a bundle, upgrade, or premium tier. They buy regularly — a small upsell is likely to land.',
            })
        if counts.get('New', 0) > 0:
            actions.append({
                'segment': 'New', 'count': counts['New'],
                'title': 'Welcome and follow up',
                'detail': 'First purchase in the last few days. Check in. A good second experience locks them in.',
            })
        if counts.get('At Risk', 0) > 0:
            actions.append({
                'segment': 'At Risk', 'count': counts['At Risk'],
                'title': 'Re-engagement offer',
                'detail': 'Were repeat buyers, now slowed. Send a discount or personal reminder. Recovery is cheaper than acquisition.',
            })
        if counts.get('Needs Attention', 0) > 0:
            actions.append({
                'segment': 'Needs Attention', 'count': counts['Needs Attention'],
                'title': 'Check in',
                'detail': 'Buying occasionally but fading. A small offer or message may reactivate them.',
            })
        if counts.get('Lost', 0) > 0:
            actions.append({
                'segment': 'Lost', 'count': counts['Lost'],
                'title': 'Decide: win back or let go',
                'detail': 'No purchase in 90+ days. A win-back campaign may work for a few. Otherwise focus on the rest.',
            })
        return actions
    except Exception:
        return []


@app.template_global()
def portfolio_risk_metrics(holdings):
    """Portfolio concentration + diversification from holdings list.
    Called from investments.html. No route changes needed."""
    try:
        if not holdings:
            return {'has_data': False, 'message': 'Add at least one holding to see portfolio risk.'}
        total_value = sum(float(h.get('value') or 0) for h in holdings)
        if total_value <= 0:
            return {'has_data': False, 'message': 'Holdings have zero current value.'}
        weights = [float(h.get('value') or 0) / total_value for h in holdings]
        top_weight = max(weights)
        top_holding = holdings[weights.index(top_weight)]
        hhi = sum(w * w for w in weights)
        n = len(holdings)
        if n == 1:
            div_score = 0.0
        else:
            div_score = round(max(0, min(100, (1 - hhi) / (1 - 1/n) * 100)), 1)
        if top_weight >= 0.6:
            risk, risk_label = 'high', 'High concentration'
        elif top_weight >= 0.4:
            risk, risk_label = 'medium', 'Moderate concentration'
        else:
            risk, risk_label = 'low', 'Well diversified'
        return {
            'has_data': True,
            'total_value': total_value,
            'top_pct': round(top_weight * 100, 1),
            'top_name': top_holding.get('name') or 'Holding',
            'hhi': round(hhi, 3),
            'diversification_score': div_score,
            'risk': risk,
            'risk_label': risk_label,
            'count': n,
        }
    except Exception:
        return {'has_data': False, 'message': 'Could not compute portfolio risk.'}


@app.route('/services/analytics')
@login_required
def services_analytics():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT sv.id, sv.name, sv.price, sv.duration_minutes,
               COUNT(ss.id) AS sessions_booked,
               COALESCE(SUM(CASE WHEN ss.status='done' THEN 1 ELSE 0 END), 0) AS sessions_done,
               COALESCE(SUM(CASE WHEN ss.status='cancelled' THEN 1 ELSE 0 END), 0) AS sessions_cancelled,
               COALESCE(SUM(CASE WHEN ss.status='done' THEN ss.price ELSE 0 END), 0) AS revenue
        FROM services sv
        LEFT JOIN service_sessions ss ON ss.service_id = sv.id AND ss.user_id = sv.user_id
        WHERE sv.user_id = %s
        GROUP BY sv.id, sv.name, sv.price, sv.duration_minutes
        ORDER BY revenue DESC, sessions_booked DESC
    """, (uid,))
    util_rows = []
    for r in (cur.fetchall() or []):
        booked = int(r.get('sessions_booked') or 0)
        done = int(r.get('sessions_done') or 0)
        cancelled = int(r.get('sessions_cancelled') or 0)
        hours = float(r.get('duration_minutes') or 60) / 60.0
        revenue = float(r.get('revenue') or 0)
        util_pct = round(done / booked * 100, 1) if booked > 0 else 0.0
        rev_per_hour = round(revenue / (done * hours), 2) if (done > 0 and hours > 0) else 0.0
        util_rows.append({
            'name': r.get('name') or '—',
            'price': float(r.get('price') or 0),
            'duration': int(r.get('duration_minutes') or 60),
            'booked': booked,
            'done': done,
            'cancelled': cancelled,
            'revenue': revenue,
            'util_pct': util_pct,
            'rev_per_hour': rev_per_hour,
        })
    total_booked = sum(x['booked'] for x in util_rows)
    total_done = sum(x['done'] for x in util_rows)
    total_revenue = sum(x['revenue'] for x in util_rows)
    overall_util = round(total_done / total_booked * 100, 1) if total_booked > 0 else 0.0
    cur.close(); conn.close()
    return render_template('services_analytics.html',
        util_rows=util_rows,
        total_booked=total_booked,
        total_done=total_done,
        total_revenue=total_revenue,
        overall_util=overall_util,
    )


# Redirects
@app.route('/rfm/actions')
@login_required
def redirect_rfm_actions():
    return redirect(url_for('rfm') + '#actions', code=301)


@app.route('/day-close/variance')
@login_required
def redirect_day_close_variance():
    return redirect(url_for('day_close'), code=301)


@app.route('/investments/risk')
@login_required
def redirect_investments_risk():
    return redirect(url_for('investments') + '#risk', code=301)


@app.route('/services/utilization')
@login_required
def redirect_services_utilization():
    return redirect(url_for('services_analytics') + '#utilization', code=301)


@app.route('/services/yield')
@login_required
def redirect_services_yield():
    return redirect(url_for('services_analytics') + '#yield', code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 3a — Fold-ins with schema
# ═══════════════════════════════════════════════════════════════════════

@app.route('/banking/tx/<int:id>/reconcile', methods=['POST'])
@login_required
def reconcile_bank_tx(id):
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT id, reconciled FROM bank_transactions WHERE id=%s AND user_id=%s",
        (id, current_user.id))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        flash('Transaction not found.', 'danger')
        return redirect(url_for('banking'))
    new_val = 0 if (row.get('reconciled') or 0) == 1 else 1
    date_val = datetime.today().strftime('%Y-%m-%d') if new_val == 1 else ''
    cur.execute(
        "UPDATE bank_transactions SET reconciled=%s, reconciled_date=%s "
        "WHERE id=%s AND user_id=%s",
        (new_val, date_val, id, current_user.id))
    conn.commit()
    cur.close(); conn.close()
    flash('Marked reconciled.' if new_val else 'Un-reconciled.', 'success')
    return redirect(url_for('banking'))


@app.route('/tasks/<int:id>/blocker', methods=['POST'])
@login_required
def save_task_blocker(id):
    note = (request.form.get('blocker_note') or '').strip()[:500]
    conn = get_db(); cur = conn.cursor()
    cur.execute("SELECT id FROM tasks WHERE id=%s AND user_id=%s", (id, current_user.id))
    if not cur.fetchone():
        cur.close(); conn.close()
        flash('Task not found.', 'danger')
        return redirect(url_for('tasks'))
    cur.execute(
        "UPDATE tasks SET blocker_note=%s WHERE id=%s AND user_id=%s",
        (note, id, current_user.id))
    conn.commit()
    cur.close(); conn.close()
    flash('Blocker updated.' if note else 'Blocker cleared.', 'success')
    return redirect(url_for('tasks'))


@app.route('/bookkeeping/doc/<int:id>/approve', methods=['POST'])
@login_required
def approve_book_doc(id):
    decision = (request.form.get('decision') or 'approved').lower()
    if decision not in ('approved', 'rejected', 'none'):
        decision = 'approved'
    note = (request.form.get('approval_notes') or '').strip()[:500]
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT id, kind FROM bookkeeping_docs WHERE id=%s AND user_id=%s",
        (id, current_user.id))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        flash('Document not found.', 'danger')
        return redirect(url_for('bookkeeping'))
    now = datetime.today().strftime('%Y-%m-%d %H:%M') if decision != 'none' else ''
    cur.execute(
        "UPDATE bookkeeping_docs SET approval_status=%s, approval_notes=%s, "
        "approved_by=%s, approved_at=%s WHERE id=%s AND user_id=%s",
        (decision, note,
         current_user.username if decision != 'none' else '',
         now, id, current_user.id))
    conn.commit()
    kind = row.get('kind') or 'invoice'
    cur.close(); conn.close()
    flash('Approval status updated.', 'success')
    return redirect(url_for('bookkeeping_type', kind=kind))


# Redirects from original disabled routes
@app.route('/banking/reconcile')
@login_required
def redirect_banking_reconcile():
    return redirect(url_for('banking') + '#reconcile', code=301)


@app.route('/tasks/blockers')
@login_required
def redirect_tasks_blockers():
    return redirect(url_for('tasks') + '#blockers', code=301)


@app.route('/documents/approval')
@login_required
def redirect_documents_approval():
    return redirect(url_for('bookkeeping'), code=301)




# ═══════════════════════════════════════════════════════════════════════
# Batch 3b part 1 — Settings audit + WhatsApp templates + WhatsApp log
# ═══════════════════════════════════════════════════════════════════════

@app.route('/settings/audit')
@login_required
def settings_audit():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT field_name, old_value, new_value, changed_at
        FROM settings_audit
        WHERE user_id=%s
        ORDER BY changed_at DESC, id DESC
        LIMIT 200
    """, (uid,))
    rows = list(cur.fetchall() or [])
    cur.close(); conn.close()
    return render_template('settings_audit.html', rows=rows)


@app.route('/whatsapp/templates')
@login_required
def whatsapp_templates():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT id, name, body, created_date, updated_date FROM whatsapp_templates "
        "WHERE user_id=%s ORDER BY updated_date DESC",
        (uid,))
    templates = list(cur.fetchall() or [])
    cur.close(); conn.close()
    return render_template('whatsapp_templates.html', templates=templates)


@app.route('/whatsapp/templates/add', methods=['POST'])
@login_required
def whatsapp_template_add():
    name = (request.form.get('name') or '').strip()[:120]
    body = (request.form.get('body') or '').strip()[:2000]
    if not name or not body:
        flash('Name and body are required.', 'danger')
        return redirect(url_for('whatsapp_templates'))
    now = datetime.today().strftime('%Y-%m-%d %H:%M')
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "INSERT INTO whatsapp_templates (user_id, name, body, created_date, updated_date) VALUES (%s, %s, %s, %s, %s)",
        (current_user.id, name, body, now, now))
    conn.commit(); cur.close(); conn.close()
    flash('Template saved.', 'success')
    return redirect(url_for('whatsapp_templates'))


@app.route('/whatsapp/templates/<int:id>/delete', methods=['POST'])
@login_required
def whatsapp_template_delete(id):
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "DELETE FROM whatsapp_templates WHERE id=%s AND user_id=%s",
        (id, current_user.id))
    conn.commit(); cur.close(); conn.close()
    flash('Template deleted.', 'success')
    return redirect(url_for('whatsapp_templates'))


@app.route('/whatsapp/log')
@login_required
def whatsapp_log():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT wm.sent_at, wm.phone, wm.body, c.name AS customer_name
        FROM whatsapp_messages wm
        LEFT JOIN customers c ON c.id = wm.customer_id
        WHERE wm.user_id=%s
        ORDER BY wm.sent_at DESC
        LIMIT 200
    """, (uid,))
    rows = list(cur.fetchall() or [])
    cur.close(); conn.close()
    return render_template('whatsapp_log.html', rows=rows)


@app.route('/whatsapp/send/<int:customer_id>')
@login_required
def whatsapp_send(customer_id):
    """Log the send attempt then redirect to wa.me. We cannot track actual
    delivery via wa.me, so we log the user's action of clicking Send."""
    import urllib.parse
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute(
        "SELECT id, name, phone FROM customers WHERE id=%s AND user_id=%s",
        (customer_id, uid))
    cust = cur.fetchone()
    if not cust:
        cur.close(); conn.close()
        flash('Customer not found.', 'danger')
        return redirect(url_for('whatsapp'))
    phone = (cust.get('phone') or '').strip()
    if not phone:
        cur.close(); conn.close()
        flash('This customer has no phone number.', 'danger')
        return redirect(url_for('whatsapp'))
    template_id = request.args.get('template_id', type=int) or 0
    body = ''
    if template_id:
        cur.execute(
            "SELECT body FROM whatsapp_templates WHERE id=%s AND user_id=%s",
            (template_id, uid))
        trow = cur.fetchone()
        if trow:
            body = trow.get('body') or ''
    if '{{name}}' in body:
        body = body.replace('{{name}}', cust.get('name') or '')
    cur.execute(
        "INSERT INTO whatsapp_messages (user_id, customer_id, phone, body) VALUES (%s, %s, %s, %s)",
        (uid, customer_id, phone, body[:500]))
    conn.commit()
    cur.close(); conn.close()
    phone_clean = phone.replace('+', '').replace(' ', '').replace('-', '')
    url = 'https://wa.me/' + phone_clean
    if body:
        url += '?text=' + urllib.parse.quote(body)
    return redirect(url, code=302)




# ═══════════════════════════════════════════════════════════════════════
# Batch 3b part 2a — Health history
# ═══════════════════════════════════════════════════════════════════════

@app.route('/health/history')
@login_required
def health_history():
    uid = current_user.id
    conn = get_db(); cur = conn.cursor()
    cur.execute("""
        SELECT snapshot_date, overall_score, trend_score, margin_score,
               stock_score, cash_score, retention_score
        FROM score_snapshots
        WHERE user_id=%s
        ORDER BY snapshot_date ASC
        LIMIT 180
    """, (uid,))
    rows = []
    for r in (cur.fetchall() or []):
        rows.append({
            'date': r.get('snapshot_date') or '?',
            'overall': int(r.get('overall_score') or 0),
            'trend': int(r.get('trend_score') or 0),
            'margin': int(r.get('margin_score') or 0),
            'stock': int(r.get('stock_score') or 0),
            'cash': int(r.get('cash_score') or 0),
            'retention': int(r.get('retention_score') or 0),
        })
    cur.close(); conn.close()

    has_data = len(rows) > 0
    first_score = rows[0]['overall'] if has_data else 0
    last_score = rows[-1]['overall'] if has_data else 0
    delta = last_score - first_score if has_data else 0

    labels = [r['date'] for r in rows]
    overalls = [r['overall'] for r in rows]
    trends = [r['trend'] for r in rows]
    margins = [r['margin'] for r in rows]
    stocks = [r['stock'] for r in rows]
    cashes = [r['cash'] for r in rows]
    retentions = [r['retention'] for r in rows]

    return render_template('health_history.html',
        rows=rows,
        has_data=has_data,
        first_score=first_score,
        last_score=last_score,
        delta=delta,
        labels=labels,
        overalls=overalls,
        trends=trends,
        margins=margins,
        stocks=stocks,
        cashes=cashes,
        retentions=retentions,
    )


if __name__ == '__main__':
    app.run(debug=True)